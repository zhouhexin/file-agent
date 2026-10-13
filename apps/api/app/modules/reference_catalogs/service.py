"""Deploy and query a shared school journal catalog without reading XLS at request time."""

from __future__ import annotations

import hashlib
import json
import os
import re
import unicodedata
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.db.models import (
    Document, DocumentVersion, FileObject, ReferenceCatalogVersion, User,
)
from app.modules.file_lifecycle.shared_workspace import get_shared_workspace_id
from app.modules.files.office_conversion import LegacyOfficeConversionService


CATALOG_KEY = "xaut-journals-2019"
CATALOG_ID = str(uuid5(NAMESPACE_URL, f"file-agent:reference-catalog:{CATALOG_KEY}"))
CATALOG_FILENAME = "2019版学术期刊A++、A+,A,B类期刊目录.xls"
ASSET_PATH = Path(__file__).parent / "assets" / "xaut-journals-2019.xls"
ASSET_SHA256 = "273f8a8cf47268a5016da26a79133a0e29ea4a6edd24c78796ccec6d02c06775"
SYSTEM_USERNAME = "__file_agent_reference_catalog__"
_ISSN_RE = re.compile(r"(?<!\d)\d{4}-?\d{3}[0-9Xx](?!\d)")


class ReferenceCatalogError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_issn(value: str) -> str:
    return re.sub(r"[^0-9X]", "", unicodedata.normalize("NFKC", value).upper())


def normalize_title(value: str) -> str:
    return "".join(unicodedata.normalize("NFKC", value).casefold().split())


def _storage_path(relative: str, *, root: Path) -> Path:
    candidate = (root / relative).resolve()
    if not candidate.is_relative_to(root.resolve()):
        raise ReferenceCatalogError("REFERENCE_CATALOG_PATH_INVALID", "参考目录存储路径越界")
    return candidate


def build_index(xlsx_path: Path) -> dict:
    """Only the two published visible sheets may determine a school journal grade."""
    book = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if "Sheet1" not in book.sheetnames or "A+、A、B类期刊目录" not in book.sheetnames:
            raise ReferenceCatalogError("REFERENCE_CATALOG_LAYOUT_INVALID", "2019 目录缺少正式可见工作表")
        rules_sheet = book["Sheet1"]
        entries_sheet = book["A+、A、B类期刊目录"]
        if rules_sheet.sheet_state != "visible" or entries_sheet.sheet_state != "visible":
            raise ReferenceCatalogError("REFERENCE_CATALOG_LAYOUT_INVALID", "正式目录工作表不可见")
        rows = entries_sheet.iter_rows(min_row=2, max_col=5, values_only=True)
        first = next(rows, None)
        header = [first[column - 1] for column in (2, 4, 5)] if first else []
        if [str(item or "").strip() for item in header] != ["期刊名称", "刊号", "所属类别"]:
            raise ReferenceCatalogError("REFERENCE_CATALOG_LAYOUT_INVALID", "逐刊目录 B/D/E 表头不匹配")
        grade_map = {"A+": "A+", "第三类A": "A", "第四类B": "B"}
        entries: list[dict] = []
        for row, cells in enumerate(rows, start=3):
            title = str(cells[1] or "").strip()
            serial = str(cells[3] or "").strip()
            category = str(cells[4] or "").strip()
            if not title and not serial and not category:
                continue
            grade = grade_map.get("".join(category.split()))
            if not title or grade is None:
                raise ReferenceCatalogError("REFERENCE_CATALOG_ROW_INVALID", f"正式目录第 {row} 行内容不完整")
            entries.append({
                "sheet": entries_sheet.title, "row": row, "journal_name": title,
                "issn": serial, "issn_keys": sorted({normalize_issn(match.group()) for match in _ISSN_RE.finditer(serial)}),
                "category": category, "school_grade": grade,
            })
        if len(entries) != 621:
            raise ReferenceCatalogError("REFERENCE_CATALOG_COUNT_INVALID", f"2019 正式目录应有 621 条，实际 {len(entries)} 条")
        rules = []
        for row in rules_sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.strip():
                    rules.append({"cell": cell.coordinate, "text": value.strip()})
        return {"catalog_key": CATALOG_KEY, "entry_count": len(entries), "entries": entries,
                "rules_sheet": rules_sheet.title, "rules": rules}
    finally:
        book.close()


def seed_catalog(db: Session, *, asset_path: Path = ASSET_PATH) -> ReferenceCatalogVersion:
    """Explicit deployment operation; never called from a user fill task."""
    source = asset_path.resolve()
    if not source.is_file() or _sha256(source) != ASSET_SHA256:
        raise ReferenceCatalogError("REFERENCE_CATALOG_ASSET_INVALID", "打包的 2019 目录缺失或哈希改变")
    existing = db.query(ReferenceCatalogVersion).filter_by(
        catalog_key=CATALOG_KEY, source_sha256=ASSET_SHA256).one_or_none()
    if existing is not None:
        if existing.status != "ACTIVE":
            raise ReferenceCatalogError("REFERENCE_CATALOG_VERSION_INACTIVE", "打包目录已经导入，但当前版本未启用")
        _load_index(existing)
        return existing

    root = Path(get_settings().file_storage_root).resolve()
    relative_original = f"reference-catalogs/{CATALOG_KEY}/{ASSET_SHA256}/original.xls"
    original = _storage_path(relative_original, root=root)
    original.parent.mkdir(parents=True, exist_ok=True)
    if original.exists():
        if _sha256(original) != ASSET_SHA256:
            raise ReferenceCatalogError("REFERENCE_CATALOG_ORIGINAL_CHANGED", "服务器已存目录原件哈希不符")
    else:
        staged = original.with_suffix(".xls.tmp")
        staged.write_bytes(source.read_bytes())
        if _sha256(staged) != ASSET_SHA256:
            staged.unlink(missing_ok=True)
            raise ReferenceCatalogError("REFERENCE_CATALOG_COPY_FAILED", "参考目录复制校验失败")
        os.replace(staged, original)

    system_user = db.query(User).filter_by(username=SYSTEM_USERNAME).one_or_none()
    if system_user is None:
        system_user = User(username=SYSTEM_USERNAME, role="system", display_name="系统参考目录", password_hash="!")
        db.add(system_user)
        db.flush()
    document = Document(user_id=system_user.id, workspace_id=get_shared_workspace_id(db),
                        original_filename=CATALOG_FILENAME, content_type="application/vnd.ms-excel",
                        size_bytes=original.stat().st_size, sha256=ASSET_SHA256,
                        status="READY", ingest_status="READY")
    db.add(document)
    db.flush()
    version = DocumentVersion(document_id=document.id, version_number=1,
                              storage_tier="REFERENCE_CATALOG", storage_path=relative_original,
                              filename=CATALOG_FILENAME, content_type="application/vnd.ms-excel",
                              size_bytes=original.stat().st_size, sha256=ASSET_SHA256,
                              source_type="SYSTEM_REFERENCE", created_by=system_user.id)
    db.add(version)
    db.add(FileObject(document_id=document.id, storage_backend="local",
                      storage_path=relative_original, size_bytes=original.stat().st_size, sha256=ASSET_SHA256))
    db.flush()
    derivative = LegacyOfficeConversionService(db=db).get_or_create_xlsx(
        document=document, document_version=version, source_path=original)
    index = build_index(derivative.file_path)
    index.update(source_sha256=ASSET_SHA256, derivative_sha256=derivative.sha256,
                 document_version_id=version.id, derivative_artifact_id=derivative.artifact_id)
    encoded = json.dumps(index, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    relative_index = f"reference-catalogs/{CATALOG_KEY}/{ASSET_SHA256}/index.json"
    index_path = _storage_path(relative_index, root=root)
    staged_index = index_path.with_suffix(".json.tmp")
    staged_index.write_bytes(encoded)
    os.replace(staged_index, index_path)
    index_sha256 = hashlib.sha256(encoded).hexdigest()
    previous = db.query(ReferenceCatalogVersion).filter_by(catalog_key=CATALOG_KEY).all()
    for item in previous:
        if item.status == "ACTIVE":
            item.status = "ARCHIVED"
    record = ReferenceCatalogVersion(catalog_key=CATALOG_KEY,
        version_number=max((item.version_number for item in previous), default=0) + 1,
        document_version_id=version.id, derivative_artifact_id=derivative.artifact_id,
        source_sha256=ASSET_SHA256, index_storage_path=relative_index,
        index_sha256=index_sha256, entry_count=len(index["entries"]), status="ACTIVE")
    db.add(record)
    db.flush()
    return record


def active_catalog(db: Session) -> ReferenceCatalogVersion:
    records = db.query(ReferenceCatalogVersion).filter_by(catalog_key=CATALOG_KEY, status="ACTIVE").all()
    if len(records) != 1:
        raise ReferenceCatalogError("REFERENCE_CATALOG_UNAVAILABLE", "2019 版共享目录未部署或启用版本不唯一")
    return records[0]


def _load_index(record: ReferenceCatalogVersion) -> dict:
    path = _storage_path(record.index_storage_path, root=Path(get_settings().file_storage_root))
    if not path.is_file() or _sha256(path) != record.index_sha256:
        raise ReferenceCatalogError("REFERENCE_CATALOG_INDEX_CHANGED", "共享目录索引缺失或已变化")
    data = json.loads(path.read_text(encoding="utf-8"))
    if (data.get("source_sha256") != record.source_sha256
            or data.get("document_version_id") != record.document_version_id
            or len(data.get("entries", [])) != record.entry_count):
        raise ReferenceCatalogError("REFERENCE_CATALOG_INDEX_CHANGED", "共享目录索引与版本记录不一致")
    return data


def catalog_info(db: Session) -> dict:
    record = active_catalog(db)
    _load_index(record)
    return {"catalog_id": CATALOG_ID, "catalog_key": CATALOG_KEY,
            "catalog_version_id": record.id, "version_number": record.version_number,
            "source_sha256": record.source_sha256, "document_version_id": record.document_version_id,
            "entry_count": record.entry_count, "status": record.status}


def query_catalog(db: Session, *, issn: str | None = None,
                  journal_name: str | None = None, include_rules: bool = False) -> dict:
    record = active_catalog(db)
    index = _load_index(record)
    issn_key = normalize_issn(issn or "")
    name_key = normalize_title(journal_name or "")
    if not issn_key and not name_key and not include_rules:
        raise ReferenceCatalogError("REFERENCE_CATALOG_QUERY_REQUIRED", "须提供刊号或完整期刊名")
    by_issn = [entry for entry in index["entries"] if issn_key in entry["issn_keys"]] if issn_key else []
    by_name = [entry for entry in index["entries"]
               if normalize_title(entry["journal_name"]) == name_key] if name_key else []
    if issn_key and name_key and {e["row"] for e in by_issn} != {e["row"] for e in by_name}:
        matches = []
        status = "CONFLICT"
    else:
        matches = by_issn if issn_key else by_name
        if issn_key and name_key and not by_issn:
            matches = []  # A supplied ISSN cannot silently fall back to a name hit.
        status = "FOUND" if len(matches) == 1 else "AMBIGUOUS" if matches else "NOT_FOUND"
    result = {"catalog_id": CATALOG_ID, "catalog_key": CATALOG_KEY,
              "catalog_version_id": record.id, "version_number": record.version_number,
              "source_sha256": record.source_sha256, "document_version_id": record.document_version_id,
              "entry_count": record.entry_count, "status": record.status,
              "match_status": status, "matches": matches[:10],
              "match_basis": "ISSN" if issn_key else "FULL_NAME" if name_key else "RULES_ONLY"}
    if include_rules:
        result["rules_sheet"] = index["rules_sheet"]
        result["rules"] = [rule for rule in index["rules"] if any(
            signal in rule["text"] for signal in ("SCI", "A++", "A+", "第三类", "第四类", "二区", "一区", "三区", "四区"))][:40]
    return result

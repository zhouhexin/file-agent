"""复杂 Excel 模板的显式布局契约与只读字段映射，不猜测数据写入区域。"""

import re
from openpyxl.utils.cell import range_boundaries, get_column_letter
from pydantic import BaseModel, ConfigDict, Field, model_validator


class WorkbookLayout(BaseModel):
    """表头和填写范围按真实坐标绑定；固定范围填满后必须由用户选择处理方式。"""

    model_config = ConfigDict(extra="forbid")
    header_range: str = Field(pattern=r"^[A-Z]{1,3}[1-9][0-9]*:[A-Z]{1,3}[1-9][0-9]*$")
    data_range: str = Field(pattern=r"^[A-Z]{1,3}[1-9][0-9]*:[A-Z]{1,3}[1-9][0-9]*$")

    @model_validator(mode="after")
    def validate_ranges(self):
        """限定宽度和行数，拒绝表头/数据重叠及反向区域。"""
        left, top, right, bottom = range_boundaries(self.header_range)
        dl, dt, dr, db = range_boundaries(self.data_range)
        if not (1 <= left <= right <= 16384 and right - left < 40 and 1 <= top <= bottom <= 100):
            raise ValueError("表头范围必须位于前 100 行且不超过 40 列")
        if (dl, dr) != (left, right) or not bottom < dt <= db <= 10000 or db - dt >= 1000:
            raise ValueError("填写区域必须与表头同列、不重叠，且不超过 1000 行")
        return self


def blank_column_label(column: int) -> str:
    """以物理列号标识无标题的保留列；不写入原工作簿表头。"""
    return f"〔空白列:{get_column_letter(column)}〕"


def is_blank_column_label(label: str) -> bool:
    return re.fullmatch(r"〔空白列:[A-Z]{1,3}〕", label) is not None


def layout_headers(sheet, layout: WorkbookLayout) -> list[str]:
    """合并表头按锚点展开，重复叶标题用实际列号消歧，不修改模板。"""
    left, top, right, bottom = range_boundaries(layout.header_range)
    labels = []
    for col in range(left, right + 1):
        parts = []
        for row in range(top, bottom + 1):
            cell = sheet.cell(row, col)
            for merged in sheet.merged_cells.ranges:
                if cell.coordinate in merged:
                    if not (top <= merged.min_row <= merged.max_row <= bottom and
                            left <= merged.min_col <= merged.max_col <= right):
                        raise ValueError("表头区域不能截断合并单元格")
                    cell = sheet.cell(merged.min_row, merged.min_col)
                    break
            if cell.data_type == "f":
                raise ValueError("表头不能使用公式")
            value = " ".join(str(cell.value or "").split())
            if value and (not parts or value != parts[-1]):
                parts.append(value)
        label = " / ".join(parts) if parts else blank_column_label(col)
        if parts and is_blank_column_label(label):
            raise ValueError("表头不能使用系统保留的空白列标记")
        labels.append(label)
    if all(is_blank_column_label(label) for label in labels):
        raise ValueError("表头范围没有可填写的字段")
    # 重复标题仍保留原文，以列坐标建立唯一字段；同名字段不再被错误合并。
    return [f"{label} [{get_column_letter(left + i)}]" if labels.count(label) > 1 else label
            for i, label in enumerate(labels)]


def sequence_offsets(labels: list[str]) -> set[int]:
    """仅原表明确的序号列可以作为预置值保留，其他有值单元格不得覆盖。"""
    return {i for i, label in enumerate(labels) if label in {"序号", "编号"}}


def read_layout(sheet, layout: WorkbookLayout):
    """读取固定表头和数据区，拒绝跨行合并的数据区域，排除仅有预置序号的空行。"""
    labels = layout_headers(sheet, layout)
    left, top, right, bottom = range_boundaries(layout.data_range)
    for merged in sheet.merged_cells.ranges:
        if not (merged.max_row < top or merged.min_row > bottom or merged.max_col < left or merged.min_col > right):
            raise ValueError("填写区域含合并单元格，不能按记录行写入")
    sequences = sequence_offsets(labels)
    rows = []
    for row in sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
        values = [c.value for c in row]
        if any(v is not None and str(v).strip() for i, v in enumerate(values) if i not in sequences):
            rows.append(values)
    return labels, rows

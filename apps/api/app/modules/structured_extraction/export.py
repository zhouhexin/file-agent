"""结构化抽取 CSV/XLSX 派生件生成与安全下载投影。"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import posixpath
import re
import tempfile
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries
from app.modules.structured_extraction.workbook_layout import WorkbookLayout, read_layout, sequence_offsets

from app.core.config import Settings, get_settings
from app.db.models import DocumentArtifact, DocumentVersion, StructuredExtractionRun
from app.modules.files.artifact_repository import DocumentArtifactRepository
from app.modules.structured_extraction.schemas import StructuredExtractionResult


class StructuredExtractionExportService:
    """只把规范化 JSON 投影成派生文件，不读取或改写原件。"""

    def __init__(self, *, db: Any, settings: Settings | None = None) -> None:
        self.db = db
        self.settings = settings or get_settings()
        self.repository = DocumentArtifactRepository(db)

    def ensure_export(
        self,
        *,
        run: StructuredExtractionRun,
        result: StructuredExtractionResult,
        force: bool = False,
    ) -> dict[str, Any] | None:
        """按运行配置幂等创建 CSV 或 XLSX，并返回不含本地路径的投影。"""

        presentation = str(run.presentation or "AUTO").upper()
        if presentation not in {"CSV", "XLSX"}:
            return None
        version = self.db.get(DocumentVersion, run.document_version_id)
        if version is None:
            raise RuntimeError("结构化导出缺少文档版本。")
        artifact_type = f"STRUCTURED_EXTRACTION_{presentation}"
        config_hash = _export_config_hash(run=run, presentation=presentation)
        existing = self.repository.get_for_document(
            document_id=run.document_id,
            artifact_type=artifact_type,
            source_sha256=version.sha256,
            converter_config_hash=config_hash,
        )
        if not force and existing is not None and self._artifact_path(existing).is_file():
            return _artifact_projection(existing, presentation=presentation, reused=True)

        suffix = ".csv" if presentation == "CSV" else ".xlsx"
        relative_dir = Path("derivatives") / "structured-extraction" / run.document_id
        storage_root = Path(self.settings.file_storage_root).resolve()
        target_dir = (storage_root / relative_dir).resolve()
        if not _is_relative_to(target_dir, storage_root):
            raise RuntimeError("结构化导出目录越界。")
        target_dir.mkdir(parents=True, exist_ok=True)
        target_path = target_dir / f"{run.id}{suffix}"
        rows = _tabular_rows(result)
        headers = [
            _safe_spreadsheet_text(str(item.get("label") or item.get("key") or ""))
            for item in result.field_schema
        ]
        if presentation == "CSV":
            _write_csv_atomic(target_path=target_path, headers=headers, rows=rows)
            content_type = "text/csv; charset=utf-8"
        else:
            _write_xlsx_atomic(target_path=target_path, headers=headers, rows=rows)
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        digest = hashlib.sha256(target_path.read_bytes()).hexdigest()
        artifact = self.repository.upsert_link(
            document_id=run.document_id,
            artifact_type=artifact_type,
            storage_path=(relative_dir / target_path.name).as_posix(),
            content_type=content_type,
            size_bytes=target_path.stat().st_size,
            sha256=digest,
            source_sha256=version.sha256,
            converter_name="structured-extraction-export",
            converter_version="1",
            converter_config_hash=config_hash,
        )
        return _artifact_projection(artifact, presentation=presentation, reused=False)

    def get_existing(self, *, run: StructuredExtractionRun) -> dict[str, Any] | None:
        """从持久化派生件恢复刷新后的下载信息。"""

        presentation = str(run.presentation or "AUTO").upper()
        if presentation not in {"CSV", "XLSX"}:
            return None
        version = self.db.get(DocumentVersion, run.document_version_id)
        if version is None:
            return None
        artifact = self.repository.get_for_document(
            document_id=run.document_id,
            artifact_type=f"STRUCTURED_EXTRACTION_{presentation}",
            source_sha256=version.sha256,
            converter_config_hash=_export_config_hash(run=run, presentation=presentation),
        )
        if artifact is None or not self._artifact_path(artifact).is_file():
            return None
        return _artifact_projection(artifact, presentation=presentation, reused=True)

    def _artifact_path(self, artifact: DocumentArtifact) -> Path:
        storage_root = Path(self.settings.file_storage_root).resolve()
        path = (storage_root / artifact.storage_path).resolve()
        if not _is_relative_to(path, storage_root):
            raise RuntimeError("结构化导出文件路径越界。")
        return path


def _tabular_rows(result: StructuredExtractionResult) -> list[list[Any]]:
    keys = [str(item.get("key") or "") for item in result.field_schema]
    rows: list[list[Any]] = []
    for record in result.records:
        fields = dict(record.get("fields") or {})
        rows.append([_cell_value((fields.get(key) or {}).get("normalized_value")) for key in keys])
    return rows


def _cell_value(value: Any) -> Any:
    if isinstance(value, str):
        return _safe_spreadsheet_text(value)
    if value is None or isinstance(value, (int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _safe_spreadsheet_text(value: str) -> str:
    """阻止用户字段和识别文本在 CSV/XLSX 中被解释为公式。"""

    return f"'{value}" if re.match(r"^[\t\r ]*[=+\-@]", value) else value


def _write_csv_atomic(*, target_path: Path, headers: list[str], rows: list[list[Any]]) -> None:
    descriptor, temporary_name = tempfile.mkstemp(prefix="structured-", suffix=".csv", dir=target_path.parent)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerows(rows)
        os.replace(temporary, target_path)
    finally:
        temporary.unlink(missing_ok=True)


def _write_xlsx_atomic(
    *,
    target_path: Path,
    headers: list[str],
    rows: list[list[Any]],
    sheet_name: str = "结构化抽取",
) -> None:
    """原子写入无宏 XLSX；工作表名由服务端校验后传入并保持用户目标结构。"""

    descriptor, temporary_name = tempfile.mkstemp(prefix="structured-", suffix=".xlsx", dir=target_path.parent)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook = Workbook(write_only=True)
        safe_sheet_name = re.sub(r"[\\/*?:\[\]]", "_", sheet_name).strip()[:31] or "结构化抽取"
        sheet = workbook.create_sheet(title=safe_sheet_name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        workbook.save(temporary)
        os.replace(temporary, target_path)
    finally:
        temporary.unlink(missing_ok=True)


def _append_xlsx_rows_atomic(
    *,
    template_path: Path,
    target_path: Path,
    headers: list[str],
    rows: list[list[Any]],
    sheet_name: str,
    layout: dict | None = None,
) -> None:
    """直接修改 XLSX 包内目标工作表，避免重存整个工作簿造成高级对象丢失。

    除目标 worksheet、与其相连且需要扩展的 table XML 外，其余 ZIP part 均按原始
    字节复制。新增单元格使用 inline string 或标量值并继承上一业务行样式，不执行
    公式、宏、外链或嵌入对象。
    """

    if not template_path.is_file():
        raise FileNotFoundError("目标 Excel 模板不存在")
    parsed_layout = WorkbookLayout.model_validate(layout) if layout else None
    if parsed_layout:
        # 写入前再次只读核对实际模板，不能单凭旧 schema 覆盖单元格。
        workbook = load_workbook(template_path, data_only=False)
        try:
            actual, _ = read_layout(workbook[sheet_name], parsed_layout)
            if actual != headers:
                raise ValueError("TARGET_EXCEL_SCHEMA_CONFLICT: 模板表头发生变化")
        finally:
            workbook.close()
    descriptor, temporary_name = tempfile.mkstemp(
        prefix="structured-", suffix=".xlsx", dir=target_path.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        with zipfile.ZipFile(template_path, "r") as source_zip:
            package = {item.filename: source_zip.read(item.filename) for item in source_zip.infolist()}
            infos = {item.filename: item for item in source_zip.infolist()}
        worksheet_path = _xlsx_worksheet_path(package, sheet_name)
        shared_strings = _xlsx_shared_strings(package.get("xl/sharedStrings.xml"))
        worksheet_bytes, header_row, final_row = _xlsx_append_rows(
            package[worksheet_path],
            headers=headers,
            rows=rows,
            shared_strings=shared_strings,
            layout=parsed_layout,
        )
        replacements = {worksheet_path: worksheet_bytes}
        if not parsed_layout:
            replacements.update(
                _xlsx_expand_linked_tables(
                    package,
                    worksheet_path=worksheet_path,
                    header_row=header_row,
                    final_row=final_row,
                    field_count=len(headers),
                )
            )
        with zipfile.ZipFile(temporary, "w") as target_zip:
            for name, data in package.items():
                target_zip.writestr(infos[name], replacements.get(name, data))
        os.replace(temporary, target_path)
    finally:
        temporary.unlink(missing_ok=True)


def _xlsx_worksheet_path(package: dict[str, bytes], sheet_name: str) -> str:
    """通过 workbook relationship 解析工作表 part，禁止根据名称拼接路径。"""

    workbook = ET.fromstring(package["xl/workbook.xml"])
    relationship_id = None
    for element in workbook.iter():
        if _xml_local_name(element.tag) != "sheet" or element.attrib.get("name") != sheet_name:
            continue
        relationship_id = next(
            (value for key, value in element.attrib.items() if key.endswith("}id") or key == "id"),
            None,
        )
        break
    if not relationship_id:
        raise ValueError("目标 Excel 工作表不存在")
    relationships = ET.fromstring(package["xl/_rels/workbook.xml.rels"])
    target = next(
        (
            element.attrib.get("Target")
            for element in relationships
            if element.attrib.get("Id") == relationship_id
        ),
        None,
    )
    if not target:
        raise ValueError("目标 Excel 工作表关系损坏")
    normalized = posixpath.normpath(
        target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target)
    )
    if normalized not in package or not normalized.startswith("xl/worksheets/"):
        raise ValueError("目标 Excel 工作表路径无效")
    return normalized


def _xlsx_shared_strings(content: bytes | None) -> list[str]:
    """只读共享字符串表，用于校验目标标题，不改写其索引。"""

    if not content:
        return []
    root = ET.fromstring(content)
    return [
        "".join(element.text or "" for element in item.iter() if _xml_local_name(element.tag) == "t")
        for item in root
        if _xml_local_name(item.tag) == "si"
    ]


def _xlsx_append_rows(
    content: bytes,
    *,
    headers: list[str],
    rows: list[list[Any]],
    shared_strings: list[str],
    layout: WorkbookLayout | None = None,
) -> tuple[bytes, int, int]:
    """修改单个 worksheet XML，并返回标题行与最终数据行。"""

    _register_xml_namespaces(content)
    root = ET.fromstring(content)
    namespace = _xml_namespace(root.tag)
    qname = lambda name: f"{{{namespace}}}{name}" if namespace else name
    sheet_data = next((element for element in root if _xml_local_name(element.tag) == "sheetData"), None)
    if sheet_data is None:
        raise ValueError("目标 Excel 工作表缺少 sheetData")
    row_elements = {
        int(element.attrib["r"]): element
        for element in sheet_data
        if _xml_local_name(element.tag) == "row" and str(element.attrib.get("r", "")).isdigit()
    }
    if layout is not None:
        return _xlsx_fill_layout(root, sheet_data, row_elements, headers, rows, shared_strings, layout, qname)
    expected = [str(value).strip() for value in headers]
    header_row = None
    last_data_row = None
    for row_number in sorted(row_elements):
        cells = _xlsx_row_cells(row_elements[row_number])
        all_values = [_xlsx_cell_value(cell, shared_strings) for cell in cells.values()]
        if header_row is None and any(value is not None and str(value).strip() for value in all_values):
            actual = [
                "" if (value := _xlsx_cell_value(cells.get(column), shared_strings)) is None
                else str(value).strip()
                for column in range(1, len(headers) + 1)
            ]
            if actual != expected:
                raise ValueError("目标 Excel 标题与数据集 schema 不一致")
            header_row = row_number
            last_data_row = row_number
            continue
        if header_row is not None and row_number > header_row and any(
            _xlsx_cell_has_value(cells.get(column)) for column in range(1, len(headers) + 1)
        ):
            last_data_row = row_number
    if header_row is None or last_data_row is None:
        raise ValueError("目标 Excel 缺少标题行")
    style_row = row_elements[last_data_row]
    style_cells = _xlsx_row_cells(style_row)
    next_row = last_data_row + 1
    for values in rows:
        destination = row_elements.get(next_row)
        if destination is None:
            destination = ET.Element(qname("row"), {"r": str(next_row)})
            for key, value in style_row.attrib.items():
                if key not in {"r", "spans"}:
                    destination.set(key, value)
            _xlsx_insert_row(sheet_data, destination, next_row)
            row_elements[next_row] = destination
        destination_cells = _xlsx_row_cells(destination)
        for column, value in enumerate(values, 1):
            cell = destination_cells.get(column)
            if cell is None:
                cell = ET.Element(qname("c"), {"r": f"{_xlsx_column_name(column)}{next_row}"})
                style = style_cells.get(column)
                if style is not None and "s" in style.attrib:
                    cell.set("s", style.attrib["s"])
                _xlsx_insert_cell(destination, cell, column)
            _xlsx_set_cell_value(cell, value, qname=qname)
        next_row += 1
    final_row = next_row - 1 if rows else last_data_row
    _xlsx_expand_range_attributes(root, header_row=header_row, final_row=final_row, field_count=len(headers))
    return ET.tostring(root, encoding="utf-8", xml_declaration=True), header_row, final_row


def _xlsx_fill_layout(root, sheet_data, row_elements, headers, rows, shared_strings, layout, qname):
    """只填充明确预留区域，不移动尾注、不覆盖公式和已有业务记录。"""
    left, top, right, bottom = range_boundaries(layout.data_range)
    sequences = sequence_offsets(headers)
    available = []
    for number in range(top, bottom + 1):
        element = row_elements.get(number)
        cells = _xlsx_row_cells(element) if element is not None else {}
        if not any(_xlsx_cell_has_value(cells.get(col)) for col in range(left, right + 1)
                   if col - left not in sequences):
            available.append(number)
    if len(rows) > len(available):
        raise ValueError(f"TARGET_EXCEL_DATA_RANGE_FULL: 填写区域剩余 {len(available)} 行，本次需要 {len(rows)} 行")
    for values, number in zip(rows, available):
        if len(values) != len(headers) or not any(v is not None and str(v).strip() for i, v in enumerate(values) if i not in sequences):
            raise ValueError("TARGET_EXCEL_EMPTY_RECORD: 记录没有可填写的业务字段")
        destination = row_elements.get(number)
        if destination is None:
            destination = ET.Element(qname("row"), {"r": str(number)})
            _xlsx_insert_row(sheet_data, destination, number)
        cells = _xlsx_row_cells(destination)
        for offset, value in enumerate(values):
            col = left + offset
            cell = cells.get(col)
            if offset in sequences and _xlsx_cell_has_value(cell):
                existing = _xlsx_cell_value(cell, shared_strings)
                if value is not None and str(value).strip() and str(existing) != str(value):
                    raise ValueError("TARGET_EXCEL_SEQUENCE_CONFLICT: 提交序号与模板预置序号不一致")
                continue
            if value is None:
                continue
            if cell is None:
                cell = ET.Element(qname("c"), {"r": f"{_xlsx_column_name(col)}{number}"})
                _xlsx_insert_cell(destination, cell, col)
            _xlsx_set_cell_value(cell, value, qname=qname)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True), range_boundaries(layout.header_range)[3], bottom


def _xlsx_expand_linked_tables(
    package: dict[str, bytes],
    *,
    worksheet_path: str,
    header_row: int,
    final_row: int,
    field_count: int,
) -> dict[str, bytes]:
    """仅扩展目标工作表关联表格的 ref，其他 OOXML part 保持原始字节。"""

    relation_path = posixpath.join(
        posixpath.dirname(worksheet_path),
        "_rels",
        f"{posixpath.basename(worksheet_path)}.rels",
    )
    if relation_path not in package:
        return {}
    relationships = ET.fromstring(package[relation_path])
    replacements: dict[str, bytes] = {}
    for relationship in relationships:
        if not str(relationship.attrib.get("Type") or "").endswith("/table"):
            continue
        target = str(relationship.attrib.get("Target") or "")
        table_path = posixpath.normpath(
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.join(posixpath.dirname(worksheet_path), target)
        )
        if table_path not in package:
            continue
        content = package[table_path]
        _register_xml_namespaces(content)
        table = ET.fromstring(content)
        reference = table.attrib.get("ref")
        expanded = _xlsx_expanded_range(
            reference,
            header_row=header_row,
            final_row=final_row,
            field_count=field_count,
        )
        if expanded == reference:
            continue
        table.set("ref", expanded)
        for element in table:
            if _xml_local_name(element.tag) == "autoFilter" and element.attrib.get("ref"):
                element.set("ref", expanded)
        replacements[table_path] = ET.tostring(table, encoding="utf-8", xml_declaration=True)
    return replacements


def _xlsx_expand_range_attributes(root, *, header_row: int, final_row: int, field_count: int) -> None:
    """扩展 worksheet dimension 和匹配的 autoFilter，不改变其他范围。"""

    for element in root.iter():
        local_name = _xml_local_name(element.tag)
        reference = element.attrib.get("ref")
        if local_name == "dimension" and reference:
            element.set("ref", _xlsx_expand_dimension(reference, final_row))
        elif local_name == "autoFilter" and reference:
            element.set(
                "ref",
                _xlsx_expanded_range(
                    reference,
                    header_row=header_row,
                    final_row=final_row,
                    field_count=field_count,
                ),
            )


def _xlsx_expand_dimension(reference: str, final_row: int) -> str:
    """只向下扩展 worksheet dimension，绝不收缩模板已有使用范围。"""

    parts = reference.split(":", 1)
    end = parts[-1]
    match = re.fullmatch(r"(\$?[A-Z]+)\$?(\d+)", end)
    if not match or int(match.group(2)) >= final_row:
        return reference
    expanded_end = f"{match.group(1)}{final_row}"
    return f"{parts[0]}:{expanded_end}" if len(parts) == 2 else expanded_end


def _xlsx_expanded_range(
    reference: str | None,
    *,
    header_row: int,
    final_row: int,
    field_count: int,
) -> str | None:
    """仅扩展覆盖目标 schema 且从标题行开始的矩形范围。"""

    if not reference:
        return reference
    match = re.fullmatch(r"\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)", reference)
    if not match:
        return reference
    start_column = _xlsx_column_number(match.group(1))
    start_row = int(match.group(2))
    end_column = _xlsx_column_number(match.group(3))
    end_row = int(match.group(4))
    if start_row != header_row or start_column > 1 or end_column < field_count or end_row >= final_row:
        return reference
    return f"{match.group(1)}{start_row}:{match.group(3)}{final_row}"


def _xlsx_row_cells(row) -> dict[int, Any]:
    """按列号索引一行中的现有单元格。"""

    result = {}
    for cell in row:
        if _xml_local_name(cell.tag) != "c":
            continue
        match = re.match(r"([A-Z]+)", str(cell.attrib.get("r") or ""))
        if match:
            result[_xlsx_column_number(match.group(1))] = cell
    return result


def _xlsx_cell_value(cell, shared_strings: list[str]) -> Any:
    """读取标题校验所需的标量值，不计算公式。"""

    if cell is None:
        return None
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        return "".join(
            element.text or "" for element in cell.iter() if _xml_local_name(element.tag) == "t"
        )
    value = next((element.text for element in cell if _xml_local_name(element.tag) == "v"), None)
    if cell_type == "s" and value is not None and value.isdigit():
        index = int(value)
        return shared_strings[index] if index < len(shared_strings) else None
    if cell_type == "b":
        return value == "1"
    return value


def _xlsx_cell_has_value(cell) -> bool:
    """公式或非空标量都视为已有业务内容。"""

    if cell is None:
        return False
    return any(
        _xml_local_name(element.tag) in {"f", "v", "is"}
        and (element.text is not None or len(element) > 0)
        for element in cell
    )


def _xlsx_set_cell_value(cell, value: Any, *, qname) -> None:
    """清除旧值后写入安全标量；字符串统一使用 inlineStr，不改共享字符串表。"""

    for child in list(cell):
        if _xml_local_name(child.tag) in {"f", "v", "is"}:
            cell.remove(child)
    if value is None:
        cell.attrib.pop("t", None)
        return
    if isinstance(value, bool):
        cell.set("t", "b")
        ET.SubElement(cell, qname("v")).text = "1" if value else "0"
        return
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.attrib.pop("t", None)
        ET.SubElement(cell, qname("v")).text = str(value)
        return
    text = str(value)
    cell.set("t", "inlineStr")
    inline = ET.SubElement(cell, qname("is"))
    text_element = ET.SubElement(inline, qname("t"))
    if text != text.strip() or "\n" in text or "\t" in text:
        text_element.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_element.text = text


def _xlsx_insert_row(sheet_data, row, row_number: int) -> None:
    """按 r 属性顺序插入新行，兼容模板预置的空白格式行。"""

    for index, existing in enumerate(sheet_data):
        value = str(existing.attrib.get("r") or "")
        if value.isdigit() and int(value) > row_number:
            sheet_data.insert(index, row)
            return
    sheet_data.append(row)


def _xlsx_insert_cell(row, cell, column: int) -> None:
    """按列顺序插入单元格，保留目标行中 schema 之外的模板内容。"""

    for index, existing in enumerate(row):
        match = re.match(r"([A-Z]+)", str(existing.attrib.get("r") or ""))
        if match and _xlsx_column_number(match.group(1)) > column:
            row.insert(index, cell)
            return
    row.append(cell)


def _xlsx_column_number(name: str) -> int:
    """把 Excel 列名转换为从 1 开始的列号。"""

    result = 0
    for character in name:
        result = result * 26 + ord(character.upper()) - 64
    return result


def _xlsx_column_name(number: int) -> str:
    """把从 1 开始的列号转换为 Excel 列名。"""

    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _xml_local_name(tag: str) -> str:
    """返回带命名空间 XML 标签的本地名。"""

    return tag.rsplit("}", 1)[-1]


def _xml_namespace(tag: str) -> str:
    """返回 ElementTree 标签的命名空间 URI。"""

    return tag[1:].split("}", 1)[0] if tag.startswith("{") else ""


def _register_xml_namespaces(content: bytes) -> None:
    """保留原 OOXML 前缀，防止 mc:Ignorable 等前缀引用失效。"""

    for _, namespace in ET.iterparse(io.BytesIO(content), events=("start-ns",)):
        prefix, uri = namespace
        if prefix and re.fullmatch(r"ns\d+", prefix):
            continue
        try:
            ET.register_namespace(prefix or "", uri)
        except ValueError:
            continue


def _export_config_hash(*, run: StructuredExtractionRun, presentation: str) -> str:
    payload = {
        "schema_fingerprint": run.schema_fingerprint,
        "prompt_version": run.prompt_version,
        "presentation": presentation,
        "export_version": "1",
    }
    return hashlib.sha256(
        json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def _artifact_projection(
    artifact: DocumentArtifact,
    *,
    presentation: str,
    reused: bool,
) -> dict[str, Any]:
    suffix = ".csv" if presentation == "CSV" else ".xlsx"
    return {
        "artifact_id": artifact.id,
        "format": presentation,
        "filename": f"structured-extraction-{artifact.document_id}{suffix}",
        "content_type": artifact.content_type,
        "size_bytes": artifact.size_bytes,
        "download_url": f"/api/files/{artifact.document_id}/artifacts/{artifact.id}",
        "reused": reused,
    }


def _is_relative_to(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
    except ValueError:
        return False
    return True

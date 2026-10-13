"""WorkBuddy 生成完整 XLSX 后的受控验收与同名工作副本发布。"""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path
from uuid import uuid4
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from sqlalchemy import func

from app.db.models import DocumentExtractionConfirmation, DocumentVersion, StructuredDatasetExport, StructuredDatasetTask, utcnow
from app.modules.structured_extraction.export import _xlsx_worksheet_path
from .assessment import excel_content_key, excel_display_text, validate_quote, validate_approval_claim
from .schemas import DatasetAppend
from .service import StructuredDatasetService, fail
from .task_schemas import TaskPayload
from .tasks import DatasetTaskService
from .writeback_schemas import WritebackCheckout, WritebackConfirm, WritebackPreview
from .xlsx_safety import XlsxSafetyError, inspect_xlsx
from .publication import stage_publication


def writeback_digest(payload: dict, file_sha256: str) -> str:
    """候选和回传文件共用一个不可变预览指纹。"""
    body = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256((body + ":" + file_sha256).encode("utf-8")).hexdigest()


def _shift_expected_footer(root: ET.Element, first_row: int, count: int) -> None:
    """独立还原允许的行插入，以便逐字节比较移位后的尾注。"""
    unsupported = {"autoFilter", "tableParts", "dataValidations", "conditionalFormatting",
                   "hyperlinks", "rowBreaks", "drawing", "legacyDrawing", "extLst"}
    if any(node.tag.rsplit("}", 1)[-1] in unsupported for node in root):
        fail("WRITEBACK_ROW_INSERT_UNSUPPORTED", "工作表包含不能安全移动的区域引用。", 422)
    data = next((node for node in root if node.tag.rsplit("}", 1)[-1] == "sheetData"), None)
    if data is None or any(node.tag.rsplit("}", 1)[-1] == "f" for node in data.iter()):
        fail("WRITEBACK_ROW_INSERT_UNSUPPORTED", "工作表公式或数据区不能安全移动。", 422)
    for row in data:
        if row.tag.rsplit("}", 1)[-1] != "row" or int(row.attrib.get("r", "0")) < first_row:
            continue
        number = int(row.attrib["r"]) + count
        row.set("r", str(number))
        for cell in row:
            if cell.tag.rsplit("}", 1)[-1] == "c":
                match = re.fullmatch(r"([A-Z]+)\d+", cell.attrib.get("r", ""))
                if match is None:
                    fail("WRITEBACK_XLSX_INVALID", "尾注单元格坐标无效。", 422)
                cell.set("r", f"{match.group(1)}{number}")
    for node in root.iter():
        if node.tag.rsplit("}", 1)[-1] != "mergeCell":
            continue
        left, top, right, bottom = range_boundaries(node.attrib.get("ref", ""))
        if top < first_row <= bottom:
            fail("WRITEBACK_ROW_INSERT_UNSUPPORTED", "合并区域跨越插入点。", 422)
        if top >= first_row:
            node.set("ref", f"{get_column_letter(left)}{top + count}:{get_column_letter(right)}{bottom + count}")
    dimension = next((node for node in root if node.tag.rsplit("}", 1)[-1] == "dimension"), None)
    if dimension is not None:
        match = re.fullmatch(r"([A-Z]+\d+:)?([A-Z]+)(\d+)", dimension.attrib.get("ref", ""))
        if match:
            previous_end = int(match.group(3))
            new_end = max(previous_end + (count if previous_end >= first_row else 0), first_row + count - 1)
            dimension.set("ref", f"{match.group(1) or ''}{match.group(2)}{new_end}")


class ExcelWritebackService:
    """只接管一个已注册目标的本轮回传；旧累计表格工具保持原契约。"""

    def __init__(self, db, user, agent_run=None):
        self.db = db
        self.user = user
        self.datasets = StructuredDatasetService(db, user, agent_run)
        self.tasks = DatasetTaskService(db, user, agent_run)

    def checkout(self, request: WritebackCheckout) -> dict:
        """按文件名唯一解析并冻结版本；读取阶段不长期占用目标。"""
        target = self.datasets.resolve_target(request.filename)
        dataset = self.datasets.dataset(target["dataset_id"], lock=True)
        working = self.datasets._dataset_working_copy(dataset)
        if working is None or not working.current_version_id:
            fail("WRITEBACK_TARGET_MISSING", "目标 Excel 尚无可读取的当前工作副本。", 409)
        self.datasets._assert_task_write_access(dataset)
        version = self.db.get(DocumentVersion, working.current_version_id)
        if version is None or Path(working.filename).suffix.lower() != ".xlsx":
            fail("WRITEBACK_XLSX_REQUIRED", "新回传链路只接受已提升并核验的 XLSX 目标。", 422)
        path = self.datasets.storage.working_copy_path(version.storage_path)
        if not path.is_file() or self.datasets.storage.sha256_file(path) != version.sha256:
            fail("WRITEBACK_TARGET_MISSING", "目标 Excel 当前版本缺失或哈希不匹配。", 409)
        if path.stat().st_size > 32 * 1024 * 1024:
            fail("WRITEBACK_TARGET_TOO_LARGE", "目标 Excel 超过新回传链路的 32 MiB 上限。", 413)
        snapshot = {
            "dataset_id": dataset.id, "working_copy_id": working.id,
            "document_id": working.document_id, "document_version_id": version.id,
            "sha256": version.sha256, "revision": dataset.revision,
            "filename": working.filename, "sheet_name": dataset.target_sheet_name,
            "layout": dataset.target_layout_json or None,
            "labels": [field["label"] for field in dataset.schema_json],
        }
        existing = self.db.query(StructuredDatasetTask).filter_by(
            user_id=self.user.id, request_key=f"writeback:{request.request_key}"
        ).one_or_none()
        if existing is not None:
            if existing.manifest_json != snapshot or existing.review_mode != request.review_mode:
                fail("WRITEBACK_REQUEST_CONFLICT", "同一请求键已绑定其他目标或版本，请使用新请求键。")
            return self.receipt(existing)
        task = StructuredDatasetTask(
            user_id=self.user.id, request_key=f"writeback:{request.request_key}",
            result_digest="0" * 64, review_mode=request.review_mode,
            status="WRITEBACK_CHECKED_OUT", manifest_json=snapshot,
            payload_json={}, progress_json={}, result_json={}, error_json={},
        )
        self.db.add(task)
        self.db.flush()
        return self.receipt(task)

    def owned(self, task_id: str, *, lock=False) -> StructuredDatasetTask:
        query = self.db.query(StructuredDatasetTask).filter_by(id=task_id, user_id=self.user.id)
        if lock:
            query = query.with_for_update()
        task = query.one_or_none()
        if (task is None or not task.request_key.startswith("writeback:")
                or (task.manifest_json or {}).get("mode") == "CORRECTION"):
            fail("WRITEBACK_TASK_NOT_FOUND", "回传任务不存在或无权访问。", 404)
        return task

    def receipt(self, task: StructuredDatasetTask) -> dict:
        snapshot = task.manifest_json or {}
        return {
            "task_id": task.id, "status": task.status, "review_mode": task.review_mode,
            "result_digest": None if task.result_digest == "0" * 64 else task.result_digest,
            "dataset_id": snapshot.get("dataset_id"),
            "working_copy_id": snapshot.get("working_copy_id"),
            "document_id": snapshot.get("document_id"),
            "document_version_id": snapshot.get("document_version_id"),
            "base_sha256": snapshot.get("sha256"), "base_revision": snapshot.get("revision"),
            "filename": snapshot.get("filename"), "sheet_name": snapshot.get("sheet_name"),
            "layout": snapshot.get("layout"), "labels": snapshot.get("labels"),
            "download_url": (f"/api/working-copies/{snapshot['working_copy_id']}/download"
                             if snapshot.get("working_copy_id") else None),
            "error": task.error_json or None, **(task.result_json or {}),
        }

    def preview(self, task_id: str, request: WritebackPreview) -> dict:
        """冻结 WorkBuddy 生成的候选和文件指纹，供可选人工复核。"""
        task = self.owned(task_id, lock=True)
        if task.status == "NEEDS_REPAIR":
            return self.receipt(task)
        if task.status == "PUBLISHING":
            if (task.result_digest != request.result_digest
                    or (task.progress_json or {}).get("file_sha256") != request.file_sha256
                    or (task.progress_json or {}).get("extraction_confirmations", []) != [
                        ref.model_dump(mode="json") for ref in request.extraction_confirmations]):
                fail("WRITEBACK_PREVIEW_CHANGED", "发布中任务不能使用不同的候选或文件摘要重放。", 409)
            return self.receipt(task)
        if task.status == "COMPLETED":
            if (task.result_digest != request.result_digest
                    or (task.progress_json or {}).get("file_sha256") != request.file_sha256
                    or (task.progress_json or {}).get("extraction_confirmations", []) != [
                        ref.model_dump(mode="json") for ref in request.extraction_confirmations]):
                fail("WRITEBACK_PREVIEW_CHANGED", "已完成任务不能使用不同的候选或文件摘要重放。", 409)
            self.tasks._verify_publication(task)
            return self.receipt(task)
        if task.status not in {"WRITEBACK_CHECKED_OUT", "WRITEBACK_WAITING_CONFIRMATION", "WRITEBACK_READY"}:
            fail("WRITEBACK_TASK_STATE", "当前任务不能登记回传预览。")
        if task.result_digest != "0" * 64 and task.result_digest != request.result_digest:
            fail("WRITEBACK_PREVIEW_CHANGED", "预览内容变化后必须创建新任务。")
        if (task.progress_json or {}).get("file_sha256") not in {None, request.file_sha256}:
            fail("WRITEBACK_PREVIEW_CHANGED", "预览文件变化后必须创建新任务。")
        confirmation_refs = [ref.model_dump(mode="json") for ref in request.extraction_confirmations]
        previous_refs = (task.progress_json or {}).get("extraction_confirmations")
        if previous_refs is not None and previous_refs != confirmation_refs:
            fail("WRITEBACK_PREVIEW_CHANGED", "来源正文复核引用变化后必须创建新任务。")
        for ref in request.extraction_confirmations:
            self.tasks._extraction(ref.task_id)
            confirmation = self.db.get(DocumentExtractionConfirmation, str(ref.confirmation_id))
            if (confirmation is None or confirmation.user_id != self.user.id
                    or confirmation.task_id != str(ref.task_id)
                    or confirmation.result_digest != ref.result_digest):
                fail("EXTRACTION_CONFIRMATION_STALE", "来源正文预览引用与本轮任务不符。")
        task.result_digest = request.result_digest
        task.progress_json = {"file_sha256": request.file_sha256,
                              "extraction_confirmations": confirmation_refs}
        task.status = "WRITEBACK_READY" if task.review_mode == "DIRECT_SAVE" else "WRITEBACK_WAITING_CONFIRMATION"
        task.updated_at = utcnow()
        return self.receipt(task)

    def confirm(self, task_id: str, request: WritebackConfirm) -> dict:
        """仅复核路径需要人工确认同一文件与候选摘要。"""
        task = self.owned(task_id, lock=True)
        if task.review_mode != "REVIEW_REQUIRED" or task.result_digest != request.result_digest:
            fail("WRITEBACK_CONFIRMATION_MISMATCH", "确认指纹与本轮回传预览不一致。")
        if task.status == "WRITEBACK_WAITING_CONFIRMATION":
            from app.modules.external_extraction.document_service import WorkBuddyDocumentExtractionService
            service = WorkBuddyDocumentExtractionService(self.db)
            for ref in (task.progress_json or {}).get("extraction_confirmations", []):
                service.confirm(task_id=ref["task_id"], confirmation_id=ref["confirmation_id"],
                                result_digest=ref["result_digest"], current_user=self.user)
            task.confirmed_at = utcnow()
            task.status = "WRITEBACK_READY"
        elif task.status not in {"WRITEBACK_READY", "COMPLETED"}:
            fail("WRITEBACK_TASK_STATE", "请先登记并展示回传预览。")
        if task.status == "COMPLETED":
            self.tasks._verify_publication(task)
        return self.receipt(task)

    def publish(self, task_id: str, *, result_digest: str, file_sha256: str,
                payload: TaskPayload, file_path: Path) -> dict:
        """正式校验候选和回传 XLSX，复用成员审计但绝不调用旧 Excel writer。"""
        task = self.owned(task_id, lock=True)
        if task.status == "NEEDS_REPAIR":
            return self.receipt(task)
        if task.status == "PUBLISHING":
            body = payload.model_dump(mode="json")
            if (task.result_digest != result_digest or (task.progress_json or {}).get("file_sha256") != file_sha256
                    or writeback_digest(body, file_sha256) != result_digest
                    or not file_path.is_file() or self.datasets.storage.sha256_file(file_path) != file_sha256):
                fail("WRITEBACK_PREVIEW_CHANGED", "发布中任务不能使用不同的候选或文件内容重放。", 409)
            return self.receipt(task)
        if task.status == "COMPLETED":
            body = payload.model_dump(mode="json")
            if (task.result_digest != result_digest or (task.progress_json or {}).get("file_sha256") != file_sha256
                    or writeback_digest(body, file_sha256) != result_digest
                    or not file_path.is_file() or self.datasets.storage.sha256_file(file_path) != file_sha256):
                fail("WRITEBACK_PREVIEW_CHANGED", "已完成任务不能使用不同的候选或文件内容重放。", 409)
            self.tasks._verify_publication(task)
            return self.receipt(task)
        if task.status != "WRITEBACK_READY" or (task.review_mode == "REVIEW_REQUIRED" and not task.confirmed_at):
            fail("WRITEBACK_CONFIRMATION_REQUIRED", "请先展示并确认完整回传预览。")
        body = payload.model_dump(mode="json")
        submitted_refs = body["manifest"].get("extraction_confirmations", [])
        frozen_refs = (task.progress_json or {}).get("extraction_confirmations", [])
        if submitted_refs != frozen_refs:
            fail("EXTRACTION_CONFIRMATION_STALE", "回传候选的来源正文确认与预览不一致。", 422)
        confirmed_task_ids = {str(ref["task_id"]) for ref in frozen_refs}
        source_task_ids = {str(source.extraction_task_id) for source in payload.manifest.sources
                           if source.extraction_task_id}
        if not confirmed_task_ids <= source_task_ids:
            fail("EXTRACTION_CONFIRMATION_STALE", "来源确认引用不属于本轮上传文件。", 422)
        for source_ref in payload.manifest.sources:
            if source_ref.extraction_task_id:
                extraction, _, batch = self.tasks._extraction(source_ref.extraction_task_id)
                direct_source = (batch.policy_json or {}).get("extraction_save_policy") == "DIRECT_ARCHIVE"
                if task.review_mode == "DIRECT_SAVE" and not direct_source:
                    fail("EXTRACTION_CONFIRMATION_REQUIRED", "来源正文需要复核，不能改走直接保存。")
                if task.review_mode == "REVIEW_REQUIRED" and not direct_source and str(extraction.id) not in confirmed_task_ids:
                    fail("EXTRACTION_CONFIRMATION_REQUIRED", "来源正文需要先由用户确认。")
        if (writeback_digest(body, file_sha256) != result_digest
                or task.result_digest != result_digest
                or (task.progress_json or {}).get("file_sha256") != file_sha256):
            fail("WRITEBACK_PREVIEW_CHANGED", "文件或候选与已冻结预览不一致。", 422)
        if not file_path.is_file() or self.datasets.storage.sha256_file(file_path) != file_sha256:
            fail("WRITEBACK_FILE_CHANGED", "回传文件与预览哈希不一致。", 422)
        try:
            inspect_xlsx(file_path)
        except XlsxSafetyError as exc:
            fail(exc.code, str(exc), 413 if "TOO_LARGE" in exc.code or "TOO_COMPLEX" in exc.code else 422)
        snapshot = task.manifest_json
        dataset = self.datasets.dataset(snapshot["dataset_id"], lock=True)
        self.datasets._assert_task_write_access(dataset)
        working = self.datasets._dataset_working_copy(dataset)
        if working is None or working.id != snapshot["working_copy_id"]:
            fail("WRITEBACK_TARGET_CHANGED", "目标工作副本已改变，请重新 checkout。")
        current = self.db.get(DocumentVersion, working.current_version_id)
        if (current is None or current.id != snapshot["document_version_id"]
                or current.sha256 != snapshot["sha256"] or dataset.revision != snapshot["revision"]
                or working.filename != snapshot["filename"]):
            fail("WRITEBACK_VERSION_CONFLICT", "目标 Excel 已有新版本，请重新读取当前文件后追加。")
        target = payload.manifest.target
        if (str(target.dataset_id or "") != dataset.id
                or str(target.document_version_id or "") != current.id
                or target.filename != working.filename
                or (target.sheet_name or "Sheet1") != (dataset.target_sheet_name or "Sheet1")
                or (target.layout or None) != (dataset.target_layout_json or None)
                or target.labels != [field["label"] for field in dataset.schema_json]):
            fail("WRITEBACK_TARGET_CHANGED", "候选目标与 checkout 的工作表、列或版本不一致。", 422)
        base_path = self.datasets.storage.working_copy_path(current.storage_path)
        if not base_path.is_file() or self.datasets.storage.sha256_file(base_path) != current.sha256:
            fail("WRITEBACK_TARGET_MISSING", "目标基线文件缺失或损坏。")
        for record_index, record in enumerate(payload.records, 1):
            for offset, field in enumerate(record.fields):
                if field.approval:
                    try:
                        validate_approval_claim(field.approval)
                    except ValueError as exc:
                        code, _, message = str(exc).partition(":")
                        fail(code, message.strip(), 422, record_index=record_index,
                             field_label=target.labels[offset])
        if [record.source_index for record in payload.records] != sorted(
                record.source_index for record in payload.records):
            fail("WRITEBACK_ROWS_MISMATCH", "多来源记录必须按来源顺序排列。", 422)
        scoped = payload.manifest.write_scope == "A:I"
        scoped_review = payload.manifest.write_scope == "A:I,L"
        limited = scoped or scoped_review
        full_row = payload.manifest.write_scope == "A:AA"
        template_all = payload.manifest.write_scope == "TEMPLATE_ALL"
        from openpyxl.utils.cell import range_boundaries
        left = range_boundaries(target.layout["header_range"])[0] if target.layout else 1
        if template_all:
            name_offset = next(index for index, label in enumerate(target.labels)
                               if re.sub(r"\s*\[[A-Z]+\]$", "", label.split(" / ")[-1]).strip() == "姓名")
            id_offset = next((index for index, label in enumerate(target.labels)
                              if re.sub(r"\s*\[[A-Z]+\]$", "", label.split(" / ")[-1]).strip() == "身份证号"), None)
            names = [str(record.fields[name_offset].raw_text).strip() for record in payload.records]
            if len(set(names)) != len(names):
                fail("TEMPLATE_PERSON_DUPLICATE", "本轮有重复姓名，须先核实身份并合并材料。", 422)
            _, top, _, bottom = range_boundaries(target.layout["data_range"])
            book = load_workbook(base_path, read_only=True, data_only=True)
            try:
                sheet = book[target.sheet_name or "Sheet1"]
                existing_names = {str(sheet.cell(row, left + name_offset).value).strip()
                                  for row in range(top, bottom + 1)
                                  if sheet.cell(row, left + name_offset).value is not None}
            finally:
                book.close()
            if existing_names.intersection(names):
                fail("TEMPLATE_PERSON_ALREADY_PRESENT", "目标表已有同名人员，请核实后再追加。", 409)
        if limited and (payload.manifest.required_pdf_columns != ("B:I,L" if scoped_review else "B:I")
                       or [record.source_index for record in payload.records]
                       != list(range(len(payload.manifest.sources)))):
            fail("SCOPED_WRITE_SCOPE_INVALID", "每份定向来源必须按顺序绑定一行及指定列。", 422)
        self.datasets.workflow_task_id = task.id
        self.datasets.source_fact_keys = {
            spec["key"] for offset, spec in enumerate(dataset.schema_json)
            if payload.manifest.source_fact_column_range == "J:P" and left + offset in {10, 11, 13, 15, 16}
        }
        self.datasets.hr_annual_key = next((spec["key"] for offset, spec in enumerate(dataset.schema_json)
            if left + offset == 12 and payload.manifest.sealed_only), None)
        dataset.active_task_id = task.id
        self.datasets.scoped_write_scope = limited
        self.datasets.scoped_review_scope = scoped_review
        self.datasets.full_row_ocr_mode = full_row
        self.datasets.template_all_mode = template_all
        start_revision = dataset.revision
        frozen_sources = [self.tasks._source(reference) for reference in payload.manifest.sources]
        if any(source is None for source in frozen_sources):
            fail("WRITEBACK_SOURCE_NOT_READY", "来源文件尚未归档完成，请稍后重试。")
        for source_index, source_ref in enumerate(payload.manifest.sources):
            source = frozen_sources[source_index]
            scoped_pdf = source["source_kind"] == "WORKBUDDY_SCOPED"
            scoped_word = (source["source_kind"] == "WORKBUDDY_DOCUMENT"
                           and Path(source["filename"]).suffix.lower() in {".doc", ".docx"})
            if limited and not (scoped_pdf or scoped_word):
                fail("SCOPED_SOURCE_NOT_READY", "来源必须是已完成的定向 PDF 或 Word 正文。", 422)
            if limited and scoped_pdf:
                resolved = source["coverage"].get("resolved_columns")
                if scoped_review:
                    actual = source["field_refs"]
                    if (not isinstance(resolved, list) or len(resolved) != len(set(resolved))
                            or set(resolved) != set(actual)
                            or not set(resolved) <= set("BCDEFGHI") | {"L"} or "L" not in resolved):
                        fail("SCOPED_SOURCE_NOT_READY", "PDF 来源的已提取列与冻结证据不一致。", 422)
                elif resolved != list("BCDEFGHI"):
                    fail("SCOPED_SOURCE_NOT_READY", "PDF 来源没有完成指定列的定向 OCR。", 422)
            if full_row and (source["source_kind"] != "WORKBUDDY_DOCUMENT"
                             or Path(source["filename"]).suffix.lower() != ".pdf"
                             or source["coverage"].get("status") != "COMPLETE"):
                fail("FULL_ROW_SOURCE_NOT_READY", "完整填表必须绑定已完成一次读取的 PDF 来源。", 422)
            records = []
            for candidate in (record for record in payload.records if record.source_index == source_index):
                group = ([frozen_sources[index] for index in candidate.source_indices]
                         if template_all else [source])
                if template_all and left == 1 and len(candidate.fields) >= 15:
                    def _numbers(column_number):
                        specification = candidate.fields[column_number - 1].teacher_render
                        return [item.ordinal for item in specification.items] if specification else []

                    paper_numbers = _numbers(10)
                    if paper_numbers and paper_numbers != list(range(1, len(paper_numbers) + 1)):
                        fail("TEACHER_RENDER_NUMBERING_INVALID", "论文题目须从 1 连续编号。", 422)
                    if any(not set(_numbers(column)) <= set(paper_numbers) for column in (11, 12)):
                        fail("TEACHER_RENDER_NUMBERING_INVALID", "刊物与收录项必须沿用论文题目的序号。", 422)
                    project_numbers = sorted(set().union(*[set(_numbers(column)) for column in (13, 14, 15)]))
                    if project_numbers and project_numbers != list(range(1, len(project_numbers) + 1)):
                        fail("TEACHER_RENDER_NUMBERING_INVALID", "项目序号须从 1 连续编号，M/N/O 同号对齐。", 422)
                if template_all and any(item["source_kind"] != "WORKBUDDY_DOCUMENT" for item in group):
                    fail("TEMPLATE_SOURCE_INVALID", "模板全列来源必须是已完成的正式 PDF/Office 正文。", 422)
                group_units = [unit for item in group for unit in item["source_units"]]
                if template_all:
                    person = str(candidate.fields[name_offset].raw_text).strip()
                    person_id = (str(candidate.fields[id_offset].raw_text).strip()
                                 if id_offset is not None and candidate.fields[id_offset].raw_text else None)
                    primary_units = source["source_units"]
                    if any(not any(reference.source_unit_id in {
                            str(unit["source_unit_id"]), str(unit.get("external_unit_id") or "")}
                            for unit in primary_units)
                           for reference in candidate.fields[name_offset].evidence_refs):
                        fail("TEMPLATE_PERSON_SOURCE_MISMATCH", "姓名引文须来自该行主来源。", 422)
                    for item in group:
                        content = "".join("".join(str(unit["text"]).split()) for unit in item["source_units"])
                        if "".join(person.split()) not in content and (not person_id or person_id not in content):
                            fail("TEMPLATE_PERSON_SOURCE_MISMATCH", "同一行的每份材料须含姓名或已核实身份证号。", 422)
                values = {}
                for offset, (spec, field) in enumerate(zip(dataset.schema_json, candidate.fields)):
                    if limited:
                        column = left + offset
                        if column == 1 and (field.raw_text or field.blank_reason != "TARGET_DERIVED"):
                            fail("SCOPED_WRITE_SCOPE_INVALID", "A 列必须沿用模板序号。", 422)
                        if column == 12 and scoped_review:
                            from app.modules.external_extraction.scoped_schemas import excellent_years, has_review_heading
                            if field.approval or field.assessment or field.fragments or field.source_only_fragments:
                                fail("SCOPED_REVIEW_INVALID", "L 列只能填写优秀年度，不进行资格或盖章判定。", 422)
                            if scoped_pdf:
                                review_ref = source["field_refs"].get("L")
                                if not review_ref:
                                    fail("SCOPED_REVIEW_MISSING", "PDF 来源缺少冻结的考核情况证据。", 422)
                                if review_ref["status"] == "EXCELLENT_FOUND":
                                    expected_refs = list({(item["start"], item["end"]): item
                                                          for item in review_ref["years"]}.values())
                                    if field.raw_text != review_ref["value"] or len(field.evidence_refs) != len(expected_refs):
                                        fail("SCOPED_REVIEW_MISMATCH", "L 列只允许冻结的优秀年度。", 422)
                                    if any(str(actual.source_unit_id) != review_ref["source_unit_id"]
                                           or actual.quote != frozen["quote"] or actual.start != frozen["start"]
                                           or actual.end != frozen["end"]
                                           for actual, frozen in zip(field.evidence_refs, expected_refs)):
                                        fail("SCOPED_REVIEW_MISMATCH", "L 列引文不属于冻结的考核情况区域。", 422)
                                else:
                                    reason = ({"NO_EXCELLENT": "NO_EXCELLENT_YEARS",
                                               "UNREADABLE": "ANNUAL_REVIEW_UNREADABLE",
                                               "NOT_FOUND": "ANNUAL_REVIEW_NOT_FOUND"}
                                              .get(review_ref["status"]))
                                    if field.raw_text or field.blank_reason != reason or field.evidence_refs:
                                        fail("SCOPED_REVIEW_MISMATCH", "L 列留空原因与考核区域结果不一致。", 422)
                            else:
                                if field.raw_text:
                                    quoted_years = sorted({year for evidence in field.evidence_refs
                                        if has_review_heading(evidence.quote)
                                        for year in excellent_years(evidence.quote)})
                                    if not quoted_years or field.raw_text != "、".join(map(str, quoted_years)):
                                        fail("SCOPED_REVIEW_MISMATCH", "Word 的 L 列须引用考核栏目中的优秀年度。", 422)
                                else:
                                    text = "\n".join(str(unit.get("text") or "") for unit in source["source_units"])
                                    heading = has_review_heading(text)
                                    if ((field.blank_reason == "ANNUAL_REVIEW_NOT_FOUND") != (not heading)
                                            or (heading and field.blank_reason == "NO_EXCELLENT_YEARS"
                                                and excellent_years(text))):
                                        fail("SCOPED_REVIEW_MISMATCH", "Word 的 L 列留空原因与正文不一致。", 422)
                        elif column >= 10 and (field.raw_text or field.blank_reason != "OUT_OF_SCOPE"
                                             or field.evidence_refs or field.fragments or field.source_only_fragments
                                             or field.approval or field.assessment):
                            fail("SCOPED_WRITE_SCOPE_INVALID", "范围外业务列不得提取或写入。", 422)
                        if 2 <= column <= 9:
                            ref = (source["field_refs"].get(get_column_letter(column))
                                   if scoped_pdf else None)
                            if field.raw_text:
                                if (len(field.evidence_refs) != 1 or field.approval or field.assessment
                                        or field.fragments or field.source_only_fragments):
                                    fail("SCOPED_EVIDENCE_REQUIRED", "B–I 已填写列必须引用本文件的正式原文。", 422,
                                         column=get_column_letter(column))
                                evidence = field.evidence_refs[0]
                                if scoped_pdf:
                                    if (not ref or str(evidence.source_unit_id) != ref["source_unit_id"]
                                            or evidence.quote != ref["quote"]
                                            or evidence.start != ref["start"] or evidence.end != ref["end"]):
                                        fail("SCOPED_EVIDENCE_MISMATCH", "字段证据与冻结 PDF OCR 结果不一致。", 422,
                                             column=get_column_letter(column))
                                if "".join(field.raw_text.split()) not in "".join(evidence.quote.split()):
                                    fail("SCOPED_EVIDENCE_MISMATCH", "字段值不在本文件原文引文中。", 422,
                                         column=get_column_letter(column))
                                if column == 2:
                                    from app.modules.external_extraction.scoped_schemas import labeled_b_column_value
                                    name = labeled_b_column_value(evidence.quote)
                                    if name is not None and field.raw_text != name:
                                        fail("SCOPED_B_VALUE_INVALID", "B 列只能填写“单位/学院”标签后的学院名称。", 422)
                            elif (not scoped_review
                                  or field.blank_reason not in {"SOURCE_FIELD_NOT_FOUND", "OCR_UNREADABLE"}
                                  or ref is not None or field.evidence_refs or field.fragments
                                  or field.source_only_fragments or field.approval or field.assessment):
                                fail("SCOPED_FIELD_BLANK_INVALID", "B–I 留空必须无冻结值并给出有效原因。", 422,
                                     column=get_column_letter(column))
                    value = field.model_dump(mode="json")
                    try:
                        value["evidence_refs"] = [validate_quote(ref, group_units)
                                                  for ref in field.evidence_refs]
                    except ValueError as exc:
                        code, _, message = str(exc).partition(":")
                        fail(code, message.strip(), 422, field_label=spec["label"])
                    values[spec["key"]] = value
                records.append({"fields": values})
            if not records:
                continue
            request_data = {key: source[key] for key in (
                "document_id", "document_version_id", "extraction_run_id", "source_sha256")}
            if template_all:
                candidate = next(record for record in payload.records if record.source_index == source_index)
                request_data["supplemental_sources"] = [
                    {key: frozen_sources[index][key] for key in (
                        "document_id", "document_version_id", "extraction_run_id", "source_sha256")}
                    for index in candidate.source_indices[1:]]
            request_data.update(records=records, content_origin="WORKBUDDY_SCOPED" if limited and scoped_pdf else "WORKBUDDY",
                                idempotency_key=f"writeback-{task.id}-{source_index}",
                                target_working_copy_id=working.id, target_document_id=working.document_id,
                                target_document_version_id=current.id)
            self.datasets.append(dataset.id, DatasetAppend.model_validate(request_data))
        added = [record for record in self.datasets._records(dataset)
                 if int(record.get("dataset_revision") or 0) > start_revision]
        if len(added) != len(payload.records) or not added or any(record["needs_review"] for record in added):
            fail("WRITEBACK_RECORDS_NEED_REVIEW", "候选存在不可写字段，回传文件未发布。", 422)
        expected = [[record["fields"][field["key"]].get(
            "display_value", record["fields"][field["key"]]["value"])
            for field in dataset.schema_json] for record in added]
        if limited or full_row or template_all:
            for row in expected:
                if limited or full_row or (template_all and left == 1 and
                        dataset.schema_json[0]["label"].strip() in {"序号", "编号"}):
                    row[0] = None
                if limited:
                    for offset in range(9, len(row)):
                        if not (scoped_review and offset == 11):
                            row[offset] = None
        inserted = self._verify_workbook(base_path=base_path, candidate_path=file_path,
                                         dataset=dataset, expected_rows=expected,
                                         scoped_a_i=limited or full_row or template_all)
        if inserted:
            updated_layout = dict(dataset.target_layout_json)
            left_col, top_row, right_col, bottom_row = range_boundaries(updated_layout["data_range"])
            updated_layout["data_range"] = (f"{get_column_letter(left_col)}{top_row}:"
                                                  f"{get_column_letter(right_col)}{bottom_row + inserted}")
            dataset.target_layout_json = updated_layout
        export_id = str(uuid4())
        version_number = int(self.db.query(func.max(DocumentVersion.version_number)).filter(
            DocumentVersion.document_id == working.document_id).scalar() or 0) + 1
        storage_path = self.datasets.storage.dataset_working_version_relative_path(
            dataset_id=dataset.id, export_id=export_id, version_number=version_number)
        self.datasets.storage.import_working_copy(
            source=file_path, relative_path=storage_path, expected_sha256=file_sha256)
        working, version = self.datasets._publish_working_version(
            dataset=dataset, working_copy=working, parent_version=current,
            storage_path=storage_path, version_number=version_number,
            sha256=file_sha256, size_bytes=file_path.stat().st_size)
        version.extraction_policy_json = {**dict(version.extraction_policy_json or {}),
                                          "structured_dataset_export_mode": "workbuddy-file-writeback-v1"}
        export = StructuredDatasetExport(
            id=export_id, dataset_id=dataset.id, dataset_revision=dataset.revision,
            storage_path=storage_path, sha256=file_sha256, size_bytes=file_path.stat().st_size,
            document_version_id=version.id)
        self.db.add(export)
        pending_result = {**self.datasets._export_result(export, reused=False),
                          "task_record_count": len(added)}
        publication = stage_publication(
            self.db, dataset=dataset, working_copy=working, version=version,
            base_version_id=current.id, task=task, result=pending_result, agent_run=self.datasets.agent_run)
        task.status = "PUBLISHING"
        task.payload_json = body if template_all else task.payload_json
        task.result_json = {}
        task.updated_at = utcnow()
        return {**self.receipt(task), "_publication_id": publication.id}

    def _verify_workbook(self, *, base_path: Path, candidate_path: Path,
                         dataset, expected_rows: list[list], scoped_a_i: bool = False) -> int:
        """只允许目标工作表的新增数据格变化，包内其他部件必须原字节保留。"""
        layout = dataset.target_layout_json or None
        labels = [field["label"] for field in dataset.schema_json]
        try:
            inspect_xlsx(base_path)
            inspect_xlsx(candidate_path)
        except XlsxSafetyError as exc:
            fail(exc.code, str(exc), 413 if "TOO_LARGE" in exc.code or "TOO_COMPLEX" in exc.code else 422)
        try:
            base_book = load_workbook(base_path, data_only=False)
            candidate_book = load_workbook(candidate_path, data_only=False)
        except Exception:
            fail("WRITEBACK_XLSX_INVALID", "回传 Excel 无法读取或已损坏。", 422)
        try:
            if base_book.sheetnames != candidate_book.sheetnames:
                fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件改变了工作表清单或顺序。", 422)
            name = dataset.target_sheet_name or "Sheet1"
            if name not in base_book.sheetnames:
                fail("WRITEBACK_TEMPLATE_CHANGED", "目标工作表不存在。", 422)
            base_sheet, candidate_sheet = base_book[name], candidate_book[name]
            if layout:
                from app.modules.structured_extraction.workbook_layout import WorkbookLayout, read_layout, sequence_offsets
                parsed = WorkbookLayout.model_validate(layout)
                base_labels, _ = read_layout(base_sheet, parsed)
                candidate_labels, _ = read_layout(candidate_sheet, parsed)
                left, top, right, bottom = range_boundaries(layout["data_range"])
                sequences = sequence_offsets(labels)
                if scoped_a_i and left == 1:
                    sequences.add(0)
                available = [row for row in range(top, bottom + 1)
                    if not any(base_sheet.cell(row, left + offset).value is not None
                               for offset in range(len(labels)) if offset not in sequences)]
                inserted = max(0, len(expected_rows) - len(available)) if scoped_a_i else 0
                if bottom + inserted > 1_048_576:
                    fail("TARGET_EXCEL_DATA_RANGE_FULL", "已达到 Excel 最大行数。", 422)
                available.extend(range(bottom + 1, bottom + 1 + inserted))
                if inserted and any(cell.data_type == "f" for sheet in base_book
                                    for sheet_row in sheet for cell in sheet_row):
                    fail("WRITEBACK_ROW_INSERT_UNSUPPORTED", "工作簿公式可能引用移动的尾注。", 422)
            else:
                base_name, base_fields, _ = self.datasets._read_excel_schema(base_path, name)
                candidate_name, candidate_fields, _ = self.datasets._read_excel_schema(candidate_path, name)
                base_labels = [field.label for field in base_fields]
                candidate_labels = [field.label for field in candidate_fields]
                if base_name != name or candidate_name != name:
                    fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件的目标工作表不一致。", 422)
                left, right = 1, len(labels)
                header = next((row for row in base_sheet.iter_rows()
                    if any(cell.value is not None for cell in row)), None)
                if header is None:
                    fail("WRITEBACK_TEMPLATE_CHANGED", "原表缺少标题行。", 422)
                top = header[0].row + 1
                last = max((row for row in range(top, base_sheet.max_row + 1)
                            if any(base_sheet.cell(row, col).value is not None
                                   for col in range(left, right + 1))), default=top - 1)
                available = list(range(last + 1, last + 1 + len(expected_rows)))
                inserted = 0
            if base_labels != labels or candidate_labels != labels:
                fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件的完整表头与基线不一致。", 422)
            if len(expected_rows) > len(available):
                fail("TARGET_EXCEL_DATA_RANGE_FULL", "模板数据区没有足够的可填写行。", 422)
            allowed = set()
            height_rows = set()
            for row_number, values in zip(available, expected_rows):
                if len(values) != len(labels):
                    fail("WRITEBACK_ROWS_MISMATCH", "回传记录与目标列数量不一致。", 422)
                if scoped_a_i and (row_number > bottom or base_sheet.cell(row_number, 1).value is None):
                    previous = next((candidate_sheet.cell(prior, 1).value
                                     for prior in range(row_number - 1, top - 1, -1)
                                     if candidate_sheet.cell(prior, 1).value is not None), None)
                    try:
                        sequence_value = int(previous) if previous is not None else 0
                        if (sequence_value < (0 if previous is None else 1)
                                or previous is not None and str(sequence_value) != str(previous)):
                            raise ValueError
                    except (TypeError, ValueError):
                        fail("TARGET_EXCEL_SEQUENCE_MISSING", "A 列无法从上一序号生成。", 422,
                             cell=f"A{row_number}")
                    values[0] = sequence_value + 1
                for offset, expected in enumerate(values):
                    col = left + offset
                    previous = base_sheet.cell(row_number, col)
                    actual = candidate_sheet.cell(row_number, col)
                    previous_value = None if inserted and row_number > bottom else previous.value
                    if previous_value is not None:
                        sequence_preserved = layout and offset in sequences and expected is None
                        if (not sequence_preserved and str(previous_value) != str(expected)) or actual.value != previous_value:
                            fail("WRITEBACK_EXISTING_CELL_CHANGED", "回传文件覆盖了模板既有单元格。", 422,
                                 cell=f"{get_column_letter(col)}{row_number}")
                    elif expected is not None and (actual.value is None or
                            str(excel_content_key(actual.value)) != str(excel_content_key(expected))):
                        fail("WRITEBACK_ROWS_MISMATCH", "回传 Excel 单元格与已核验字段不一致。", 422,
                             cell=f"{get_column_letter(col)}{row_number}")
                    elif expected is None and actual.value is not None:
                        fail("WRITEBACK_ROWS_MISMATCH", "留空字段被回传文件填入了内容。", 422,
                             cell=f"{get_column_letter(col)}{row_number}")
                    if previous_value is None and isinstance(actual.value, str) and (
                            excel_display_text(actual.value) != actual.value):
                        fail("WRITEBACK_BREAK_TAG_REMAINING", "新增单元格仍含 HTML 换行标签。", 422,
                             cell=f"{get_column_letter(col)}{row_number}")
                    if any(isinstance(value, str) and "\n" in excel_display_text(value)
                           for value in (expected, actual.value)):
                        height_rows.add(row_number)
                    if expected is not None and actual.data_type == "f":
                        fail("WRITEBACK_FORMULA_NOT_ALLOWED", "新数据格不能写入公式。", 422)
                    allowed.add(f"{get_column_letter(col)}{row_number}")
        finally:
            base_book.close()
            candidate_book.close()
        try:
            with zipfile.ZipFile(base_path) as base_zip, zipfile.ZipFile(candidate_path) as result_zip:
                names = set(base_zip.namelist())
                if names != set(result_zip.namelist()):
                    fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件增删了工作簿部件。", 422)
                lookup = {name: b"" for name in names}
                lookup["xl/workbook.xml"] = base_zip.read("xl/workbook.xml")
                lookup["xl/_rels/workbook.xml.rels"] = base_zip.read("xl/_rels/workbook.xml.rels")
                worksheet_path = _xlsx_worksheet_path(lookup, dataset.target_sheet_name or "Sheet1")
                for name in names - {worksheet_path}:
                    hashes = []
                    for archive in (base_zip, result_zip):
                        part_hash = hashlib.sha256()
                        with archive.open(name) as stream:
                            while chunk := stream.read(1024 * 1024):
                                part_hash.update(chunk)
                        hashes.append(part_hash.digest())
                    if hashes[0] != hashes[1]:
                        fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件改变了非目标工作表、样式或其他部件。", 422)
                base_xml = base_zip.read(worksheet_path)
                result_xml = result_zip.read(worksheet_path)
        except (OSError, zipfile.BadZipFile, RuntimeError, KeyError, ValueError):
            fail("WRITEBACK_XLSX_INVALID", "回传 Excel 包无法完整读取。", 422)
        base_root = ET.fromstring(base_xml)
        result_root = ET.fromstring(result_xml)
        if inserted:
            _shift_expected_footer(base_root, bottom + 1, inserted)
        if base_root.attrib != result_root.attrib:
            fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件改变了工作表根属性。", 422)
        def local(tag: str) -> str:
            return tag.rsplit("}", 1)[-1]
        def parts(root):
            nodes = list(root)
            if len({local(node.tag) for node in nodes}) != len(nodes):
                fail("WRITEBACK_XLSX_INVALID", "工作表包含重复结构节点。", 422)
            return {local(node.tag): node for node in nodes}
        base_nodes, result_nodes = parts(base_root), parts(result_root)
        if base_nodes.keys() != result_nodes.keys():
            fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件改变了工作表结构。", 422)
        for name in base_nodes:
            if name not in {"sheetData", "dimension"} and ET.tostring(base_nodes[name]) != ET.tostring(result_nodes[name]):
                fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件改变了目标表非数据区域。", 422)
        if inserted:
            before_dimension, after_dimension = base_nodes.get("dimension"), result_nodes.get("dimension")
            if ((before_dimension is None) != (after_dimension is None)
                    or (before_dimension is not None and
                        ET.tostring(before_dimension) != ET.tostring(after_dimension))):
                fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件的工作表范围与行插入不一致。", 422)
        def xml_rows(root):
            sheet_data = parts(root).get("sheetData")
            if sheet_data is None:
                fail("WRITEBACK_XLSX_INVALID", "回传文件缺少工作表数据区。", 422)
            return {int(row.attrib["r"]): row for row in sheet_data if local(row.tag) == "row"}
        base_rows, result_rows = xml_rows(base_root), xml_rows(result_root)
        if not base_rows.keys() <= result_rows.keys():
            fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件删除了原有行。", 422)
        for number, result_row in result_rows.items():
            base_row = base_rows.get(number)
            if (base_row is not None and
                    [ET.tostring(child) for child in base_row if local(child.tag) != "c"] !=
                    [ET.tostring(child) for child in result_row if local(child.tag) != "c"]):
                fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件改变了原有行的非单元格内容。", 422)
            if base_row is None and any(local(child.tag) != "c" for child in result_row):
                fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件新行包含未授权结构。", 422)
            result_cells = {cell.attrib.get("r"): cell for cell in result_row if local(cell.tag) == "c"}
            base_cells = ({cell.attrib.get("r"): cell for cell in base_row if local(cell.tag) == "c"}
                          if base_row is not None else {})
            if (len(result_cells) != sum(local(child.tag) == "c" for child in result_row)
                    or (base_row is not None and len(base_cells) != sum(local(child.tag) == "c" for child in base_row))):
                fail("WRITEBACK_XLSX_INVALID", "工作表存在重复或缺少坐标的单元格。", 422)
            prior_attributes = dict(base_row.attrib) if base_row is not None else {}
            current_attributes = dict(result_row.attrib)
            if number in height_rows:
                for attribute in ("ht", "customHeight"):
                    prior_attributes.pop(attribute, None)
                    current_attributes.pop(attribute, None)
            if base_row is not None and (prior_attributes != current_attributes
                    or not base_cells.keys() <= result_cells.keys()):
                fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件改变了原有行结构。", 422)
            if base_row is None and not any(coordinate in allowed for coordinate in result_cells):
                fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件增加了未授权空行。", 422)
            for coordinate, cell in result_cells.items():
                before = base_cells.get(coordinate)
                if coordinate not in allowed:
                    if before is None or ET.tostring(before) != ET.tostring(cell):
                        fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件改变了未授权单元格。", 422,
                             cell=coordinate)
                elif before is not None and before.attrib.get("s") != cell.attrib.get("s"):
                    fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件改变了数据格原有样式。", 422,
                         cell=coordinate)
            if base_row is None and any(coordinate not in allowed for coordinate in result_cells):
                fail("WRITEBACK_TEMPLATE_CHANGED", "回传文件增加了未授权单元格。", 422)
        return inserted

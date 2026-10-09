"""共享累计表格 Demo：复用工作副本、结构化事实、证据规范化和 XLSX writer。"""

from __future__ import annotations

import hashlib
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath
from uuid import uuid4

from fastapi import HTTPException
from fastapi.responses import FileResponse
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from sqlalchemy import func, or_

from app.core.config import get_settings
from app.db.models import (
    ChangeItem, Document, DocumentArtifact, DocumentElement, DocumentExtractionRun, DocumentPage,
    DocumentVersion, FileObject, ManagedFile, ManagedRoot, StructuredDataset,
    StructuredDatasetCorrection, StructuredDatasetExport, StructuredDatasetMember, StructuredDatasetTask, StructuredExtractionRun,
    WorkingCopy, WorkingCopyPathRecord, WorkingCopyRoot, utcnow,
)
from app.modules.agent.tool_schemas import StructuredFieldSpec, StructuredImageExtractionInput
from app.modules.file_lifecycle.shared_access import CanonicalWorkingFileResolver, CanonicalWorkingFileError
from app.modules.file_lifecycle.shared_workspace import get_shared_workspace_id
from app.modules.file_lifecycle.storage import FileLifecycleStorageService
from app.modules.files.office_conversion import (
    CONVERTED_XLSX_ARTIFACT_TYPE,
    OfficeConversionError,
    OfficeDerivativeService,
    XLSX_CONTENT_TYPE,
    XLS_TO_XLSX,
)
from app.modules.managed_files.jobs import FilesystemJobQueue
from app.modules.structured_extraction.export import (
    _append_xlsx_rows_atomic,
    _safe_spreadsheet_text,
    _write_xlsx_atomic,
)
from app.modules.structured_extraction.repository import StructuredExtractionRepository
from app.modules.structured_extraction.normalization import normalize_field_value
from app.modules.structured_extraction.schemas import (
    CandidateExtraction, CandidateRecord, CandidateFieldValue,
    LayoutParseResult, LayoutPage, LayoutElement,
)
from app.modules.structured_extraction.service import StructuredExtractionService
from .schemas import DatasetCreate, DatasetCreateFromExcel, DatasetAppend, DatasetQuery
from .assessment import assess_field, display_value, validate_quote, validate_approval, validate_scope
from .xlsx_safety import XlsxSafetyError, inspect_xlsx
from .publication import assert_publication_ready, pending_publication, pending_replacement, stage_publication
from app.modules.structured_extraction.evidence import EvidenceElement
from app.modules.structured_extraction.workbook_layout import WorkbookLayout, read_layout


def fail(code: str, message: str, status: int = 409, **details):
    """通过现有异常响应转换输出稳定业务错误。"""
    raise HTTPException(
        status_code=status,
        detail={"code": code, "message": message, **details},
    )


def digest(value) -> str:
    """对完整请求语义生成摘要，正文不进入审计字段。"""
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


class StructuredDatasetService:
    """所有登录用户共享数据集；写入由 runtime 白名单节点调用并提交事务。"""

    def __init__(self, db, user, agent_run=None):
        """依赖限于本次请求，不缓存用户或数据库会话。"""
        self.db, self.user, self.agent_run = db, user, agent_run
        self.settings = get_settings()
        self.repository = StructuredExtractionRepository(db)
        self.storage = FileLifecycleStorageService(self.settings)
        self.workflow_task_id = None
        self.source_fact_keys: set[str] = set()
        self.full_row_ocr_mode = False

    def _assert_task_write_access(self, dataset):
        """跨请求表格任务占用目标时，旧粒度接口也不能穿透并发发布。"""
        pending = pending_publication(self.db, dataset_id=dataset.id)
        if pending is not None and pending.agent_run_id != getattr(self.agent_run, "id", None):
            assert_publication_ready(self.db, dataset_id=dataset.id)
        if (dataset.active_task_id and dataset.active_task_id != self.workflow_task_id
                and not (pending is not None and pending.id == dataset.active_task_id
                         and pending.agent_run_id == getattr(self.agent_run, "id", None))):
            fail("DATASET_TASK_IN_PROGRESS", "该表格有尚未发布完成的任务，请先恢复该任务。")

    def dataset(self, dataset_id: str, *, lock=False):
        """按共享工作区读取，创建人不构成访问限制。"""
        query = self.db.query(StructuredDataset).filter(
            StructuredDataset.id == dataset_id,
            StructuredDataset.workspace_id == get_shared_workspace_id(self.db),
        )
        result = (query.with_for_update() if lock else query).one_or_none()
        if result is None:
            fail("DATASET_NOT_FOUND", "累计表格不存在。", 404)
        if pending_replacement(self.db, dataset_id=result.id) is not None:
            fail("DATASET_TARGET_REPLACEMENT_PENDING", "同名新目标正在发布，请等待发布完成。", 423)
        return result

    def create(self, request: DatasetCreate):
        """固定共享 schema，同用户同创建键重试复用原对象。"""
        key = digest([self.user.id, "dataset-create", request.idempotency_key])
        payload_digest = digest(request.model_dump(mode="json", exclude={"idempotency_key"}))
        existing = self.db.query(StructuredDataset).filter_by(creation_key=key).one_or_none()
        if existing:
            if existing.creation_digest != payload_digest:
                fail("IDEMPOTENCY_CONFLICT", "此创建键已用于不同表格。")
            return {**self.get(existing.id), "reused": True}
        target_filename = self._normalize_target_filename(request.filename or request.name)
        self._assert_target_filename_available(target_filename)
        row = StructuredDataset(
            workspace_id=get_shared_workspace_id(self.db), created_by=self.user.id,
            name=request.name, target_filename=target_filename, target_sheet_name="Sheet1",
            schema_json=[f.model_dump() for f in request.fields], base_rows_json=[],
            creation_key=key, creation_digest=payload_digest,
        )
        self.db.add(row)
        self.db.flush()
        return {**self.get(row.id), "reused": False}

    def create_from_excel(self, request: DatasetCreateFromExcel):
        """读取当前工作副本中的真实 XLSX，并把它注册为新的持续追加目标。"""

        key = digest([self.user.id, "dataset-create-from-excel", request.idempotency_key])
        payload_digest = digest(request.model_dump(mode="json", exclude={"idempotency_key"}))
        existing = self.db.query(StructuredDataset).filter_by(creation_key=key).one_or_none()
        if existing:
            if existing.creation_digest != payload_digest:
                fail("IDEMPOTENCY_CONFLICT", "此创建键已用于不同 Excel 目标。")
            assert_publication_ready(self.db, dataset_id=existing.id)
            source_version = (
                self.db.get(DocumentVersion, existing.target_source_version_id)
                if existing.target_source_version_id
                else None
            )
            if source_version is not None and Path(source_version.filename).suffix.lower() == ".xls":
                self._retire_reused_legacy_visible_file(existing, source_version=source_version)
                return self._legacy_target_result(existing, reused=True)
            return {**self.get(existing.id), "reused": True}
        try:
            source = CanonicalWorkingFileResolver(self.db).resolve_document(
                document_id=str(request.document_id),
                document_version_id=(str(request.document_version_id) if request.document_version_id else None),
            )
        except CanonicalWorkingFileError as exc:
            fail(exc.code, str(exc))
        working_copy = source.working_copy
        version = source.document_version
        if working_copy.status != "ACTIVE" or version.id != working_copy.current_version_id:
            fail("TARGET_EXCEL_VERSION_CONFLICT", "上传 Excel 已不是当前活动版本，请重新选择。")
        source_suffix = Path(working_copy.filename).suffix.lower()
        if source_suffix == ".xls":
            return self._enqueue_legacy_excel_target(
                request=request,
                source=source,
                creation_key=key,
                creation_digest=payload_digest,
            )
        if source_suffix != ".xlsx":
            fail("TARGET_EXCEL_TYPE_UNSUPPORTED", "当前只支持将 .xlsx 注册为目标表。", 422)
        if Path(working_copy.relative_path).name != working_copy.filename:
            fail("TARGET_EXCEL_SCHEMA_INVALID", "工作副本路径与展示文件名不一致，需要先修复文件。")
        already_bound = self.db.query(StructuredDataset).filter_by(
            workspace_id=get_shared_workspace_id(self.db), working_copy_id=working_copy.id
        ).one_or_none()
        if already_bound is not None:
            fail("TARGET_EXCEL_ALREADY_REGISTERED", "该 Excel 已经是累计目标表，请按文件名继续追加。")
        archive_target = self._replacement_for_uploaded_target(
            working_copy.filename, exclude_working_copy_id=working_copy.id)
        source_path = self.storage.working_copy_path(version.storage_path)
        if not source_path.is_file() or self.storage.sha256_file(source_path) != version.sha256:
            fail("TARGET_EXCEL_NOT_FOUND", "目标 Excel 内容缺失或校验失败。", 404)
        sheet_name, fields, base_rows = self._read_excel_schema(source_path, request.sheet_name, request.layout)
        if request.extraction_run_id is not None:
            # WorkBuddy 模板模式不仅信任字段清单，还逐格核对上传的当前 XLSX。
            expected = self._validate_workbuddy_target_extraction(
                document_id=version.document_id, document_version_id=version.id,
                extraction_run_id=str(request.extraction_run_id), requested_sheet=sheet_name)
            self._require_workbuddy_header_grid(expected=expected, sheet_name=sheet_name,
                                                layout=request.layout)
            self._assert_workbuddy_schema_matches(
                expected=expected, sheet_name=sheet_name,
                actual_labels=[field.label for field in fields], layout=request.layout,
                converted_path=source_path)
        row = StructuredDataset(
            workspace_id=get_shared_workspace_id(self.db),
            created_by=self.user.id,
            name=Path(working_copy.filename).stem,
            target_filename=working_copy.filename,
            target_sheet_name=sheet_name,
            target_layout_json=request.layout.model_dump() if request.layout else {},
            schema_json=[field.model_dump() for field in fields],
            base_rows_json=base_rows,
            target_source_version_id=version.id,
            creation_key=key,
            creation_digest=payload_digest,
            working_copy_id=working_copy.id,
        )
        self.db.add(row)
        self.db.flush()
        # 注册时先把上传版本复制进隐藏历史目录；之后可见同名文件可以安全更新。
        export_id = str(uuid4())
        history_path = self.storage.dataset_working_version_relative_path(
            dataset_id=row.id,
            export_id=export_id,
            version_number=version.version_number,
        )
        self.storage.import_working_copy(
            source=source_path,
            relative_path=history_path,
            expected_sha256=version.sha256,
        )
        version.storage_path = history_path
        file_object = self.db.query(FileObject).filter_by(
            document_id=working_copy.document_id,
            storage_backend="working_copy_local",
        ).one_or_none()
        if file_object is not None:
            file_object.storage_path = self._working_copy_storage_path(working_copy)
        self.db.add(StructuredDatasetExport(
            id=export_id,
            dataset_id=row.id,
            dataset_revision=row.revision,
            storage_path=history_path,
            sha256=version.sha256,
            size_bytes=version.size_bytes,
            document_version_id=version.id,
        ))
        cleanup = self._relocate_dataset_target(
            dataset=row,
            working_copy=working_copy,
            current_version=version,
        )
        self.db.add(WorkingCopyPathRecord(
            working_copy_id=working_copy.id,
            sequence_number=self._next_path_sequence(working_copy.id),
            operation_type="DATASET_TARGET_REGISTERED",
            before_relative_path=working_copy.relative_path,
            after_relative_path=working_copy.relative_path,
            before_filename=working_copy.filename,
            after_filename=working_copy.filename,
            document_version_id=version.id,
            content_sha256=version.sha256,
            status="COMPLETED",
            executed_by=self.user.id,
        ))
        self.db.flush()
        result = {**self.get(row.id), "reused": False, "source": "UPLOADED_EXCEL"}
        publication = stage_publication(
            self.db, dataset=row, working_copy=working_copy, version=version,
            result=result, cleanup=cleanup, archive_target=archive_target, agent_run=self.agent_run)
        if archive_target:
            working_copy.status = "PUBLISHING"
            self.db.get(ManagedFile, working_copy.managed_file_id).status = "PUBLISHING"
        row.active_task_id = publication.id
        result["_publication_id"] = publication.id
        if cleanup:
            result["_visible_cleanup"] = cleanup
        return result

    def _enqueue_legacy_excel_target(
        self,
        *,
        request: DatasetCreateFromExcel,
        source,
        creation_key: str,
        creation_digest: str,
    ) -> dict:
        """校验 WorkBuddy 的 XLS 读取快照并幂等提交服务端格式提升任务。"""

        if request.legacy_target_policy != "CONVERT_TO_XLSX":
            fail(
                "TARGET_EXCEL_LEGACY_POLICY_REQUIRED",
                "旧版 .xls 目标表必须明确使用 CONVERT_TO_XLSX 格式提升策略。",
                422,
            )
        if request.extraction_run_id is None:
            fail(
                "TARGET_EXCEL_EXTRACTION_REQUIRED",
                "旧版 .xls 必须先由 WorkBuddy 完成电子表格读取并提供 extraction_run_id。",
                422,
            )
        expected = self._validate_workbuddy_target_extraction(
            document_id=source.document_version.document_id,
            document_version_id=source.document_version.id,
            extraction_run_id=str(request.extraction_run_id),
            requested_sheet=request.sheet_name,
        )
        # 复杂模板没有真实单元格网格时立即拒绝，不能先排入耗时转换队列。
        self._require_workbuddy_header_grid(expected=expected, sheet_name=request.sheet_name,
                                            layout=request.layout)
        target_filename = f"{Path(source.working_copy.filename).stem}.xlsx"
        self._replacement_for_uploaded_target(
            target_filename, exclude_working_copy_id=source.working_copy.id)
        job = FilesystemJobQueue(self.db).create_job(
            job_type="STRUCTURED_DATASET_TARGET_PREPARE",
            queue_name="STRUCTURED_EXTRACTION",
            root_id=None,
            created_by=self.user.id,
            deduplication_key=f"structured-dataset-target:{creation_key}",
            priority=20,
            max_attempts=3,
            payload={
                "request": request.model_dump(mode="json"),
                "creation_key": creation_key,
                "creation_digest": creation_digest,
                "structured_dataset_task_id": self.workflow_task_id,
            },
        )
        stored_digest = str((job.payload_json or {}).get("creation_digest") or "")
        if stored_digest and stored_digest != creation_digest:
            fail("IDEMPOTENCY_CONFLICT", "此创建键已用于不同 Excel 目标。")
        if job.status == "FAILED":
            fail(
                "TARGET_EXCEL_CONVERSION_FAILED",
                job.error_message or "旧版 Excel 格式提升失败，原文件未修改。",
                filesystem_job_id=job.id,
            )
        return {
            "kind": "filesystem_job",
            "status": "PREPARING_TARGET",
            "filesystem_job_id": job.id,
            "source_filename": source.working_copy.filename,
            "target_filename": target_filename,
            "source_document_id": source.document_version.document_id,
            "source_document_version_id": source.document_version.id,
            "converted_from_legacy": False,
            "original_unchanged": True,
            "reused": job.status != "PENDING" or job.attempt_count > 0,
        }

    def prepare_legacy_excel_target(self, request: DatasetCreateFromExcel) -> dict:
        """由异步 worker 将已验收的 XLS 提升为同一逻辑文件的 XLSX 当前版本。"""

        key = digest([self.user.id, "dataset-create-from-excel", request.idempotency_key])
        payload_digest = digest(request.model_dump(mode="json", exclude={"idempotency_key"}))
        existing = self.db.query(StructuredDataset).filter_by(creation_key=key).one_or_none()
        if existing is not None:
            if existing.creation_digest != payload_digest:
                fail("IDEMPOTENCY_CONFLICT", "此创建键已用于不同 Excel 目标。")
            return self._legacy_target_result(existing, reused=True)
        try:
            source = CanonicalWorkingFileResolver(self.db).resolve_document(
                document_id=str(request.document_id),
                document_version_id=(str(request.document_version_id) if request.document_version_id else None),
            )
        except CanonicalWorkingFileError as exc:
            fail(exc.code, str(exc))
        working_copy = (
            self.db.query(WorkingCopy)
            .filter(WorkingCopy.id == source.working_copy.id)
            .with_for_update()
            .one()
        )
        source_version = (
            self.db.query(DocumentVersion)
            .filter(DocumentVersion.id == source.document_version.id)
            .with_for_update()
            .one()
        )
        if working_copy.current_version_id != source_version.id or working_copy.status != "ACTIVE":
            fail("TARGET_EXCEL_VERSION_CONFLICT", "目标 XLS 在格式提升前已发生变化，请重新读取并确认。")
        if Path(working_copy.filename).suffix.lower() != ".xls":
            fail("TARGET_EXCEL_VERSION_CONFLICT", "目标文件已不再是待提升的 .xls 当前版本。")
        if request.legacy_target_policy != "CONVERT_TO_XLSX" or request.extraction_run_id is None:
            fail("TARGET_EXCEL_EXTRACTION_REQUIRED", "格式提升缺少固定策略或 WorkBuddy 提取运行。", 422)
        expected = self._validate_workbuddy_target_extraction(
            document_id=source_version.document_id,
            document_version_id=source_version.id,
            extraction_run_id=str(request.extraction_run_id),
            requested_sheet=request.sheet_name,
        )
        # 异步执行时再检查一次固定证据，防止旧任务或绕过入口的调用写入近似模板。
        self._require_workbuddy_header_grid(expected=expected, sheet_name=request.sheet_name,
                                            layout=request.layout)
        old_filename = working_copy.filename
        old_relative_path = working_copy.relative_path
        target_filename = f"{Path(old_filename).stem}.xlsx"
        archive_target = self._replacement_for_uploaded_target(
            target_filename, exclude_working_copy_id=working_copy.id)
        old_current_storage_path = self._working_copy_storage_path(working_copy)
        source_path = self.storage.working_copy_path(old_current_storage_path)
        if not source_path.is_file() or self.storage.sha256_file(source_path) != source_version.sha256:
            fail("TARGET_EXCEL_NOT_FOUND", "旧版 Excel 当前文件缺失或内容校验失败。", 404)
        document = self.db.get(Document, working_copy.document_id)
        managed_file = self.db.get(ManagedFile, working_copy.managed_file_id)
        if document is None or managed_file is None:
            fail("DATASET_WORKING_COPY_INVALID", "目标 Excel 的文件关联不完整。")
        try:
            artifact = OfficeDerivativeService(db=self.db).get_or_create(
                document=document,
                document_version=source_version,
                source_path=source_path,
                spec=XLS_TO_XLSX,
            )
        except OfficeConversionError as exc:
            fail(exc.code, exc.message, 503 if exc.retryable else 422)
        sheet_name, fields, base_rows = self._read_excel_schema(artifact.file_path, request.sheet_name, request.layout)
        self._assert_workbuddy_schema_matches(
            expected=expected,
            sheet_name=sheet_name,
            actual_labels=[field.label for field in fields],
            layout=request.layout,
            converted_path=artifact.file_path,
        )
        row = StructuredDataset(
            workspace_id=get_shared_workspace_id(self.db),
            created_by=self.user.id,
            name=Path(target_filename).stem,
            target_filename=target_filename,
            target_sheet_name=sheet_name,
            target_layout_json=request.layout.model_dump() if request.layout else {},
            schema_json=[field.model_dump() for field in fields],
            base_rows_json=base_rows,
            target_source_version_id=source_version.id,
            creation_key=key,
            creation_digest=payload_digest,
            working_copy_id=working_copy.id,
        )
        self.db.add(row)
        self.db.flush()
        generated_managed_root, generated_working_root = self._ensure_generated_root(dataset=row)
        target_relative_path = f"{row.id}/{target_filename}"
        export_id = str(uuid4())
        version_number = int(
            self.db.query(func.max(DocumentVersion.version_number))
            .filter(DocumentVersion.document_id == document.id)
            .scalar()
            or 0
        ) + 1
        source_history_path = self.storage.dataset_source_version_relative_path(
            dataset_id=row.id,
            document_version_id=source_version.id,
            version_number=source_version.version_number,
            suffix=".xls",
        )
        xlsx_history_path = self.storage.dataset_working_version_relative_path(
            dataset_id=row.id,
            export_id=export_id,
            version_number=version_number,
        )
        self.storage.import_working_copy(
            source=source_path,
            relative_path=source_history_path,
            expected_sha256=source_version.sha256,
        )
        self.storage.import_working_copy(
            source=artifact.file_path,
            relative_path=xlsx_history_path,
            expected_sha256=artifact.sha256,
        )
        new_version = DocumentVersion(
            document_id=document.id,
            version_number=version_number,
            parent_version_id=source_version.id,
            working_copy_id=working_copy.id,
            storage_tier="WORKING_COPY",
            storage_path=xlsx_history_path,
            filename=target_filename,
            content_type=XLSX_CONTENT_TYPE,
            size_bytes=artifact.file_path.stat().st_size,
            sha256=artifact.sha256,
            source_type="STRUCTURED_DATASET_XLS_CONVERSION",
            extraction_policy_json=dict(source_version.extraction_policy_json or {}),
            created_by=self.user.id,
        )
        self.db.add(new_version)
        self.db.flush()
        root_path = generated_working_root.relative_storage_path.strip("/\\")
        normalized_target_relative_path = target_relative_path.strip("/\\")
        new_current_storage_path = f"{root_path}/{normalized_target_relative_path}"
        new_current_path = self.storage.working_copy_path(new_current_storage_path)
        if new_current_path.exists() and self.storage.sha256_file(new_current_path) != artifact.sha256:
            fail("TARGET_EXCEL_FILENAME_CONFLICT", f"文件名“{target_filename}”对应的物理文件已存在。")
        source_version.storage_path = source_history_path
        now = utcnow()
        document.size_bytes = new_version.size_bytes
        document.sha256 = artifact.sha256
        document.status = document.ingest_status = "READY"
        managed_file.root_id = generated_managed_root.id
        managed_file.relative_path = target_relative_path
        managed_file.relative_path_hash = hashlib.sha256(target_relative_path.encode()).hexdigest()
        managed_file.filename = target_filename
        managed_file.extension = ".xlsx"
        managed_file.size_bytes = new_version.size_bytes
        managed_file.modified_at = now
        managed_file.fingerprint = artifact.sha256
        managed_file.content_sha256 = artifact.sha256
        working_copy.working_copy_root_id = generated_working_root.id
        working_copy.current_version_id = new_version.id
        working_copy.relative_path = target_relative_path
        working_copy.relative_path_hash = hashlib.sha256(target_relative_path.encode()).hexdigest()
        working_copy.filename = target_filename
        working_copy.extension = ".xlsx"
        working_copy.size_bytes = new_version.size_bytes
        working_copy.content_sha256 = artifact.sha256
        working_copy.imported_source_sha256 = artifact.sha256
        working_copy.sync_status = "SYNCED"
        working_copy.revision += 1
        file_object = self.db.query(FileObject).filter_by(
            document_id=document.id,
            storage_backend="working_copy_local",
        ).one_or_none()
        if file_object is None:
            file_object = FileObject(
                document_id=document.id,
                storage_backend="working_copy_local",
                storage_path=new_current_storage_path,
                size_bytes=new_version.size_bytes,
                sha256=artifact.sha256,
            )
            self.db.add(file_object)
        else:
            file_object.storage_path = new_current_storage_path
            file_object.size_bytes = new_version.size_bytes
            file_object.sha256 = artifact.sha256
        self.db.add(StructuredDatasetExport(
            id=export_id,
            dataset_id=row.id,
            dataset_revision=0,
            storage_path=xlsx_history_path,
            sha256=artifact.sha256,
            size_bytes=new_version.size_bytes,
            document_version_id=new_version.id,
        ))
        self.db.add(WorkingCopyPathRecord(
            working_copy_id=working_copy.id,
            sequence_number=self._next_path_sequence(working_copy.id),
            operation_type="DATASET_TARGET_FORMAT_PROMOTED",
            before_relative_path=old_relative_path,
            after_relative_path=target_relative_path,
            before_filename=old_filename,
            after_filename=target_filename,
            document_version_id=new_version.id,
            content_sha256=artifact.sha256,
            status="COMPLETED",
            executed_by=self.user.id,
        ))
        self.db.flush()
        result = self._legacy_target_result(
            row,
            reused=False,
            conversion_artifact_id=artifact.artifact_id,
            source_version_id=source_version.id,
            current_version=new_version,
        )
        cleanup = {
            "legacy_relative_path": old_current_storage_path,
            "expected_legacy_sha256": source_version.sha256,
            "current_relative_path": new_current_storage_path,
            "expected_current_sha256": artifact.sha256,
        }
        publication = stage_publication(
            self.db, dataset=row, working_copy=working_copy, version=new_version,
            base_version_id=source_version.id, result=result, cleanup=cleanup,
            archive_target=archive_target)
        if archive_target:
            working_copy.status = "PUBLISHING"
            managed_file.status = "PUBLISHING"
        row.active_task_id = publication.id
        return {
            **result,
            "_publication_id": publication.id,
            "_legacy_cleanup": cleanup,
        }

    def _validate_workbuddy_target_extraction(
        self,
        *,
        document_id: str,
        document_version_id: str,
        extraction_run_id: str,
        requested_sheet: str | None,
    ) -> dict:
        """固定 WorkBuddy 的完整 SHEET 快照，并提取可核对的标题定位。"""

        run = self.db.query(DocumentExtractionRun).filter_by(
            id=extraction_run_id,
            document_id=document_id,
            document_version_id=document_version_id,
            status="COMPLETED",
            extractor="workbuddy-external-document",
        ).one_or_none()
        if run is None:
            # WorkBuddy submits against the immutable upload (staging) version. Archiving
            # creates a new document/version and clones that run; callers still hold the
            # original run ID. Resolve only the clone explicitly linked to this upload.
            version = self.db.get(DocumentVersion, document_version_id)
            working_copy = self.db.query(WorkingCopy).filter_by(
                document_id=document_id, current_version_id=document_version_id,
            ).one_or_none()
            managed_file = self.db.get(ManagedFile, working_copy.managed_file_id) if working_copy else None
            policy = dict(version.extraction_policy_json or {}) if version else {}
            source_version_id = str(policy.get("origin_upload_version_id") or "")
            source_version = self.db.get(DocumentVersion, source_version_id) if source_version_id else None
            source_run = self.db.get(DocumentExtractionRun, extraction_run_id)
            if (
                version is not None and version.document_id == document_id
                and managed_file is not None
                and managed_file.source_upload_version_id == source_version_id
                and source_version is not None and source_version.sha256 == version.sha256
                and policy.get("origin_upload_extraction_run_id") == extraction_run_id
                and source_run is not None
                and source_run.document_id == source_version.document_id
                and source_run.document_version_id == source_version.id
                and source_run.status == "COMPLETED"
                and source_run.extractor == "workbuddy-external-document"
            ):
                clones = self.db.query(DocumentExtractionRun).filter_by(
                    document_id=document_id, document_version_id=document_version_id,
                    status="COMPLETED", extractor="workbuddy-external-document",
                ).all()
                source_page_ids = {
                    page.id for page in self.db.query(DocumentPage).filter_by(extraction_run_id=source_run.id).all()
                }
                for clone in clones:
                    clone_pages = self.db.query(DocumentPage).filter_by(extraction_run_id=clone.id).all()
                    if source_page_ids and len(clone_pages) == len(source_page_ids) and {
                        (page.metadata_json or {}).get("reused_from_upload_page_id")
                        for page in clone_pages
                        if (page.metadata_json or {}).get("reused_from_upload_extraction") == source_run.id
                    } == source_page_ids:
                        run = clone
                        break
        if run is None:
            fail("TARGET_EXCEL_EXTRACTION_INVALID", "extraction_run_id 不是该 XLS 当前版本的已完成 WorkBuddy 读取结果。")
        pages = self.db.query(DocumentPage).filter_by(extraction_run_id=run.id).all()
        sheets = {}
        for page in pages:
            metadata = dict(page.metadata_json or {})
            if metadata.get("unit_kind") != "SHEET" or metadata.get("coverage_status") != "COMPLETE":
                fail("TARGET_EXCEL_EXTRACTION_INCOMPLETE", "WorkBuddy 必须完整提交 XLS 的 SHEET 读取结果。")
            sheet = str(page.sheet_name or metadata.get("sheet_name") or "").strip()
            if not sheet:
                fail("TARGET_EXCEL_EXTRACTION_INCOMPLETE", "WorkBuddy SHEET 结果缺少真实工作表名称。")
            sheets[sheet] = page
        if not sheets:
            fail("TARGET_EXCEL_EXTRACTION_INCOMPLETE", "WorkBuddy 没有提交可用的工作表内容。")
        if requested_sheet and requested_sheet not in sheets:
            fail("TARGET_EXCEL_SCHEMA_CONFLICT", "指定工作表不在 WorkBuddy 固定读取范围内。")
        return {"run_id": run.id, "sheets": sheets}

    def _assert_workbuddy_schema_matches(
        self,
        *,
        expected: dict,
        sheet_name: str,
        actual_labels: list[str],
        layout: WorkbookLayout | None = None,
        converted_path: Path | None = None,
    ) -> None:
        """用 WorkBuddy 单元格坐标或制表符首行核对转换后的标题，拒绝静默错列。"""

        page = expected["sheets"].get(sheet_name)
        if page is None:
            fail("TARGET_EXCEL_CONVERSION_SCHEMA_MISMATCH", "转换后的工作表与 WorkBuddy 读取结果不一致。")
        elements = self.db.query(DocumentElement).filter_by(
            extraction_run_id=expected["run_id"],
        ).all()
        if layout is not None:
            # 多层表头不能再与第一行大标题比较；按真实坐标核对宿主表头单元格。
            from openpyxl.utils.cell import range_boundaries
            left, top, right, bottom = range_boundaries(layout.header_range)
            workbook = load_workbook(converted_path, data_only=False)
            try:
                sheet = workbook[sheet_name]
                cells = {}
                conflicts = 0
                for element in elements:
                    meta = element.metadata_json or {}
                    if element.label == "table_cell" and meta.get("sheet_name") == sheet_name:
                        coordinate = (meta.get("row_index"), meta.get("column_index"))
                        value = " ".join(element.text_content.split())
                        if coordinate in cells and cells[coordinate] != value:
                            conflicts += 1
                        cells[coordinate] = value
                missing = mismatched = unexpected = 0
                actual_coordinates = set()
                for row in sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
                    for cell in row:
                        if cell.value is not None and str(cell.value).strip():
                            coordinate = (cell.row - 1, cell.column - 1)
                            actual_coordinates.add(coordinate)
                            if coordinate not in cells:
                                missing += 1
                            elif cells[coordinate] != " ".join(str(cell.value).split()):
                                mismatched += 1
                for coordinate, value in cells.items():
                    row_index, column_index = coordinate
                    if (isinstance(row_index, int) and isinstance(column_index, int)
                            and top <= row_index + 1 <= bottom and left <= column_index + 1 <= right
                            and value and coordinate not in actual_coordinates):
                        unexpected += 1
                if missing or mismatched or unexpected or conflicts:
                    fail("TARGET_EXCEL_CONVERSION_SCHEMA_MISMATCH",
                         "WorkBuddy 表头网格与目标 Excel 不一致；缺失、错位或冲突单元格需要重新读取。",
                         422, missing_cells=missing, mismatched_cells=mismatched,
                         unexpected_cells=unexpected, conflicting_cells=conflicts)
            finally:
                workbook.close()
            return
        located = []
        for element in elements:
            metadata = dict(element.metadata_json or {})
            if str(metadata.get("sheet_name") or "") != sheet_name:
                continue
            if element.label == "table_cell" and metadata.get("row_index") == 0:
                located.append((int(metadata.get("column_index") or 0), element.text_content.strip()))
        if located:
            for column_index, label in sorted(located):
                if column_index >= len(actual_labels) or actual_labels[column_index] != label:
                    fail("TARGET_EXCEL_CONVERSION_SCHEMA_MISMATCH", "转换后的标题与 WorkBuddy 单元格证据不一致。")
            return
        first_line = next((line.strip() for line in page.text_content.splitlines() if line.strip()), "")
        labels = [value.strip() for value in first_line.split("\t") if value.strip()]
        if not labels:
            fail("TARGET_EXCEL_HEADER_CELLS_MISSING", "WorkBuddy 没有提交可核对的标题，请重新读取目标 Excel。", 422)
        if labels != actual_labels:
            fail("TARGET_EXCEL_CONVERSION_SCHEMA_MISMATCH", "转换后的标题与 WorkBuddy 读取结果不一致。")

    def _require_workbuddy_header_grid(
        self, *, expected: dict, sheet_name: str | None, layout: WorkbookLayout | None,
    ) -> None:
        """复杂表头必须有宿主真实网格；普通正文或 OCR 不得推算 Excel 坐标。"""

        if layout is None:
            selected = sheet_name or (next(iter(expected["sheets"])) if len(expected["sheets"]) == 1 else None)
            if selected is None:
                return
            rows: dict[int, set[int]] = {}
            for element in self.db.query(DocumentElement).filter_by(extraction_run_id=expected["run_id"]).all():
                metadata = element.metadata_json or {}
                if element.label != "table_cell" or metadata.get("sheet_name") != selected:
                    continue
                row, column = metadata.get("row_index"), metadata.get("column_index")
                if (isinstance(row, int) and not isinstance(row, bool) and row >= 0
                        and isinstance(column, int) and not isinstance(column, bool) and column >= 0
                        and element.text_content.strip()):
                    rows.setdefault(row, set()).add(column)
            if rows:
                first_row = min(rows)
                first_columns = rows[first_row]
                if first_row != 0 or min(first_columns) != 0 or any(
                    max(columns) > max(first_columns) for columns in rows.values()
                ):
                    fail("TARGET_EXCEL_LAYOUT_REQUIRED",
                         "WorkBuddy 网格显示首行不是完整字段表头；请根据真实单元格和合并区域提供 sheet_name、header_range、data_range。",
                         422)
            return
        from openpyxl.utils.cell import range_boundaries

        selected = sheet_name or (next(iter(expected["sheets"])) if len(expected["sheets"]) == 1 else None)
        if selected is None or selected not in expected["sheets"]:
            fail("TARGET_EXCEL_SHEET_REQUIRED", "多个工作表时必须明确指定本轮目标工作表。", 422)
        left, top, right, bottom = range_boundaries(layout.header_range)
        cells = set()
        for element in self.db.query(DocumentElement).filter_by(extraction_run_id=expected["run_id"]).all():
            metadata = element.metadata_json or {}
            if element.label != "table_cell" or metadata.get("sheet_name") != selected:
                continue
            row, column = metadata.get("row_index"), metadata.get("column_index")
            if isinstance(row, int) and isinstance(column, int) and top <= row + 1 <= bottom and left <= column + 1 <= right:
                cells.add((row, column))
        if not cells:
            fail("TARGET_EXCEL_HEADER_CELLS_MISSING",
                 "当前表格读取结果没有表头单元格坐标；请用支持结构化网格的 WorkBuddy 能力重新读取，或上传可读取的 .xlsx。", 422)

    def _legacy_target_result(
        self,
        dataset: StructuredDataset,
        *,
        reused: bool,
        conversion_artifact_id: str | None = None,
        source_version_id: str | None = None,
        current_version: DocumentVersion | None = None,
    ) -> dict:
        """返回格式提升的稳定对象绑定与版本信息。"""

        result = self.get(dataset.id)
        working_copy = self._dataset_working_copy(dataset)
        current = current_version or (
            self.db.get(DocumentVersion, working_copy.current_version_id)
            if working_copy and working_copy.current_version_id
            else None
        )
        artifact_id = conversion_artifact_id
        if artifact_id is None and dataset.target_source_version_id:
            artifact = (
                self.db.query(DocumentArtifact)
                .filter_by(
                    document_version_id=dataset.target_source_version_id,
                    artifact_type=CONVERTED_XLSX_ARTIFACT_TYPE,
                )
                .order_by(DocumentArtifact.created_at.desc())
                .first()
            )
            artifact_id = artifact.id if artifact is not None else None
        return {
            **result,
            "kind": "structured_dataset",
            "status": "READY",
            "reused": reused,
            "source": "UPLOADED_EXCEL",
            "converted_from_legacy": True,
            "conversion_artifact_id": artifact_id,
            "source_document_version_id": source_version_id or dataset.target_source_version_id,
            "document_version_id": current.id if current else result.get("document_version_id"),
            "version_number": current.version_number if current else None,
            "original_unchanged": True,
        }

    def _retire_reused_legacy_visible_file(
        self,
        dataset: StructuredDataset,
        *,
        source_version: DocumentVersion,
    ) -> None:
        """幂等重试时补偿清理可能因进程中断残留的旧 XLS 可见入口。"""

        working_copy = self._dataset_working_copy(dataset)
        if working_copy is None or not working_copy.current_version_id:
            fail("DATASET_WORKING_COPY_INVALID", "格式提升后的工作副本关联无效。")
        current_version = self.db.get(DocumentVersion, working_copy.current_version_id)
        root = self.db.get(WorkingCopyRoot, working_copy.working_copy_root_id)
        if current_version is None or root is None:
            fail("DATASET_WORKING_COPY_INVALID", "格式提升后的版本链不完整。")
        legacy_relative_path = str(
            PurePosixPath(working_copy.relative_path).with_name(source_version.filename)
        )
        root_path = root.relative_storage_path.strip("/\\")
        normalized_legacy_relative_path = legacy_relative_path.strip("/\\")
        try:
            self.storage.retire_dataset_legacy_current_file(
                legacy_relative_path=f"{root_path}/{normalized_legacy_relative_path}",
                expected_legacy_sha256=source_version.sha256,
                current_relative_path=self._working_copy_storage_path(working_copy),
                expected_current_sha256=current_version.sha256,
            )
        except ValueError:
            fail(
                "TARGET_EXCEL_VISIBLE_CLEANUP_FAILED",
                "目标 XLSX 已注册，但旧 XLS 可见入口尚未安全清理，请稍后使用原提交键重试。",
            )

    def resolve_target(self, filename: str, *, manual_upload_sha256: str | None = None):
        """按完整文件名解析唯一活动目标，返回稳定 ID 和当前 Excel schema。"""

        normalized = self._normalize_target_filename(filename, require_extension=True)
        workspace_id = get_shared_workspace_id(self.db)
        matches = []
        rows = (
            self.db.query(StructuredDataset, WorkingCopy)
            .join(WorkingCopy, WorkingCopy.id == StructuredDataset.working_copy_id)
            .filter(
                StructuredDataset.workspace_id == workspace_id,
                WorkingCopy.status == "ACTIVE",
            )
            .all()
        )
        for dataset, working_copy in rows:
            if working_copy.filename.casefold() == normalized.casefold():
                matches.append((dataset, working_copy))
        if not matches:
            fail("TARGET_EXCEL_NOT_FOUND", f"没有找到名为“{normalized}”的累计 Excel。", 404)
        if len(matches) > 1:
            fail(
                "TARGET_EXCEL_AMBIGUOUS",
                f"找到多个名为“{normalized}”的活动 Excel，请指定具体文件。",
                candidates=[{
                    "dataset_id": dataset.id,
                    "working_copy_id": working_copy.id,
                    "document_id": working_copy.document_id,
                    "document_version_id": working_copy.current_version_id,
                    "filename": working_copy.filename,
                } for dataset, working_copy in matches],
            )
        dataset, working_copy = matches[0]
        if pending_replacement(self.db, dataset_id=dataset.id) is not None:
            fail("DATASET_TARGET_REPLACEMENT_PENDING", "同名新目标正在发布，请等待发布完成。", 423)
        assert_publication_ready(self.db, dataset_id=dataset.id)
        current_fields = self._read_and_validate_current_target_schema(
            dataset=dataset,
            working_copy=working_copy,
            manual_upload_sha256=manual_upload_sha256,
        )
        result = self.get(dataset.id)
        return {
            **result,
            "fields": current_fields,
            "resolved_by": "EXACT_FILENAME",
            "filename": working_copy.filename,
        }

    def list_datasets(self):
        """列出小 Demo 的共享数据集，重连后仍可找到稳定 ID。"""
        rows = (self.db.query(StructuredDataset)
                .outerjoin(WorkingCopy, WorkingCopy.id == StructuredDataset.working_copy_id)
                .filter(StructuredDataset.workspace_id == get_shared_workspace_id(self.db),
                        or_(StructuredDataset.working_copy_id.is_(None), WorkingCopy.status == "ACTIVE"))
                .order_by(StructuredDataset.created_at.desc()).limit(100).all())
        return {"datasets": [{
            "dataset_id": r.id,
            "name": r.name,
            "filename": self._dataset_filename(r),
            "revision": r.revision,
        } for r in rows]}

    def get(self, dataset_id):
        """读取 schema、累计状态及最新导出，不暴露操作用户和本地路径。"""
        row = self.dataset(dataset_id)
        members = self._members(row)
        versions = {
            version.id: version
            for version in self.db.query(DocumentVersion).filter(
                DocumentVersion.id.in_([member.document_version_id for member, _ in members])
            ).all()
        } if members else {}
        format_counts = {"image": 0, "word": 0, "pdf": 0}
        for member, _ in members:
            suffix = Path(versions[member.document_version_id].filename).suffix.lower()
            category = "image" if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp"} else "pdf" if suffix == ".pdf" else "word"
            format_counts[category] += 1
        working_copy = self._dataset_working_copy(row, allow_archived=True)
        return {
            "dataset_id": row.id, "name": row.name, "revision": row.revision,
            "status": "ARCHIVED" if working_copy and working_copy.status == "ARCHIVED" else "ACTIVE",
            "filename": working_copy.filename if working_copy else self._dataset_filename(row),
            "sheet_name": row.target_sheet_name or "Sheet1",
            "layout": row.target_layout_json or None,
            "fields": row.schema_json, "document_count": len(members),
            "image_count": format_counts["image"], "format_counts": format_counts,
            "record_count": len(row.base_rows_json or []) + sum(run.record_count for _, run in members),
            "base_record_count": len(row.base_rows_json or []),
            "working_copy_id": working_copy.id if working_copy else None,
            "document_id": working_copy.document_id if working_copy else None,
            "document_version_id": working_copy.current_version_id if working_copy else None,
        }

    def _normalize_target_filename(self, value: str, *, require_extension: bool = False) -> str:
        """规范化用户可见 Excel 名称，禁止路径和非 xlsx 扩展名进入存储。"""

        raw = str(value or "").strip()
        sanitized = self.storage.sanitize_filename(raw)
        if not raw or sanitized != raw:
            fail(
                "TARGET_EXCEL_FILENAME_INVALID",
                "目标 Excel 必须是单个有效文件名，不能包含路径、控制字符或被静默改名。",
                422,
            )
        suffix = Path(sanitized).suffix.lower()
        if suffix and suffix != ".xlsx":
            fail("TARGET_EXCEL_TYPE_UNSUPPORTED", "目标文件名必须使用 .xlsx 扩展名。", 422)
        if not suffix:
            if require_extension:
                fail("TARGET_EXCEL_TYPE_UNSUPPORTED", "请提供完整的 .xlsx 文件名。", 422)
            sanitized = f"{sanitized}.xlsx"
        if sanitized in {".xlsx", ""}:
            fail("TARGET_EXCEL_SCHEMA_INVALID", "目标 Excel 文件名不能为空。", 422)
        return sanitized

    def _dataset_filename(self, dataset: StructuredDataset) -> str:
        """兼容迁移前数据集，始终返回可向用户展示的完整 xlsx 文件名。"""

        if dataset.working_copy_id:
            working_copy = self.db.get(WorkingCopy, dataset.working_copy_id)
            if working_copy is not None:
                return self._normalize_target_filename(working_copy.filename)
        if dataset.target_filename:
            return self._normalize_target_filename(dataset.target_filename)
        return self._normalize_target_filename(dataset.name)

    def _assert_target_filename_available(
        self,
        filename: str,
        *,
        exclude_working_copy_id: str | None = None,
    ) -> None:
        """新目标禁止与现有活动工作副本或尚未发布的数据集静默同名。"""

        workspace_id = get_shared_workspace_id(self.db)
        working_conflicts = [
            row for row in self.db.query(WorkingCopy).filter_by(
                workspace_id=workspace_id,
                status="ACTIVE",
            ).all()
            if row.id != exclude_working_copy_id
            and row.filename.casefold() == filename.casefold()
        ]
        dataset_conflicts = [
            row for row in self.db.query(StructuredDataset).filter_by(
                workspace_id=workspace_id,
            ).all()
            if row.working_copy_id != exclude_working_copy_id
            and (row.working_copy_id is None or (
                (bound := self.db.get(WorkingCopy, row.working_copy_id)) is None
                or bound.status != "ARCHIVED"))
            and self._dataset_filename(row).casefold() == filename.casefold()
        ]
        if working_conflicts or dataset_conflicts:
            fail(
                "TARGET_EXCEL_FILENAME_CONFLICT",
                f"文件名“{filename}”已存在，请指定其他名称或明确选择已有文件。",
            )

    def _replacement_for_uploaded_target(
        self, filename: str, *, exclude_working_copy_id: str,
    ) -> dict | None:
        """本轮上传表优先；仅允许取代一个已发布、未被任务占用的同名目标。"""

        workspace_id = get_shared_workspace_id(self.db)
        copies = [copy for copy in self.db.query(WorkingCopy).filter_by(
            workspace_id=workspace_id, status="ACTIVE").order_by(
                WorkingCopy.id).with_for_update().all()
            if copy.id != exclude_working_copy_id
            and copy.filename.casefold() == filename.casefold()]
        unbound = [row for row in self.db.query(StructuredDataset).filter_by(
            workspace_id=workspace_id, working_copy_id=None).all()
            if self._dataset_filename(row).casefold() == filename.casefold()]
        if unbound or len(copies) > 1:
            fail("TARGET_EXCEL_FILENAME_CONFLICT", "同名目标不唯一，无法自动归档，请先检查当前文件。")
        if not copies:
            return None
        old_copy = copies[0]
        old_dataset = (self.db.query(StructuredDataset)
                       .filter_by(workspace_id=workspace_id, working_copy_id=old_copy.id)
                       .with_for_update().one_or_none())
        if old_dataset is None:
            fail("TARGET_EXCEL_FILENAME_CONFLICT", "同名文件不是累计目标表，无法自动归档。")
        if old_dataset.active_task_id or pending_publication(self.db, dataset_id=old_dataset.id):
            fail("DATASET_TASK_IN_PROGRESS", "同名旧目标仍有未完成任务，请先完成或恢复。", 423)
        if pending_replacement(self.db, dataset_id=old_dataset.id) is not None:
            fail("DATASET_TARGET_REPLACEMENT_PENDING", "同名新目标正在发布，请等待发布完成。", 423)
        return {
            "dataset_id": old_dataset.id,
            "working_copy_id": old_copy.id,
            "document_version_id": old_copy.current_version_id,
            "sha256": old_copy.content_sha256,
            "revision": old_copy.revision,
        }

    def _read_excel_schema(
        self,
        path: Path,
        requested_sheet: str | None,
        layout: WorkbookLayout | None = None,
    ) -> tuple[str, list[StructuredFieldSpec], list[dict]]:
        """从受控 XLSX 的唯一业务工作表读取标题与已有标量行，不执行公式或宏。"""

        try:
            inspect_xlsx(path)
        except XlsxSafetyError as exc:
            fail(exc.code, str(exc), 413 if "TOO_LARGE" in exc.code or "TOO_COMPLEX" in exc.code else 422)
        try:
            workbook = load_workbook(path, read_only=layout is None, data_only=layout is None)
        except Exception:
            fail("TARGET_EXCEL_SCHEMA_INVALID", "Excel 无法读取或文件已损坏。", 422)
        try:
            if requested_sheet:
                if requested_sheet not in workbook.sheetnames:
                    fail("TARGET_EXCEL_SCHEMA_INVALID", "指定的工作表不存在。", 422)
                candidates = [workbook[requested_sheet]]
            else:
                candidates = []
                for sheet in workbook.worksheets:
                    if any(
                        any(value is not None and str(value).strip() for value in row)
                        for row in sheet.iter_rows(
                            min_row=1,
                            max_row=min(sheet.max_row or 1001, 1001),
                            values_only=True,
                        )
                    ):
                        candidates.append(sheet)
                if len(candidates) > 1:
                    fail(
                        "TARGET_EXCEL_SHEET_AMBIGUOUS",
                        "Excel 中存在多个非空工作表，请明确指定目标工作表。",
                        sheets=[sheet.title for sheet in candidates],
                    )
            if len(candidates) != 1:
                fail("TARGET_EXCEL_SCHEMA_INVALID", "Excel 中没有可读取的业务工作表。", 422)
            sheet = candidates[0]
            rows = list(
                sheet.iter_rows(
                    min_row=1,
                    max_row=min(sheet.max_row or 1002, 1002),
                    values_only=True,
                )
            )
            header_index = next((index for index, row in enumerate(rows) if any(
                value is not None and str(value).strip() for value in row
            )), None)
            if header_index is None:
                fail("TARGET_EXCEL_SCHEMA_INVALID", "Excel 缺少标题行。", 422)
            raw_headers = list(rows[header_index])
            while raw_headers and (raw_headers[-1] is None or not str(raw_headers[-1]).strip()):
                raw_headers.pop()
            if layout is not None:
                try:
                    raw_headers, layout_rows = read_layout(sheet, layout)
                except ValueError as exc:
                    fail("TARGET_EXCEL_LAYOUT_INVALID", str(exc), 422)
            elif any(any(v is not None and str(v).strip() for v in row[len(raw_headers):])
                     for row in rows[header_index + 1:]):
                fail("TARGET_EXCEL_LAYOUT_REQUIRED", "首行不是完整字段表头，请提供真实 header_range 和 data_range，不能把大标题或尾注当数据。", 422)
            if not raw_headers or len(raw_headers) > 40 or any(
                value is None or not str(value).strip() for value in raw_headers
            ):
                fail("TARGET_EXCEL_SCHEMA_INVALID", "标题必须连续、非空且不超过 40 列；复杂表头请指定 layout。", 422)
            labels = [" ".join(str(value).split()) for value in raw_headers]
            if len({label.casefold() for label in labels}) != len(labels):
                fail("TARGET_EXCEL_SCHEMA_INVALID", "Excel 标题不能重复。", 422)
            data_rows = []
            for values in (layout_rows if layout is not None else rows[header_index + 1:]):
                selected = list(values[:len(labels)])
                # write_only 生成的 XLSX 可以省略尾部空单元格；读取时按真实表头补齐，
                # 避免可选空列导致发布回读或后续按文件名追加出现 IndexError。
                selected.extend([None] * (len(labels) - len(selected)))
                if not any(value is not None and str(value).strip() for value in selected):
                    continue
                if len(data_rows) >= 1000:
                    fail("DEMO_RECORD_LIMIT", "Excel 已有数据超过 1000 行限制。", 422)
                data_rows.append(selected)
            fields = []
            used_keys: set[str] = set()
            for offset, label in enumerate(labels, 1):
                key = re.sub(r"[^a-z0-9_]+", "_", label.casefold()).strip("_")
                if not re.fullmatch(r"[a-z][a-z0-9_]{0,63}", key or ""):
                    key = f"field_{offset}"
                candidate = key
                sequence = 2
                while candidate in used_keys:
                    candidate = f"{key[:60]}_{sequence}"
                    sequence += 1
                used_keys.add(candidate)
                column_values = [row[offset - 1] for row in data_rows if row[offset - 1] is not None]
                fields.append(StructuredFieldSpec(
                    key=candidate,
                    label=label,
                    field_type=self._infer_excel_field_type(column_values),
                    required=False,
                ))
            base_rows = [
                {
                    field.key: self._json_excel_value(values[index])
                    for index, field in enumerate(fields)
                }
                for values in data_rows
            ]
            return sheet.title, fields, base_rows
        finally:
            workbook.close()

    @staticmethod
    def _infer_excel_field_type(values: list) -> str:
        """仅依据真实已有标量推断基础类型；混合列安全降级为字符串。"""

        if values and all(isinstance(value, (datetime, date)) for value in values):
            return "date"
        if values and all(isinstance(value, int) and not isinstance(value, bool) for value in values):
            return "integer"
        if values and all(
            isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)
            for value in values
        ):
            return "decimal"
        return "string"

    @staticmethod
    def _json_excel_value(value):
        """把 openpyxl 标量转换为 JSON 可持久化值，拒绝非有限数字。"""

        if value is None:
            return None
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        if isinstance(value, bool):
            return "true" if value else "false"
        if isinstance(value, (int, float, Decimal)):
            number = Decimal(str(value))
            if not number.is_finite():
                fail("TARGET_EXCEL_SCHEMA_INVALID", "Excel 包含非有限数值。", 422)
            return int(number) if number == number.to_integral_value() else str(number)
        return str(value)

    def _validate_target_binding(self, dataset: StructuredDataset, request: DatasetAppend) -> None:
        """把按文件名解析得到的稳定目标冻结到追加请求，检测确认期间的版本变化。"""

        supplied = request.target_working_copy_id is not None
        working_copy = self._dataset_working_copy(dataset)
        if working_copy is not None and not supplied:
            fail(
                "TARGET_EXCEL_BINDING_REQUIRED",
                "目标 Excel 已发布，请先按文件名重新解析并携带当前稳定 ID。",
            )
        if not supplied:
            return
        if working_copy is None:
            fail("TARGET_EXCEL_VERSION_CONFLICT", "目标 Excel 尚未发布，请重新获取目标。")
        if (
            working_copy.id != str(request.target_working_copy_id)
            or working_copy.document_id != str(request.target_document_id)
            or working_copy.current_version_id != str(request.target_document_version_id)
        ):
            fail("TARGET_EXCEL_VERSION_CONFLICT", "目标 Excel 在确认期间已发生变化，请重新读取字段并确认。")

    def _read_and_validate_current_target_schema(
        self,
        *,
        dataset: StructuredDataset,
        working_copy: WorkingCopy,
        manual_upload_sha256: str | None = None,
    ) -> list[dict]:
        """读取当前可见 Excel 的真实标题，并拒绝与数据集冻结 schema 不一致的文件。"""

        version = (
            self.db.get(DocumentVersion, working_copy.current_version_id)
            if working_copy.current_version_id
            else None
        )
        if version is None or version.document_id != working_copy.document_id:
            fail("TARGET_EXCEL_VERSION_CONFLICT", "目标 Excel 缺少有效当前版本，请重新选择。")
        path = self.storage.working_copy_path(self._working_copy_storage_path(working_copy))
        if not path.is_file() or version.sha256 != working_copy.content_sha256:
            fail("TARGET_EXCEL_VERSION_CONFLICT", "目标 Excel 当前文件与版本记录不一致，请重新同步。")
        visible_sha256 = self.storage.sha256_file(path)
        if visible_sha256 != version.sha256 and visible_sha256 != manual_upload_sha256:
            if manual_upload_sha256 is not None:
                fail("MANUAL_EXCEL_VISIBLE_CONFLICT", "服务器同名 Excel 与上传文件不一致，请先核对两份修改。", 409)
            fail("TARGET_EXCEL_VERSION_CONFLICT", "目标 Excel 当前文件与版本记录不一致，请重新同步。")
        sheet_name, actual_fields, _ = self._read_excel_schema(
            path,
            dataset.target_sheet_name or None,
            WorkbookLayout.model_validate(dataset.target_layout_json) if dataset.target_layout_json else None,
        )
        expected_fields = [StructuredFieldSpec.model_validate(item) for item in dataset.schema_json]
        if (
            sheet_name != (dataset.target_sheet_name or "Sheet1")
            or [item.label for item in actual_fields] != [item.label for item in expected_fields]
        ):
            fail(
                "TARGET_EXCEL_SCHEMA_CONFLICT",
                "目标 Excel 的工作表或标题字段已经变化，请先确认新的目标结构。",
            )
        # Excel 标题本身不声明 person_name 等业务类型；标题以真实文件为准，
        # 类型与必填约束继续使用已确认的数据集 schema。
        return [item.model_dump() for item in expected_fields]

    def source(self, document_id, *, version_id=None, extraction_run_id=None, content_origin=None,
               mode="FULL", unit_orders=None):
        """读取固定提取运行的索引或指定区段；内部校验仍默认使用完整正文。"""
        if mode not in {"FULL", "INDEX"}:
            fail("INVALID_SOURCE_MODE", "来源读取模式无效。", 422)
        if unit_orders is not None and (not unit_orders or len(unit_orders) > 100
                                        or any(not isinstance(order, int) or order < 0 for order in unit_orders)
                                        or len(set(unit_orders)) != len(unit_orders)):
            fail("INVALID_SOURCE_UNIT_ORDERS", "来源区段序号必须是不重复的非负整数，最多 100 项。", 422)
        if content_origin == "WORKBUDDY_SCOPED":
            return self._scoped_source(document_id, version_id=version_id,
                extraction_run_id=extraction_run_id, mode=mode, unit_orders=unit_orders)
        try:
            source = CanonicalWorkingFileResolver(self.db).resolve_document(
                document_id=document_id, document_version_id=version_id,
            )
        except CanonicalWorkingFileError as exc:
            fail(exc.code, str(exc))
        version = source.document_version
        if version.document_id != document_id:
            fail("USE_WORKING_DOCUMENT", "请使用已整理工作副本的 Document ID。")
        suffix = Path(version.filename).suffix.lower()
        is_image = suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
        if suffix not in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".pdf", ".doc", ".docx"}:
            fail("DEMO_DOCUMENT_REQUIRED", "Demo 仅支持图片、PDF、DOC 或 DOCX 工作副本。", 422)
        query = self.db.query(DocumentExtractionRun).filter_by(
            document_id=document_id, document_version_id=version.id, status="COMPLETED",
        )
        if extraction_run_id:
            query = query.filter(DocumentExtractionRun.id == extraction_run_id)
        elif content_origin == "WORKBUDDY" or not is_image:
            query = query.filter(DocumentExtractionRun.extractor == "workbuddy-external-document")
        elif is_image:
            query = query.filter(DocumentExtractionRun.extractor != "workbuddy-dataset-reuse")
        extraction = query.order_by(DocumentExtractionRun.created_at.desc()).first()
        pages = self.db.query(DocumentPage).filter_by(extraction_run_id=extraction.id).all() if extraction else []
        if not pages or not any(page.text_content.strip() for page in pages):
            fail("EXTRACTION_NOT_READY", "请先完成宿主内容提取和工作副本整理流程。")
        if is_image and (len(pages) != 1 or pages[0].page_number != 1):
            fail("OCR_NOT_READY", "图片必须具有现有单页 OCR 结果。")
        if content_origin == "WORKBUDDY" and extraction.extractor != "workbuddy-external-document":
            fail("WORKBUDDY_CONTENT_REQUIRED", "该来源没有合格的 WorkBuddy 文档提取正文。")
        if sum(len(page.text_content) for page in pages) > 100_000:
            fail("DEMO_SOURCE_LIMIT", "文件正文超过 Demo 处理范围。", 422)
        ordered_pages = sorted(
            pages,
            key=lambda page: (
                int((page.metadata_json or {}).get("unit_order", 0)),
                page.created_at,
            ),
        )
        all_orders = {int((page.metadata_json or {}).get("unit_order", offset))
                      for offset, page in enumerate(ordered_pages)}
        requested_orders = set(unit_orders) if unit_orders is not None else None
        if requested_orders is not None and not requested_orders <= all_orders:
            fail("SOURCE_UNIT_NOT_FOUND", "所选区段不属于该来源版本。", 422)
        selected_pages = [page for offset, page in enumerate(ordered_pages)
                          if requested_orders is None or
                          int((page.metadata_json or {}).get("unit_order", offset)) in requested_orders]
        if mode == "INDEX":
            # 先返回无正文区段目录，避免 WorkBuddy 每轮把整份 PDF 重复送进模型上下文。
            return {
                "document_id": document_id, "document_version_id": version.id,
                "filename": source.working_copy.filename, "source_sha256": version.sha256,
                "extraction_run_id": extraction.id,
                "content_origin": "WORKBUDDY" if extraction.extractor == "workbuddy-external-document" else "LEGACY_IMAGE",
                "source_kind": "IMAGE" if is_image else "WORKBUDDY_DOCUMENT",
                "coverage": dict((ordered_pages[0].metadata_json or {}).get("coverage") or {}),
                "total_source_unit_count": len(ordered_pages),
                "source_units": [{
                    "source_unit_id": page.id,
                    "external_unit_id": str((page.metadata_json or {}).get("source_unit_id") or ""),
                    "kind": str((page.metadata_json or {}).get("unit_kind") or ("IMAGE_PAGE" if is_image else "TEXT_SECTION")),
                    "order": int((page.metadata_json or {}).get("unit_order", offset)),
                    "page_number": page.page_number,
                    "sheet_name": (page.metadata_json or {}).get("sheet_name"),
                    "text_length": len(page.text_content),
                } for offset, page in enumerate(selected_pages)],
                "pages": [],
            }
        source_elements = self.db.query(DocumentElement).filter_by(extraction_run_id=extraction.id).all()
        elements_by_unit: dict[str, list[str]] = {}
        seal_blocks_by_unit: dict[str, list[dict]] = {}
        for element in source_elements:
            external_unit_id = str((element.metadata_json or {}).get("source_unit_id") or "")
            if external_unit_id:
                elements_by_unit.setdefault(external_unit_id, []).append(element.id)
                if element.label == "seal" and element.bbox_json:
                    seal_blocks_by_unit.setdefault(external_unit_id, []).append({
                        "text": element.text_content, "bbox": element.bbox_json,
                        "start": (element.metadata_json or {}).get("start"),
                        "end": (element.metadata_json or {}).get("end"),
                        "element_id": element.id,
                    })
        source_units = [
            {
                "source_unit_id": page.id,
                "external_unit_id": str((page.metadata_json or {}).get("source_unit_id") or ""),
                "kind": str((page.metadata_json or {}).get("unit_kind") or ("IMAGE_PAGE" if is_image else "TEXT_SECTION")),
                "order": int((page.metadata_json or {}).get("unit_order", offset)),
                "page_number": page.page_number,
                "sheet_name": (page.metadata_json or {}).get("sheet_name"),
                "text": page.text_content,
                "source_element_ids": elements_by_unit.get(str((page.metadata_json or {}).get("source_unit_id") or ""), []),
                "seal_blocks": seal_blocks_by_unit.get(str((page.metadata_json or {}).get("source_unit_id") or ""), []),
            }
            for offset, page in enumerate(selected_pages)
        ]
        return {
            "document_id": document_id, "document_version_id": version.id,
            "filename": source.working_copy.filename, "source_sha256": version.sha256,
            "extraction_run_id": extraction.id,
            "content_origin": "WORKBUDDY" if extraction.extractor == "workbuddy-external-document" else "LEGACY_IMAGE",
            "source_kind": "IMAGE" if is_image else "WORKBUDDY_DOCUMENT",
            "coverage": dict((ordered_pages[0].metadata_json or {}).get("coverage") or {}) if ordered_pages else {},
            "total_source_unit_count": len(ordered_pages),
            "source_units": source_units,
            # 旧图片调用方继续得到原 pages 形状；新 Skill 使用 source_units 的正式 UUID。
            "pages": [{"page_number": page.page_number, "text": page.text_content} for page in selected_pages] if is_image else [],
        }

    def _scoped_source(self, document_id, *, version_id, extraction_run_id, mode, unit_orders):
        """只向来源用户的明确表格任务暴露冻结的身份和考核 OCR 区域。"""
        from app.db.models import ExternalExtractionTask, IngestBatch, IngestItem
        from app.modules.external_extraction.scoped_schemas import CONTRACT, DISCOVERY_CONTRACT, FIXED_25_CONTRACT, LEGACY_CONTRACT

        if not version_id or not extraction_run_id:
            fail("SCOPED_SOURCE_BINDING_REQUIRED", "定向来源必须绑定版本和提取运行。", 422)
        version = self.db.get(DocumentVersion, str(version_id))
        run = self.db.get(DocumentExtractionRun, str(extraction_run_id))
        task = self.db.query(ExternalExtractionTask).filter_by(
            source_version_id=str(version_id), extraction_run_id=str(extraction_run_id),
            status="SCOPE_COMPLETE").filter(
            ExternalExtractionTask.provider_contract_version.in_([
                LEGACY_CONTRACT, CONTRACT, DISCOVERY_CONTRACT, FIXED_25_CONTRACT])).one_or_none()
        item = self.db.get(IngestItem, task.ingest_item_id) if task else None
        batch = self.db.get(IngestBatch, item.batch_id) if item else None
        if (version is None or run is None or task is None or batch is None or batch.user_id != self.user.id
                or version.document_id != str(document_id) or run.document_id != str(document_id)
                or run.document_version_id != version.id or run.status != "SCOPE_COMPLETE"
                or run.extractor != "workbuddy-scoped-pdf" or version.sha256 != task.source_sha256):
            fail("SCOPED_SOURCE_NOT_READY", "定向来源未完成或不属于当前用户和版本。", 409)
        pages = self.db.query(DocumentPage).filter_by(extraction_run_id=run.id).all()
        region_specs = (task.source_manifest_json or {}).get("regions", [])
        if (len(pages) != len(region_specs) or
                {((p.metadata_json or {}).get("region_id"), p.page_number) for p in pages}
                != {(region["region_id"], region["page_number"]) for region in region_specs}):
            fail("SCOPED_SOURCE_NOT_READY", "定向区域证据不完整。", 409)
        pages.sort(key=lambda p: int((p.metadata_json or {}).get("unit_order", 0)))
        if unit_orders is not None and not set(unit_orders) <= set(range(len(pages))):
            fail("SOURCE_UNIT_NOT_FOUND", "所选区域不属于本次冻结范围。", 422)
        units = []
        for page in pages:
            metadata = page.metadata_json or {}
            order = int(metadata.get("unit_order", 0))
            if unit_orders is not None and order not in unit_orders:
                continue
            units.append({"source_unit_id": page.id, "external_unit_id": metadata.get("region_id"),
                "kind": "PDF_PAGE", "order": order, "page_number": page.page_number,
                "sheet_name": None, "text_length": len(page.text_content),
                **({"text": page.text_content, "source_element_ids": [], "seal_blocks": []} if mode == "FULL" else {})})
        field_refs = {}
        for column, ref in (task.result_summary_json or {}).get("field_refs", {}).items():
            page = next((p for p in pages if (p.metadata_json or {}).get("region_id") == ref.get("region_id")), None)
            if column == "L" and ref.get("status") == "NOT_FOUND":
                field_refs[column] = {"source_unit_id": None, "page_number": None,
                    "region_id": None, "status": "NOT_FOUND", "value": None, "years": []}
                continue
            if page:
                if column == "L":
                    field_refs[column] = {"source_unit_id": page.id, "page_number": page.page_number,
                        "region_id": ref["region_id"], "status": ref["status"], "value": ref["value"],
                        "years": [{"year": item["year"], "start": item["start"], "end": item["end"],
                                   **({"quote": page.text_content[item["start"]:item["end"]]}
                                      if mode == "FULL" else {})} for item in ref["years"]]}
                else:
                    field_refs[column] = {"source_unit_id": page.id, "page_number": page.page_number,
                        "region_id": ref["region_id"], "start": ref["start"], "end": ref["end"],
                        **({"quote": page.text_content[ref["start"]:ref["end"]]} if mode == "FULL" else {})}
        return {"document_id": str(document_id), "document_version_id": version.id,
            "filename": version.filename, "source_sha256": version.sha256, "extraction_run_id": run.id,
            "content_origin": "WORKBUDDY_SCOPED", "source_kind": "WORKBUDDY_SCOPED",
            "scope_digest": (task.source_manifest_json or {}).get("scope_digest"),
            "coverage": (task.result_summary_json or {}).get("coverage", {}),
            "total_source_unit_count": len(pages), "source_units": units, "field_refs": field_refs, "pages": []}

    def append(self, dataset_id, request: DatasetAppend):
        """在同一事务锁定数据集，保存单图事实及成员，再推进 revision。"""
        dataset = self.dataset(dataset_id, lock=True)
        self._assert_task_write_access(dataset)
        payload_for_digest = request.model_dump(mode="json", exclude={"idempotency_key"})
        if payload_for_digest.get("content_origin") is None:
            payload_for_digest.pop("content_origin", None)
        for record in payload_for_digest.get("records", []):
            for value in record.get("fields", {}).values():
                # 新可选字段不改变旧请求指纹，保证升级后原幂等键仍可恢复。
                for optional in ("fragments", "blank_reason", "approval", "assessment", "ocr_status", "ocr_unit_id"):
                    if not value.get(optional):
                        value.pop(optional, None)
                if not value.get("evidence_refs"):
                    value.pop("evidence_refs", None)
        payload_digest = digest(payload_for_digest)
        existing = self.db.query(StructuredDatasetMember).filter_by(
            dataset_id=dataset.id, append_key=request.idempotency_key,
        ).one_or_none()
        if existing:
            if existing.payload_digest != payload_digest:
                fail("IDEMPOTENCY_CONFLICT", "此提交键已用于不同结构化结果。")
            return {**self.get(dataset.id), "appended_revision": existing.appended_revision, "reused": True}
        self._validate_target_binding(dataset, request)
        if self.db.query(StructuredDatasetMember.id).filter_by(
            dataset_id=dataset.id, document_version_id=str(request.document_version_id)
        ).first():
            fail("SOURCE_ALREADY_APPENDED", "此文件版本已经加入该表格。")
        if self.get(dataset.id)["record_count"] + len(request.records) > 1000:
            fail("DEMO_RECORD_LIMIT", "Demo 每个表格最多累计 1000 条记录。", 422)
        source = self.source(str(request.document_id), version_id=str(request.document_version_id),
                             extraction_run_id=str(request.extraction_run_id), content_origin=request.content_origin)
        if source["source_sha256"] != request.source_sha256:
            fail("SOURCE_SNAPSHOT_CHANGED", "文件内容哈希与提交时不一致。")
        if source["source_kind"] == "WORKBUDDY_DOCUMENT" and request.content_origin != "WORKBUDDY":
            fail("WORKBUDDY_CONTENT_REQUIRED", "Word/PDF 追加必须明确固定 WorkBuddy 正文来源。")
        if not request.records:
            return {**self.get(dataset.id), "appended_revision": None, "review_field_count": 0,
                    "record_count_added": 0, "message": "未找到匹配记录", "original_unchanged": True, "reused": False}
        fields = [StructuredFieldSpec.model_validate(f) for f in dataset.schema_json]
        keys = {f.key for f in fields}
        first_column = (range_boundaries(dataset.target_layout_json["header_range"])[0]
                        if dataset.target_layout_json else 1)
        field_locations = {
            field.key: (field.label, get_column_letter(first_column + offset))
            for offset, field in enumerate(fields)
        }
        if source["source_kind"] == "WORKBUDDY_SCOPED":
            if not (self.workflow_task_id and getattr(self, "scoped_write_scope", False)
                    and request.content_origin == "WORKBUDDY_SCOPED"
                    and first_column == 1 and len(fields) >= 9):
                fail("SCOPED_WRITE_SCOPE_REQUIRED", "定向来源只能用于受控限列回传任务。", 422)
            required_keys = {field.key for index, field in enumerate(fields) if 1 <= index <= 8}
            for row in request.records:
                if not required_keys <= set(row.fields):
                    fail("SCOPED_EVIDENCE_REQUIRED", "缺少 B–I 来源字段。", 422)
                allow_review = getattr(self, "scoped_review_scope", False)
                if any(row.fields.get(field.key) and row.fields[field.key].raw_text
                       for index, field in enumerate(fields)
                       if index == 0 or (index >= 9 and not (allow_review and index == 11))):
                    fail("SCOPED_WRITE_SCOPE_INVALID", "定向来源只能写入 B–I 与获授权的 L 列。", 422)
        if any(set(row.fields) - keys for row in request.records):
            fail("UNKNOWN_FIELD", "提交包含当前表格未定义的字段。", 422)
        run = self.repository.create_run(
            tool_input=StructuredImageExtractionInput(document_id=source["document_id"],
                schema_mode="EXPLICIT_FIELDS", fields=fields, record_mode="TABLE_ROWS", presentation="TABLE"),
            document_version_id=source["document_version_id"], schema_fingerprint=digest(dataset.schema_json),
            provider="workbuddy-dataset-demo", model_name="external-field-mapping",
            prompt_version="workbuddy-dataset-demo-v1", agent_run_id=self.agent_run.id,
        )
        # 完整原文来自已有事实；新运行只表达结构化证据派生，不覆盖基础 OCR 和索引。
        # Word 无真实页码时这里只生成内部 layout 序号，字段证据仍按 source unit_id 校验。
        layout_pages = [
            LayoutPage(page_number=index + 1, elements=[LayoutElement(element_index=index, text=unit["text"])])
            for index, unit in enumerate(source["source_units"])
        ]
        layout, elements = self.repository.create_layout_extraction(
            run=run, parser_config_hash=digest([source["extraction_run_id"], "dataset-demo-v1"]),
            layout=LayoutParseResult(provider="workbuddy-dataset-reuse", provider_version="1",
                pages=layout_pages),
        )
        persisted_elements = {
            element.id: element
            for element in self.db.query(DocumentElement).filter_by(extraction_run_id=layout.id).all()
        }
        derived_pages = self.db.query(DocumentPage).filter_by(extraction_run_id=layout.id).order_by(DocumentPage.page_number).all()
        for index, page in enumerate(derived_pages):
            page.metadata_json = {**page.metadata_json, "source_extraction_run_id": source["extraction_run_id"],
                                  "source_unit_id": source["source_units"][index]["source_unit_id"]}
        candidates = []
        presentations = {}
        unscored_fields = set()
        for index, record in enumerate(request.records, 1):
            values = {}
            for key, value in record.fields.items():
                presentation = {"blank_reason": value.blank_reason} if value.blank_reason else {}
                if value.ocr_status:
                    if (not self.workflow_task_id or not self.full_row_ocr_mode
                            or source["source_kind"] != "WORKBUDDY_DOCUMENT"
                            or Path(source["filename"]).suffix.lower() != ".pdf"):
                        fail("OCR_MARKER_NOT_ALLOWED", "OCR 待核对标记只允许完整教授表 PDF 任务。", 422)
                    matching_units = [(offset, unit) for offset, unit in enumerate(source["source_units"])
                                      if unit["source_unit_id"] == str(value.ocr_unit_id)]
                    if (len(matching_units) != 1 or matching_units[0][1]["kind"] != "PDF_PAGE"
                            or matching_units[0][1]["page_number"] != value.page_number):
                        fail("OCR_SOURCE_MISMATCH", "待核对字段未绑定本份 PDF 的真实页面。", 422)
                    if value.ocr_status == "UNCERTAIN" and any(
                            str(ref.source_unit_id) != str(value.ocr_unit_id) for ref in value.evidence_refs):
                        fail("OCR_SOURCE_MISMATCH", "OCR 候选引文不属于指定页面。", 422)
                    if fields[[field.key for field in fields].index(key)].field_type != "string":
                        fail("OCR_MARKER_TEXT_FIELD_REQUIRED", "OCR 待核对标记只能写入文本列。", 422)
                    presentation.update(ocr_status=value.ocr_status,
                                        ocr_source_unit_id=str(value.ocr_unit_id))
                    if value.ocr_status == "UNCERTAIN":
                        presentation["ocr_candidate"] = value.raw_text
                if value.source_only_fragments and key not in self.source_fact_keys:
                    fail("SOURCE_ONLY_NOT_ALLOWED", "无章来源事实只允许在专用任务指定的业务列写入。", 422)
                if value.source_only_fragments and Path(source["filename"]).suffix.lower() != ".pdf":
                    fail("SOURCE_ONLY_PDF_REQUIRED", "无章来源事实仅接受当前 WorkBuddy PDF 原文。", 422)
                if self.workflow_task_id and value.confidence is None:
                    # 已授权的固定候选仍必须通过逐字证据与类型校验；缺失评分不是低分，不能伪造 1.0。
                    unscored_fields.add((index, key))
                    presentation["confidence_origin"] = "UNSCORED"
                try:
                    if value.approval:
                        if value.approval.approval_kind == "HR_ANNUAL_EXCELLENCE" and not (
                            self.workflow_task_id and key == getattr(self, "hr_annual_key", None)
                        ):
                            fail("HR_SEAL_SCOPE_MISMATCH", "年度考核方章仅允许绑定专用模板 L 列。", 422)
                        validate_approval(value.approval, source["source_units"])
                        for ref in [ref for ref in value.fragments if ref not in value.source_only_fragments] or (
                            [] if value.source_only_fragments else value.evidence_refs
                        ):
                            from .assessment import SourceQuote
                            validate_scope(SourceQuote.model_validate(ref.model_dump(mode="json")),
                                           value.approval, source["source_units"])
                        presentation["approval"] = value.approval.model_dump()
                    if value.assessment:
                        spec = next(field for field in fields if field.key == key)
                        if spec.field_type != "string":
                            fail("CHECKMARK_TEXT_FIELD_REQUIRED", "只允许在业务文本列生成对勾。", 422)
                        presentation.update(assess_field(label=spec.label, assessment=value.assessment,
                            approval=value.approval, units=source["source_units"]))
                    fragment_refs = [validate_quote(ref, source["source_units"]) for ref in value.fragments]
                    if value.source_only_fragments:
                        indexes = [value.fragments.index(ref) for ref in value.source_only_fragments]
                        presentation["source_only_fragment_indexes"] = indexes
                        presentation["source_status"] = "SOURCE_ONLY" if len(indexes) == len(value.fragments) else "MIXED"
                except ValueError as exc:
                    code, _, message = str(exc).partition(":")
                    label, column = field_locations[key]
                    fail(code, message.strip(), 422, record_index=index,
                         field_key=key, field_label=label, column=column)
                presentations[index, key] = presentation
                if fragment_refs:
                    # 多段原文先逐一校验，再形成仅供本次结构化运行的派生证据。
                    # 不覆盖原 DocumentPage；每段仍保留正式 unit 与字符位置。
                    element = DocumentElement(document_id=run.document_id, extraction_run_id=layout.id,
                        element_index=len(elements), label="VERIFIED_FRAGMENTS", text_content=value.raw_text,
                        metadata_json={"source_evidence_refs": fragment_refs,
                                       "source_extraction_run_id": source["extraction_run_id"]}, bbox_json={})
                    self.db.add(element)
                    self.db.flush()
                    evidence = EvidenceElement(id=element.id, document_id=run.document_id,
                        extraction_run_id=layout.id, text=value.raw_text, page_number=None,
                        bbox={}, metadata=element.metadata_json)
                    elements.append(evidence)
                    values[key] = CandidateFieldValue(raw_text=value.raw_text, value=None,
                        confidence=value.confidence or 0, evidence_element_ids=[element.id])
                    continue
                if source["source_kind"] == "WORKBUDDY_DOCUMENT" and value.raw_text and not value.evidence_refs:
                    fail("EVIDENCE_SCOPE_MISMATCH", "Word/PDF 字段必须提供正式 source_unit_id 和原文 quote。", 422)
                matched_offsets = []
                validated_refs = []
                for reference in value.evidence_refs:
                    matching = [(offset, unit) for offset, unit in enumerate(source["source_units"])
                                if unit["source_unit_id"] == str(reference.source_unit_id)]
                    if len(matching) != 1:
                        fail("EVIDENCE_SCOPE_MISMATCH", "字段证据不属于指定正文运行。", 422)
                    source_offset, source_unit = matching[0]
                    quote = reference.quote
                    if reference.start is not None:
                        if reference.end > len(source_unit["text"]) or source_unit["text"][reference.start:reference.end] != quote:
                            fail("EVIDENCE_SCOPE_MISMATCH", "字段引文范围与正式正文不一致。", 422)
                    elif source_unit["text"].count(quote) != 1:
                        fail("EVIDENCE_AMBIGUOUS", "字段引文在指定区段中不唯一，请提供字符范围。")
                    if value.raw_text and value.raw_text not in quote:
                        fail("EVIDENCE_SCOPE_MISMATCH", "字段原值没有被证据引文完整支持。", 422)
                    matched_offsets.append(source_offset)
                    validated_refs.append({
                        "_source_offset": source_offset,
                        "source_unit_id": str(reference.source_unit_id),
                        "quote": quote,
                        "start": reference.start,
                        "end": reference.end,
                        "source_element_ids": source_unit.get("source_element_ids", []),
                    })
                raw = re.sub(r"\s+", "", value.raw_text or "").casefold()
                if matched_offsets:
                    supported = all(raw in re.sub(r"\s+", "", source["source_units"][offset]["text"]).casefold()
                                    for offset in matched_offsets) if raw else False
                    evidence_ids = [elements[offset].id for offset in matched_offsets] if supported else []
                    if supported:
                        for offset in set(matched_offsets):
                            persisted_element = persisted_elements[elements[offset].id]
                            current = list((persisted_element.metadata_json or {}).get("source_evidence_refs") or [])
                            for evidence_ref in (
                                {key: item for key, item in evidence_ref.items() if key != "_source_offset"}
                                for evidence_ref in validated_refs
                                if evidence_ref["_source_offset"] == offset
                            ):
                                if evidence_ref not in current:
                                    current.append(evidence_ref)
                            persisted_element.metadata_json = {
                                **(persisted_element.metadata_json or {}),
                                "source_extraction_run_id": source["extraction_run_id"],
                                "source_evidence_refs": current,
                            }
                else:
                    # 旧图片 payload 沿用 page_number=1 的兼容定位。
                    matches = [(offset, unit) for offset, unit in enumerate(source["source_units"])
                               if unit["page_number"] == value.page_number]
                    evidence_ids = [elements[matches[0][0]].id] if len(matches) == 1 and raw in re.sub(r"\s+", "", matches[0][1]["text"]).casefold() else []
                values[key] = CandidateFieldValue(
                    raw_text=value.raw_text, value=None, confidence=value.confidence or 0,
                    evidence_element_ids=evidence_ids,
                )
            candidates.append(CandidateRecord(record_index=index, fields=values))
        # 禁用 Provider 配置，仅调用公共 finalizer；不触发网络和第二次 OCR。
        settings = self.settings.model_copy(update={
            "structured_extraction_layout_provider": "pp_structure_v3",
            "structured_extraction_llm_provider": "disabled", "structured_extraction_vision_provider": "disabled",
        })
        finalizer = StructuredExtractionService(db=self.db, user_id=self.user.id, settings=settings)
        result = finalizer.finalize_candidates(run=run, layout_extraction_run_id=layout.id,
            fields=fields, candidates=CandidateExtraction(records=candidates), elements=elements,
            unscored_fields=unscored_fields)
        for persisted in self.repository.load_fields(run.id):
            presentation = presentations.get((persisted.record_index, persisted.field_key), {})
            persisted.presentation_json = presentation
            if presentation.get("ocr_status"):
                # 已读但不清晰是可发布的显式待核对状态；规范化值保持空，查询与合计不能误用候选。
                persisted.status = "EXTRACTED"
                persisted.normalized_value_json = None
                persisted.warning_codes_json = [*(persisted.warning_codes_json or []),
                                                f"OCR_{presentation['ocr_status']}"]
                if presentation["ocr_status"] == "UNREADABLE":
                    persisted.page_number = next(
                        (unit["page_number"] for unit in source["source_units"]
                         if unit["source_unit_id"] == presentation["ocr_source_unit_id"]),
                        persisted.page_number,
                    )
        changeset_id = self.repository.record_changeset(
            run=run, agent_run_id=self.agent_run.id, conversation_id=self.agent_run.conversation_id,
            user_id=self.user.id, review_count=result.review_count,
            field_count=result.field_count, record_count=result.record_count,
        )
        dataset.revision += 1
        member = StructuredDatasetMember(dataset_id=dataset.id, structured_extraction_run_id=run.id,
            document_version_id=source["document_version_id"], append_key=request.idempotency_key,
            payload_digest=payload_digest, appended_revision=dataset.revision)
        self.db.add(member)
        self.db.add(ChangeItem(changeset_id=changeset_id, target_type="DATASET", target_id=dataset.id,
            target_document_id=source["document_id"], change_type="STRUCTURED_DATASET_APPENDED",
            after_value_json={"revision": dataset.revision, "record_count": result.record_count},
            source="workbuddy-dataset-demo", execution_status="COMPLETED"))
        self.db.flush()
        return {**self.get(dataset.id), "appended_revision": dataset.revision,
                "review_field_count": result.review_count, "changeset_id": changeset_id,
                "record_count_added": result.record_count,
                "original_unchanged": True, "reused": False}

    def _members(self, dataset):
        """固定 revision，历史运行不可被后续追加改写。"""
        return self.db.query(StructuredDatasetMember, StructuredExtractionRun).join(
            StructuredExtractionRun, StructuredExtractionRun.id == StructuredDatasetMember.structured_extraction_run_id
        ).filter(StructuredDatasetMember.dataset_id == dataset.id,
                 StructuredDatasetMember.appended_revision <= dataset.revision
        ).order_by(StructuredDatasetMember.appended_revision).all()

    def _records(self, dataset):
        """逐文件聚合字段，缺失可选字段不阻止整条记录进入正式结果。"""
        required = {f["key"] for f in dataset.schema_json if f.get("required")}
        source_version = (
            self.db.get(DocumentVersion, dataset.target_source_version_id)
            if dataset.target_source_version_id
            else None
        )
        records = [
            {
                "record_id": f"base:{dataset.id}:{index}",
                "document_id": source_version.document_id if source_version else None,
                "document_version_id": source_version.id if source_version else None,
                "fields": {
                    field["key"]: {
                        "value": row.get(field["key"]),
                        "raw_text": None if row.get(field["key"]) is None else str(row.get(field["key"])),
                        "status": "IMPORTED_FROM_TARGET_EXCEL",
                        "confidence": 1.0,
                        "page_number": None,
                        "bbox": {},
                    }
                    for field in dataset.schema_json
                },
                "needs_review": False,
                "dataset_revision": 0,
            }
            for index, row in enumerate(dataset.base_rows_json or [], 1)
        ]
        for member, run in self._members(dataset):
            grouped = {}
            for field in self.repository.load_fields(run.id):
                row = grouped.setdefault(field.record_index, {
                    "record_id": f"{member.id}:{field.record_index}", "document_id": run.document_id,
                    "document_version_id": run.document_version_id, "fields": {}, "needs_review": False,
                    "dataset_revision": member.appended_revision,
                })
                invalid = field.status not in {"EXTRACTED", "NORMALIZED"} and not (
                    field.status == "MISSING" and field.field_key not in required)
                row["needs_review"] |= invalid
                row["fields"][field.field_key] = {
                    "value": field.normalized_value_json, "raw_text": field.raw_text,
                    "display_value": display_value(field.normalized_value_json, field.presentation_json or {}),
                    "assessment": field.presentation_json or {},
                    "status": field.status, "confidence": (
                        None if (field.presentation_json or {}).get("confidence_origin") == "UNSCORED" else field.confidence),
                    "page_number": field.page_number, "bbox": field.bbox_json,
                }
            records.extend(grouped.values())
        # 更正只投影当前值；原始提取字段、成员和旧 Excel 版本均保持不可变。
        by_id = {row["record_id"]: row for row in records}
        corrections = self.db.query(StructuredDatasetCorrection).filter(
            StructuredDatasetCorrection.dataset_id == dataset.id,
            StructuredDatasetCorrection.dataset_revision <= dataset.revision,
        ).order_by(StructuredDatasetCorrection.dataset_revision).all()
        for correction in corrections:
            row = by_id.get(correction.record_id)
            if row is None:
                continue
            row["fields"].update(correction.changes_json)
            row["correction_revision"] = correction.dataset_revision
        return records

    def query(self, dataset_id, request: DatasetQuery):
        """过滤和聚合使用完整合格集合，再截取返回页。"""
        dataset = self.dataset(dataset_id)
        assert_publication_ready(self.db, dataset_id=dataset.id)
        fields = {f["key"]: f for f in dataset.schema_json}
        for condition in request.filters:
            if condition.field not in fields:
                fail("UNKNOWN_FIELD", "查询字段不存在。", 422)
            self._comparable(fields[condition.field], condition.value)
        if request.sum_field and (request.sum_field not in fields or fields[request.sum_field]["field_type"] not in {"integer", "decimal"}):
            fail("INVALID_SUM_FIELD", "只能对整数或小数字段求和。", 422)
        records = self._records(dataset)
        self._overlay_manual_revision_values(dataset, records)
        excluded = sum(row["needs_review"] for row in records)
        matched = []
        for row in records:
            if row["needs_review"] and not request.include_review:
                continue
            if all(self._matches(fields[c.field], row["fields"][c.field]["value"], c.op, c.value) for c in request.filters):
                matched.append(row)
        total = None
        if request.sum_field:
            # 即使查看待复核记录，合计也只纳入合格记录，避免未验证金额进入结论。
            total = str(sum((Decimal(str(r["fields"][request.sum_field]["value"])) for r in matched
                if not r["needs_review"] and r["fields"][request.sum_field]["value"] is not None), Decimal(0)))
        return {"dataset_id": dataset.id, "revision": dataset.revision, "total_records": len(records),
                "review_records": excluded, "matched_records": len(matched), "sum": total,
                "records": matched[request.offset:request.offset + request.limit],
                "truncated": request.offset + request.limit < len(matched)}

    def _overlay_manual_revision_values(self, dataset, records):
        """结构化查询展示已发布的人工补充值，保留原来源证据字段。"""
        by_id = {record["record_id"]: record for record in records}
        tasks = self.db.query(StructuredDatasetTask).filter(
            StructuredDatasetTask.status == "COMPLETED",
            StructuredDatasetTask.request_key.like("manual-revision:%"),
        ).all()
        ordered = sorted((task for task in tasks
                          if (task.manifest_json or {}).get("dataset_id") == dataset.id
                          and int((task.result_json or {}).get("revision") or 0) <= dataset.revision),
                         key=lambda task: int((task.result_json or {}).get("revision") or 0))
        for task in ordered:
            revision = int((task.result_json or {}).get("revision") or 0)
            for change in (task.payload_json or {}).get("changes", []):
                record = by_id.get(change.get("record_id"))
                if record is None:
                    continue
                field = record["fields"].get(change.get("field_key"))
                if field is None or int(record.get("correction_revision") or 0) > revision:
                    continue
                field.update(value=change["after"], display_value=change["after"],
                             status="MANUAL_EDIT", manual_revision=revision)

    def _comparable(self, field, value):
        """字符串金额不能按字典序比较；非法数值返回契约错误。"""
        if field["field_type"] == "date":
            normalized, status, _ = normalize_field_value(
                field=StructuredFieldSpec.model_validate(field), raw_text=str(value), candidate_value=None,
            )
            if status != "NORMALIZED":
                fail("INVALID_FILTER_VALUE", "日期过滤条件必须包含明确年月日。", 422)
            return normalized
        if field["field_type"] in {"integer", "decimal"}:
            try:
                number = Decimal(str(value))
                if not number.is_finite():
                    raise InvalidOperation
                return number
            except InvalidOperation:
                fail("INVALID_FILTER_VALUE", "数值过滤条件必须是有限数字。", 422)
        return str(value)

    def _matches(self, field, value, op, expected):
        """仅执行已校验的固定比较，不计算表达式。"""
        if value is None:
            return False
        left, right = self._comparable(field, value), self._comparable(field, expected)
        return {"eq": lambda: left == right, "contains": lambda: str(right) in str(left),
                "gte": lambda: left >= right, "lte": lambda: left <= right}[op]()

    def export(self, dataset_id):
        """把当前 revision 发布为同一普通工作副本的一个不可变版本。"""
        dataset = self.dataset(dataset_id, lock=True)
        self._assert_task_write_access(dataset)
        self._assert_dataset_revision_integrity(dataset)
        all_records = self._records(dataset)
        records = [record for record in all_records if not record["needs_review"]]
        if dataset.revision > 0 and not any(int(r.get("dataset_revision") or 0) > 0 for r in records):
            fail(
                "DATASET_EXPORT_NO_CONFIRMED_RECORDS",
                "当前累计版本没有任何已确认记录，拒绝生成只有标题行的 Excel。请先复核或修正提取结果。",
                review_record_count=len(all_records),
            )
        export = self.db.query(StructuredDatasetExport).filter_by(
            dataset_id=dataset.id, dataset_revision=dataset.revision).one_or_none()
        export_version = (
            self.db.get(DocumentVersion, export.document_version_id)
            if export is not None and export.document_version_id
            else None
        )
        rebuild_legacy_export = bool(
            export_version is not None
            and dataset.revision > 0
            and (export_version.extraction_policy_json or {}).get(
                "structured_dataset_export_mode"
            ) != "package-append-v3"
        )
        if export and export.document_version_id and not rebuild_legacy_export:
            path = self.storage.working_copy_path(export.storage_path)
            if path.is_file() and self.storage.sha256_file(path) == export.sha256:
                working_copy = self._dataset_working_copy(dataset)
                cleanup = None
                if working_copy is not None:
                    version = self.db.get(DocumentVersion, export.document_version_id)
                    if version is None:
                        fail("DATASET_WORKING_COPY_INVALID", "累计表格版本链不完整，需要管理员修复。")
                    cleanup = self._relocate_dataset_target(
                        dataset=dataset,
                        working_copy=working_copy,
                        current_version=version,
                    )
                result = self._export_result(export, reused=True)
                if working_copy is not None:
                    visible = self.storage.working_copy_path(self._working_copy_storage_path(working_copy))
                    if cleanup or not visible.is_file() or self.storage.sha256_file(visible) != export.sha256:
                        publication = stage_publication(
                            self.db, dataset=dataset, working_copy=working_copy, version=version,
                            base_version_id=version.parent_version_id, result=result,
                            cleanup=cleanup, agent_run=self.agent_run)
                        if not self.workflow_task_id:
                            dataset.active_task_id = publication.id
                        result["_publication_id"] = publication.id
                if cleanup:
                    result["_visible_cleanup"] = cleanup
                return result
            fail("EXPORT_MISSING", "已发布快照缺失或损坏，需要检查存储。")
        export_id = export.id if export else str(uuid4())
        working_copy = self._dataset_working_copy(dataset)
        parent_version = (
            self.db.get(DocumentVersion, working_copy.current_version_id)
            if working_copy is not None and working_copy.current_version_id
            else None
        )
        cleanup = None
        previous_export = None
        template_version = parent_version
        if working_copy is not None and parent_version is not None:
            previous_export = self.db.query(StructuredDatasetExport).filter_by(
                dataset_id=dataset.id,
                document_version_id=parent_version.id,
            ).one_or_none()
            if previous_export is None:
                fail("DATASET_WORKING_COPY_INVALID", "累计表格当前版本缺少对应导出记录，需要管理员修复。")
            cleanup = self._relocate_dataset_target(
                dataset=dataset,
                working_copy=working_copy,
                current_version=parent_version,
            )
        if rebuild_legacy_export and dataset.target_source_version_id:
            base_export = self.db.query(StructuredDatasetExport).filter_by(
                dataset_id=dataset.id,
                dataset_revision=0,
            ).one_or_none()
            if base_export is None or not base_export.document_version_id:
                fail("DATASET_WORKING_COPY_INVALID", "累计表格原始模板版本缺失，需要管理员修复。")
            template_version = self.db.get(DocumentVersion, base_export.document_version_id)
            if template_version is None:
                fail("DATASET_WORKING_COPY_INVALID", "累计表格原始模板版本缺失，需要管理员修复。")
            previous_export = base_export
        elif rebuild_legacy_export:
            previous_export = None
            template_version = None
        version_number = (
            int(
                self.db.query(func.max(DocumentVersion.version_number))
                .filter(DocumentVersion.document_id == working_copy.document_id)
                .scalar()
                or 0
            )
            + 1
            if working_copy is not None
            else 1
        )
        storage_path = self.storage.dataset_working_version_relative_path(
            dataset_id=dataset.id,
            export_id=export_id,
            version_number=version_number,
        )
        path = self.storage.working_copy_path(storage_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        keys = [f["key"] for f in dataset.schema_json]
        if previous_export is not None:
            records = [
                record for record in records
                if int(record.get("dataset_revision") or 0) > previous_export.dataset_revision
            ]
        rows = [
            [
                _safe_spreadsheet_text(value) if isinstance(value, str) else value
                for value in [record["fields"][key].get("display_value", record["fields"][key]["value"]) for key in keys]
            ]
            for record in records
        ]
        if template_version is not None:
            template_path = self.storage.working_copy_path(template_version.storage_path)
            if not template_path.is_file() or self.storage.sha256_file(template_path) != template_version.sha256:
                fail("EXPORT_MISSING", "目标 Excel 当前版本缺失或损坏，需要检查存储。")
            try:
                _append_xlsx_rows_atomic(
                    template_path=template_path,
                    target_path=path,
                    headers=[f["label"] for f in dataset.schema_json],
                    rows=rows,
                    sheet_name=dataset.target_sheet_name or "Sheet1",
                    layout=dataset.target_layout_json or None,
                )
            except ValueError as exc:
                code, _, message = str(exc).partition(":")
                fail(code if code.startswith("TARGET_EXCEL_") else "TARGET_EXCEL_WRITE_FAILED",
                     message.strip() or str(exc), 422)
        else:
            _write_xlsx_atomic(
                target_path=path,
                headers=[_safe_spreadsheet_text(f["label"]) for f in dataset.schema_json],
                rows=rows,
                sheet_name=dataset.target_sheet_name or "Sheet1",
            )
        sha256 = self.storage.sha256_file(path)
        size_bytes = path.stat().st_size
        working_copy, version = self._publish_working_version(
            dataset=dataset,
            working_copy=working_copy,
            parent_version=parent_version,
            storage_path=storage_path,
            version_number=version_number,
            sha256=sha256,
            size_bytes=size_bytes,
        )
        version.extraction_policy_json = {
            **dict(version.extraction_policy_json or {}),
            "structured_dataset_export_mode": "package-append-v3",
        }
        if export is None:
            export = StructuredDatasetExport(
                id=export_id,
                dataset_id=dataset.id,
                dataset_revision=dataset.revision,
                sha256=sha256,
                size_bytes=size_bytes,
                storage_path=storage_path,
                document_version_id=version.id,
            )
            self.db.add(export)
        else:
            # 兼容旧 derivative 导出：同一 revision 首次重新导出时转为工作副本版本。
            export.sha256 = sha256
            export.size_bytes = size_bytes
            export.storage_path = storage_path
            export.document_version_id = version.id
        self.db.flush()
        result = self._export_result(export, reused=False)
        publication = stage_publication(
            self.db, dataset=dataset, working_copy=working_copy, version=version,
            base_version_id=parent_version.id if parent_version else None,
            result=result, cleanup=({**cleanup, "expected_current_sha256": sha256} if cleanup else None),
            agent_run=self.agent_run)
        if not self.workflow_task_id:
            dataset.active_task_id = publication.id
        result["_publication_id"] = publication.id
        if cleanup:
            # 迁移后本次导出又发布了新内容，提交后清理必须校验最终版本哈希，
            # 不能继续使用迁移前 parent version 的哈希，否则旧入口会被错误保留。
            result["_visible_cleanup"] = {
                **cleanup,
                "expected_current_sha256": sha256,
            }
        return result

    def _assert_dataset_revision_integrity(self, dataset: StructuredDataset) -> None:
        """拒绝 revision 与成员链不一致的历史数据，避免把事实缺失误报为成功导出。"""

        revisions = [
            value
            for (value,) in self.db.query(StructuredDatasetMember.appended_revision)
            .filter_by(dataset_id=dataset.id)
            .order_by(StructuredDatasetMember.appended_revision)
            .all()
        ]
        expected = list(range(1, dataset.revision + 1))
        if revisions != expected:
            fail(
                "DATASET_REVISION_INCONSISTENT",
                "累计表格 revision 与已保存成员不一致，拒绝发布不完整 Excel；请重新提交缺失来源。",
                dataset_revision=dataset.revision,
                member_revisions=revisions,
            )

    def _dataset_working_copy(self, dataset: StructuredDataset, *, allow_archived=False) -> WorkingCopy | None:
        """读取数据集唯一活动工作副本，关联损坏时关闭式失败。"""

        if not dataset.working_copy_id:
            return None
        working_copy = self.db.get(WorkingCopy, dataset.working_copy_id)
        if (
            working_copy is None
            or working_copy.workspace_id != dataset.workspace_id
            or working_copy.status not in ({"ACTIVE", "ARCHIVED"} if allow_archived else {"ACTIVE"})
        ):
            fail("DATASET_WORKING_COPY_INVALID", "累计表格工作副本关联无效，需要管理员修复。")
        return working_copy

    def _ensure_generated_root(
        self,
        *,
        dataset: StructuredDataset,
    ) -> tuple[ManagedRoot, WorkingCopyRoot]:
        """取得服务器生成文件专用根，不把任意客户端路径带入存储。"""

        root_key = "system.structured-datasets"
        managed_root = self.db.query(ManagedRoot).filter_by(root_key=root_key).one_or_none()
        if managed_root is None:
            managed_root = ManagedRoot(
                root_key=root_key,
                display_name="累计表格",
                container_path="generated://structured-datasets",
                classification_mode="NONE",
                enabled=True,
                read_only=True,
                archive_write_enabled=False,
                allowed_operations_json=["READ", "DOWNLOAD"],
                created_by=self.user.id,
            )
            self.db.add(managed_root)
            self.db.flush()
        working_root = (
            self.db.query(WorkingCopyRoot)
            .filter_by(workspace_id=dataset.workspace_id, managed_root_id=managed_root.id)
            .one_or_none()
        )
        if working_root is None:
            working_root = WorkingCopyRoot(
                workspace_id=dataset.workspace_id,
                managed_root_id=managed_root.id,
                root_key=root_key,
                relative_storage_path="structured-datasets",
                status="ACTIVE",
            )
            self.db.add(working_root)
            self.db.flush()
        return managed_root, working_root

    def _relocate_dataset_target(
        self,
        *,
        dataset: StructuredDataset,
        working_copy: WorkingCopy,
        current_version: DocumentVersion,
    ) -> dict | None:
        """把目标表的唯一可见入口收敛到 structured-datasets 受控目录。

        不可变版本始终作为发布源；旧可见入口只在数据库提交成功后由运行时移除。
        """

        filename = self._dataset_filename(dataset)
        # 兼容 target_filename 字段引入前已经发布的累计表：工作副本名称才是
        # 用户实际看到和继续引用的名称，迁移或再次导出时把它回填到稳定绑定。
        if dataset.target_filename != filename:
            dataset.target_filename = filename
        managed_root, working_root = self._ensure_generated_root(dataset=dataset)
        desired_relative_path = f"{dataset.id}/{filename}"
        current_root = self.db.get(WorkingCopyRoot, working_copy.working_copy_root_id)
        if current_root is None:
            fail("DATASET_WORKING_COPY_INVALID", "累计表格工作副本根不存在。")
        old_storage_path = self._working_copy_storage_path(working_copy)
        root_path = working_root.relative_storage_path.strip("/\\")
        relative_path = desired_relative_path.strip("/\\")
        desired_storage_path = f"{root_path}/{relative_path}"
        if (
            working_copy.working_copy_root_id == working_root.id
            and working_copy.relative_path == desired_relative_path
        ):
            return None
        version_path = self.storage.working_copy_path(current_version.storage_path)
        if not version_path.is_file() or self.storage.sha256_file(version_path) != current_version.sha256:
            fail("EXPORT_MISSING", "累计表格当前版本缺失或损坏，需要检查存储。")
        destination = self.storage.working_copy_path(desired_storage_path)
        if destination.exists() and (
            not destination.is_file()
            or self.storage.sha256_file(destination) != current_version.sha256
        ):
            fail("TARGET_EXCEL_FILENAME_CONFLICT", f"文件名“{filename}”对应的物理文件已存在。")
        managed_file = self.db.get(ManagedFile, working_copy.managed_file_id)
        if managed_file is None:
            fail("DATASET_WORKING_COPY_INVALID", "累计表格文件关联不完整。")
        before_relative_path = working_copy.relative_path
        before_filename = working_copy.filename
        managed_file.root_id = managed_root.id
        managed_file.relative_path = desired_relative_path
        managed_file.relative_path_hash = hashlib.sha256(desired_relative_path.encode()).hexdigest()
        managed_file.filename = filename
        managed_file.extension = ".xlsx"
        working_copy.working_copy_root_id = working_root.id
        working_copy.relative_path = desired_relative_path
        working_copy.relative_path_hash = hashlib.sha256(desired_relative_path.encode()).hexdigest()
        working_copy.filename = filename
        working_copy.extension = ".xlsx"
        working_copy.revision += 1
        file_object = self.db.query(FileObject).filter_by(
            document_id=working_copy.document_id,
            storage_backend="working_copy_local",
        ).one_or_none()
        if file_object is None:
            file_object = FileObject(
                document_id=working_copy.document_id,
                storage_backend="working_copy_local",
                storage_path=desired_storage_path,
                size_bytes=current_version.size_bytes,
                sha256=current_version.sha256,
            )
            self.db.add(file_object)
        else:
            file_object.storage_path = desired_storage_path
            file_object.size_bytes = current_version.size_bytes
            file_object.sha256 = current_version.sha256
        self.db.add(WorkingCopyPathRecord(
            working_copy_id=working_copy.id,
            sequence_number=self._next_path_sequence(working_copy.id),
            operation_type="DATASET_TARGET_RELOCATED",
            before_relative_path=before_relative_path,
            after_relative_path=desired_relative_path,
            before_filename=before_filename,
            after_filename=filename,
            document_version_id=current_version.id,
            content_sha256=current_version.sha256,
            status="COMPLETED",
            executed_by=self.user.id,
        ))
        self.db.flush()
        return {
            "legacy_relative_path": old_storage_path,
            "expected_legacy_sha256": current_version.sha256,
            "current_relative_path": desired_storage_path,
            "expected_current_sha256": current_version.sha256,
        }

    def _publish_working_version(
        self,
        *,
        dataset: StructuredDataset,
        working_copy: WorkingCopy | None,
        parent_version: DocumentVersion | None,
        storage_path: str,
        version_number: int,
        sha256: str,
        size_bytes: int,
    ) -> tuple[WorkingCopy, DocumentVersion]:
        """首次创建一个逻辑文件，后续只更新其当前版本。"""

        filename = self._dataset_filename(dataset)
        if dataset.target_filename != filename:
            dataset.target_filename = filename
        logical_path = working_copy.relative_path if working_copy is not None else f"{dataset.id}/{filename}"
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        now = utcnow()
        if working_copy is None:
            managed_root, working_root = self._ensure_generated_root(dataset=dataset)
            managed_file = ManagedFile(
                root_id=managed_root.id,
                relative_path=logical_path,
                relative_path_hash=hashlib.sha256(logical_path.encode()).hexdigest(),
                filename=filename,
                extension=".xlsx",
                size_bytes=size_bytes,
                modified_at=now,
                fingerprint=sha256,
                content_sha256=sha256,
                source_type="GENERATED_DATASET",
                status="ACTIVE",
            )
            document = Document(
                user_id=dataset.created_by,
                workspace_id=dataset.workspace_id,
                original_filename=filename,
                content_type=content_type,
                size_bytes=size_bytes,
                sha256=sha256,
                status="READY",
                ingest_status="READY",
            )
            self.db.add_all([managed_file, document])
            self.db.flush()
            version = DocumentVersion(
                document_id=document.id,
                version_number=version_number,
                parent_version_id=None,
                storage_tier="WORKING_COPY",
                storage_path=storage_path,
                filename=filename,
                content_type=content_type,
                size_bytes=size_bytes,
                sha256=sha256,
                source_type="STRUCTURED_DATASET_EXPORT",
                created_by=self.user.id,
            )
            self.db.add(version)
            self.db.flush()
            working_copy = WorkingCopy(
                working_copy_root_id=working_root.id,
                workspace_id=dataset.workspace_id,
                managed_file_id=managed_file.id,
                document_id=document.id,
                current_version_id=version.id,
                relative_path=logical_path,
                relative_path_hash=hashlib.sha256(logical_path.encode()).hexdigest(),
                filename=filename,
                extension=".xlsx",
                size_bytes=size_bytes,
                content_sha256=sha256,
                imported_source_sha256=sha256,
                is_primary_import=True,
                status="ACTIVE",
                sync_status="SYNCED",
                placement_status="GENERATED",
                placement_policy_version="structured-dataset-v1",
            )
            self.db.add(working_copy)
            self.db.flush()
            version.working_copy_id = working_copy.id
            dataset.working_copy_id = working_copy.id
            before_path = ""
            before_filename = ""
        else:
            document = self.db.get(Document, working_copy.document_id)
            managed_file = self.db.get(ManagedFile, working_copy.managed_file_id)
            if document is None or managed_file is None or parent_version is None:
                fail("DATASET_WORKING_COPY_INVALID", "累计表格版本链不完整，需要管理员修复。")
            if working_copy.filename != filename or Path(working_copy.relative_path).name != filename:
                fail("DATASET_WORKING_COPY_INVALID", "累计表格工作副本名称与用户可见名称不一致。")
            version = DocumentVersion(
                document_id=document.id,
                version_number=version_number,
                parent_version_id=parent_version.id,
                working_copy_id=working_copy.id,
                storage_tier="WORKING_COPY",
                storage_path=storage_path,
                filename=working_copy.filename,
                content_type=content_type,
                size_bytes=size_bytes,
                sha256=sha256,
                source_type="STRUCTURED_DATASET_EXPORT",
                created_by=self.user.id,
            )
            self.db.add(version)
            self.db.flush()
            before_path = working_copy.relative_path
            before_filename = working_copy.filename
            document.size_bytes = size_bytes
            document.sha256 = sha256
            document.status = "READY"
            document.ingest_status = "READY"
            managed_file.size_bytes = size_bytes
            managed_file.modified_at = now
            managed_file.fingerprint = sha256
            managed_file.content_sha256 = sha256
            working_copy.current_version_id = version.id
            working_copy.size_bytes = size_bytes
            working_copy.content_sha256 = sha256
            working_copy.imported_source_sha256 = sha256
            working_copy.sync_status = "SYNCED"
            working_copy.revision += 1
        file_object = (
            self.db.query(FileObject)
            .filter_by(document_id=document.id, storage_backend="working_copy_local")
            .one_or_none()
        )
        if file_object is None:
            file_object = FileObject(
                document_id=document.id,
                storage_backend="working_copy_local",
                storage_path=self._working_copy_storage_path(working_copy),
                size_bytes=size_bytes,
                sha256=sha256,
            )
            self.db.add(file_object)
        else:
            # FileObject 表示当前普通文件，而不可变历史内容由 DocumentVersion 保存。
            file_object.storage_path = self._working_copy_storage_path(working_copy)
            file_object.size_bytes = size_bytes
            file_object.sha256 = sha256
        self.db.add(WorkingCopyPathRecord(
            working_copy_id=working_copy.id,
            sequence_number=self._next_path_sequence(working_copy.id),
            operation_type="DATASET_VERSION_PUBLISHED",
            before_relative_path=before_path,
            after_relative_path=working_copy.relative_path,
            before_filename=before_filename,
            after_filename=working_copy.filename,
            document_version_id=version.id,
            content_sha256=sha256,
            status="COMPLETED",
            executed_by=self.user.id,
        ))
        return working_copy, version

    def _next_path_sequence(self, working_copy_id: str) -> int:
        """在当前事务内取得下一条工作副本路径审计序号。"""

        return int(
            self.db.query(func.max(WorkingCopyPathRecord.sequence_number))
            .filter(WorkingCopyPathRecord.working_copy_id == working_copy_id)
            .scalar()
            or 0
        ) + 1

    def _working_copy_storage_path(self, working_copy: WorkingCopy) -> str:
        """把工作副本根和业务相对路径组合为 StorageService 使用的受控路径。"""

        root = self.db.get(WorkingCopyRoot, working_copy.working_copy_root_id)
        if root is None or not root.relative_storage_path.strip("/\\"):
            fail("DATASET_WORKING_COPY_INVALID", "累计表格工作副本根不存在。")
        root_path = root.relative_storage_path.strip("/\\")
        relative_path = working_copy.relative_path.strip("/\\")
        return f"{root_path}/{relative_path}"

    def _export_result(self, export, *, reused):
        """导出只返回稳定对象引用，不暴露存储路径。"""
        version = self.db.get(DocumentVersion, export.document_version_id) if export.document_version_id else None
        dataset = self.dataset(export.dataset_id)
        working_copy = self._dataset_working_copy(dataset)
        # 回执只统计此导出 revision 已确认并写入的记录，不把待复核记录算成成功。
        records = [record for record in self._records(dataset)
                   if not record["needs_review"] and 0 < int(record.get("dataset_revision") or 0) <= export.dataset_revision]
        return {"dataset_id": export.dataset_id, "revision": export.dataset_revision,
                "excel_published": True, "saved_to_server": True,
                "confirmed_record_count": len(records),
                "base_record_count": len(dataset.base_rows_json or []),
                "export_id": export.id, "sha256": export.sha256, "size_bytes": export.size_bytes,
                "working_copy_id": dataset.working_copy_id,
                "document_id": version.document_id if version else None,
                "document_version_id": export.document_version_id,
                "version_number": version.version_number if version else None,
                "filename": working_copy.filename if working_copy else self._dataset_filename(dataset),
                "download_url": f"/api/structured-datasets/{export.dataset_id}/exports/{export.id}/content",
                "reused": reused}

    def download(self, dataset_id, export_id):
        """登录用户均可下载共享快照，路径只通过 StorageService 生成。"""
        self.dataset(dataset_id)
        export = self.db.query(StructuredDatasetExport).filter_by(id=export_id, dataset_id=dataset_id).one_or_none()
        if export is None:
            fail("EXPORT_NOT_FOUND", "导出不存在。", 404)
        pending = pending_publication(self.db, dataset_id=dataset_id)
        if pending is not None and export.document_version_id == pending.candidate_document_version_id:
            assert_publication_ready(self.db, dataset_id=dataset_id)
        path = (
            self.storage.working_copy_path(export.storage_path)
            if export.document_version_id
            else self.storage.upload_path(export.storage_path)
        )
        if not path.is_file() or self.storage.sha256_file(path) != export.sha256:
            fail("EXPORT_MISSING", "导出文件缺失或校验失败。", 404)
        working_copy = self._dataset_working_copy(self.dataset(dataset_id), allow_archived=True)
        filename = working_copy.filename if working_copy else self._dataset_filename(self.dataset(dataset_id))
        return FileResponse(path, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            filename=filename)

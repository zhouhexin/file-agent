"""用户手工补全已发布 Excel 后，将完整文件发布为同一工作副本的新版本。"""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from uuid import uuid4
import re
import tempfile
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from sqlalchemy import func
from fastapi import HTTPException

from app.db.models import DocumentVersion, StructuredDatasetExport, StructuredDatasetTask, utcnow
from app.modules.structured_extraction.workbook_layout import WorkbookLayout

from .publication import stage_publication
from .service import StructuredDatasetService, fail
from .xlsx_safety import XlsxSafetyError, inspect_xlsx


def _json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _funding_text_value(value, *, field_key: str, cell: str):
    """保留 Excel 标量原意，将汇总页文本列的非文本值转为可保存的文本。"""
    if value is None or isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if field_key == "date" and value.time() == time.min:
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, (int, float, Decimal)):
        try:
            if Decimal(str(value)).is_finite():
                return str(value)
        except InvalidOperation:
            pass
    fail("MANUAL_EXCEL_FIELD_TYPE", "登记、日期或状态含无法安全转成文本的值。", 422,
         cell=cell)


class ManualRevisionService:
    def __init__(self, db, user, agent_run=None):
        self.db = db
        self.user = user
        self.datasets = StructuredDatasetService(db, user, agent_run)

    def publish(self, request, path: Path) -> dict:
        existing = self.db.query(StructuredDatasetTask).filter_by(
            user_id=self.user.id, request_key=f"manual-revision:{request.request_key}").one_or_none()
        if existing is not None:
            snapshot = existing.manifest_json or {}
            if (snapshot.get("filename") != request.filename
                    or snapshot.get("base_document_version_id") != str(request.base_document_version_id)
                    or snapshot.get("base_sha256") != request.base_sha256
                    or snapshot.get("base_revision") != request.base_revision
                    or (existing.progress_json or {}).get("file_sha256") != request.file_sha256):
                fail("MANUAL_EXCEL_REQUEST_CONFLICT", "请求键已用于其他目标、版本或文件。", 409)
            return self.receipt(existing)
        if not path.is_file() or self.datasets.storage.sha256_file(path) != request.file_sha256:
            fail("MANUAL_EXCEL_FILE_CHANGED", "上传文件与声明的哈希不一致。", 422)
        target = self.datasets.resolve_target(
            request.filename, manual_upload_sha256=request.file_sha256)
        dataset = self.datasets.dataset(target["dataset_id"], lock=True)
        working = self.datasets._dataset_working_copy(dataset)
        if working is None or not working.current_version_id:
            fail("MANUAL_EXCEL_TARGET_MISSING", "目标 Excel 尚未发布。", 409)
        self.datasets._assert_task_write_access(dataset)
        current = self.db.get(DocumentVersion, working.current_version_id)
        if (current is None or current.id != str(request.base_document_version_id)
                or current.sha256 != request.base_sha256 or dataset.revision != request.base_revision
                or working.filename != request.filename):
            fail("MANUAL_EXCEL_VERSION_CONFLICT", "目标 Excel 已变化，请下载最新版本后重新补充。", 409)
        visible = self.datasets.storage.working_copy_path(
            self.datasets._working_copy_storage_path(working))
        visible_sha256 = self.datasets.storage.sha256_file(visible) if visible.is_file() else None
        if visible_sha256 not in {current.sha256, request.file_sha256}:
            fail("MANUAL_EXCEL_VISIBLE_CONFLICT", "服务器同名 Excel 与上传文件不一致，请先核对两份修改。", 409)
        base_path = self.datasets.storage.working_copy_path(current.storage_path)
        if not base_path.is_file() or self.datasets.storage.sha256_file(base_path) != current.sha256:
            fail("MANUAL_EXCEL_BASE_MISSING", "目标基线缺失或损坏。", 409)
        origin = self.db.query(DocumentVersion).filter_by(
            document_id=working.document_id, version_number=1).one_or_none()
        funding = bool(origin and (origin.extraction_policy_json or {}).get(
            "structured_dataset_export_mode") == "funding-usage-v1" and any(
            field.get("key") == "date" for field in dataset.schema_json or []))
        if not funding and current.sha256 == request.file_sha256:
            fail("MANUAL_EXCEL_UNCHANGED", "上传文件与当前版本相同。", 422)
        if funding:
            changes, normalized_bytes, normalized_rows, balance_overrides, source_policy, deleted_sequences = (
                self._prepare_funding_revision(dataset, base_path, path, current, origin, request.request_key))
            output_sha = sha256(normalized_bytes).hexdigest()
            output_size = len(normalized_bytes)
            dataset.base_rows_json = normalized_rows
            dataset.target_layout_json = {"header_range": "A1:G1",
                "data_range": f"A2:G{max(2, len(normalized_rows) + 1)}"}
        else:
            changes = self._validate_and_diff(dataset, base_path, path)
            normalized_bytes = None
            balance_overrides = None
            source_policy = {}
            deleted_sequences = []
            output_sha = request.file_sha256
            output_size = path.stat().st_size
        if funding and not changes:
            fail("MANUAL_EXCEL_UNCHANGED", "汇总页没有需要发布的业务修改。", 422)
        task = StructuredDatasetTask(
            user_id=self.user.id, request_key=f"manual-revision:{request.request_key}",
            result_digest=request.file_sha256, review_mode="DIRECT_SAVE", status="PUBLISHING",
            manifest_json={"mode": "MANUAL_REVISION", "dataset_id": dataset.id,
                           "working_copy_id": working.id, "base_document_version_id": current.id,
                           "base_sha256": current.sha256, "base_revision": dataset.revision,
                           "filename": working.filename},
            payload_json={"changes": changes}, progress_json={
                "file_sha256": request.file_sha256,
                "published_sha256": output_sha,
                "visible_sha256_before_publish": visible_sha256,
            },
            result_json={}, error_json={},
        )
        self.db.add(task)
        self.db.flush()
        dataset.active_task_id = task.id
        dataset.revision += 1
        export_id = str(uuid4())
        version_number = int(self.db.query(func.max(DocumentVersion.version_number)).filter(
            DocumentVersion.document_id == working.document_id).scalar() or 0) + 1
        relative_path = self.datasets.storage.dataset_working_version_relative_path(
            dataset_id=dataset.id, export_id=export_id, version_number=version_number)
        if normalized_bytes is None:
            self.datasets.storage.import_working_copy(
                source=path, relative_path=relative_path, expected_sha256=output_sha)
        else:
            with tempfile.TemporaryDirectory(prefix="file-agent-funding-revision-") as directory:
                normalized_path = Path(directory) / request.filename
                normalized_path.write_bytes(normalized_bytes)
                self.datasets.storage.import_working_copy(
                    source=normalized_path, relative_path=relative_path, expected_sha256=output_sha)
        working, version = self.datasets._publish_working_version(
            dataset=dataset, working_copy=working, parent_version=current,
            storage_path=relative_path, version_number=version_number,
            sha256=output_sha, size_bytes=output_size)
        version.extraction_policy_json = {**dict(current.extraction_policy_json or {}),
            "structured_dataset_export_mode": "funding-usage-v2" if funding else "package-append-v3",
            "manual_revision": True,
            **({"funding_usage_balance_overrides": balance_overrides, **source_policy} if funding else {})}
        export = StructuredDatasetExport(
            id=export_id, dataset_id=dataset.id, dataset_revision=dataset.revision,
            storage_path=relative_path, sha256=output_sha,
            size_bytes=output_size, document_version_id=version.id)
        self.db.add(export)
        result = {**self.datasets._export_result(export, reused=False),
                  "manual_changed_cell_count": sum(item.get("field_key") not in {"__row__", "__sequence__"}
                                                   for item in changes),
                  "deleted_row_sequences": deleted_sequences,
                  "deleted_record_count": len(deleted_sequences),
                  "added_record_count": len(source_policy.get("funding_usage_added_rows", [])),
                  "added_row_sequences": [item["sequence"] for item in
                                          source_policy.get("funding_usage_added_rows", [])],
                  "renumbered_sequences": source_policy.get("funding_usage_sequence_changes", []),
                  "task_record_count": max(1, len(changes))}
        stage = stage_publication(self.db, dataset=dataset, working_copy=working,
            version=version, base_version_id=current.id, task=task,
            result=result, agent_run=self.datasets.agent_run)
        return {**self.receipt(task), "_publication_id": stage.id}

    def receipt(self, task) -> dict:
        return {"task_id": task.id, "status": task.status,
                "dataset_id": (task.manifest_json or {}).get("dataset_id"),
                "filename": (task.manifest_json or {}).get("filename"),
                "excel_published": task.status == "COMPLETED",
                "error": task.error_json or None, **(task.result_json or {})}

    def _prepare_funding_revision(self, dataset, base_path: Path, path: Path, current, origin, request_key):
        """首表是唯一可编辑的业务来源，个人页由服务端重新生成。"""
        from app.modules.funding_usage.service import FundingUsageService
        from app.modules.funding_usage.workbook import (KEYS, money, read_summary_rows,
            rebuild_personal_sheets, repair_summary_row_styles)

        try:
            inspect_xlsx(base_path)
            inspect_xlsx(path)
            with zipfile.ZipFile(path) as archive:
                if any(name.lower().startswith("xl/externallinks/") for name in archive.namelist()):
                    fail("MANUAL_EXCEL_EXTERNAL_LINK", "补全文件不能包含外部工作簿链接。", 422)
            baseline = load_workbook(base_path, data_only=False)
            candidate = load_workbook(path, data_only=False)
        except XlsxSafetyError as exc:
            fail(exc.code, str(exc), 422)
        except HTTPException:
            raise
        except Exception:
            fail("MANUAL_EXCEL_INVALID", "Excel 无法安全读取。", 422)
        try:
            if baseline.sheetnames != candidate.sheetnames:
                fail("MANUAL_EXCEL_STRUCTURE_CHANGED", "不能增删或重命名工作表。", 422)
            for old_sheet, new_sheet in zip(baseline.worksheets, candidate.worksheets):
                if {str(x) for x in old_sheet.merged_cells.ranges} != {
                        str(x) for x in new_sheet.merged_cells.ranges}:
                    fail("MANUAL_EXCEL_STRUCTURE_CHANGED", "不能改变工作表合并区域。", 422)
                if any(cell.data_type == "f" for cell in new_sheet._cells.values()):
                    fail("MANUAL_EXCEL_FORMULA_CHANGED", "补全文件不能包含公式。", 422)
            old_rows = read_summary_rows(baseline)
            new_rows = read_summary_rows(candidate, allow_new_rows=True)
            records = self.datasets._records(dataset)
            if len(old_rows) != len(records):
                fail("FUNDING_ROW_ID_INVALID", "当前登记行与服务器记录不一致。", 409)
            old_by_sequence = {row["sequence"]: row for row in old_rows}
            # 空 A 列只表示本次人工新增；旧行的 A 列仍用于匹配冻结来源。
            for row in new_rows:
                if row["sequence"] is None:
                    row["sequence"] = -row["_manual_input_row"]
            new_sequences = [row["sequence"] for row in new_rows]
            retained = [sequence for sequence in new_sequences if sequence > 0]
            if (retained != [row["sequence"] for row in old_rows
                             if row["sequence"] in retained]
                    or any(old_by_sequence[row["sequence"]]["name"] != row["name"]
                           for row in new_rows if row["sequence"] in old_by_sequence)
                    or any(sequence not in old_by_sequence for sequence in retained)):
                fail("FUNDING_ROW_ID_INVALID", "已有行的序号、姓名和顺序不能改变；新增行的 A 列须留空。", 422)
            deleted_sequences = [row["sequence"] for row in old_rows
                                 if row["sequence"] not in new_sequences]
            old_summary, new_summary = baseline.worksheets[0], candidate.worksheets[0]
            for coordinate in set(old_summary._cells) | set(new_summary._cells):
                row_number, column_number = coordinate
                if (column_number > 7 or row_number < 2):
                    if old_summary.cell(*coordinate).value != new_summary.cell(*coordinate).value:
                        fail("FUNDING_SUMMARY_REVISION_UNSUPPORTED",
                             "只能修改首个汇总页 A–G 数据行；不能改动其他列或表头。", 422,
                             cell=new_summary.cell(*coordinate).coordinate)
            # 兼容用户已按旧说明同步修改的个人页，但首表仍是唯一业务准据。
            legacy = (current.extraction_policy_json or {}).get(
                "structured_dataset_export_mode") != "funding-usage-v2"
            for index, (old_sheet, new_sheet) in enumerate(zip(
                    baseline.worksheets[1:], candidate.worksheets[1:]), 2):
                allowed = {(2, 2): (index, 3), (2, 3): (index, 5),
                           (3, 3): (index, 6), (2, 4): (index, 7)} if legacy else {}
                for coordinate in set(old_sheet._cells) | set(new_sheet._cells):
                    before = old_sheet.cell(*coordinate).value
                    after = new_sheet.cell(*coordinate).value
                    if before == after:
                        continue
                    target = allowed.get(coordinate)
                    if target is None or after != candidate.worksheets[0].cell(*target).value:
                        fail("FUNDING_PERSONAL_SHEET_REVISION_UNSUPPORTED",
                             "只需修改首个汇总页；个人页其他修改不能合并。", 422)
            source_version = self.db.get(DocumentVersion, dataset.target_source_version_id)
            if source_version is None:
                fail("FUNDING_BASELINE_MISSING", "首次旧汇总表版本缺失。", 409)
            source_path = self.datasets.storage.working_copy_path(source_version.storage_path)
            if not source_path.is_file() or self.datasets.storage.sha256_file(source_path) != source_version.sha256:
                fail("FUNDING_BASELINE_MISSING", "首次旧汇总表文件缺失或校验失败。", 409)
            totals = FundingUsageService(self.db, self.user)._baseline(source_path)
            overrides = dict((current.extraction_policy_json or {}).get(
                "funding_usage_balance_overrides") or {})
            previous = {}
            old_previous = {}
            pending_names = set()
            for old in old_rows:
                name = old["name"]
                if name not in old_previous:
                    matches = totals.get("".join(name.split()), [])
                    old_previous[name] = matches[0] if len(matches) == 1 else None
                old_amount = money(old["usage_amount"], cell="E")
                old_balance = money(old["balance"], cell="F")
                if old_amount is not None and old_balance is not None and old_previous[name] is not None:
                    if old_balance != old_previous[name] - old_amount and str(old["sequence"]) not in overrides:
                        overrides[str(old["sequence"])] = str(old_balance)
                old_previous[name] = old_balance
            for sequence in deleted_sequences:
                overrides.pop(str(sequence), None)
            for row in new_rows:
                old = old_by_sequence.get(row["sequence"])
                sequence = str(row["sequence"])
                name = row["name"]
                for key in ("registration", "date", "status"):
                    value = row[key]
                    coordinate = f"{ {'registration': 'C', 'date': 'D', 'status': 'G'}[key]}{row['_sheet_row_number']}"
                    if new_summary[coordinate].data_type == "e":
                        fail("MANUAL_EXCEL_FIELD_TYPE", "登记、日期或状态含 Excel 错误值。", 422,
                             cell=coordinate)
                    row[key] = _funding_text_value(value, field_key=key, cell=coordinate)
                if name not in previous:
                    matches = totals.get("".join(name.split()), [])
                    previous[name] = matches[0] if len(matches) == 1 else None
                old_amount = money(old["usage_amount"], cell="E") if old else None
                old_balance = money(old["balance"], cell="F") if old else None
                amount = money(row["usage_amount"], cell="E")
                input_balance = money(row["balance"], cell="F")
                if old_amount != amount and input_balance == old_balance:
                    overrides.pop(sequence, None)
                expected = previous[name] - amount if previous[name] is not None and amount is not None else None
                manually_changed = input_balance != old_balance
                if manually_changed and input_balance is not None and input_balance != expected:
                    status = str(row["status"] or "")
                    if not re.search(r"(?:人工更正|余额更正|更正余额|更正)[：:\s]\s*\S{2,}", status):
                        fail("FUNDING_BALANCE_CORRECTION_REASON_REQUIRED",
                             "手工更改余额须在核对状态写明更正原因。", 422,
                             cell=f"F{row['_sheet_row_number']}")
                    if amount is None or name in pending_names:
                        fail("FUNDING_BALANCE_PENDING_AMOUNT", "先确认该人较早的待核对金额，再更正余额。", 422)
                    overrides[sequence] = str(input_balance)
                elif manually_changed and input_balance == expected:
                    overrides.pop(sequence, None)
                override = money(overrides[sequence], cell="F") if sequence in overrides else None
                result = override if override is not None and name not in pending_names and amount is not None else expected
                row["balance"] = float(result) if result is not None else None
                previous[name] = result
                if amount is None:
                    pending_names.add(name)
                    if "手写金额待核对" not in str(row["status"] or ""):
                        row["status"] = (str(row["status"] or "") + "；手写金额待核对").strip("；")
                else:
                    row["status"] = (str(row["status"] or "").replace("；手写金额待核对", "")
                        .replace("手写金额待核对；", "").replace("手写金额待核对", "")
                        .replace("；当前余额待核对", "").replace("当前余额待核对；", "")
                        .replace("当前余额待核对", "")
                        .replace("；可能超额，待核对", "").replace("可能超额，待核对；", "")
                        .replace("可能超额，待核对", "")
                        .strip("；") or "已核对")
                    if result is None and "当前余额待核对" not in row["status"]:
                        row["status"] = ("当前余额待核对" if row["status"] == "已核对"
                                         else row["status"] + "；当前余额待核对")
                    if result is not None and result < 0 and "可能超额" not in row["status"]:
                        row["status"] += "；可能超额，待核对"
            changes = []
            records_by_sequence = {record["fields"]["sequence"]["value"]: record
                                   for record in records}
            for index, row in enumerate(new_rows, 2):
                old = old_by_sequence.get(row["sequence"])
                if old is None:
                    continue
                record = records_by_sequence[row["sequence"]]
                for key in KEYS[2:]:
                    if old[key] != row[key]:
                        if len(changes) >= 1000:
                            fail("MANUAL_EXCEL_TOO_MANY_CHANGES", "一次最多补充 1000 个现有数据格。", 422)
                        changes.append({"record_id": record["record_id"],
                            "row_number": index, "field_key": key,
                            "before": _json_value(old[key]), "after": _json_value(row[key])})
            for index, old in enumerate(old_rows, 2):
                if old["sequence"] in deleted_sequences:
                    changes.append({"record_id": records_by_sequence[old["sequence"]]["record_id"],
                                    "row_number": index, "field_key": "__row__",
                                    "change_type": "ROW_DELETED",
                                    "before": {key: _json_value(value) for key, value in old.items()},
                                    "after": None})
            if len(changes) > 1000:
                fail("MANUAL_EXCEL_TOO_MANY_CHANGES", "一次最多修改或删除 1000 条内容。", 422)
            stored_ids = {row["sequence"]: row.get("_funding_row_id")
                          for row in dataset.base_rows_json or []}
            source_policy = self._funding_source_policy(
                current, origin, old_rows, new_rows, totals, request_key, stored_ids,
                self.datasets.storage.sha256_file(path))
            source_ids = {entry["sequence"]: entry["source_row_id"]
                          for entry in source_policy["funding_usage_entries"]}
            old_to_new = {row["sequence"]: index for index, row in enumerate(new_rows, 1)}
            sequence_changes = []
            for row in new_rows:
                original_sequence = row["sequence"]
                row["_funding_row_id"] = source_ids[original_sequence]
                row["sequence"] = old_to_new[original_sequence]
                if original_sequence < 0:
                    changes.append({"record_id": f"base:{dataset.id}:{row['_funding_row_id']}",
                                    "row_number": row["sequence"] + 1, "field_key": "__row__",
                                    "change_type": "ROW_ADDED", "before": None,
                                    "after": {key: _json_value(row[key]) for key in KEYS}})
                if original_sequence > 0 and original_sequence != row["sequence"]:
                    sequence_changes.append({"before": original_sequence, "after": row["sequence"],
                                             "source_row_id": row["_funding_row_id"]})
                    changes.append({"record_id": records_by_sequence[original_sequence]["record_id"],
                                    "row_number": row["sequence"] + 1, "field_key": "__sequence__",
                                    "change_type": "ROW_RENUMBERED", "before": original_sequence,
                                    "after": row["sequence"]})
            if len(changes) > 1000:
                fail("MANUAL_EXCEL_TOO_MANY_CHANGES", "一次最多修改或重排 1000 条内容。", 422)
            for entry in source_policy["funding_usage_entries"]:
                original_sequence = entry["sequence"]
                entry["sequence"] = old_to_new[original_sequence]
                entry["current_row"]["sequence"] = entry["sequence"]
            for evidence in source_policy["funding_usage_evidence"]:
                evidence["display_sequence"] = old_to_new[evidence["display_sequence"]]
            source_policy["funding_usage_added_rows"] = [
                {"sequence": row["sequence"], "source_row_id": row["_funding_row_id"],
                 "uploaded_row": row["_manual_input_row"]}
                for row in new_rows if "_manual_input_row" in row]
            source_policy["funding_usage_sequence_changes"] = sequence_changes
            source_policy["funding_usage_next_sequence"] = len(new_rows) + 1
            overrides = {str(old_to_new[int(sequence)]): value for sequence, value in overrides.items()
                         if int(sequence) in old_to_new}
            for row in new_rows:
                row.pop("_manual_input_row", None)
                row.pop("_sheet_row_number", None)
            summary = candidate.worksheets[0]
            for index, row in enumerate(new_rows, 2):
                for column, key in enumerate(KEYS, 1):
                    summary.cell(index, column).value = row[key]
            if summary.max_row > len(new_rows) + 1:
                summary.delete_rows(len(new_rows) + 2, summary.max_row - len(new_rows) - 1)
            if new_rows:
                repair_summary_row_styles(candidate, new_rows)
            rebuild_personal_sheets(candidate, new_rows, allow_removed_people=bool(deleted_sequences),
                                    original_rows=old_rows)
            output = BytesIO()
            candidate.save(output)
            return changes, output.getvalue(), new_rows, overrides, source_policy, deleted_sequences
        finally:
            baseline.close()
            candidate.close()

    def _funding_source_policy(self, current, origin, old_rows, new_rows, totals, request_key,
                               stored_ids, upload_sha256):
        """冻结原始来源，删除时只撤销活跃登记，保留不可变版本和删除审计。"""
        from .service import digest

        policy = current.extraction_policy_json or {}
        origin_policy = origin.extraction_policy_json or {}
        entries = {int(item["sequence"]): dict(item)
                   for item in policy.get("funding_usage_entries") or []}
        image_version_id = origin_policy.get("image_document_version_id")
        image_version = self.db.get(DocumentVersion, image_version_id) if image_version_id else None
        if not policy.get("funding_usage_source_active_complete"):
            for evidence in origin_policy.get("funding_usage_evidence") or []:
                sequence = int(evidence["row"])
                if sequence in entries:
                    continue
                if image_version is None:
                    fail("FUNDING_SOURCE_HISTORY_MISSING", "首次图片版本缺失，不能更新来源记录。", 409)
                registration = evidence.get("registration_ref") or evidence.get("name_ref") or {}
                if "start" not in registration or "end" not in registration:
                    fail("FUNDING_SOURCE_HISTORY_MISSING", "首次登记引文缺失，不能更新来源记录。", 409)
                entries[sequence] = {**evidence, "sequence": sequence,
                    "registration_key": digest([image_version.sha256,
                        evidence.get("page_number", 1), registration["start"], registration["end"]]),
                    "image_document_id": image_version.document_id,
                    "image_document_version_id": image_version.id,
                    "image_sha256": image_version.sha256,
                    "image_extraction_run_id": origin_policy.get("image_extraction_run_id"),
                    "name": next((row["name"] for row in old_rows
                                  if row["sequence"] == sequence), None)}
        old_sequences = {row["sequence"] for row in old_rows}
        if not old_sequences.issubset(entries):
            fail("FUNDING_SOURCE_HISTORY_MISSING", "已有登记缺少可定位来源，不能安全回传。", 409)
        for sequence in old_sequences:
            entries[sequence]["source_row_id"] = (stored_ids.get(sequence)
                or entries[sequence].get("source_row_id") or f"legacy:{sequence}")
        for row in new_rows:
            if row["sequence"] >= 0:
                continue
            source_row_id = str(uuid4())
            uploaded_row = row["_manual_input_row"]
            entries[row["sequence"]] = {
                "registration_key": digest(["manual-excel-revision", upload_sha256,
                                            uploaded_row, request_key]),
                "sequence": row["sequence"], "source_row_id": source_row_id,
                "source_kind": "MANUAL_EXCEL_REVISION", "name": row["name"],
                "manual_revision_request_key": request_key,
                "manual_upload_sha256": upload_sha256,
                "manual_upload_sheet": "使用登记汇总",
                "manual_upload_row": uploaded_row,
            }
        active = {row["sequence"] for row in new_rows}
        deleted = [row for row in old_rows if row["sequence"] not in active]
        tombstones = list(policy.get("funding_usage_deleted_entries") or [])
        for row in deleted:
            tombstones.append({**entries[row["sequence"]],
                "deleted_at": utcnow().isoformat(), "deletion_request_key": request_key,
                "deleted_row": {key: _json_value(value) for key, value in row.items()}})
        prior = {}
        for row in new_rows:
            sequence = row["sequence"]
            entry = entries[sequence]
            name = row["name"]
            if name not in prior:
                matches = totals.get("".join(name.split()), [])
                prior[name] = matches[0] if len(matches) == 1 else None
            entry["balance_before"] = float(prior[name]) if prior[name] is not None else None
            entry["balance_after"] = row["balance"]
            entry["status"] = row["status"]
            entry["name"] = name
            entry["current_row"] = {key: _json_value(row[key]) for key in
                                    ("sequence", "name", "registration", "date",
                                     "usage_amount", "balance", "status")}
            prior[name] = Decimal(str(row["balance"])) if row["balance"] is not None else None
        active_entries = [entries[row["sequence"]] for row in new_rows]
        active_by_key = {entry["registration_key"]: entry for entry in active_entries}
        active_initial_evidence = []
        if image_version is not None:
            for original in origin_policy.get("funding_usage_evidence") or []:
                registration = original.get("registration_ref") or original.get("name_ref") or {}
                if "start" not in registration or "end" not in registration:
                    continue
                key = digest([image_version.sha256, original.get("page_number", 1),
                              registration["start"], registration["end"]])
                if key in active_by_key:
                    active_initial_evidence.append({**original,
                        "display_sequence": active_by_key[key]["sequence"]})
        return {"funding_usage_entries": active_entries,
                "funding_usage_history": [entry["registration_key"] for entry in active_entries],
                "funding_usage_deleted_entries": tombstones,
                "funding_usage_next_sequence": len(new_rows) + 1,
                "funding_usage_evidence": active_initial_evidence,
                "funding_usage_source_active_complete": True}

    def _validate_and_diff(self, dataset, base_path: Path, path: Path) -> list[dict]:
        try:
            inspect_xlsx(base_path)
            inspect_xlsx(path)
            with zipfile.ZipFile(path) as archive:
                if any(name.lower().startswith("xl/externallinks/") for name in archive.namelist()):
                    fail("MANUAL_EXCEL_EXTERNAL_LINK", "补全文件不能包含外部工作簿链接。", 422)
            baseline = load_workbook(base_path, read_only=False, data_only=False)
            candidate = load_workbook(path, read_only=False, data_only=False)
        except XlsxSafetyError as exc:
            fail(exc.code, str(exc), 413 if "TOO_LARGE" in exc.code or "TOO_COMPLEX" in exc.code else 422)
        except HTTPException:
            raise
        except Exception:
            fail("MANUAL_EXCEL_INVALID", "Excel 无法安全读取。", 422)
        try:
            if baseline.sheetnames != candidate.sheetnames:
                fail("MANUAL_EXCEL_STRUCTURE_CHANGED", "不能增删或重命名工作表。", 422)
            name = dataset.target_sheet_name or "Sheet1"
            if name not in candidate.sheetnames:
                fail("MANUAL_EXCEL_STRUCTURE_CHANGED", "目标工作表不存在。", 422)
            for old_sheet, new_sheet in zip(baseline.worksheets, candidate.worksheets):
                if {str(item) for item in old_sheet.merged_cells.ranges} != {
                        str(item) for item in new_sheet.merged_cells.ranges}:
                    fail("MANUAL_EXCEL_STRUCTURE_CHANGED", "不能改变工作表合并区域。", 422)
                for coordinate, old_cell in old_sheet._cells.items():
                    new_cell = new_sheet._cells.get(coordinate)
                    if old_cell.value is not None and (new_cell is None or new_cell.value is None):
                        fail("MANUAL_EXCEL_CELL_CLEARED", "不能清空原有非空单元格。", 422,
                             cell=old_cell.coordinate)
                    if old_cell.data_type == "f" and (new_cell is None or new_cell.value != old_cell.value):
                        fail("MANUAL_EXCEL_FORMULA_CHANGED", "不能修改原有公式。", 422,
                             cell=old_cell.coordinate)
                for coordinate, new_cell in new_sheet._cells.items():
                    old_cell = old_sheet._cells.get(coordinate)
                    if new_cell.data_type == "f" and (old_cell is None or old_cell.value != new_cell.value):
                        fail("MANUAL_EXCEL_FORMULA_CHANGED", "不能新增或修改公式。", 422,
                             cell=new_cell.coordinate)
            working = self.datasets._dataset_working_copy(dataset)
            origin_version = (self.db.query(DocumentVersion).filter_by(
                document_id=working.document_id, version_number=1).one_or_none()
                if working else None)
            funding_origin = bool(origin_version and
                (origin_version.extraction_policy_json or {}).get("structured_dataset_export_mode")
                == "funding-usage-v1")
            if funding_origin:
                # 个人页仅是汇总页的镜像：用户可同时改两处，服务端逐格核对，
                # 审计仍以汇总页字段为唯一业务记录，避免一处改动被静默漏记。
                old_summary, new_summary = baseline[name], candidate[name]
                if len(baseline.worksheets) != 1 + len(dataset.base_rows_json or []):
                    fail("FUNDING_PERSONAL_SHEET_REVISION_UNSUPPORTED",
                         "个人页数量与汇总人数不一致。", 422)
                for index, (old_personal, new_personal) in enumerate(zip(
                        baseline.worksheets[1:], candidate.worksheets[1:]), 2):
                    if (old_summary.cell(index, 1).value != new_summary.cell(index, 1).value
                            or old_summary.cell(index, 2).value != new_summary.cell(index, 2).value):
                        fail("FUNDING_PERSONAL_SHEET_REVISION_UNSUPPORTED",
                             "人员序号和姓名不能通过手工回传修改。", 422)
                    date_offset = 1 if any(
                        field.get("key") == "date" for field in dataset.schema_json or []) else 0
                    mirrors = {(2, 2): (index, 3), (2, 3): (index, 4 + date_offset),
                               (3, 3): (index, 5 + date_offset),
                               (2, 4): (index, 6 + date_offset)}
                    if (old_personal.max_row != new_personal.max_row
                            or old_personal.max_column != new_personal.max_column):
                        fail("FUNDING_PERSONAL_SHEET_REVISION_UNSUPPORTED",
                             "个人页不能增删行列。", 422)
                    for (row_number, column_number), old_cell in old_personal._cells.items():
                        new_cell = new_personal.cell(row_number, column_number)
                        if old_cell.value != new_cell.value and (row_number, column_number) not in mirrors:
                            fail("FUNDING_PERSONAL_SHEET_REVISION_UNSUPPORTED",
                                 "个人页只能同步修改汇总页的对应值。", 422)
                    for (row_number, column_number), new_cell in new_personal._cells.items():
                        if ((row_number, column_number) not in mirrors
                                and old_personal.cell(row_number, column_number).value != new_cell.value):
                            fail("FUNDING_PERSONAL_SHEET_REVISION_UNSUPPORTED",
                                 "个人页不能增加未审计的业务内容。", 422)
                    for (row_number, column_number), (summary_row, summary_column) in mirrors.items():
                        if (new_personal.cell(row_number, column_number).value
                                != new_summary.cell(summary_row, summary_column).value):
                            fail("FUNDING_PERSONAL_SHEET_REVISION_UNSUPPORTED",
                                 "个人页与汇总页内容不一致，请同时修改两处。", 422)
            old_sheet = baseline[name]
            new_sheet = candidate[name]
            records = self.datasets._records(dataset)
            left, top, right, bottom = (range_boundaries(dataset.target_layout_json["data_range"])
                if dataset.target_layout_json else
                (1, 2, len(dataset.schema_json), 1 + len(records)))
            if right - left + 1 != len(dataset.schema_json):
                fail("MANUAL_EXCEL_STRUCTURE_CHANGED", "目标表数据列范围与 schema 不一致。", 422)
            if len(records) > bottom - top + 1:
                fail("MANUAL_EXCEL_STRUCTURE_CHANGED", "目标表数据行范围与已登记记录不一致。", 422)
            for (row, column), cell in new_sheet._cells.items():
                if (row >= top + len(records) and left <= column <= right and cell.value is not None
                        and old_sheet.cell(row, column).value is None):
                    fail("MANUAL_EXCEL_NEW_ROW", "新增人员须使用来源追加流程。", 422)
            changes = []
            for index, record in enumerate(records):
                row = top + index
                for offset, spec in enumerate(dataset.schema_json):
                    column = left + offset
                    before = old_sheet.cell(row, column).value
                    after = new_sheet.cell(row, column).value
                    if before == after:
                        continue
                    field_type = spec.get("field_type")
                    if after is not None and field_type in {"integer", "decimal"}:
                        try:
                            number = Decimal(str(after))
                            if isinstance(after, bool) or not number.is_finite() or (
                                    field_type == "integer" and number != number.to_integral_value()):
                                raise InvalidOperation
                        except (InvalidOperation, ValueError):
                            fail("MANUAL_EXCEL_FIELD_TYPE", "补充值不符合目标列的数值类型。", 422,
                                 cell=new_sheet.cell(row, column).coordinate)
                    if after is not None and field_type == "date":
                        try:
                            self.datasets._comparable(spec, after)
                        except HTTPException:
                            fail("MANUAL_EXCEL_FIELD_TYPE", "补充值不符合目标列的日期类型。", 422,
                                 cell=new_sheet.cell(row, column).coordinate)
                    if len(changes) >= 1000:
                        fail("MANUAL_EXCEL_TOO_MANY_CHANGES", "一次最多补充 1000 个现有数据格。", 422)
                    changes.append({"record_id": record["record_id"], "row_number": row,
                                    "field_key": spec["key"], "before": _json_value(before),
                                    "after": _json_value(after)})
            # 标题需与服务器冻结的字段完全一致；只读 schema 解析也会验证模板定位。
            sheet_name, actual_fields, _ = self.datasets._read_excel_schema(
                path, name, WorkbookLayout.model_validate(dataset.target_layout_json)
                if dataset.target_layout_json else None)
            if (sheet_name != name or [field.label for field in actual_fields]
                    != [field["label"] for field in dataset.schema_json]):
                fail("MANUAL_EXCEL_STRUCTURE_CHANGED", "目标表标题或字段已变化。", 422)
            return changes
        finally:
            baseline.close()
            candidate.close()

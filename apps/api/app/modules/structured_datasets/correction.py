"""已发布 Excel 数据行的受控更正，复用正文与版本发布而不重复追加。"""

from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path
from uuid import uuid4
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from sqlalchemy import func

from app.db.models import (DocumentPage, DocumentVersion, StructuredDatasetCorrection, StructuredDatasetExport,
                           StructuredDatasetMember, StructuredDatasetTask, StructuredExtractionRun, utcnow)
from app.modules.structured_extraction.export import _xlsx_worksheet_path
from .assessment import (SourceQuote, assess_field, display_value, validate_approval,
                         validate_quote, validate_scope)
from .correction_schemas import CorrectionCheckout, CorrectionPayload
from .publication import stage_publication
from .service import fail
from .writeback import ExcelWritebackService, writeback_digest
from .writeback_schemas import WritebackConfirm, WritebackPreview
from .xlsx_safety import XlsxSafetyError, inspect_xlsx


def _same(left, right) -> bool:
    """Excel 数值和文本类型可能不同，但空格与非空格不能混同。"""
    return left is right if left is None or right is None else str(left) == str(right)


class ExcelCorrectionService(ExcelWritebackService):
    """更正只接受已发布成员行，追加入口的非空格保护保持原样。"""

    def owned(self, task_id: str, *, lock=False) -> StructuredDatasetTask:
        """独立任务命名空间，不能被旧 append 回传端点接管。"""
        query = self.db.query(StructuredDatasetTask).filter_by(id=task_id, user_id=self.user.id)
        if lock:
            query = query.with_for_update()
        task = query.one_or_none()
        if task is None or not task.request_key.startswith("correction:") or (task.manifest_json or {}).get("mode") != "CORRECTION":
            fail("CORRECTION_TASK_NOT_FOUND", "更正任务不存在或无权访问。", 404)
        return task

    def receipt(self, task: StructuredDatasetTask) -> dict:
        """只返回固定目标、预览状态及已发布版本，不泄漏服务器路径。"""
        base = super().receipt(task)
        snapshot = task.manifest_json or {}
        return {**base, "record_id": snapshot.get("record_id"), "row_number": snapshot.get("row_number"),
                "source_document_id": snapshot.get("source_document_id"),
                "source_document_version_id": snapshot.get("source_document_version_id"),
                "source_extraction_run_id": snapshot.get("source_extraction_run_id"),
                "source_sha256": snapshot.get("source_sha256"),
                "base_values": snapshot.get("base_values"), "field_keys": snapshot.get("field_keys"),
                "mode": "CORRECTION"}

    def checkout(self, request: CorrectionCheckout) -> dict:
        """以稳定记录 ID 和真实行号绑定当前版本，直接复用已归档 PDF 提取。"""
        target = self.datasets.resolve_target(request.filename)
        dataset = self.datasets.dataset(target["dataset_id"], lock=True)
        self.datasets._assert_task_write_access(dataset)
        working = self.datasets._dataset_working_copy(dataset)
        if working is None or not working.current_version_id:
            fail("CORRECTION_TARGET_MISSING", "目标 Excel 没有当前工作副本。")
        version = self.db.get(DocumentVersion, working.current_version_id)
        if version is None or Path(working.filename).suffix.lower() != ".xlsx":
            fail("WRITEBACK_XLSX_REQUIRED", "更正只接受已发布 XLSX 工作副本。", 422)
        path = self.datasets.storage.working_copy_path(version.storage_path)
        if not path.is_file() or self.datasets.storage.sha256_file(path) != version.sha256:
            fail("WRITEBACK_TARGET_MISSING", "当前工作副本缺失或哈希不符。")
        if path.stat().st_size > 32 * 1024 * 1024:
            fail("WRITEBACK_TARGET_TOO_LARGE", "目标 Excel 超过 32 MiB。", 413)
        record = next((item for item in self.datasets._records(dataset) if item["record_id"] == request.record_id), None)
        member_id, _, index = request.record_id.partition(":")
        member = self.db.query(StructuredDatasetMember).filter_by(id=member_id, dataset_id=dataset.id).one_or_none()
        if record is None or member is None or not index.isdigit():
            fail("CORRECTION_RECORD_NOT_FOUND", "只能更正该数据集已发布的来源记录。", 404)
        source_run_id = str(request.source_extraction_run_id) if request.source_extraction_run_id else None
        if source_run_id is None:
            # 默认固定当初追加时使用的完整 WorkBuddy 正文，而非随意选择后来产生的局部运行。
            original_run = self.db.get(StructuredExtractionRun, member.structured_extraction_run_id)
            page = (self.db.query(DocumentPage).filter_by(
                extraction_run_id=original_run.layout_extraction_run_id).first()
                if original_run and original_run.layout_extraction_run_id else None)
            source_run_id = (page.metadata_json or {}).get("source_extraction_run_id") if page else None
        source = self.datasets.source(record["document_id"], version_id=record["document_version_id"],
                                      extraction_run_id=source_run_id, content_origin="WORKBUDDY")
        left = range_boundaries(dataset.target_layout_json["data_range"])[0] if dataset.target_layout_json else 1
        if dataset.target_layout_json:
            _, top, _, bottom = range_boundaries(dataset.target_layout_json["data_range"])
            if not top <= request.row_number <= bottom:
                fail("CORRECTION_ROW_MISMATCH", "行号不在目标模板数据区域。", 422)
        try:
            inspect_xlsx(path)
            book = load_workbook(path, data_only=False, read_only=True)
            sheet = book[dataset.target_sheet_name or "Sheet1"]
            values = {spec["key"]: sheet.cell(request.row_number, left + offset).value
                      for offset, spec in enumerate(dataset.schema_json)}
            book.close()
        except (XlsxSafetyError, KeyError, ValueError, OSError):
            fail("WRITEBACK_XLSX_INVALID", "当前 Excel 不能安全读取。", 422)
        for spec in dataset.schema_json:
            key = spec["key"]
            expected = record["fields"].get(key, {}).get("display_value", record["fields"].get(key, {}).get("value"))
            if spec["label"] in {"序号", "编号"} and left + [item["key"] for item in dataset.schema_json].index(key) == 1:
                continue
            if not _same(values[key], expected):
                fail("CORRECTION_ROW_MISMATCH", "指定行与已发布记录不一致，不能猜测绑定。", 409,
                     field_label=spec["label"])
        snapshot = {
            "mode": "CORRECTION", "dataset_id": dataset.id, "working_copy_id": working.id,
            "document_id": working.document_id, "document_version_id": version.id,
            "sha256": version.sha256, "revision": dataset.revision,
            "filename": working.filename, "sheet_name": dataset.target_sheet_name,
            "layout": dataset.target_layout_json or None,
            "labels": [field["label"] for field in dataset.schema_json],
            "field_keys": [field["key"] for field in dataset.schema_json],
            "record_id": request.record_id, "row_number": request.row_number,
            "base_values": values, "source_document_id": source["document_id"],
            "source_document_version_id": source["document_version_id"],
            "source_extraction_run_id": source["extraction_run_id"],
            "source_sha256": source["source_sha256"],
            "review_mode": request.review_mode,
        }
        existing = self.db.query(StructuredDatasetTask).filter_by(
            user_id=self.user.id, request_key=f"correction:{request.request_key}").one_or_none()
        if existing is not None:
            if existing.manifest_json != snapshot:
                fail("CORRECTION_REQUEST_CONFLICT", "请求键已绑定另一目标或版本。")
            return self.receipt(existing)
        task = StructuredDatasetTask(user_id=self.user.id, request_key=f"correction:{request.request_key}",
            result_digest="0" * 64, review_mode=request.review_mode, status="WRITEBACK_CHECKED_OUT",
            manifest_json=snapshot, payload_json={}, progress_json={}, result_json={}, error_json={})
        self.db.add(task)
        self.db.flush()
        return self.receipt(task)

    def preview(self, task_id: str, request: WritebackPreview) -> dict:
        """按任务冻结的模式登记候选；显式复核时等待用户确认。"""
        if request.extraction_confirmations:
            fail("CORRECTION_SOURCE_ALREADY_ARCHIVED", "更正只复用已归档来源，不接受新的来源确认引用。", 422)
        return super().preview(task_id, request)

    def confirm(self, task_id: str, request: WritebackConfirm) -> dict:
        """沿用同一预览指纹确认，不放宽已发布格覆盖权限。"""
        return super().confirm(task_id, request)

    def _validated_changes(self, dataset, snapshot: dict, payload: CorrectionPayload) -> dict:
        """重用正式来源页，逐字核验改动；仅专用模板允许无章事实。"""
        source_ref = payload.source
        if source_ref.extraction_task_id or (
            str(source_ref.document_id) != snapshot["source_document_id"] or
            str(source_ref.document_version_id) != snapshot["source_document_version_id"] or
            str(source_ref.extraction_run_id) != snapshot["source_extraction_run_id"] or
            source_ref.source_sha256 != snapshot["source_sha256"]
        ):
            fail("CORRECTION_SOURCE_CHANGED", "更正来源与冻结的已归档文档版本不一致。", 422)
        source = self.datasets.source(snapshot["source_document_id"],
            version_id=snapshot["source_document_version_id"],
            extraction_run_id=snapshot["source_extraction_run_id"], content_origin="WORKBUDDY")
        if source["source_sha256"] != snapshot["source_sha256"]:
            fail("CORRECTION_SOURCE_CHANGED", "来源 PDF 内容版本已变化。", 409)
        left = range_boundaries(dataset.target_layout_json["header_range"])[0] if dataset.target_layout_json else 1
        specs = {spec["key"]: (spec, left + offset) for offset, spec in enumerate(dataset.schema_json)}
        checked = {}
        for change in payload.changes:
            if change.field_key not in specs:
                fail("UNKNOWN_FIELD", "更正字段不属于目标表。", 422)
            spec, column = specs[change.field_key]
            if column == 1 and spec["label"] in {"序号", "编号"}:
                fail("CORRECTION_SEQUENCE_PROTECTED", "模板预置序号不可更正。", 422)
            if spec["field_type"] != "string":
                fail("CORRECTION_FIELD_TYPE_UNSUPPORTED", "当前更正只接受专用模板文字列。", 422)
            if not _same(change.before, snapshot["base_values"].get(change.field_key)):
                fail("CORRECTION_BEFORE_CHANGED", "更正前值与冻结目标格不一致。", 409,
                     field_label=spec["label"])
            field = change.field
            is_identity = 2 <= column <= 9
            is_source_fact = column in {10, 11, 13, 15, 16}
            if field.raw_text is None or field.raw_text == "":
                if spec.get("required"):
                    fail("CORRECTION_REQUIRED_FIELD", "必填身份或业务列不能清空。", 422,
                         field_label=spec["label"])
                if change.after is not None or not field.blank_reason:
                    fail("CORRECTION_VALUE_MISMATCH", "清空数据格必须明确留空原因。", 422)
                presentation = {"blank_reason": field.blank_reason}
            else:
                if not field.fragments and not field.evidence_refs:
                    fail("EVIDENCE_SCOPE_MISMATCH", "更正文字必须引用已有 PDF 原文。", 422)
                for ref in [*field.fragments, *field.evidence_refs]:
                    validate_quote(ref, source["source_units"])
                if field.evidence_refs and not field.fragments and not any(
                    field.raw_text in ref.quote for ref in field.evidence_refs):
                    fail("EVIDENCE_SCOPE_MISMATCH", "更正文字未被引用原文支持。", 422)
                if field.source_only_fragments and not is_source_fact:
                    fail("SOURCE_ONLY_NOT_ALLOWED", "无章事实只允许专用模板 J/K/M/O/P。", 422)
                if not is_identity and not field.approval and not field.source_only_fragments:
                    fail("APPROVAL_SEAL_REQUIRED", "业务内容缺少审批印章；不能借更正绕过。", 422)
                presentation = {}
                if field.approval:
                    if field.approval.approval_kind == "HR_ANNUAL_EXCELLENCE" and column != 12:
                        fail("HR_SEAL_SCOPE_MISMATCH", "年度优秀方章只能用于 L 列。", 422)
                    validate_approval(field.approval, source["source_units"])
                    for ref in [ref for ref in field.fragments if ref not in field.source_only_fragments] or (
                            [] if field.source_only_fragments else field.evidence_refs):
                        validate_scope(SourceQuote.model_validate(ref.model_dump(mode="json")),
                                       field.approval, source["source_units"])
                    presentation["approval"] = field.approval.model_dump(mode="json")
                if field.source_only_fragments:
                    presentation["source_only_fragment_indexes"] = [
                        field.fragments.index(ref) for ref in field.source_only_fragments]
                    presentation["source_status"] = (
                        "SOURCE_ONLY" if len(field.source_only_fragments) == len(field.fragments) else "MIXED")
                if field.assessment:
                    if is_identity or spec["field_type"] != "string":
                        fail("CHECKMARK_TEXT_FIELD_REQUIRED", "身份格或非文字格不能生成资格对勾。", 422)
                    presentation.update(assess_field(label=spec["label"], assessment=field.assessment,
                        approval=field.approval, units=source["source_units"]))
                if column == 12 and not field.approval:
                    fail("HR_SEAL_SCOPE_MISMATCH", "L 列考核优秀必须核对人事部门方章。", 422)
                expected = display_value(field.raw_text, presentation)
                if not isinstance(change.after, str) or change.after != expected:
                    fail("CORRECTION_VALUE_MISMATCH", "回传值与已核验的展示值或整列对勾不一致。", 422,
                         field_label=spec["label"])
            checked[change.field_key] = {
                "value": field.raw_text, "raw_text": field.raw_text,
                "display_value": change.after, "assessment": presentation,
                "status": "MISSING" if field.raw_text is None else "EXTRACTED",
                "confidence": field.confidence, "page_number": field.page_number,
                "bbox": {},
            }
        return checked

    def publish(self, task_id: str, *, result_digest: str, file_sha256: str,
                payload: CorrectionPayload, file_path: Path) -> dict:
        """先核对逐格证据与 OOXML 差异，再发布同名新版本及更正事件。"""
        task = self.owned(task_id, lock=True)
        body = payload.model_dump(mode="json")
        if task.status in {"PUBLISHING", "COMPLETED", "NEEDS_REPAIR"}:
            if (task.result_digest != result_digest or
                    (task.progress_json or {}).get("file_sha256") != file_sha256 or
                    writeback_digest(body, file_sha256) != result_digest or
                    not file_path.is_file() or self.datasets.storage.sha256_file(file_path) != file_sha256):
                fail("CORRECTION_PREVIEW_CHANGED", "已发布任务不能用不同候选重放。", 409)
            if task.status == "COMPLETED":
                self.tasks._verify_publication(task)
            return self.receipt(task)
        if task.status != "WRITEBACK_READY" or (task.review_mode == "REVIEW_REQUIRED" and not task.confirmed_at):
            fail("CORRECTION_CONFIRMATION_REQUIRED", "请先展示逐格差异并确认。")
        if (task.result_digest != result_digest or
                (task.progress_json or {}).get("file_sha256") != file_sha256 or
                writeback_digest(body, file_sha256) != result_digest or
                not file_path.is_file() or self.datasets.storage.sha256_file(file_path) != file_sha256):
            fail("CORRECTION_PREVIEW_CHANGED", "候选文件或更正内容与已确认预览不一致。", 422)
        snapshot = task.manifest_json
        if payload.record_id != snapshot["record_id"] or payload.row_number != snapshot["row_number"]:
            fail("CORRECTION_RECORD_CHANGED", "候选指向了不同的记录或数据行。", 422)
        dataset = self.datasets.dataset(snapshot["dataset_id"], lock=True)
        self.datasets._assert_task_write_access(dataset)
        working = self.datasets._dataset_working_copy(dataset)
        current = self.db.get(DocumentVersion, working.current_version_id) if working else None
        if (working is None or working.id != snapshot["working_copy_id"] or current is None or
                current.id != snapshot["document_version_id"] or current.sha256 != snapshot["sha256"] or
                dataset.revision != snapshot["revision"] or working.filename != snapshot["filename"]):
            fail("WRITEBACK_VERSION_CONFLICT", "目标 Excel 已变化，请重新 checkout。", 409)
        base_path = self.datasets.storage.working_copy_path(current.storage_path)
        if not base_path.is_file() or self.datasets.storage.sha256_file(base_path) != current.sha256:
            fail("WRITEBACK_TARGET_MISSING", "目标基线缺失或损坏。")
        self._verify_correction_workbook(base_path, file_path, dataset, snapshot, payload)
        try:
            checked = self._validated_changes(dataset, snapshot, payload)
        except ValueError as exc:
            code, _, message = str(exc).partition(":")
            fail(code if message else "CORRECTION_EVIDENCE_INVALID", message.strip() or str(exc), 422)
        dataset.active_task_id = task.id
        dataset.revision += 1
        self.db.add(StructuredDatasetCorrection(dataset_id=dataset.id, task_id=task.id,
            record_id=payload.record_id, row_number=payload.row_number,
            dataset_revision=dataset.revision, changes_json=checked))
        self.db.flush()
        export_id = str(uuid4())
        version_number = int(self.db.query(func.max(DocumentVersion.version_number)).filter(
            DocumentVersion.document_id == working.document_id).scalar() or 0) + 1
        storage_path = self.datasets.storage.dataset_working_version_relative_path(
            dataset_id=dataset.id, export_id=export_id, version_number=version_number)
        self.datasets.storage.import_working_copy(source=file_path, relative_path=storage_path,
                                                  expected_sha256=file_sha256)
        working, version = self.datasets._publish_working_version(dataset=dataset, working_copy=working,
            parent_version=current, storage_path=storage_path, version_number=version_number,
            sha256=file_sha256, size_bytes=file_path.stat().st_size)
        version.extraction_policy_json = {**dict(version.extraction_policy_json or {}),
            "structured_dataset_export_mode": "workbuddy-file-correction-v1"}
        export = StructuredDatasetExport(id=export_id, dataset_id=dataset.id,
            dataset_revision=dataset.revision, storage_path=storage_path, sha256=file_sha256,
            size_bytes=file_path.stat().st_size, document_version_id=version.id)
        self.db.add(export)
        result = {**self.datasets._export_result(export, reused=False),
                  "corrected_record_id": payload.record_id, "corrected_cell_count": len(checked),
                  "task_record_count": 1}
        publication = stage_publication(self.db, dataset=dataset, working_copy=working,
            version=version, base_version_id=current.id, task=task,
            result=result, agent_run=self.datasets.agent_run)
        task.status = "PUBLISHING"
        task.updated_at = utcnow()
        return {**self.receipt(task), "_publication_id": publication.id}

    def _verify_correction_workbook(self, base_path: Path, candidate_path: Path,
                                    dataset, snapshot: dict, payload: CorrectionPayload) -> None:
        """只允许冻结行的指定值变化；所有未授权 XML 格和 ZIP 部件原样保留。"""
        try:
            inspect_xlsx(base_path)
            inspect_xlsx(candidate_path)
            base_book = load_workbook(base_path, data_only=False, read_only=True)
            result_book = load_workbook(candidate_path, data_only=False, read_only=True)
        except XlsxSafetyError as exc:
            fail(exc.code, str(exc), 413 if "TOO_LARGE" in exc.code or "TOO_COMPLEX" in exc.code else 422)
        except (OSError, ValueError, zipfile.BadZipFile):
            fail("WRITEBACK_XLSX_INVALID", "更正候选 Excel 无法安全读取。", 422)
        try:
            if base_book.sheetnames != result_book.sheetnames:
                fail("WRITEBACK_TEMPLATE_CHANGED", "更正不能更改工作表清单。", 422)
            name = dataset.target_sheet_name or "Sheet1"
            left = range_boundaries(dataset.target_layout_json["header_range"])[0] if dataset.target_layout_json else 1
            offsets = {spec["key"]: left + index for index, spec in enumerate(dataset.schema_json)}
            allowed = {}
            for change in payload.changes:
                coordinate = f"{get_column_letter(offsets[change.field_key])}{payload.row_number}"
                before = base_book[name][coordinate]
                after = result_book[name][coordinate]
                if (not _same(before.value, change.before) or not _same(after.value, change.after)
                        or before.data_type == "f" or after.data_type == "f"):
                    fail("CORRECTION_CELL_MISMATCH", "更正格值、公式或样式与预览不一致。", 422, cell=coordinate)
                allowed[coordinate] = change
        finally:
            base_book.close()
            result_book.close()
        try:
            with zipfile.ZipFile(base_path) as old_zip, zipfile.ZipFile(candidate_path) as new_zip:
                names = set(old_zip.namelist())
                if names != set(new_zip.namelist()):
                    fail("WRITEBACK_TEMPLATE_CHANGED", "更正文件增删了 OOXML 部件。", 422)
                lookup = {name: b"" for name in names}
                lookup["xl/workbook.xml"] = old_zip.read("xl/workbook.xml")
                lookup["xl/_rels/workbook.xml.rels"] = old_zip.read("xl/_rels/workbook.xml.rels")
                sheet_path = _xlsx_worksheet_path(lookup, dataset.target_sheet_name or "Sheet1")
                for name in names - {sheet_path}:
                    digests = []
                    for archive in (old_zip, new_zip):
                        sha = hashlib.sha256()
                        with archive.open(name) as stream:
                            while chunk := stream.read(1024 * 1024):
                                sha.update(chunk)
                        digests.append(sha.digest())
                    if digests[0] != digests[1]:
                        fail("WRITEBACK_TEMPLATE_CHANGED", "更正改动了非目标表或工作簿部件。", 422)
                old_root = ET.fromstring(old_zip.read(sheet_path))
                new_root = ET.fromstring(new_zip.read(sheet_path))
        except (OSError, zipfile.BadZipFile, KeyError, ValueError, ET.ParseError):
            fail("WRITEBACK_XLSX_INVALID", "更正 Excel 包损坏。", 422)
        local = lambda tag: tag.rsplit("}", 1)[-1]
        if old_root.attrib != new_root.attrib or [local(x.tag) for x in old_root] != [local(x.tag) for x in new_root]:
            fail("WRITEBACK_TEMPLATE_CHANGED", "更正改变了目标表结构。", 422)
        for old_node, new_node in zip(old_root, new_root):
            if local(old_node.tag) != "sheetData":
                if ET.tostring(old_node) != ET.tostring(new_node):
                    fail("WRITEBACK_TEMPLATE_CHANGED", "更正改变了目标表非数据区域。", 422)
                continue
            old_rows = {row.attrib.get("r"): row for row in old_node}
            new_rows = {row.attrib.get("r"): row for row in new_node}
            if old_rows.keys() != new_rows.keys():
                fail("WRITEBACK_TEMPLATE_CHANGED", "更正增删了工作表行。", 422)
            for number, old_row in old_rows.items():
                new_row = new_rows[number]
                if old_row.attrib != new_row.attrib:
                    fail("WRITEBACK_TEMPLATE_CHANGED", "更正改变了行属性。", 422)
                if ([ET.tostring(child) for child in old_row if local(child.tag) != "c"] !=
                        [ET.tostring(child) for child in new_row if local(child.tag) != "c"]):
                    fail("WRITEBACK_TEMPLATE_CHANGED", "更正改变了行内非单元格结构。", 422)
                old_cells = {cell.attrib.get("r"): cell for cell in old_row if local(cell.tag) == "c"}
                new_cells = {cell.attrib.get("r"): cell for cell in new_row if local(cell.tag) == "c"}
                if (not old_cells.keys() <= new_cells.keys()
                        or any(key not in allowed for key in new_cells.keys() - old_cells.keys())
                        or len(old_cells) != sum(local(child.tag) == "c" for child in old_row)
                        or len(new_cells) != sum(local(child.tag) == "c" for child in new_row)):
                    fail("WRITEBACK_TEMPLATE_CHANGED", "更正增删了单元格或行内结构。", 422)
                for coordinate, old_cell in old_cells.items():
                    new_cell = new_cells[coordinate]
                    if coordinate not in allowed and ET.tostring(old_cell) != ET.tostring(new_cell):
                        fail("WRITEBACK_TEMPLATE_CHANGED", "更正改变了未授权单元格。", 422, cell=coordinate)
                    if coordinate in allowed:
                        before_attrs = {key: value for key, value in old_cell.attrib.items() if key != "t"}
                        after_attrs = {key: value for key, value in new_cell.attrib.items() if key != "t"}
                        before_extra = [ET.tostring(child) for child in old_cell
                            if local(child.tag) not in {"f", "v", "is"}]
                        after_extra = [ET.tostring(child) for child in new_cell
                            if local(child.tag) not in {"f", "v", "is"}]
                        if before_attrs != after_attrs or before_extra != after_extra:
                            fail("WRITEBACK_TEMPLATE_CHANGED", "更正改变了目标格样式或附加结构。", 422,
                                 cell=coordinate)
                for coordinate in new_cells.keys() - old_cells.keys():
                    new_cell = new_cells[coordinate]
                    if set(new_cell.attrib) - {"r", "t"}:
                        fail("WRITEBACK_TEMPLATE_CHANGED", "更正给新增数据格加入了未授权属性。", 422,
                             cell=coordinate)

"""可恢复累计表任务；每次白名单调用提交一个阶段，避免确认后遗漏发布。"""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter

from fastapi import HTTPException

from app.db.models import (StructuredDatasetTask, ExternalExtractionTask, IngestItem, IngestBatch,
    WorkingCopy, DocumentVersion, DocumentExtractionRun, DocumentExtractionConfirmation,
    StructuredDatasetMember, FilesystemJob, utcnow)
from .schemas import DatasetAppend, DatasetCreateFromExcel
from .task_schemas import TaskPreview, TaskConfirm, TaskCommit, TaskPayload, TaskManifest
from .assessment import SourceQuote, validate_quote, validate_approval_claim
from .service import StructuredDatasetService, fail


def mark_target_preparation_failed(db, job, *, code: str, message: str, needs_input: bool) -> None:
    """异步转换终态同步写回父任务，不能让预览任务永久停在准备中。"""

    task_id = str((job.payload_json or {}).get("structured_dataset_task_id") or "")
    row = db.get(StructuredDatasetTask, task_id) if task_id else None
    if row is None:
        # 兼容升级前已排队的任务：仅按固定 job ID 关联，不按文件名猜测。
        row = next((candidate for candidate in db.query(StructuredDatasetTask).filter_by(status="PREPARING_TARGET")
                    if (candidate.progress_json or {}).get("filesystem_job_id") == job.id), None)
    if row is None or row.status in {"COMPLETED", "FAILED", "NEEDS_INPUT"}:
        return
    row.status = "NEEDS_INPUT" if needs_input else "FAILED"
    row.error_json = {"code": code, "message": message, "filesystem_job_id": job.id}
    row.updated_at = utcnow()


def payload_digest(payload: dict) -> str:
    """与宿主规范 JSON 一致的确认摘要，传输租约不属于表格候选。"""
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True,
        separators=(",", ":")).encode("utf-8")).hexdigest()


class DatasetTaskService:
    """任务按发起用户隔离，共享目标访问仍由原数据集服务校验。"""

    def __init__(self, db, user, agent_run=None):
        """依赖仅在本次白名单调用中使用，不存入业务状态。"""
        self.db, self.user = db, user
        self.datasets = StructuredDatasetService(db, user, agent_run)

    def owned(self, task_id, *, lock=False):
        """禁止读取其他人的预览、确认和任务候选。"""
        query = self.db.query(StructuredDatasetTask).filter_by(id=str(task_id), user_id=self.user.id)
        row = (query.with_for_update() if lock else query).one_or_none()
        if row is None:
            fail("DATASET_TASK_NOT_FOUND", "表格任务不存在。", 404)
        return row

    def _extraction(self, task_id):
        """仅解析当前用户的固定提取任务，不暴露其他人的导入上下文。"""
        task = self.db.get(ExternalExtractionTask, str(task_id))
        item = self.db.get(IngestItem, task.ingest_item_id) if task else None
        batch = self.db.get(IngestBatch, item.batch_id) if item else None
        if batch is None or batch.user_id != self.user.id:
            fail("EXTRACTION_TASK_NOT_FOUND", "来源提取任务不可访问。", 404)
        return task, item, batch

    def preview(self, request: TaskPreview):
        """只登记有界元数据和摘要，不把未确认字段正文写入服务器。"""
        existing = self.db.query(StructuredDatasetTask).filter_by(
            user_id=self.user.id, request_key=request.request_key).with_for_update().one_or_none()
        manifest = request.manifest.model_dump(mode="json")
        if existing:
            if (existing.result_digest != request.result_digest or existing.manifest_json != manifest
                    or existing.review_mode != request.review_mode):
                fail("IDEMPOTENCY_CONFLICT", "同一任务键对应的预览已经不同，请生成新预览。")
            return self.receipt(existing)
        target = request.manifest.target
        if target.dataset_id:
            current = self.datasets.resolve_target(target.filename)
            if (current["dataset_id"] != str(target.dataset_id)
                    or current["document_version_id"] != str(target.document_version_id)):
                fail("TARGET_EXCEL_VERSION_CONFLICT", "目标版本已改变，请重新预览。")
        snapshots = {}
        task_ids = [s.extraction_task_id for s in request.manifest.sources if s.extraction_task_id]
        if target.uploaded_task_id:
            task_ids.append(target.uploaded_task_id)
        for source in request.manifest.sources:
            if not source.extraction_task_id:
                self._source(source)
        for task_id in task_ids:
            extraction, _, batch = self._extraction(task_id)
            if extraction.provider_contract_version != "workbuddy-document-v1":
                fail("WORKBUDDY_CONTENT_REQUIRED", "表格任务仅接受 WorkBuddy 文档提取。")
            snapshots[str(task_id)] = {"version_id": extraction.source_version_id, "sha256": extraction.source_sha256}
            if request.review_mode == "DIRECT_SAVE" and (batch.policy_json or {}).get("extraction_save_policy") != "DIRECT_ARCHIVE":
                fail("EXTRACTION_CONFIRMATION_REQUIRED", "来源要求复核，不能改走直接保存。")
        for source in request.manifest.sources:
            if source.extraction_task_id and snapshots[str(source.extraction_task_id)]["sha256"] != source.source_sha256:
                fail("SOURCE_SNAPSHOT_CHANGED", "来源哈希与提取任务不符。")
        for ref in request.manifest.extraction_confirmations:
            self._extraction(ref.task_id)
            confirmation = self.db.get(DocumentExtractionConfirmation, str(ref.confirmation_id))
            if (confirmation is None or confirmation.user_id != self.user.id or confirmation.task_id != str(ref.task_id)
                    or confirmation.result_digest != ref.result_digest):
                fail("EXTRACTION_CONFIRMATION_STALE", "文档预览引用与本轮来源不符。")
        row = StructuredDatasetTask(user_id=self.user.id, request_key=request.request_key,
            result_digest=request.result_digest, review_mode=request.review_mode,
            status="WAITING_CONFIRMATION" if request.review_mode == "REVIEW_REQUIRED" else "READY",
            manifest_json=manifest, progress_json={"source_snapshots": snapshots})
        self.db.add(row)
        self.db.flush()
        return self.receipt(row)

    def confirm(self, task_id, request: TaskConfirm):
        """确认固定表格与文档指纹，不接收新的字段值。"""
        row = self.owned(task_id, lock=True)
        self._digest(row, request.result_digest)
        if row.status == "WAITING_CONFIRMATION":
            from app.modules.external_extraction.document_service import WorkBuddyDocumentExtractionService
            service = WorkBuddyDocumentExtractionService(self.db)
            for ref in TaskManifest.model_validate(row.manifest_json).extraction_confirmations:
                service.confirm(task_id=str(ref.task_id), confirmation_id=str(ref.confirmation_id),
                    result_digest=ref.result_digest, current_user=self.user)
            row.confirmed_at = utcnow()
            row.status = "READY"
        return self.receipt(row)

    def _digest(self, row, value):
        """任何阶段都固定同一用户已查看的指纹。"""
        if value != row.result_digest:
            fail("DATASET_PREVIEW_CHANGED", "表格预览内容已变化，必须重新预览确认。")

    def _working(self, task_id, snapshots):
        """只从后端导入映射解析已发布工作副本，等待归档时不猜测 ID。"""
        task, item, batch = self._extraction(task_id)
        expected = snapshots.get(str(task_id))
        if expected and expected != {"version_id": task.source_version_id, "sha256": task.source_sha256}:
            fail("SOURCE_SNAPSHOT_CHANGED", "提取来源版本已改变。")
        if task.status == "COMPLETE" and not item.final_working_copy_id:
            from app.modules.ingestion.workflow import IngestionWorkflow

            # worker 已发布副本而条目投影尚未刷新时，沿正式工作流同步；不能让表格提交永久等旧状态。
            IngestionWorkflow(self.db).synchronize_batch(batch)
        if task.status in {"FAILED", "CANCELLED"} or item.status in {"FAILED", "CANCELLED"}:
            fail("DATASET_SOURCE_FAILED", "来源归档或提取失败，请先修复来源。")
        working = self.db.get(WorkingCopy, item.final_working_copy_id) if item.final_working_copy_id else None
        if task.status != "COMPLETE" or not working or working.status != "ACTIVE":
            return None
        version = self.db.get(DocumentVersion, working.current_version_id)
        if version is None or version.sha256 != task.source_sha256:
            fail("SOURCE_SNAPSHOT_CHANGED", "工作副本已不是已确认的来源内容。")
        run = self.db.query(DocumentExtractionRun).filter_by(document_id=working.document_id,
            document_version_id=version.id, status="COMPLETED", extractor="workbuddy-external-document").order_by(
                DocumentExtractionRun.created_at.desc()).first()
        if not run:
            return None
        return {"document_id": working.document_id, "document_version_id": version.id,
                "extraction_run_id": run.id, "source_sha256": version.sha256}

    def _source(self, source, snapshots=None):
        """复用正式 WorkBuddy 正文，不触发 OCR 或解析。"""
        binding = (self._working(source.extraction_task_id, snapshots or {}) if source.extraction_task_id
                   else source.model_dump(mode="json", exclude={"extraction_task_id"}))
        if binding is None:
            return None
        run = self.db.get(DocumentExtractionRun, binding["extraction_run_id"])
        origin = "WORKBUDDY_SCOPED" if run is not None and run.extractor == "workbuddy-scoped-pdf" else "WORKBUDDY"
        result = self.datasets.source(binding["document_id"], version_id=binding["document_version_id"],
            extraction_run_id=binding["extraction_run_id"], content_origin=origin)
        if result["source_sha256"] != source.source_sha256:
            fail("SOURCE_SNAPSHOT_CHANGED", "来源内容与预览哈希不一致。")
        return result

    def commit(self, task_id, request: TaskCommit):
        """执行一个可提交阶段；宿主短循环推进，异步依赖未好时原任务等待。"""
        row = self.owned(task_id, lock=True)
        self._digest(row, request.result_digest)
        if row.status in {"FAILED", "NEEDS_INPUT", "NEEDS_REPAIR"}:
            return self.receipt(row)
        if row.status == "WAITING_CONFIRMATION":
            fail("DATASET_CONFIRMATION_REQUIRED", "请先展示逐列预览并取得用户确认。")
        if request.payload:
            payload = request.payload.model_dump(mode="json")
            if payload_digest(payload) != row.result_digest or payload["manifest"] != row.manifest_json:
                fail("DATASET_PREVIEW_CHANGED", "提交字段与已确认预览不一致。")
            row.payload_json = payload
        if not row.payload_json:
            fail("DATASET_DRAFT_REQUIRED", "首次提交需要已确认的宿主草稿。")
        if row.status == "COMPLETED":
            self._verify_publication(row)
            return self.receipt(row)
        if row.status == "PUBLISHING":
            from app.modules.structured_datasets.publication import pending_publication
            dataset_id = (row.progress_json or {}).get("dataset_id")
            pending = pending_publication(self.db, dataset_id=dataset_id) if dataset_id else None
            if pending is not None and pending.task_id == row.id:
                return self.receipt(row)
        payload = TaskPayload.model_validate(row.payload_json)
        self.datasets.workflow_task_id = row.id
        # 例外只在该任务内按真实物理列映射到本次数据集字段 key，旧 append 接口不继承权限。
        from openpyxl.utils.cell import get_column_letter, range_boundaries
        left = range_boundaries(payload.manifest.target.layout["header_range"])[0] if payload.manifest.target.layout else 1
        for record_index, record in enumerate(payload.records, 1):
            for offset, field in enumerate(record.fields):
                if field.approval is None:
                    continue
                try:
                    validate_approval_claim(field.approval)
                except ValueError as exc:
                    code, _, message = str(exc).partition(":")
                    fail(code, message.strip(), 422, record_index=record_index,
                         field_label=payload.manifest.target.labels[offset],
                         column=get_column_letter(left + offset))
        self.datasets.source_fact_keys = set()
        self.datasets.hr_annual_key = None
        progress = dict(row.progress_json)
        row.error_json = {}
        if not progress.get("dataset_id"):
            result = self._prepare_target(row, payload)
            progress = dict(row.progress_json)
            if result is None:
                row.status = "WAITING_DEPENDENCY"
                return self.receipt(row)
            if result.get("kind") == "terminal":
                return self.receipt(row)
            if result.get("kind") == "filesystem_job":
                row.status = "PREPARING_TARGET"
                progress["filesystem_job_id"] = result["filesystem_job_id"]
                row.progress_json = progress
                return {**self.receipt(row), "kind": "filesystem_job", "filesystem_job_id": result["filesystem_job_id"]}
            self._check_target(payload, result)
            target_dataset = self.datasets.dataset(result["dataset_id"], lock=True)
            self.datasets._assert_task_write_access(target_dataset)
            target_dataset.active_task_id = row.id
            progress.update(dataset_id=result["dataset_id"], target_binding={key: result[key] for key in
                ("working_copy_id", "document_id", "document_version_id")}, next_source=0, revisions=[])
            row.progress_json = progress
            row.status = "APPENDING"
            return {**self.receipt(row),
                    **({"_publication_id": result["_publication_id"]} if result.get("_publication_id") else {}),
                    **({"_visible_cleanup": result["_visible_cleanup"]} if result.get("_visible_cleanup") else {})}
        dataset = self.datasets.dataset(progress["dataset_id"], lock=True)
        self.datasets.source_fact_keys = {
            spec["key"] for index, spec in enumerate(dataset.schema_json)
            if payload.manifest.source_fact_column_range == "J:P" and left + index in {10, 11, 13, 15, 16}
        }
        self.datasets.hr_annual_key = next((spec["key"] for index, spec in enumerate(dataset.schema_json)
            if left + index == 12 and payload.manifest.sealed_only), None)
        self._check_target(payload, self.datasets.get(dataset.id))
        source_index = progress.get("next_source", 0)
        self.datasets._assert_task_write_access(dataset)
        dataset.active_task_id = row.id
        if source_index < len(payload.manifest.sources):
            # 先等齐所有来源，再在同一事务追加；某一来源校验失败不能留下半批占用。
            sources = [self._source(source, progress.get("source_snapshots"))
                       for source in payload.manifest.sources[source_index:]]
            if any(source is None for source in sources):
                row.status = "WAITING_DEPENDENCY"
                return self.receipt(row)
            for source_index, source in enumerate(sources, source_index):
                binding = progress["target_binding"]
                records = []
                for candidate in (r for r in payload.records if r.source_index == source_index):
                    fields = {}
                    for spec, field in zip(dataset.schema_json, candidate.fields):
                        value = field.model_dump(mode="json")
                        try:
                            value["evidence_refs"] = [validate_quote(ref, source["source_units"]) for ref in field.evidence_refs]
                        except ValueError as exc:
                            fail("EVIDENCE_SCOPE_MISMATCH", str(exc), 422)
                        fields[spec["key"]] = value
                    records.append({"fields": fields})
                if not records or not any(f.get("raw_text") for r in records for f in r["fields"].values()):
                    fail("DATASET_NO_WRITABLE_RECORDS", "本轮来源没有可写内容，不能报告完成。")
                request_data = {key: source[key] for key in ("document_id", "document_version_id", "extraction_run_id", "source_sha256")}
                request_data.update(records=records, content_origin="WORKBUDDY", idempotency_key=f"task-{row.id}-{source_index}")
                if binding["working_copy_id"]:
                    request_data.update({f"target_{key}": value for key, value in binding.items()})
                result = self.datasets.append(dataset.id, DatasetAppend.model_validate(request_data))
                members = self.datasets._records(dataset)
                written = [r for r in members if r["dataset_revision"] == result["appended_revision"]]
                if not written or any(r["needs_review"] for r in written):
                    fail("DATASET_RECORDS_NEED_REVIEW", "本轮候选存在无法写入的字段，请重新预览修正。")
                progress["revisions"] = [*progress["revisions"], result["appended_revision"]]
                progress["next_source"] = source_index + 1
                row.progress_json = progress
            row.status = "PUBLISHING"
            return self.receipt(row)
        # 导出必须仍作用于注册时绑定的版本，防止两个任务确认后交叉覆盖。
        current = self.datasets.get(dataset.id)
        if current["document_version_id"] != progress["target_binding"]["document_version_id"]:
            fail("TARGET_EXCEL_VERSION_CONFLICT", "发布前目标已被其他任务更新。")
        records = [r for r in self.datasets._records(dataset) if r["dataset_revision"] in progress["revisions"] and not r["needs_review"]]
        if not records:
            fail("DATASET_NO_WRITABLE_RECORDS", "本轮没有可发布记录。")
        exported = self.datasets.export(dataset.id)
        cleanup = exported.pop("_visible_cleanup", None)
        publication_id = exported.pop("_publication_id", None)
        final_result = {**exported, "task_record_count": len(records),
            "task_cell_count": sum(bool(f.get("value") is not None) for r in records for f in r["fields"].values()),
            "checkmark_count": sum(bool(f.get("assessment", {}).get("checkmark")) for r in records for f in r["fields"].values())}
        if publication_id:
            from app.db.models import StructuredDatasetPublication
            publication = self.db.get(StructuredDatasetPublication, publication_id)
            publication.task_id = row.id
            publication.result_json = final_result
            row.status = "PUBLISHING"
            row.result_json = {}
            return {**self.receipt(row), "_publication_id": publication_id}
        row.result_json = {**final_result, "excel_published": True}
        row.status = "COMPLETED"
        self._verify_publication(row)
        dataset.active_task_id = None
        return self.receipt(row)

    def _prepare_target(self, row, payload):
        """三种目标策略共用既有注册、格式提升与版本校验。"""
        target = payload.manifest.target
        if target.new:
            return self.datasets.create(target.new.model_copy(update={"idempotency_key": f"task-{row.id}", "filename": target.filename}))
        if target.dataset_id:
            current = self.datasets.resolve_target(target.filename)
            if current["dataset_id"] != str(target.dataset_id) or current["document_version_id"] != str(target.document_version_id):
                fail("TARGET_EXCEL_VERSION_CONFLICT", "目标已变化，请重新预览。")
            return current
        progress = dict(row.progress_json)
        if progress.get("filesystem_job_id"):
            job = self.db.get(FilesystemJob, progress["filesystem_job_id"])
            if not job or job.status == "FAILED":
                fail("DATASET_TARGET_PREPARATION_FAILED", "目标转换失败，请用真实表格网格重新预览或检查转换任务。",
                     filesystem_job_id=progress["filesystem_job_id"])
            if job.status != "COMPLETED":
                return {"kind": "filesystem_job", "filesystem_job_id": job.id}
        if progress.get("target_request"):
            request = DatasetCreateFromExcel.model_validate(progress["target_request"])
        elif target.uploaded:
            request = target.uploaded.model_copy(update={"idempotency_key": f"task-{row.id}"})
        else:
            binding = self._working(target.uploaded_task_id, progress.get("source_snapshots", {}))
            if binding is None:
                return None
            request = DatasetCreateFromExcel(document_id=binding["document_id"],
                document_version_id=binding["document_version_id"], extraction_run_id=binding["extraction_run_id"],
                sheet_name=target.sheet_name, layout=target.layout, legacy_target_policy="CONVERT_TO_XLSX",
                idempotency_key=f"task-{row.id}")
        progress["target_request"] = request.model_dump(mode="json")
        row.progress_json = progress
        try:
            return self.datasets.create_from_excel(request)
        except HTTPException as exc:
            detail = exc.detail if isinstance(exc.detail, dict) else {}
            code = str(detail.get("code") or "")
            if code not in {"TARGET_EXCEL_HEADER_CELLS_MISSING", "TARGET_EXCEL_EXTRACTION_INCOMPLETE",
                            "TARGET_EXCEL_CONVERSION_SCHEMA_MISMATCH"}:
                raise
            # 确定性模板证据错误保留在原任务上，不让 MCP 重试同一失败指纹。
            row.status = "NEEDS_INPUT"
            row.error_json = {"code": code, "message": str(detail.get("message") or "请重新读取目标 Excel。")}
            row.updated_at = utcnow()
            return {"kind": "terminal"}

    def _check_target(self, payload, result):
        """模板列顺序、完整标题、布局和名称必须与预览一致。"""
        target = payload.manifest.target
        if (result["filename"] != target.filename
                or not self._target_labels_match(target.labels, [f["label"] for f in result["fields"]], target.layout)
                or (target.sheet_name and result["sheet_name"] != target.sheet_name)
                or (target.layout or None) != (result.get("layout") or None)):
            fail("TARGET_EXCEL_SCHEMA_CONFLICT", "目标字段、布局或文件名与预览不一致。")

    @staticmethod
    def _target_labels_match(expected, actual, layout):
        """只容忍服务器合并表头分隔符与重复叶标题列号；字段文字和物理列仍须逐项一致。"""

        if expected == actual:
            return True
        if not layout or len(expected) != len(actual):
            return False
        from openpyxl.utils.cell import get_column_letter, range_boundaries

        left = range_boundaries(layout["header_range"])[0]
        bases = []
        for index, label in enumerate(actual):
            suffix = f" [{get_column_letter(left + index)}]"
            bases.append(label[:-len(suffix)] if label.endswith(suffix) else label)
        counts = Counter(bases)
        normalized_actual = []
        for label, base in zip(actual, bases):
            # 列号仅在同名叶标题确实重复、且后缀与物理列相符时才可剥离。
            display = base if base != label and counts[base] > 1 else label
            normalized_actual.append(" ".join(re.sub(r"\s+/\s+", " ", display).split()))
        return [" ".join(label.split()) for label in expected] == normalized_actual

    def _verify_publication(self, row):
        """已完成任务也核实历史下载及当前工作副本存在，不把缓存链接当发布事实。"""
        from app.modules.structured_datasets.publication import assert_publication_ready
        from app.db.models import FileObject
        result = row.result_json
        if not result.get("excel_published") or not result.get("task_record_count"):
            fail("DATASET_PUBLICATION_INCOMPLETE", "缺少本轮 Excel 发布事实。")
        self.datasets.download(result["dataset_id"], result["export_id"])
        dataset = self.datasets.dataset(result["dataset_id"])
        assert_publication_ready(self.db, dataset_id=dataset.id)
        working = self.datasets._dataset_working_copy(dataset)
        if working is None:
            fail("DATASET_PUBLICATION_MISSING", "当前工作副本不存在。")
        current = self.db.get(DocumentVersion, working.current_version_id)
        visible = self.datasets.storage.working_copy_path(self.datasets._working_copy_storage_path(working))
        file_object = self.db.query(FileObject).filter_by(
            document_id=working.document_id, storage_backend="working_copy_local").one_or_none()
        if (current is None or not visible.is_file() or self.datasets.storage.sha256_file(visible) != current.sha256
                or working.content_sha256 != current.sha256 or file_object is None
                or file_object.sha256 != current.sha256):
            fail("DATASET_PUBLICATION_INCOMPLETE", "当前 Excel 与数据库版本或同名文件不一致。")
        self.datasets._read_and_validate_current_target_schema(dataset=dataset, working_copy=working)

    def receipt(self, row):
        """返回轻量任务状态及发布回执，不输出候选正文到日志或审计。"""
        status, error = row.status, row.error_json or {}
        job_id = (row.progress_json or {}).get("filesystem_job_id")
        if status == "PREPARING_TARGET" and job_id:
            job = self.db.get(FilesystemJob, job_id)
            if job is None or job.status == "FAILED":
                # 兼容升级前未同步父任务的失败记录；GET 仍必须返回真实终态。
                status = "FAILED"
                error = {"code": "DATASET_TARGET_PREPARATION_FAILED",
                         "message": "目标格式提升失败；请用真实表格网格重新预览或检查转换任务。",
                         "filesystem_job_id": job_id}
        from app.modules.structured_datasets.publication import pending_publication
        dataset_id = (row.progress_json or {}).get("dataset_id")
        pending = pending_publication(self.db, dataset_id=dataset_id) if dataset_id and status == "PUBLISHING" else None
        next_action = ("contact_admin" if status == "NEEDS_REPAIR" else
                       "wait_for_publication" if pending is not None and pending.task_id == row.id else
                       "structured_dataset_confirm" if status == "WAITING_CONFIRMATION" else
                       "create_new_preview" if status in {"FAILED", "NEEDS_INPUT"} else
                       "structured_dataset_commit")
        return {"task_id": row.id, "status": status, "result_digest": row.result_digest,
                "review_mode": row.review_mode, "excel_published": status == "COMPLETED",
                "draft_saved": bool(row.payload_json),
                "next_action": next_action,
                "filesystem_job_id": job_id,
                "error": error, **(row.result_json if status == "COMPLETED" else {})}

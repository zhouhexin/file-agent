"""受控提示词修订：以已发布个人登记为基线，重算余额表并发布同名版本。"""

from __future__ import annotations

import json
import re
import tempfile
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from types import SimpleNamespace
from uuid import NAMESPACE_URL, uuid4, uuid5

from openpyxl import load_workbook
from sqlalchemy import func

from app.db.models import DocumentVersion, StructuredDatasetExport
from app.modules.structured_datasets.publication import stage_publication
from app.modules.structured_datasets.service import fail
from app.modules.structured_datasets.xlsx_safety import XlsxSafetyError, inspect_xlsx

from .schemas import FundingUsagePromptEdit, FundingUsagePromptPublish
from .service import FundingUsageService, _detail_matches_frozen, _excel_number
from .summary_workbook import person_order, read_detail_book, render_book, summaries


def _digest(request: FundingUsagePromptEdit) -> str:
    return sha256(json.dumps(request.model_dump(mode="json", exclude={"preview_digest"}),
        ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()


class FundingPromptEditService(FundingUsageService):
    def _target(self, filename: str, *, lock=False):
        dataset, working, current, origin, path = self._continue_target(
            SimpleNamespace(filename=filename), lock=lock)
        if (origin.extraction_policy_json or {}).get("structured_dataset_export_mode") != "funding-usage-v3":
            fail("FUNDING_PROMPT_TARGET_INVALID", "只有新版个人明细余额表支持提示词修订。", 422)
        policy = current.extraction_policy_json or {}
        rows = list(policy.get("funding_usage_rows_v3") or [])
        funding = {name: dict(value) for name, value in
                   (policy.get("funding_usage_funding_v3") or {}).items()}
        order = person_order(dataset, funding, policy)
        try:
            inspect_xlsx(path)
            book = load_workbook(path, read_only=False, data_only=False)
            try:
                published_rows, published_funding = read_detail_book(
                    book, order, required_names={row["name"] for row in rows})
                expected = summaries(rows, funding, order)
                summary = book.worksheets[0]
                summary_matches = all(
                    [summary.cell(i + 2, col).value for col in range(3, 7)]
                    == [person[key] for key in ("research_total", "other_total", "total_spent", "remaining")]
                    for i, person in enumerate(expected))
            finally:
                book.close()
        except (OSError, ValueError, XlsxSafetyError):
            fail("FUNDING_TARGET_INVALID", "目标余额表无法安全读取。", 422)
        frozen = {(row["name"], row["sequence"]): row for row in rows}
        row_ids = [row.get("_funding_row_id") for row in rows]
        if (published_funding != funding or not summary_matches
                or len(rows) != len(published_rows) or len(frozen) != len(rows)
                or not all(row_ids) or len(set(row_ids)) != len(row_ids)
                or any((row["name"], row["sequence"]) not in frozen
                       or not _detail_matches_frozen(
                           row, frozen[(row["name"], row["sequence"])])
                       for row in published_rows)):
            fail("FUNDING_TARGET_CHANGED", "服务器已发布 Excel 与冻结登记不一致。", 409)
        entries = {item.get("source_row_id"): dict(item)
                   for item in policy.get("funding_usage_entries") or []}
        if (len(entries) != len(rows) or set(entries) != set(row_ids)
                or any(not item.get("registration_key") for item in entries.values())):
            fail("FUNDING_SOURCE_HISTORY_MISSING", "已有登记缺少完整来源记录。", 409)
        return dataset, working, current, policy, rows, funding, order, entries

    def checkout(self, filename: str):
        dataset, working, current, _, rows, funding, order, _ = self._target(filename)
        return {"filename": filename, "dataset_id": dataset.id,
                "base_document_version_id": current.id, "base_sha256": current.sha256,
                "base_revision": dataset.revision, "working_copy_id": working.id,
                "people": summaries(rows, funding, order),
                "rows": [{key: row.get(key) for key in (
                    "_funding_row_id", "sequence", "name", "registration", "date",
                    "usage_amount", "status")} for row in rows]}

    def _evaluate(self, request: FundingUsagePromptEdit, *, lock=False):
        dataset, working, current, policy, old_rows, funding, order, old_entries = self._target(
            request.filename, lock=lock)
        if (current.id != str(request.base_document_version_id)
                or current.sha256 != request.base_sha256
                or dataset.revision != request.base_revision):
            fail("FUNDING_TARGET_VERSION_CONFLICT", "余额表已有新版本，请重新读取当前登记。", 409)
        groups = {name: [dict(row) for row in old_rows if row["name"] == name]
                  for name in order}
        by_id = {row["_funding_row_id"]: row for group in groups.values() for row in group}
        entries = {key: dict(value) for key, value in old_entries.items()}
        deleted_entries = list(policy.get("funding_usage_deleted_entries") or [])
        changes = []
        touched_ids = set()
        inserted_after = {}
        for index, edit in enumerate(request.row_edits):
            name = edit.name.strip()
            if name not in groups or name != edit.name:
                fail("FUNDING_PROMPT_PERSON_INVALID", "登记姓名必须与首页已有教师完全一致。", 422)
            if edit.action == "ADD":
                anchor = str(edit.after_row_id) if edit.after_row_id else None
                if anchor is not None:
                    anchor = inserted_after.get(anchor, anchor)
                    if anchor not in by_id or by_id[anchor]["name"] != name:
                        fail("FUNDING_PROMPT_ROW_INVALID", "插入位置不属于该教师的当前明细。", 422)
                row_id = str(uuid5(NAMESPACE_URL,
                    f"funding-prompt:{self.user.id}:{dataset.id}:{request.request_key}:{index}"))
                if row_id in by_id:
                    fail("FUNDING_REQUEST_CONFLICT", "请求键与已有登记冲突。", 409)
                amount = _excel_number(edit.usage_amount)
                row = {"_funding_row_id": row_id, "sequence": None, "name": name,
                       "registration": edit.registration, "date": edit.date,
                       "usage_amount": amount,
                       "status": edit.status or ("已核对" if amount is not None else "金额待核对")}
                if anchor is None:
                    groups[name].append(row)
                else:
                    position = next(i for i, item in enumerate(groups[name])
                                    if item["_funding_row_id"] == anchor)
                    groups[name].insert(position + 1, row)
                    inserted_after[str(edit.after_row_id)] = row_id
                by_id[row_id] = row
                entries[row_id] = {"source_row_id": row_id, "source_kind": "USER_PROMPT",
                    "registration_key": sha256(
                        f"funding-prompt:{request.request_key}:{row_id}".encode()).hexdigest(),
                    "prompt_request_key": request.request_key}
                changes.append({"action": "ADD", "name": name, "source_row_id": row_id,
                                "before": None, "after": dict(row)})
                continue
            row_id = str(edit.source_row_id)
            if row_id in touched_ids or row_id not in by_id or by_id[row_id]["name"] != name:
                fail("FUNDING_PROMPT_ROW_INVALID", "待修改的登记行不存在、归属不符或被重复操作。", 422)
            touched_ids.add(row_id)
            row = by_id[row_id]
            before = dict(row)
            if edit.action == "DELETE":
                groups[name].remove(row)
                del by_id[row_id]
                deleted_entries.append({**entries.pop(row_id),
                    "deleted_at": datetime.now(timezone.utc).isoformat(),
                    "deletion_request_key": request.request_key})
                changes.append({"action": "DELETE", "name": name,
                                "source_row_id": row_id, "before": before, "after": None})
                continue
            for field in ("registration", "date", "usage_amount", "status"):
                if field in edit.model_fields_set:
                    row[field] = (_excel_number(edit.usage_amount)
                                  if field == "usage_amount" else getattr(edit, field))
            if row["registration"] is None:
                fail("FUNDING_PROMPT_ROW_INVALID", "登记原文不能清空。", 422)
            if row["usage_amount"] is None:
                row["status"] = row["status"] or "金额待核对"
                if "待核对" not in row["status"]:
                    row["status"] += "；金额待核对"
            elif "status" not in edit.model_fields_set and (
                    row["status"] is None or "金额待核对" in row["status"]):
                row["status"] = (row["status"] or "").replace("手写金额待核对", "")
                row["status"] = row["status"].replace("金额待核对", "").strip("； ") or "已核对"
            if row != before:
                entries[row_id]["prompt_edit_request_keys"] = [
                    *(entries[row_id].get("prompt_edit_request_keys") or []), request.request_key]
                changes.append({"action": "UPDATE", "name": name,
                                "source_row_id": row_id, "before": before, "after": dict(row)})
        seen_totals = set()
        for edit in request.total_edits:
            name = edit.name.strip()
            key = (name, edit.field)
            if name not in funding or name != edit.name or key in seen_totals:
                fail("FUNDING_PROMPT_TOTAL_INVALID", "资助金额的教师或字段重复、无效。", 422)
            seen_totals.add(key)
            before = funding[name].get(edit.field)
            after = _excel_number(edit.value)
            if before != after:
                funding[name][edit.field] = after
                changes.append({"action": "TOTAL_UPDATE", "name": name,
                    "field": edit.field, "before": before, "after": after})
        renames = {}
        new_names = set()
        for edit in request.name_edits:
            old_name, new_name = edit.old_name, edit.new_name
            if old_name not in funding or old_name != old_name.strip():
                fail("FUNDING_PROMPT_PERSON_INVALID", "更名前姓名必须与首页已有教师完全一致。", 422)
            if (not new_name.strip() or new_name != new_name.strip()
                    or re.search(r'[\x00-\x1f\[\]:*?/\\]', new_name)
                    or new_name == "使用登记汇总"):
                fail("FUNDING_PROMPT_NAME_INVALID", "新姓名不能用作个人工作表名称。", 422)
            compact_new = "".join(new_name.split())
            if (new_name == old_name or old_name in renames or compact_new in new_names
                    or any(compact_new == "".join(name.split())
                           for name in funding if name != old_name)):
                fail("FUNDING_PROMPT_NAME_CONFLICT", "更名重复或新姓名与已有教师冲突。", 422)
            renames[old_name] = new_name
            new_names.add(compact_new)
        for old_name, new_name in renames.items():
            funding[new_name] = funding.pop(old_name)
            groups[new_name] = groups.pop(old_name)
            for row in groups[new_name]:
                row["name"] = new_name
            changes.append({"action": "RENAME", "name": new_name,
                            "before": old_name, "after": new_name})
        order = [renames.get(name, name) for name in order]
        if not changes:
            fail("FUNDING_PROMPT_UNCHANGED", "提示词没有产生需要发布的变更。", 422)
        if len(changes) > 1000:
            fail("FUNDING_PROMPT_TOO_MANY_CHANGES", "单次变更数量过多。", 422)
        normalized = []
        for name in order:
            for sequence, row in enumerate(groups[name], 1):
                row["sequence"] = sequence
                row_id = row["_funding_row_id"]
                entries[row_id].update({"name": name, "sequence": sequence,
                                        "current_row": dict(row)})
                normalized.append(row)
        final_by_id = {row["_funding_row_id"]: row for row in normalized}
        for change in changes:
            if change["action"] in {"ADD", "UPDATE"}:
                change["after"] = dict(final_by_id[change["source_row_id"]])
        new_summary = summaries(normalized, funding, order)
        affected = {item["name"] for item in changes}
        preview = {"preview_digest": _digest(request), "filename": request.filename,
            "added_record_count": sum(item["action"] == "ADD" for item in changes),
            "updated_record_count": sum(item["action"] == "UPDATE" for item in changes),
            "deleted_record_count": sum(item["action"] == "DELETE" for item in changes),
            "total_change_count": sum(item["action"] == "TOTAL_UPDATE" for item in changes),
            "renamed_person_count": len(renames),
            "changes": changes,
            "people": [row for row in new_summary if row["name"] in affected]}
        return dataset, working, current, policy, normalized, funding, order, entries, deleted_entries, new_summary, preview

    def preview(self, request: FundingUsagePromptEdit):
        return self._evaluate(request)[-1]

    def publish(self, request: FundingUsagePromptPublish):
        request_digest = _digest(request)
        if request.preview_digest != request_digest:
            fail("FUNDING_PREVIEW_CHANGED", "提示词变更与预览内容不一致。", 409)
        dataset, working, _, _, _, _, _, _ = self._target(request.filename, lock=True)
        for previous in self.db.query(DocumentVersion).filter_by(document_id=working.document_id).all():
            saved = previous.extraction_policy_json or {}
            if saved.get("funding_usage_prompt_request_key") == request.request_key:
                if saved.get("funding_usage_prompt_digest") != request_digest:
                    fail("FUNDING_REQUEST_CONFLICT", "请求键已用于不同的提示词变更。", 409)
                export = self.db.query(StructuredDatasetExport).filter_by(
                    dataset_id=dataset.id, document_version_id=previous.id).one_or_none()
                if export is None:
                    fail("FUNDING_PUBLICATION_PENDING", "上次提示词变更尚未发布。", 423)
                return {**self.datasets._export_result(export, reused=True),
                        "status": "COMPLETED", "excel_published": True,
                        "added_record_count": saved.get("funding_usage_prompt_added_count", 0),
                        "updated_record_count": saved.get("funding_usage_prompt_updated_count", 0),
                        "deleted_record_count": saved.get("funding_usage_prompt_deleted_count", 0),
                        "total_change_count": saved.get("funding_usage_prompt_total_count", 0),
                        "renamed_person_count": saved.get("funding_usage_prompt_renamed_count", 0)}
        (dataset, working, current, policy, rows, funding, order, entries,
         deleted_entries, summary_rows, preview) = self._evaluate(request, lock=True)
        self.datasets._assert_task_write_access(dataset)
        with tempfile.TemporaryDirectory(prefix="file-agent-funding-prompt-") as directory:
            output = Path(directory) / request.filename
            book = render_book(rows, funding, order=order)
            try:
                book.save(output)
            finally:
                book.close()
            file_sha = self.datasets.storage.sha256_file(output)
            rename_map = {edit.old_name: edit.new_name for edit in request.name_edits}
            previous_ids = {rename_map.get(row["name"], row["name"]): row.get("_funding_row_id")
                            for row in dataset.base_rows_json or []}
            dataset.base_rows_json = [
                {**row, "_funding_row_id": previous_ids.get(row["name"]) or str(uuid4())}
                for row in summary_rows]
            dataset.target_layout_json = {"header_range": "A1:F1",
                "data_range": f"A2:F{max(2, len(summary_rows) + 1)}"}
            dataset.revision += 1
            export_id = str(uuid4())
            version_number = int(self.db.query(func.max(DocumentVersion.version_number)).filter(
                DocumentVersion.document_id == working.document_id).scalar() or 0) + 1
            storage_path = self.datasets.storage.dataset_working_version_relative_path(
                dataset_id=dataset.id, export_id=export_id, version_number=version_number)
            self.datasets.storage.import_working_copy(source=output, relative_path=storage_path,
                                                       expected_sha256=file_sha)
            working, version = self.datasets._publish_working_version(
                dataset=dataset, working_copy=working, parent_version=current,
                storage_path=storage_path, version_number=version_number,
                sha256=file_sha, size_bytes=output.stat().st_size)
            version.extraction_policy_json = {**policy,
                "structured_dataset_export_mode": "funding-usage-v3",
                "funding_usage_rows_v3": rows,
                "funding_usage_funding_v3": funding,
                "funding_usage_person_order_v3": order,
                "funding_usage_entries": list(entries.values()),
                "funding_usage_deleted_entries": deleted_entries,
                "funding_usage_history": [item["registration_key"] for item in entries.values()],
                "funding_usage_prompt_request_key": request.request_key,
                "funding_usage_prompt_digest": request_digest,
                "funding_usage_prompt_added_count": preview["added_record_count"],
                "funding_usage_prompt_updated_count": preview["updated_record_count"],
                "funding_usage_prompt_deleted_count": preview["deleted_record_count"],
                "funding_usage_prompt_total_count": preview["total_change_count"],
                "funding_usage_prompt_renamed_count": preview["renamed_person_count"],
                "funding_usage_prompt_history": [*(policy.get("funding_usage_prompt_history") or []),
                    {"request_key": request.request_key, "user_request": request.user_request,
                     "changes": preview["changes"]}]}
            export = StructuredDatasetExport(id=export_id, dataset_id=dataset.id,
                dataset_revision=dataset.revision, storage_path=storage_path,
                sha256=file_sha, size_bytes=output.stat().st_size,
                document_version_id=version.id)
            self.db.add(export)
            self.db.flush()
            result = {**self.datasets._export_result(export, reused=False),
                "status": "PUBLISHING", "excel_published": False,
                "added_record_count": preview["added_record_count"],
                "updated_record_count": preview["updated_record_count"],
                "deleted_record_count": preview["deleted_record_count"],
                "total_change_count": preview["total_change_count"],
                "renamed_person_count": preview["renamed_person_count"]}
            publication = stage_publication(self.db, dataset=dataset,
                working_copy=working, version=version, base_version_id=current.id,
                result=result, agent_run=self.datasets.agent_run)
        return {**result, "people": preview["people"],
                "changes": preview["changes"], "_publication_id": publication.id}

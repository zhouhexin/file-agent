"""成果资助余额表第三版：个人明细是登记来源，首页只保留人员汇总。"""

from __future__ import annotations

from copy import copy
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.structured_datasets.service import fail
from app.modules.structured_extraction.export import _safe_spreadsheet_text

from .workbook import money


SUMMARY_HEADERS = [
    "序号", "姓名", "总科研成果资助金额", "其他资助金额（院优青、专委、其他资助）",
    "总支出", "剩余总资助金额",
]
DETAIL_HEADERS = ["序号", "使用登记原文", "日期", "确认使用金额（元）", "核对说明"]
TOTAL_LABELS = SUMMARY_HEADERS[2:]
SUMMARY_KEYS = ["sequence", "name", "research_total", "other_total", "total_spent", "remaining"]


def _number(value: Decimal | None):
    return float(value.quantize(Decimal("0.01"))) if value is not None else None


def person_totals(rows: list[dict], research_total, other_total):
    research = money(research_total, cell="总科研成果资助金额")
    other = money(other_total, cell="其他资助金额")
    amounts = [money(row.get("usage_amount"), cell="确认使用金额") for row in rows]
    spent = sum(amounts, Decimal("0.00")) if all(value is not None for value in amounts) else None
    remaining = (research + (other or Decimal("0.00")) - spent
                 if research is not None and spent is not None else None)
    return {"research_total": _number(research), "other_total": _number(other),
            "total_spent": _number(spent), "remaining": _number(remaining)}


def person_order(dataset, funding: dict[str, dict], policy: dict | None = None,
                 *, new_names=()) -> list[str]:
    """从有序汇总行恢复人员顺序；JSONB 对象的键顺序不能用于 Excel 行定位。"""
    base_rows = list(getattr(dataset, "base_rows_json", None) or [])
    sequences = [row.get("sequence") for row in base_rows]
    if (not base_rows or any(not isinstance(value, int) or isinstance(value, bool)
                             for value in sequences)
            or sorted(sequences) != list(range(1, len(base_rows) + 1))):
        fail("FUNDING_PERSON_ORDER_INVALID", "余额表缺少有效的首页人员顺序。", 409)
    ordered = [row.get("name") for row in sorted(base_rows, key=lambda row: row["sequence"])]
    if (any(not isinstance(name, str) or not name for name in ordered)
            or len(ordered) != len(set(ordered))):
        fail("FUNDING_PERSON_ORDER_INVALID", "余额表首页人员姓名或顺序无效。", 409)
    saved = (policy or {}).get("funding_usage_person_order_v3")
    if saved is not None and saved != ordered:
        fail("FUNDING_PERSON_ORDER_INVALID", "余额表保存的人员顺序与首页记录不一致。", 409)
    for name in new_names:
        if name not in ordered:
            ordered.append(name)
    if set(ordered) != set(funding):
        fail("FUNDING_PERSON_ORDER_INVALID", "余额表人员名单与资助记录不一致。", 409)
    return ordered


def summaries(rows: list[dict], funding: dict[str, dict], order=None):
    people = (list(order) if order is not None else
              list(dict.fromkeys([*funding, *(row["name"] for row in rows)])))
    if order is not None and (len(people) != len(set(people))
                              or set(people) != set(funding) | {row["name"] for row in rows}):
        fail("FUNDING_PERSON_ORDER_INVALID", "汇总人员顺序与登记或资助记录不一致。", 409)
    return [{"sequence": index, "name": name,
             **person_totals([row for row in rows if row["name"] == name],
                             funding.get(name, {}).get("research_total"),
                             funding.get(name, {}).get("other_total"))}
            for index, name in enumerate(people, 1)]


def _display(value):
    return _safe_spreadsheet_text(value) if isinstance(value, str) else value


def render_book(rows: list[dict], funding: dict[str, dict], *, style_sheet=None,
                order=None) -> Workbook:
    book = Workbook()
    summary = book.active
    summary.title = "使用登记汇总"
    for col, label in enumerate(SUMMARY_HEADERS, 1):
        cell = summary.cell(1, col, label)
        if style_sheet is not None and style_sheet.cell(1, min(col, style_sheet.max_column)).has_style:
            sample = style_sheet.cell(1, min(col, style_sheet.max_column))
            cell.font, cell.fill, cell.border = copy(sample.font), copy(sample.fill), copy(sample.border)
            cell.alignment = copy(sample.alignment)
        else:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
        summary.column_dimensions[get_column_letter(col)].width = (9, 18, 25, 43, 20, 23)[col - 1]
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    summary.row_dimensions[1].height = 32
    summary.freeze_panes = "A2"
    for summary_row in summaries(rows, funding, order):
        row_number = summary_row["sequence"] + 1
        for col, key in enumerate(SUMMARY_KEYS, 1):
            cell = summary.cell(row_number, col, _display(summary_row[key]))
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if col >= 3:
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
        summary.row_dimensions[row_number].height = 27

    used_titles = {summary.title}
    detail_names = {row["name"] for row in rows}
    for person in summaries(rows, funding, order):
        name = person["name"]
        if name not in detail_names:
            continue
        title = name[:31]
        if not title or any(char in title for char in '[]:*?/\\') or title in used_titles:
            title = f"人员{person['sequence']}"
        used_titles.add(title)
        sheet = book.create_sheet(title)
        sheet.append(DETAIL_HEADERS)
        own_rows = [row for row in rows if row["name"] == name]
        for index, row in enumerate(own_rows, 1):
            sheet.append([index, _display(row.get("registration")), _display(row.get("date")),
                          row.get("usage_amount"), _display(row.get("status"))])
            lines = max(str(row.get("registration") or "").count("\n"),
                        str(row.get("date") or "").count("\n"),
                        str(row.get("status") or "").count("\n")) + 1
            sheet.row_dimensions[index + 1].height = min(180, max(28, 18 * lines))
        footer_start = len(own_rows) + 3
        for offset, (label, key) in enumerate(zip(TOTAL_LABELS, SUMMARY_KEYS[2:])):
            sheet.cell(footer_start + offset, 1, label)
            sheet.cell(footer_start + offset, 2, person[key])
            sheet.cell(footer_start + offset, 1).font = Font(bold=True)
            sheet.cell(footer_start + offset, 2).number_format = '#,##0.00;[Red]-#,##0.00'
        sheet.freeze_panes = "A2"
        for col, width in enumerate((43, 48, 24, 24, 30), 1):
            sheet.column_dimensions[get_column_letter(col)].width = width
            sheet.cell(1, col).font = Font(bold=True)
            sheet.cell(1, col).fill = PatternFill("solid", fgColor="DDEBF7")
            sheet.cell(1, col).alignment = Alignment(wrap_text=True, vertical="center")
        for row_number in range(2, footer_start):
            for col in range(1, 6):
                sheet.cell(row_number, col).alignment = Alignment(vertical="top", wrap_text=True)
            sheet.cell(row_number, 4).number_format = '#,##0.00;[Red]-#,##0.00'
    return book


def read_detail_book(book, expected_names: list[str], *, required_names=None):
    required_names = set(expected_names if required_names is None else required_names)
    if not required_names.issubset(expected_names):
        fail("FUNDING_PERSONAL_SHEET_INVALID", "登记人员不在汇总名单中。", 422)
    compact_names = [name for name in expected_names if name in required_names]
    if (len(book.worksheets) not in {len(compact_names) + 1, len(expected_names) + 1}
            or book.worksheets[0].title != "使用登记汇总"):
        fail("FUNDING_PERSONAL_SHEET_INVALID", "个人工作表数量或汇总页不一致。", 422)
    # 兼容此前已发布的“每人一页”版本；新版本只为有登记的人建个人页。
    detail_names = expected_names if len(book.worksheets) == len(expected_names) + 1 else compact_names
    summary = book.worksheets[0]
    if [summary.cell(1, col).value for col in range(1, 7)] != SUMMARY_HEADERS:
        fail("FUNDING_SUMMARY_INVALID", "汇总页表头与当前版本不一致。", 422)
    result = []
    funding = {}
    detail_index = 1
    for person_index, name in enumerate(expected_names, 1):
        if summary.cell(person_index + 1, 2).value != name:
            fail("FUNDING_ROW_ID_INVALID", "汇总页姓名或人员顺序不能改变。", 422)
        if name not in detail_names:
            research = money(summary.cell(person_index + 1, 3).value,
                             cell=f"使用登记汇总!C{person_index + 1}")
            other = money(summary.cell(person_index + 1, 4).value,
                          cell=f"使用登记汇总!D{person_index + 1}")
            funding[name] = {"research_total": _number(research), "other_total": _number(other)}
            continue
        sheet = book.worksheets[detail_index]
        detail_index += 1
        if [sheet.cell(1, col).value for col in range(1, 6)] != DETAIL_HEADERS:
            fail("FUNDING_PERSONAL_SHEET_INVALID", f"{name} 的明细表头不能改变。", 422)
        footer = next((r for r in range(2, sheet.max_row + 1)
                       if sheet.cell(r, 1).value == TOTAL_LABELS[0]), None)
        if footer is None or [sheet.cell(footer + i, 1).value for i in range(4)] != TOTAL_LABELS:
            fail("FUNDING_PERSONAL_SHEET_INVALID", f"{name} 的金额汇总区域缺失或标签被修改。", 422)
        if (any(sheet.cell(row_number, col).value is not None
                for row_number in range(2, sheet.max_row + 1)
                for col in range(6, sheet.max_column + 1))
                or any(sheet.cell(footer + offset, col).value is not None
                       for offset in range(4) for col in range(3, 6))
                or any(sheet.cell(row_number, col).value is not None
                       for row_number in range(footer + 4, sheet.max_row + 1)
                       for col in range(1, 6))):
            fail("FUNDING_PERSONAL_SHEET_INVALID", f"{name} 的明细或金额汇总区域之外含未受支持的数据。", 422)
        research = money(sheet.cell(footer, 2).value, cell=f"{sheet.title}!B{footer}")
        other = money(sheet.cell(footer + 1, 2).value, cell=f"{sheet.title}!B{footer + 1}")
        if any(value is not None and (value < 0 or value > Decimal("1000000000"))
               for value in (research, other)):
            fail("FUNDING_AMOUNT_INVALID", "资助总额须为非负且不超过十亿元的金额。", 422)
        funding[name] = {"research_total": _number(research), "other_total": _number(other)}
        seen = set()
        for row_number in range(2, footer):
            values = [sheet.cell(row_number, col).value for col in range(1, 6)]
            if all(value is None for value in values):
                continue
            sequence = values[0]
            if sequence is not None and (not isinstance(sequence, int) or isinstance(sequence, bool)
                                         or sequence < 1 or sequence in seen):
                fail("FUNDING_ROW_ID_INVALID", f"{name} 的明细序号无效。", 422)
            if sequence is not None:
                seen.add(sequence)
            amount = money(values[3], cell=f"{sheet.title}!D{row_number}")
            if amount is not None and (amount < 0 or amount > Decimal("1000000000")):
                fail("FUNDING_AMOUNT_INVALID", "确认使用金额须为非负且不超过十亿元。", 422)
            result.append({"sequence": sequence, "name": name,
                           "registration": values[1], "date": values[2],
                           "usage_amount": _number(amount), "status": values[4],
                           "_sheet_title": sheet.title,
                           "_sheet_row_number": row_number})
    return result, funding


def to_bytes(book: Workbook) -> bytes:
    output = BytesIO()
    book.save(output)
    return output.getvalue()

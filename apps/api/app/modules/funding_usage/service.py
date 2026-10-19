"""从已归档的图片 OCR 和旧汇总表生成新的经费使用 Excel。"""

from __future__ import annotations

from copy import copy
from datetime import date as calendar_date
from decimal import Decimal, InvalidOperation
from hashlib import sha256
import json
import re
import tempfile
from pathlib import Path
from uuid import uuid4
from sqlalchemy import func

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.db.models import Document, DocumentPage, DocumentVersion, DocumentExtractionRun, StructuredDataset, StructuredDatasetExport
from app.modules.file_lifecycle.shared_access import CanonicalWorkingFileResolver, CanonicalWorkingFileError
from app.modules.file_lifecycle.shared_workspace import get_shared_workspace_id
from app.modules.structured_datasets.publication import assert_publication_ready, stage_publication
from app.modules.structured_datasets.service import StructuredDatasetService, digest, fail
from app.modules.structured_datasets.xlsx_safety import XlsxSafetyError, inspect_xlsx
from app.modules.structured_extraction.export import _safe_spreadsheet_text

from .schemas import (FundingUsageRequest, FundingUsagePublish,
                      FundingUsageContinue, FundingUsageContinuePublish,
                      FundingUsageRefresh, FundingUsageRefreshPublish,
                      FundingSourceSummarySync, FundingSourceSummarySyncPublish, TextSpan)
from .workbook import append_summary_rows, money, read_summary_rows, rebuild_personal_sheets
from .summary_workbook import (SUMMARY_HEADERS, SUMMARY_KEYS, person_order, read_detail_book,
                               render_book, summaries)


_MONEY = re.compile(r"\s*[¥￥]?\s*(\d{1,3}(?:,\d{3})+|\d+)(?:\.(\d{1,2}))?\s*(?:元)?\s*")
_NUMBER = re.compile(r"\d+(?:[.,]\d+)*")
_YEAR = re.compile(
    r"(?<!\d)(?:19|20)\d{2}(?=\s*年(?:度)?|[.．/\-]\s*\d{1,2}\s*(?:月|[.．/\-])\s*\d{1,2})")
_FULL_DATE = re.compile(
    r"(?<!\d)(?P<year>(?:19|20)\d{2})\s*(?:年|[.．/\-])\s*"
    r"(?P<month>\d{1,2})\s*(?:月|[.．/\-])\s*(?P<day>\d{1,2})\s*日?(?!\d)")
_MONTH_DAY = re.compile(
    r"(?<![\d.．/\-])(?P<month>\d{1,2})\s*(?P<separator>月|[.．/\-])\s*"
    r"(?P<day>\d{1,2})\s*日?(?!\d)")
_HEADERS = ["序号", "姓名", "使用登记原文", "日期", "本次使用金额合计（元）", "扣减后余额（元）", "核对状态"]
_KEYS = ["sequence", "name", "registration", "date", "usage_amount", "balance", "status"]


def _copy_cell_style(source, target) -> None:
    """跨工作簿复制样式属性，避免沿用旧工作簿的样式索引。"""
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def _span(text: str, ref: TextSpan) -> None:
    if ref.end > len(text) or text[ref.start:ref.end] != ref.quote:
        fail("FUNDING_EVIDENCE_INVALID", "图片 OCR 引文与固定来源原文不一致。", 422)


def _amount(quote: str) -> Decimal:
    # OCR 原文中的明确加法式可作为一条引文；每个加数仍必须是独立的正金额。
    parts = re.split(r"[+＋]", quote)
    if len(parts) > 30:
        fail("FUNDING_AMOUNT_INVALID", "一条使用登记最多可核对 30 个加数。", 422)
    total = Decimal("0.00")
    for index, part in enumerate(parts):
        term = part.strip()
        # 仅在加号前允许句末式孤立点，例如 OCR 原文的“2030.+3885”。
        if index < len(parts) - 1 and term.endswith((".", "．")):
            term = term[:-1].rstrip()
        match = _MONEY.fullmatch(term)
        if match is None:
            fail("FUNDING_AMOUNT_INVALID", "确认支出的引文必须是单个金额或仅由加号连接的明确金额。", 422)
        try:
            number = Decimal(match.group(1).replace(",", "") + ("." + match.group(2) if match.group(2) else ""))
        except InvalidOperation:
            fail("FUNDING_AMOUNT_INVALID", "使用金额不是有效十进制金额。", 422)
        if not number.is_finite() or number <= 0 or number > Decimal("1000000000"):
            fail("FUNDING_AMOUNT_INVALID", "使用金额超出允许范围。", 422)
        total += number
    if total > Decimal("1000000000"):
        fail("FUNDING_AMOUNT_INVALID", "使用金额超出允许范围。", 422)
    return total.quantize(Decimal("0.01"))


def _excel_number(value: Decimal | None) -> float | None:
    return None if value is None else float(value.quantize(Decimal("0.01")))


def _published_detail_value(value):
    return _safe_spreadsheet_text(value) if isinstance(value, str) else value


def _detail_matches_frozen(published: dict, frozen: dict) -> bool:
    return all(published.get(key) == _published_detail_value(frozen.get(key))
               for key in ("registration", "date", "usage_amount", "status"))


def _unique_image_year(text: str) -> int | None:
    years = {int(match.group()) for match in _YEAR.finditer(text)}
    return next(iter(years)) if len(years) == 1 else None


def _registration_dates(registration_ref: TextSpan | None, image_text: str,
                        amount_candidates) -> tuple[str | None, bool]:
    if registration_ref is None:
        return None, False
    text = registration_ref.quote
    entries: list[tuple[int, str, bool]] = []
    date_spans: list[tuple[int, int]] = []
    image_years = {int(match.group()) for match in _YEAR.finditer(image_text)}
    year_reason = "图片年份不唯一" if len(image_years) > 1 else "图片未识别到年份"

    def add(match, year: int | None) -> None:
        date_spans.append(match.span())
        raw = match.group().strip()
        start = registration_ref.start + match.start()
        end = registration_ref.start + match.end()
        if any(candidate.decision == "UNCERTAIN"
               and candidate.start < end and candidate.end > start
               for candidate in amount_candidates):
            entries.append((match.start(), f"待核对：{raw}（日期识别不清）", True))
            return
        month, day = int(match["month"]), int(match["day"])
        try:
            # 闰年仅用于判定月日是否可能有效；真实年份仍必须来自图片 OCR。
            calendar_date(2024, month, day)
        except ValueError:
            entries.append((match.start(), f"待核对：{raw}（日期无效）", True))
            return
        if year is None:
            entries.append((match.start(), f"待核对：{raw}（{year_reason}）", True))
            return
        try:
            formatted = calendar_date(year, month, day).isoformat()
        except ValueError:
            entries.append((match.start(), f"待核对：{raw}（日期无效）", True))
            return
        entries.append((match.start(), formatted, False))

    for match in _FULL_DATE.finditer(text):
        add(match, int(match["year"]))

    year = _unique_image_year(text) or _unique_image_year(image_text)
    for match in _MONTH_DAY.finditer(text):
        if any(match.start() < end and match.end() > start for start, end in date_spans):
            continue
        if match["separator"] != "月":
            start = registration_ref.start + match.start()
            end = registration_ref.start + match.end()
            if not any(candidate.start < end and candidate.end > start
                       and re.search(r"日期|时间|年月|月日", candidate.reason)
                       and candidate.decision in {"EXCLUDED", "UNCERTAIN"}
                       for candidate in amount_candidates):
                # 裸数字 12.8 也可能是金额；需要登记候选明确认定为日期。
                continue
        add(match, year)

    for candidate in amount_candidates:
        if (candidate.decision not in {"EXCLUDED", "UNCERTAIN"}
                or not re.search(r"日期|时间|年月|月日", candidate.reason)
                or any(candidate.start < registration_ref.start + end
                       and candidate.end > registration_ref.start + start
                       for start, end in date_spans)):
            continue
        reason = "日期识别不清" if candidate.decision == "UNCERTAIN" else "日期不完整"
        entries.append((candidate.start - registration_ref.start,
                        f"待核对：{candidate.quote}（{reason}）", True))

    entries.sort(key=lambda entry: entry[0])
    values = list(dict.fromkeys(value for _, value, _ in entries))
    return ("\n".join(values) if values else None,
            any(pending for _, _, pending in entries))


class FundingUsageService:
    _legacy_initial_layout = False

    def __init__(self, db, user, agent_run=None):
        self.db = db
        self.user = user
        self.datasets = StructuredDatasetService(db, user, agent_run)

    def _sources(self, request: FundingUsageRequest):
        try:
            image_source = CanonicalWorkingFileResolver(self.db).resolve_document(
                document_id=str(request.image.document_id),
                document_version_id=str(request.image.document_version_id))
        except CanonicalWorkingFileError as exc:
            fail(exc.code, str(exc), 422)
        image_version = image_source.document_version
        image_document = self.db.get(Document, str(request.image.document_id))
        run = self.db.get(DocumentExtractionRun, str(request.image.extraction_run_id))
        if (image_version is None or image_document is None or run is None
                or image_version.document_id != image_document.id
                or image_document.workspace_id != get_shared_workspace_id(self.db)
                or image_source.working_copy.status != "ACTIVE"
                or image_source.working_copy.current_version_id != image_version.id
                or image_version.sha256 != request.image.source_sha256
                or run.document_id != image_document.id
                or run.document_version_id != image_version.id
                or run.status != "COMPLETED" or run.extractor != "workbuddy-external-document"
                or Path(image_version.filename).suffix.lower() not in {".jpg", ".jpeg", ".png", ".webp"}):
            fail("FUNDING_IMAGE_SOURCE_INVALID", "图片必须是本轮已归档并完成正式 WorkBuddy OCR 的固定版本。", 422)
        image_path = self.datasets.storage.working_copy_path(image_version.storage_path)
        if not image_path.is_file() or self.datasets.storage.sha256_file(image_path) != image_version.sha256:
            fail("FUNDING_IMAGE_SOURCE_CHANGED", "已归档图片字节与固定版本不一致。", 409)
        pages = {page.page_number: page for page in self.db.query(DocumentPage).filter_by(
            extraction_run_id=run.id).all() if page.page_number is not None}
        if not pages:
            fail("FUNDING_IMAGE_OCR_MISSING", "图片没有可核对的正式 OCR 正文。", 422)
        try:
            source = CanonicalWorkingFileResolver(self.db).resolve_document(
                document_id=str(request.workbook.document_id),
                document_version_id=str(request.workbook.document_version_id))
        except CanonicalWorkingFileError as exc:
            fail(exc.code, str(exc), 422)
        version = source.document_version
        if (source.working_copy.status != "ACTIVE" or source.working_copy.current_version_id != version.id
                or version.sha256 != request.workbook.source_sha256
                or Path(version.filename).suffix.lower() != ".xlsx"):
            fail("FUNDING_WORKBOOK_VERSION_INVALID", "旧汇总表不是当前已归档的 .xlsx 版本。", 409)
        path = self.datasets.storage.working_copy_path(version.storage_path)
        if not path.is_file() or self.datasets.storage.sha256_file(path) != version.sha256:
            fail("FUNDING_WORKBOOK_MISSING", "旧汇总表文件缺失或哈希不一致。", 409)
        return pages, path

    def _baseline(self, path: Path):
        try:
            inspect_xlsx(path)
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                if "汇总" not in book.sheetnames:
                    fail("FUNDING_SUMMARY_MISSING", "旧 Excel 没有“汇总”工作表。", 422)
                sheet = book["汇总"]
                headers = [str(cell.value or "").replace(" ", "").strip() for cell in sheet[1]]
                names = [i + 1 for i, value in enumerate(headers) if value == "姓名"]
                totals = [i + 1 for i, value in enumerate(headers) if value == "资助总额（元）"]
                if len(names) != 1 or len(totals) != 1:
                    fail("FUNDING_SUMMARY_SCHEMA_INVALID", "旧表必须有唯一的“姓名”和“资助总额（元）”列。", 422)
                result: dict[str, list[Decimal | None]] = {}
                for values in sheet.iter_rows(min_row=2, values_only=True):
                    raw_name = values[names[0] - 1] if len(values) >= names[0] else None
                    if raw_name is None or not str(raw_name).strip():
                        continue
                    name = re.sub(r"\s+", "", str(raw_name))
                    # 样本末行 A=“总计”、B=“50 人”；这是统计行，不是教师。
                    row_label = re.sub(r"\s+", "", str(values[0] or ""))
                    if row_label in {"总计", "合计", "小计", "总数"} or re.fullmatch(r"\d+人", name):
                        continue
                    raw_total = values[totals[0] - 1] if len(values) >= totals[0] else None
                    try:
                        total = Decimal(str(raw_total)) if raw_total is not None and not isinstance(raw_total, bool) else None
                        if total is not None and (not total.is_finite() or total < 0):
                            total = None
                    except InvalidOperation:
                        total = None
                    result.setdefault(name, []).append(total)
                return result
            finally:
                book.close()
        except (OSError, ValueError, KeyError, XlsxSafetyError):
            fail("FUNDING_WORKBOOK_INVALID", "旧汇总表无法安全读取。", 422)

    def _evaluate(self, request: FundingUsageRequest, *, totals_override=None,
                  allow_duplicate_names=False, starting_sequence=0):
        pages, path = self._sources(request)
        totals = self._baseline(path) if totals_override is None else {
            name: list(values) for name, values in totals_override.items()}
        people = []
        seen_names = set()
        for index, row in enumerate(request.rows, starting_sequence + 1):
            page = pages.get(row.page_number)
            if page is None:
                fail("FUNDING_EVIDENCE_INVALID", "登记行引用的图片页不存在。", 422)
            text = page.text_content
            _span(text, row.name_ref)
            name = re.sub(r"\s+", "", row.name)
            if name != re.sub(r"\s+", "", row.name_ref.quote):
                fail("FUNDING_NAME_EVIDENCE_INVALID", "姓名必须与图片 OCR 姓名引文一致。", 422)
            if name in seen_names and not allow_duplicate_names:
                fail("FUNDING_NAME_DUPLICATE", "图片中同名人员重复，需先核对并合并。", 422)
            seen_names.add(name)
            if row.registration_ref:
                _span(text, row.registration_ref)
            if row.no_usage and row.registration_ref and not re.search(
                    r"未使用|无使用|未支出|无支出|未报销|无登记|未发生使用|暂未使用",
                    row.registration_ref.quote):
                fail("FUNDING_NO_USAGE_EVIDENCE_INVALID", "有手写登记时，只有明确写出未使用才能按零支出计算。", 422)
            spent = Decimal("0.00")
            uncertain = False
            ranges = []
            for candidate in row.amounts:
                _span(text, candidate)
                if (row.registration_ref is None or candidate.start < row.registration_ref.start
                        or candidate.end > row.registration_ref.end):
                    fail("FUNDING_AMOUNT_SOURCE_INVALID", "使用金额只能来自同一人的手写使用登记引文。", 422)
                if any(candidate.start < end and candidate.end > start for start, end in ranges):
                    fail("FUNDING_AMOUNT_DUPLICATE", "金额引文重叠，不能重复扣减。", 422)
                ranges.append((candidate.start, candidate.end))
                if candidate.decision == "CONFIRMED_SPEND":
                    nearby_after = text[candidate.end:candidate.end + 12]
                    addition_dot = bool(re.match(r"[.．]\s*[+＋]\s*\d", nearby_after))
                    if (re.fullmatch(r"(?:19|20)\d{2}", candidate.quote.strip())
                            and nearby_after[:1] in {"年", ".", "．", "-", "/", "月"}
                            and not addition_dot):
                        fail("FUNDING_DATE_NOT_SPEND", "日期中的年份不能作为确认支出。", 422)
                    spent += _amount(candidate.quote)
                elif candidate.decision == "UNCERTAIN":
                    uncertain = True
            if row.registration_ref and not row.amounts and not row.no_usage:
                uncertain = True
            if row.registration_ref and not row.no_usage and not any(
                    candidate.decision == "CONFIRMED_SPEND" for candidate in row.amounts):
                uncertain = True
            if row.registration_ref:
                # 登记中每个阿拉伯数字都必须解释为支出、排除或不确定；
                # 漏掉一笔时不能把部分和伪装成最终余额。
                for number in _NUMBER.finditer(row.registration_ref.quote):
                    start = row.registration_ref.start + number.start()
                    end = row.registration_ref.start + number.end()
                    if not any(left <= start and end <= right for left, right in ranges):
                        uncertain = True
            if row.no_usage:
                spent = Decimal("0.00")
            matches = totals.get(name, [])
            if not matches:
                status = "旧表未找到姓名"
            elif len(matches) != 1:
                status = "旧表同名多行"
            elif matches[0] is None:
                status = "旧表资助总额无效"
            elif uncertain:
                status = "手写金额待核对"
            else:
                status = "已核对"
            if uncertain and status != "手写金额待核对":
                status += "；手写金额待核对"
            balance = matches[0] - spent if status == "已核对" else None
            if balance is not None and balance < 0:
                status = "可能超额，待核对"
            date_value, date_pending = _registration_dates(row.registration_ref, text, row.amounts)
            if date_pending:
                status = ("金额已核对；日期待核对" if status == "已核对"
                          else f"{status}；日期待核对")
            if totals_override is not None and name in totals and len(totals[name]) == 1:
                totals[name] = [balance]
            people.append({"sequence": index, "name": name,
                           "registration": row.registration_ref.quote if row.registration_ref else None,
                           "date": date_value,
                           "usage_amount": None if uncertain else _excel_number(spent),
                           "balance": _excel_number(balance), "status": status,
                           "page_number": row.page_number,
                           "name_start": row.name_ref.start,
                           "registration_start": row.registration_ref.start if row.registration_ref else None,
                           "amounts": [{"quote": c.quote, "decision": c.decision,
                                        "reason": c.reason, "start": c.start, "end": c.end}
                                       for c in row.amounts]})
        return people, path

    def _digest(self, request: FundingUsageRequest) -> str:
        body = request.model_dump(mode="json", exclude={"preview_digest"})
        return sha256(json.dumps(body, ensure_ascii=False, sort_keys=True,
                                 separators=(",", ":")).encode("utf-8")).hexdigest()

    def preview(self, request: FundingUsageRequest):
        people, _ = self._evaluate(request)
        return {"preview_digest": self._digest(request), "filename": request.filename,
                "people": people, "person_count": len(people),
                "ready_count": sum(row["balance"] is not None for row in people),
                "requires_review": any(row["status"] != "已核对" for row in people)}

    def publish(self, request: FundingUsagePublish):
        if request.preview_digest != self._digest(request):
            fail("FUNDING_PREVIEW_CHANGED", "预览与发布候选不一致，请重新预览。", 409)
        people, baseline_path = self._evaluate(request)
        filename = self.datasets._normalize_target_filename(request.filename, require_extension=True)
        if filename.casefold() == Path(baseline_path).name.casefold():
            fail("FUNDING_OUTPUT_FILENAME_INVALID", "新表不能与旧汇总表同名。", 422)
        creation_key = digest([self.user.id, "funding-usage", request.request_key])
        creation_digest = self._digest(request)
        existing = self.db.query(StructuredDataset).filter_by(creation_key=creation_key).one_or_none()
        if existing is not None:
            if existing.creation_digest != creation_digest:
                fail("FUNDING_REQUEST_CONFLICT", "请求键已用于另一份图片、旧表或候选。", 409)
            assert_publication_ready(self.db, dataset_id=existing.id)
            exported = self.db.query(StructuredDatasetExport).filter_by(dataset_id=existing.id).one_or_none()
            if exported is None:
                fail("FUNDING_PUBLICATION_PENDING", "上次生成尚未正式发布。", 423)
            return {**self.datasets._export_result(exported, reused=True), "status": "COMPLETED",
                    "people": people, "person_count": len(people)}
        self.datasets._assert_target_filename_available(filename)
        # 只供已有旧版链路的兼容回归测试构造 v1 文件；正式入口始终发布 v3。
        legacy_layout = getattr(self, "_legacy_initial_layout", False)
        if legacy_layout:
            schema = [{"key": key, "label": label,
                       "field_type": "decimal" if key in {"usage_amount", "balance"} else
                           "integer" if key == "sequence" else "string",
                       "required": False, "multiple": False}
                      for key, label in zip(_KEYS, _HEADERS)]
            base_rows = [{**{key: row[key] for key in _KEYS}, "_funding_row_id": str(uuid4())}
                         for row in people]
            layout = {"header_range": "A1:G1", "data_range": f"A2:G{max(2, len(people) + 1)}"}
            funding = detail_rows = None
        else:
            schema = [{"key": key, "label": label,
                       "field_type": "integer" if key == "sequence" else "string" if key == "name" else "decimal",
                       "required": False, "multiple": False}
                      for key, label in zip(SUMMARY_KEYS, SUMMARY_HEADERS)]
            baseline_totals = self._baseline(baseline_path)
            # 旧表“汇总”的姓名顺序决定首页和个人页顺序；图片只提供登记明细。
            funding = {
                name: {"research_total": float(matches[0])
                       if len(matches) == 1 and matches[0] is not None else None,
                       "other_total": None}
                for name, matches in baseline_totals.items()
            }
            for row in people:
                funding.setdefault(row["name"], {"research_total": None,
                                                   "other_total": None})
            detail_rows = [{**{key: row[key] for key in ("sequence", "name", "registration",
                            "date", "usage_amount", "status")}, "sequence": 1,
                            "_funding_row_id": str(uuid4())} for row in people]
            base_rows = [{**row, "_funding_row_id": str(uuid4())}
                         for row in summaries(detail_rows, funding)]
            layout = {"header_range": "A1:F1", "data_range": f"A2:F{max(2, len(base_rows) + 1)}"}
        dataset = StructuredDataset(
            workspace_id=get_shared_workspace_id(self.db), created_by=self.user.id,
            name=Path(filename).stem[:100], target_filename=filename,
            target_sheet_name="使用登记汇总",
            target_layout_json=layout,
            schema_json=schema, base_rows_json=base_rows,
            target_source_version_id=str(request.workbook.document_version_id),
            creation_key=creation_key, creation_digest=creation_digest)
        self.db.add(dataset)
        self.db.flush()
        with tempfile.TemporaryDirectory(prefix="file-agent-funding-") as directory:
            output = Path(directory) / filename
            self._render(output, baseline_path, people if legacy_layout else detail_rows, funding)
            file_sha = self.datasets.storage.sha256_file(output)
            export_id = str(uuid4())
            storage_path = self.datasets.storage.dataset_working_version_relative_path(
                dataset_id=dataset.id, export_id=export_id, version_number=1)
            self.datasets.storage.import_working_copy(
                source=output, relative_path=storage_path, expected_sha256=file_sha)
            working, version = self.datasets._publish_working_version(
                dataset=dataset, working_copy=None, parent_version=None,
                storage_path=storage_path, version_number=1,
                sha256=file_sha, size_bytes=output.stat().st_size)
            v3_policy = ({"funding_usage_rows_v3": detail_rows,
                "funding_usage_funding_v3": funding,
                "funding_usage_person_order_v3": [row["name"] for row in base_rows],
                "funding_usage_history": [self._registration_key(request.image.source_sha256, row)
                                          for row in request.rows],
                "funding_usage_entries": [{
                    "registration_key": self._registration_key(request.image.source_sha256, source),
                    "source_row_id": detail["_funding_row_id"], "name": detail["name"],
                    "sequence": detail["sequence"], "image_document_id": str(request.image.document_id),
                    "image_document_version_id": str(request.image.document_version_id),
                    "image_sha256": request.image.source_sha256,
                    "image_extraction_run_id": str(request.image.extraction_run_id),
                    "name_ref": source.name_ref.model_dump(),
                    "registration_ref": source.registration_ref.model_dump() if source.registration_ref else None,
                    "amounts": [candidate.model_dump() for candidate in source.amounts],
                } for source, detail in zip(request.rows, detail_rows)]}
                if not legacy_layout else {})
            version.extraction_policy_json = {**dict(version.extraction_policy_json or {}),
                "structured_dataset_export_mode": "funding-usage-v1" if legacy_layout else "funding-usage-v3",
                **v3_policy,
                "image_document_version_id": str(request.image.document_version_id),
                "image_extraction_run_id": str(request.image.extraction_run_id),
                "workbook_document_version_id": str(request.workbook.document_version_id),
                "funding_usage_evidence": [{
                    "row": index, "name_ref": row.name_ref.model_dump(),
                    "page_number": row.page_number,
                    "registration_ref": row.registration_ref.model_dump() if row.registration_ref else None,
                    "amounts": [candidate.model_dump() for candidate in row.amounts],
                    "no_usage": row.no_usage,
                } for index, row in enumerate(request.rows, 1)]}
            export = StructuredDatasetExport(id=export_id, dataset_id=dataset.id,
                dataset_revision=0, storage_path=storage_path,
                sha256=file_sha, size_bytes=output.stat().st_size,
                document_version_id=version.id)
            self.db.add(export)
            self.db.flush()
            result = {**self.datasets._export_result(export, reused=False),
                "status": "PUBLISHING", "excel_published": False,
                "person_count": len(people)}
            publication = stage_publication(self.db, dataset=dataset,
                working_copy=working, version=version, result=result,
                agent_run=self.datasets.agent_run)
        return {**result, "people": people, "_publication_id": publication.id}

    @staticmethod
    def _registration_key(image_sha256, row) -> str:
        registration = row.registration_ref or row.name_ref
        return digest([image_sha256, row.page_number,
                       registration.start, registration.end])

    def _continue_target(self, request: FundingUsageContinue, *, lock=False):
        target = self.datasets.resolve_target(request.filename)
        dataset = self.datasets.dataset(target["dataset_id"], lock=lock)
        working = self.datasets._dataset_working_copy(dataset)
        if working is None or not working.current_version_id:
            fail("FUNDING_TARGET_NOT_PUBLISHED", "目标余额表尚未正式发布。", 409)
        current = self.db.get(DocumentVersion, working.current_version_id)
        origin = (self.db.query(DocumentVersion).filter_by(document_id=working.document_id,
                  version_number=1).one_or_none())
        if not origin or (origin.extraction_policy_json or {}).get(
                "structured_dataset_export_mode") not in {"funding-usage-v1", "funding-usage-v3"}:
            fail("FUNDING_TARGET_INVALID", "同名文件不是成果资助使用登记余额表。", 422)
        if current is None or current.filename != request.filename:
            fail("FUNDING_TARGET_INVALID", "目标余额表版本链不完整。", 409)
        path = self.datasets.storage.working_copy_path(current.storage_path)
        if not path.is_file() or self.datasets.storage.sha256_file(path) != current.sha256:
            fail("FUNDING_TARGET_CHANGED", "目标余额表文件缺失或内容与版本不一致。", 409)
        return dataset, working, current, origin, path

    def _refresh_source(self, request: FundingUsageRefresh) -> Path:
        try:
            source = CanonicalWorkingFileResolver(self.db).resolve_document(
                document_id=str(request.workbook.document_id),
                document_version_id=str(request.workbook.document_version_id))
        except CanonicalWorkingFileError as exc:
            fail(exc.code, str(exc), 422)
        version = source.document_version
        if (version is None or source.working_copy.status != "ACTIVE"
                or source.working_copy.current_version_id != version.id
                or version.sha256 != request.workbook.source_sha256
                or Path(version.filename).suffix.lower() != ".xlsx"):
            fail("FUNDING_REFRESH_SOURCE_INVALID", "新版成果明细表必须是当前已归档的 .xlsx 版本。", 409)
        path = self.datasets.storage.working_copy_path(version.storage_path)
        if not path.is_file() or self.datasets.storage.sha256_file(path) != version.sha256:
            fail("FUNDING_REFRESH_SOURCE_CHANGED", "新版成果明细表文件缺失或校验失败。", 409)
        return path

    @staticmethod
    def _refresh_detail_totals(path: Path, names: set[str]) -> dict[str, dict]:
        """个人成果页的 J 列合计优先；逐项金额必须能复算该合计。"""
        def amount(value):
            if value is None or isinstance(value, bool):
                return None
            try:
                result = Decimal(str(value))
            except InvalidOperation:
                return None
            if (not result.is_finite() or result < 0 or result > Decimal("1000000000")
                    or result.as_tuple().exponent < -2):
                return None
            return result

        try:
            book = load_workbook(path, read_only=True, data_only=False)
            try:
                detail: dict[str, dict] = {}
                for sheet in book.worksheets[1:]:
                    name = re.sub(r"\s+", "", sheet.title)
                    if name not in names:
                        continue
                    if name in detail:
                        detail[name] = {"status": "DETAIL_DUPLICATE"}
                        continue
                    if sheet.cell(2, 10).value != "资助金额（元）":
                        detail[name] = {"status": "DETAIL_STRUCTURE_INVALID"}
                        continue
                    footers = []
                    item_amounts = []
                    for values in sheet.iter_rows(min_row=3, max_col=10, values_only=True):
                        if str(values[0] or "").strip() == "合计":
                            footers.append(values[9])
                        elif not footers and values[9] is not None:
                            item_amounts.append(values[9])
                    if len(footers) != 1:
                        detail[name] = {"status": "DETAIL_STRUCTURE_INVALID"}
                        continue
                    total = amount(footers[0])
                    if total is None:
                        detail[name] = {"status": "DETAIL_TOTAL_INVALID"}
                        continue
                    parsed = [amount(value) for value in item_amounts]
                    if any(value is None for value in parsed):
                        detail[name] = {"status": "DETAIL_ITEM_INVALID"}
                    elif sum(parsed, Decimal("0.00")) != total:
                        detail[name] = {"status": "DETAIL_TOTAL_MISMATCH"}
                    else:
                        detail[name] = {"status": "DETAIL_VALID", "total": total}
                return detail
            finally:
                book.close()
        except (OSError, ValueError, KeyError):
            fail("FUNDING_REFRESH_WORKBOOK_INVALID", "新版成果明细表个人页无法安全读取。", 422)

    @staticmethod
    def _source_summary_digest(request: FundingSourceSummarySync) -> str:
        return sha256(json.dumps(request.model_dump(mode="json", exclude={"preview_digest"}),
            ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()

    def _source_summary_evaluate(self, request: FundingSourceSummarySync):
        path = self._refresh_source(request)
        summary = self._baseline(path)
        detail = self._refresh_detail_totals(path, set(summary))
        decisions = []
        try:
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet = book["汇总"]
                headers = [str(cell.value or "").replace(" ", "").strip() for cell in sheet[1]]
                name_col = headers.index("姓名") + 1
                total_col = headers.index("资助总额（元）") + 1
                for row_number in range(2, sheet.max_row + 1):
                    raw = sheet.cell(row_number, name_col).value
                    if raw is None or not str(raw).strip():
                        continue
                    name = re.sub(r"\s+", "", str(raw))
                    label = re.sub(r"\s+", "", str(sheet.cell(row_number, 1).value or ""))
                    if label in {"总计", "合计", "小计", "总数"} or re.fullmatch(r"\d+人", name):
                        continue
                    item = detail.get(name)
                    old = summary.get(name, [])
                    status = ("SOURCE_DUPLICATE" if len(old) != 1 else
                              item["status"] if item and item["status"] != "DETAIL_VALID" else
                              "NO_PERSONAL_DETAIL" if item is None else
                              "UPDATED" if old[0] != item["total"] else "UNCHANGED")
                    decisions.append({"name": name, "row_number": row_number, "status": status,
                                      "old_research_total": _excel_number(old[0]) if len(old) == 1 and old[0] is not None else None,
                                      "new_research_total": _excel_number(item["total"]) if item and item["status"] == "DETAIL_VALID" else None})
            finally:
                book.close()
        except (OSError, ValueError, KeyError):
            fail("FUNDING_REFRESH_WORKBOOK_INVALID", "成果明细表汇总页无法安全读取。", 422)
        return path, decisions, total_col

    def source_summary_preview(self, request: FundingSourceSummarySync):
        _, decisions, _ = self._source_summary_evaluate(request)
        return {"preview_digest": self._source_summary_digest(request),
                "changed_count": sum(item["status"] == "UPDATED" for item in decisions),
                "review_count": sum(item["status"].startswith(("SOURCE_", "DETAIL_")) for item in decisions),
                "people": decisions}

    def source_summary_publish(self, request: FundingSourceSummarySyncPublish):
        request_digest = self._source_summary_digest(request)
        if request.preview_digest != request_digest:
            fail("FUNDING_PREVIEW_CHANGED", "成果明细表汇总页预览与发布候选不一致。", 409)
        # 幂等检查先于当前版本检查；成功后的旧附件版本仍可返回同一发布回执。
        for previous in self.db.query(DocumentVersion).filter_by(
                source_type="FUNDING_SOURCE_SUMMARY_SYNC").all():
            policy = previous.extraction_policy_json or {}
            if (policy.get("funding_source_summary_request_key") != request.request_key
                    or policy.get("funding_source_summary_source_document_id") != str(request.workbook.document_id)):
                continue
            if policy.get("funding_source_summary_digest") != request_digest:
                fail("FUNDING_REQUEST_CONFLICT", "请求键已用于不同的成果明细表修订。", 409)
            export = self.db.query(StructuredDatasetExport).filter_by(
                document_version_id=previous.id).one_or_none()
            if export is None:
                fail("FUNDING_PUBLICATION_PENDING", "成果明细表上次发布尚未完成。", 423)
            assert_publication_ready(self.db, dataset_id=export.dataset_id)
            return {**self.datasets._export_result(export, reused=True),
                    "status": "COMPLETED", "changed_count": policy.get("funding_source_summary_changed_count", 0)}
        path, decisions, total_col = self._source_summary_evaluate(request)
        changed = [item for item in decisions if item["status"] == "UPDATED"]
        if not changed:
            fail("FUNDING_SOURCE_UNCHANGED", "个人明细页没有需要同步到汇总页的有效金额差异。", 422)
        source = CanonicalWorkingFileResolver(self.db).resolve_document(
            document_id=str(request.workbook.document_id),
            document_version_id=str(request.workbook.document_version_id))
        working = source.working_copy
        current = source.document_version
        assert_publication_ready(self.db, working_copy_id=working.id)
        visible_relative = self.datasets._working_copy_storage_path(working)
        visible = self.datasets.storage.working_copy_path(visible_relative)
        if not visible.is_file() or self.datasets.storage.sha256_file(visible) != current.sha256:
            fail("FUNDING_REFRESH_SOURCE_CHANGED", "成果明细表当前可见文件与版本不一致。", 409)
        dataset = self.db.query(StructuredDataset).filter_by(working_copy_id=working.id).with_for_update().one_or_none()
        if dataset is not None and not dataset.creation_key == digest(["funding-source-summary", working.id]):
            fail("FUNDING_SOURCE_DATASET_CONFLICT", "成果明细表已由其他数据集管理，不能通过本流程改写。", 409)
        if dataset is None:
            dataset = StructuredDataset(workspace_id=get_shared_workspace_id(self.db),
                created_by=self.user.id, name=Path(working.filename).stem[:100],
                target_filename=working.filename, target_sheet_name="汇总",
                target_layout_json={"funding_source_summary_sync": True},
                schema_json=[{"key": "name", "label": "姓名", "field_type": "string",
                              "required": False, "multiple": False},
                             {"key": "research_total", "label": "资助总额（元）",
                              "field_type": "decimal", "required": False, "multiple": False}],
                base_rows_json=[], target_source_version_id=current.id,
                creation_key=digest(["funding-source-summary", working.id]),
                creation_digest=digest([working.id]), working_copy_id=working.id)
            self.db.add(dataset)
            self.db.flush()
        self.datasets._assert_task_write_access(dataset)
        with tempfile.TemporaryDirectory(prefix="file-agent-funding-source-") as directory:
            output = Path(directory) / working.filename
            book = load_workbook(path, read_only=False, data_only=False)
            try:
                if any(cell.data_type == "f" for worksheet in book.worksheets
                       for row in worksheet for cell in row):
                    fail("FUNDING_SOURCE_FORMULA_UNSUPPORTED",
                         "成果明细表含公式，当前同步方式会清除公式缓存；请先保存为数值版再上传。", 422)
                sheet = book["汇总"]
                for item in changed:
                    sheet.cell(item["row_number"], total_col).value = item["new_research_total"]
                book.save(output)
            finally:
                book.close()
            inspect_xlsx(output)
            file_sha = self.datasets.storage.sha256_file(output)
            export_id = str(uuid4())
            version_number = int(self.db.query(func.max(DocumentVersion.version_number)).filter(
                DocumentVersion.document_id == working.document_id).scalar() or 0) + 1
            if current.storage_path == visible_relative:
                history_path = self.datasets.storage.dataset_source_version_relative_path(
                    dataset_id=dataset.id, document_version_id=current.id,
                    version_number=current.version_number, suffix=".xlsx")
                self.datasets.storage.import_working_copy(
                    source=visible, relative_path=history_path, expected_sha256=current.sha256)
                current.storage_path = history_path
            storage_path = self.datasets.storage.dataset_working_version_relative_path(
                dataset_id=dataset.id, export_id=export_id, version_number=version_number)
            self.datasets.storage.import_working_copy(
                source=output, relative_path=storage_path, expected_sha256=file_sha)
            dataset.revision += 1
            _, version = self.datasets._publish_working_version(
                dataset=dataset, working_copy=working, parent_version=current,
                storage_path=storage_path, version_number=version_number,
                sha256=file_sha, size_bytes=output.stat().st_size)
            version.source_type = "FUNDING_SOURCE_SUMMARY_SYNC"
            version.extraction_policy_json = {**dict(current.extraction_policy_json or {}),
                "funding_source_summary_request_key": request.request_key,
                "funding_source_summary_source_document_id": str(request.workbook.document_id),
                "funding_source_summary_digest": request_digest,
                "funding_source_summary_changed_count": len(changed)}
            dataset.base_rows_json = [{"name": item["name"],
                "research_total": item["new_research_total"] if item["status"] == "UPDATED"
                    else item["old_research_total"]} for item in decisions]
            export = StructuredDatasetExport(id=export_id, dataset_id=dataset.id,
                dataset_revision=dataset.revision, storage_path=storage_path,
                sha256=file_sha, size_bytes=output.stat().st_size,
                document_version_id=version.id)
            self.db.add(export)
            self.db.flush()
            result = {**self.datasets._export_result(export, reused=False),
                "status": "PUBLISHING", "excel_published": False,
                "changed_count": len(changed), "review_count": sum(
                    item["status"].startswith(("SOURCE_", "DETAIL_")) for item in decisions)}
            publication = stage_publication(self.db, dataset=dataset,
                working_copy=working, version=version, base_version_id=current.id,
                result=result, agent_run=self.datasets.agent_run)
        return {**result, "people": decisions, "_publication_id": publication.id}

    def _refresh_evaluate(self, request: FundingUsageRefresh, *, lock=False):
        dataset, working, current, origin, target_path = self._continue_target(request, lock=lock)
        if (origin.extraction_policy_json or {}).get("structured_dataset_export_mode") != "funding-usage-v3":
            fail("FUNDING_REFRESH_TARGET_INVALID", "只有新版个人明细驱动的余额表支持资助总额刷新。", 422)
        if (current.id != str(request.base_document_version_id)
                or current.sha256 != request.base_sha256
                or dataset.revision != request.base_revision):
            fail("FUNDING_TARGET_VERSION_CONFLICT", "目标余额表已有新版本，请重新预览。", 409)
        if working.document_id == str(request.workbook.document_id):
            fail("FUNDING_REFRESH_SOURCE_INVALID", "成果明细表不能是目标余额表本身。", 422)
        source_path = self._refresh_source(request)
        source_totals = self._baseline(source_path)
        policy = current.extraction_policy_json or {}
        rows = list(policy.get("funding_usage_rows_v3") or [])
        funding = dict(policy.get("funding_usage_funding_v3") or {})
        order = person_order(dataset, funding, policy)
        detail_totals = self._refresh_detail_totals(source_path, set(order))
        try:
            inspect_xlsx(target_path)
            book = load_workbook(target_path, read_only=False, data_only=False)
            try:
                published_rows, published_funding = read_detail_book(
                    book, order, required_names={row["name"] for row in rows})
            finally:
                book.close()
        except (OSError, ValueError, XlsxSafetyError):
            fail("FUNDING_TARGET_INVALID", "目标余额表无法安全读取。", 422)
        frozen = {(row["name"], row["sequence"]): row for row in rows}
        if (published_funding != funding or len(rows) != len(published_rows)
                or len(frozen) != len(rows)
                or any((row["name"], row["sequence"]) not in frozen
                       or not _detail_matches_frozen(
                           row, frozen[(row["name"], row["sequence"])])
                       for row in published_rows)):
            fail("FUNDING_TARGET_CHANGED", "服务器已发布 Excel 与冻结登记不一致。", 409)
        updated_funding = {name: dict(values) for name, values in funding.items()}
        decisions = []
        for name in order:
            summary_matches = source_totals.get(name, [])
            detail = detail_totals.get(name)
            matches = ([detail["total"]] if detail and detail["status"] == "DETAIL_VALID"
                       else summary_matches if detail is None else [])
            old_total = funding[name].get("research_total")
            if detail and detail["status"] != "DETAIL_VALID":
                status = detail["status"]
            elif not matches:
                status = "SOURCE_MISSING"
            elif len(matches) != 1:
                status = "SOURCE_DUPLICATE"
            elif matches[0] is None or matches[0] > Decimal("1000000000") or matches[0].as_tuple().exponent < -2:
                status = "SOURCE_TOTAL_INVALID"
            else:
                new_total = _excel_number(matches[0])
                status = "UNCHANGED" if old_total == new_total else "UPDATED"
                updated_funding[name]["research_total"] = new_total
            decisions.append({"name": name, "status": status,
                              "old_research_total": old_total,
                              "new_research_total": updated_funding[name].get("research_total"),
                              "source_basis": "PERSONAL_DETAIL" if detail else "SUMMARY",
                              "source_summary_total": _excel_number(summary_matches[0])
                                  if len(summary_matches) == 1 and summary_matches[0] is not None else None})
        calculated = {row["name"]: row for row in summaries(rows, updated_funding, order)}
        for decision in decisions:
            person = calculated[decision["name"]]
            decision.update({key: person[key] for key in ("other_total", "total_spent", "remaining")})
        source_only_names = [name for name in source_totals if name not in funding]
        return dataset, working, current, rows, updated_funding, order, decisions, source_only_names

    def refresh_preview(self, request: FundingUsageRefresh):
        _, _, _, _, _, _, decisions, source_only_names = self._refresh_evaluate(request)
        return {"preview_digest": self._continue_digest(request),
                "filename": request.filename,
                "changed_count": sum(row["status"] == "UPDATED" for row in decisions),
                "review_count": sum(row["status"].startswith(("SOURCE_", "DETAIL_")) for row in decisions),
                "people": decisions, "source_only_names": source_only_names}

    def refresh_publish(self, request: FundingUsageRefreshPublish):
        request_digest = self._continue_digest(request)
        if request.preview_digest != request_digest:
            fail("FUNDING_PREVIEW_CHANGED", "资助总额刷新预览与发布候选不一致。", 409)
        dataset, working, current, _, _ = self._continue_target(request, lock=True)
        for previous in self.db.query(DocumentVersion).filter_by(document_id=working.document_id).all():
            saved = previous.extraction_policy_json or {}
            if saved.get("funding_usage_refresh_request_key") == request.request_key:
                if saved.get("funding_usage_refresh_digest") != request_digest:
                    fail("FUNDING_REQUEST_CONFLICT", "请求键已用于其他资助总额刷新。", 409)
                export = self.db.query(StructuredDatasetExport).filter_by(
                    dataset_id=dataset.id, document_version_id=previous.id).one_or_none()
                if export is None:
                    fail("FUNDING_PUBLICATION_PENDING", "上次刷新尚未正式发布。", 423)
                return {**self.datasets._export_result(export, reused=True),
                        "status": "COMPLETED", "changed_count": saved.get("funding_usage_refresh_changed_count", 0)}
        dataset, working, current, rows, funding, order, decisions, source_only_names = self._refresh_evaluate(
            request, lock=True)
        changed_count = sum(row["status"] == "UPDATED" for row in decisions)
        if not changed_count:
            fail("FUNDING_REFRESH_UNCHANGED", "新版成果明细表没有可更新的资助总额。", 422)
        self.datasets._assert_task_write_access(dataset)
        policy = current.extraction_policy_json or {}
        with tempfile.TemporaryDirectory(prefix="file-agent-funding-refresh-") as directory:
            output = Path(directory) / request.filename
            book = render_book(rows, funding, order=order)
            try:
                book.save(output)
            finally:
                book.close()
            file_sha = self.datasets.storage.sha256_file(output)
            prior_ids = {row["name"]: row.get("_funding_row_id")
                         for row in dataset.base_rows_json or []}
            summary_rows = summaries(rows, funding, order)
            dataset.base_rows_json = [{**row, "_funding_row_id": prior_ids.get(row["name"]) or str(uuid4())}
                                      for row in summary_rows]
            dataset.target_layout_json = {"header_range": "A1:F1",
                "data_range": f"A2:F{max(2, len(summary_rows) + 1)}"}
            dataset.revision += 1
            export_id = str(uuid4())
            version_number = int(self.db.query(func.max(DocumentVersion.version_number)).filter(
                DocumentVersion.document_id == working.document_id).scalar() or 0) + 1
            storage_path = self.datasets.storage.dataset_working_version_relative_path(
                dataset_id=dataset.id, export_id=export_id, version_number=version_number)
            self.datasets.storage.import_working_copy(source=output, relative_path=storage_path,
                                                       expected_sha256=file_sha)
            working, version = self.datasets._publish_working_version(
                dataset=dataset, working_copy=working, parent_version=current,
                storage_path=storage_path, version_number=version_number,
                sha256=file_sha, size_bytes=output.stat().st_size)
            version.extraction_policy_json = {**policy,
                "structured_dataset_export_mode": "funding-usage-v3",
                "funding_usage_funding_v3": funding,
                "funding_usage_person_order_v3": order,
                "funding_usage_refresh_request_key": request.request_key,
                "funding_usage_refresh_digest": request_digest,
                "funding_usage_refresh_changed_count": changed_count,
                "workbook_document_version_id": str(request.workbook.document_version_id),
                "funding_usage_refresh_history": [*(policy.get("funding_usage_refresh_history") or []),
                    {"source_document_id": str(request.workbook.document_id),
                     "source_document_version_id": str(request.workbook.document_version_id),
                     "source_sha256": request.workbook.source_sha256,
                     "request_key": request.request_key,
                     "changed_names": [item["name"] for item in decisions if item["status"] == "UPDATED"]}]}
            export = StructuredDatasetExport(id=export_id, dataset_id=dataset.id,
                dataset_revision=dataset.revision, storage_path=storage_path,
                sha256=file_sha, size_bytes=output.stat().st_size,
                document_version_id=version.id)
            self.db.add(export)
            self.db.flush()
            result = {**self.datasets._export_result(export, reused=False),
                "status": "PUBLISHING", "excel_published": False,
                "changed_count": changed_count,
                "review_count": sum(item["status"].startswith(("SOURCE_", "DETAIL_")) for item in decisions),
                "source_document_id": str(request.workbook.document_id),
                "source_document_version_id": str(request.workbook.document_version_id)}
            publication = stage_publication(self.db, dataset=dataset,
                working_copy=working, version=version, base_version_id=current.id,
                result=result, agent_run=self.datasets.agent_run)
        return {**result, "people": decisions, "source_only_names": source_only_names,
                "_publication_id": publication.id}

    @staticmethod
    def _continue_digest(request: FundingUsageContinue) -> str:
        return sha256(json.dumps(request.model_dump(mode="json", exclude={"preview_digest"}),
            ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()

    def _continue_evaluate(self, request: FundingUsageContinue, *, lock=False):
        dataset, working, current, origin, path = self._continue_target(request, lock=lock)
        if (origin.extraction_policy_json or {}).get("structured_dataset_export_mode") == "funding-usage-v3":
            return self._continue_evaluate_v3(request, dataset, working, current, path)
        if (current.id != str(request.base_document_version_id)
                or current.sha256 != request.base_sha256
                or dataset.revision != request.base_revision):
            fail("FUNDING_TARGET_VERSION_CONFLICT", "目标余额表已有新版本，请重新预览。", 409)
        try:
            inspect_xlsx(path)
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                existing_rows = read_summary_rows(book)
            finally:
                book.close()
        except (OSError, ValueError, XlsxSafetyError):
            fail("FUNDING_TARGET_INVALID", "目标余额表无法安全读取。", 422)
        latest = {}
        for row in existing_rows:
            name = re.sub(r"\s+", "", row["name"])
            latest[name] = [money(row["balance"], cell="F")]
        deleted_names = {re.sub(r"\s+", "", str(item.get("name") or "")) for item in
            (current.extraction_policy_json or {}).get("funding_usage_deleted_entries") or []}
        missing_deleted_names = {re.sub(r"\s+", "", row.name) for row in request.rows
                                 if re.sub(r"\s+", "", row.name) not in latest
                                 and re.sub(r"\s+", "", row.name) in deleted_names}
        if missing_deleted_names:
            baseline_version = self.db.get(DocumentVersion, dataset.target_source_version_id)
            if baseline_version is None:
                fail("FUNDING_BASELINE_MISSING", "首次旧汇总表版本缺失。", 409)
            baseline_path = self.datasets.storage.working_copy_path(baseline_version.storage_path)
            if (not baseline_path.is_file()
                    or self.datasets.storage.sha256_file(baseline_path) != baseline_version.sha256):
                fail("FUNDING_BASELINE_MISSING", "首次旧汇总表文件缺失或校验失败。", 409)
            baseline_totals = self._baseline(baseline_path)
            latest.update({name: baseline_totals[name] for name in missing_deleted_names
                           if name in baseline_totals})
        adapted = FundingUsageRequest.model_validate({
            "image": request.image.model_dump(mode="json"),
            "workbook": {"document_id": working.document_id,
                         "document_version_id": current.id,
                         "source_sha256": current.sha256},
            "filename": request.filename, "request_key": request.request_key,
            "rows": [row.model_dump(mode="json") for row in request.rows],
        })
        additions, _ = self._evaluate(adapted, totals_override=latest,
            allow_duplicate_names=True,
            starting_sequence=max(max((row["sequence"] for row in existing_rows), default=0),
                int((current.extraction_policy_json or {}).get("funding_usage_next_sequence") or 1) - 1))
        for item in additions:
            if item["status"].startswith("旧表未找到姓名"):
                item["status"] = item["status"].replace("旧表未找到姓名", "当前表未找到姓名", 1)
            elif item["status"].startswith("旧表资助总额无效"):
                item["status"] = item["status"].replace("旧表资助总额无效", "当前余额待核对", 1)
        policy = current.extraction_policy_json or {}
        history = list(policy.get("funding_usage_history") or [])
        if "funding_usage_history" not in policy:
            source = origin.extraction_policy_json or {}
            image_version = (self.db.get(DocumentVersion,
                source.get("image_document_version_id")) if source.get("image_document_version_id") else None)
            if image_version is None:
                fail("FUNDING_SOURCE_HISTORY_MISSING", "首次图片版本缺失，不能安全判断重复登记。", 409)
            for item in source.get("funding_usage_evidence") or []:
                name_ref = item.get("name_ref") or {}
                registration = item.get("registration_ref") or name_ref
                if name_ref and registration:
                    history.append(digest([image_version.sha256, item.get("page_number", 1),
                                           registration["start"], registration["end"]]))
        new_keys = [self._registration_key(request.image.source_sha256, row)
                    for row in request.rows]
        if len(new_keys) != len(set(new_keys)) or set(new_keys).intersection(history):
            fail("FUNDING_REGISTRATION_DUPLICATE", "这条图片登记已经追加过，不能重复扣款。", 409)
        return dataset, working, current, path, existing_rows, additions, history + new_keys

    def continue_preview(self, request: FundingUsageContinue):
        _, _, _, _, _, additions, _ = self._continue_evaluate(request)
        return {"preview_digest": self._continue_digest(request),
                "filename": request.filename, "added_record_count": len(additions),
                "people": additions,
                "requires_review": any(row["balance"] is None or "待核对" in row["status"]
                                       for row in additions)}

    def _continue_evaluate_v3(self, request, dataset, working, current, path):
        if (current.id != str(request.base_document_version_id)
                or current.sha256 != request.base_sha256
                or dataset.revision != request.base_revision):
            fail("FUNDING_TARGET_VERSION_CONFLICT", "目标余额表已有新版本，请重新预览。", 409)
        policy = current.extraction_policy_json or {}
        existing_rows = list(policy.get("funding_usage_rows_v3") or [])
        funding = dict(policy.get("funding_usage_funding_v3") or {})
        order = person_order(dataset, funding, policy)
        try:
            inspect_xlsx(path)
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                read_detail_book(book, order,
                                 required_names={row["name"] for row in existing_rows})
            finally:
                book.close()
        except (OSError, ValueError, XlsxSafetyError):
            fail("FUNDING_TARGET_INVALID", "目标余额表无法安全读取。", 422)
        latest = {row["name"]: [money(row["remaining"], cell="F")]
                  for row in summaries(existing_rows, funding, order)}
        adapted = FundingUsageRequest.model_validate({
            "image": request.image.model_dump(mode="json"),
            "workbook": {"document_id": working.document_id,
                         "document_version_id": current.id,
                         "source_sha256": current.sha256},
            "filename": request.filename, "request_key": request.request_key,
            "rows": [row.model_dump(mode="json") for row in request.rows],
        })
        additions, _ = self._evaluate(adapted, totals_override=latest,
            allow_duplicate_names=True)
        history = list(policy.get("funding_usage_history") or [])
        keys = [self._registration_key(request.image.source_sha256, row)
                for row in request.rows]
        if len(keys) != len(set(keys)) or set(keys).intersection(history):
            fail("FUNDING_REGISTRATION_DUPLICATE", "这条图片登记已经追加过，不能重复扣款。", 409)
        return dataset, working, current, path, existing_rows, additions, history + keys

    def _continue_publish_v3(self, request, dataset, working, current):
        request_digest = self._continue_digest(request)
        for previous in self.db.query(DocumentVersion).filter_by(document_id=working.document_id).all():
            policy = previous.extraction_policy_json or {}
            if policy.get("funding_usage_continue_request_key") == request.request_key:
                if policy.get("funding_usage_continue_digest") != request_digest:
                    fail("FUNDING_REQUEST_CONFLICT", "请求键已经用于其他登记或基准版本。", 409)
                export = self.db.query(StructuredDatasetExport).filter_by(
                    dataset_id=dataset.id, document_version_id=previous.id).one_or_none()
                if export is None:
                    fail("FUNDING_PUBLICATION_PENDING", "上次追加尚未正式发布。", 423)
                return {**self.datasets._export_result(export, reused=True),
                        "status": "COMPLETED", "added_record_count": len(request.rows)}
        dataset, working, current, path, existing_rows, additions, history = (
            self._continue_evaluate(request, lock=True))
        self.datasets._assert_task_write_access(dataset)
        policy = current.extraction_policy_json or {}
        funding = dict(policy.get("funding_usage_funding_v3") or {})
        combined = list(existing_rows)
        entries = list(policy.get("funding_usage_entries") or [])
        keys = [self._registration_key(request.image.source_sha256, row)
                for row in request.rows]
        for source, addition, key in zip(request.rows, additions, keys):
            name = addition["name"]
            funding.setdefault(name, {"research_total": None, "other_total": None})
            sequence = sum(row["name"] == name for row in combined) + 1
            detail = {key: addition[key] for key in ("name", "registration", "date",
                                                        "usage_amount", "status")}
            detail.update({"sequence": sequence, "_funding_row_id": str(uuid4())})
            combined.append(detail)
            entries.append({"registration_key": key, "source_row_id": detail["_funding_row_id"],
                "sequence": sequence, "name": name,
                "image_document_id": str(request.image.document_id),
                "image_document_version_id": str(request.image.document_version_id),
                "image_sha256": request.image.source_sha256,
                "image_extraction_run_id": str(request.image.extraction_run_id),
                "name_ref": source.name_ref.model_dump(),
                "registration_ref": source.registration_ref.model_dump() if source.registration_ref else None,
                "amounts": [candidate.model_dump() for candidate in source.amounts]})
        # 个人页按首页顺序分组；JSONB 资助对象的键顺序不能用于排列。
        order = person_order(dataset, funding, policy,
                             new_names=[row["name"] for row in additions])
        combined = [row for name in order for row in combined if row["name"] == name]
        with tempfile.TemporaryDirectory(prefix="file-agent-funding-continue-") as directory:
            output = Path(directory) / request.filename
            book = render_book(combined, funding, order=order)
            try:
                book.save(output)
            finally:
                book.close()
            file_sha = self.datasets.storage.sha256_file(output)
            previous_ids = {row["name"]: row.get("_funding_row_id")
                            for row in dataset.base_rows_json or []}
            summary_rows = summaries(combined, funding, order)
            dataset.base_rows_json = [{**row, "_funding_row_id": previous_ids.get(row["name"]) or str(uuid4())}
                                      for row in summary_rows]
            dataset.target_layout_json = {"header_range": "A1:F1",
                "data_range": f"A2:F{max(2, len(summary_rows) + 1)}"}
            dataset.revision += 1
            export_id = str(uuid4())
            version_number = int(self.db.query(func.max(DocumentVersion.version_number)).filter(
                DocumentVersion.document_id == working.document_id).scalar() or 0) + 1
            storage_path = self.datasets.storage.dataset_working_version_relative_path(
                dataset_id=dataset.id, export_id=export_id, version_number=version_number)
            self.datasets.storage.import_working_copy(source=output, relative_path=storage_path,
                                                       expected_sha256=file_sha)
            working, version = self.datasets._publish_working_version(
                dataset=dataset, working_copy=working, parent_version=current,
                storage_path=storage_path, version_number=version_number,
                sha256=file_sha, size_bytes=output.stat().st_size)
            version.extraction_policy_json = {**policy,
                "structured_dataset_export_mode": "funding-usage-v3",
                "funding_usage_rows_v3": combined,
                "funding_usage_funding_v3": funding,
                "funding_usage_person_order_v3": order,
                "funding_usage_history": history,
                "funding_usage_entries": entries,
                "funding_usage_continue_request_key": request.request_key,
                "funding_usage_continue_digest": request_digest,
                "image_document_version_id": str(request.image.document_version_id),
                "image_extraction_run_id": str(request.image.extraction_run_id)}
            export = StructuredDatasetExport(id=export_id, dataset_id=dataset.id,
                dataset_revision=dataset.revision, storage_path=storage_path,
                sha256=file_sha, size_bytes=output.stat().st_size,
                document_version_id=version.id)
            self.db.add(export)
            self.db.flush()
            result = {**self.datasets._export_result(export, reused=False),
                      "status": "PUBLISHING", "excel_published": False,
                      "added_record_count": len(additions)}
            publication = stage_publication(self.db, dataset=dataset,
                working_copy=working, version=version, base_version_id=current.id,
                result=result, agent_run=self.datasets.agent_run)
        return {**result, "people": additions, "_publication_id": publication.id}

    def continue_publish(self, request: FundingUsageContinuePublish):
        if request.preview_digest != self._continue_digest(request):
            fail("FUNDING_PREVIEW_CHANGED", "追加预览与发布候选不一致。", 409)
        dataset, working, current, origin, _ = self._continue_target(request, lock=True)
        if (origin.extraction_policy_json or {}).get("structured_dataset_export_mode") == "funding-usage-v3":
            return self._continue_publish_v3(request, dataset, working, current)
        request_digest = self._continue_digest(request)
        for previous in self.db.query(DocumentVersion).filter_by(document_id=working.document_id).all():
            policy = previous.extraction_policy_json or {}
            if policy.get("funding_usage_continue_request_key") == request.request_key:
                if policy.get("funding_usage_continue_digest") != request_digest:
                    fail("FUNDING_REQUEST_CONFLICT", "请求键已经用于其他登记或基准版本。", 409)
                export = self.db.query(StructuredDatasetExport).filter_by(
                    dataset_id=dataset.id, document_version_id=previous.id).one_or_none()
                if export is None:
                    fail("FUNDING_PUBLICATION_PENDING", "上次追加尚未正式发布。", 423)
                return {**self.datasets._export_result(export, reused=True),
                        "status": "COMPLETED", "added_record_count": len(request.rows)}
        dataset, working, current, path, existing_rows, additions, history = self._continue_evaluate(
            request, lock=True)
        self.datasets._assert_task_write_access(dataset)
        stored_by_sequence = {row["sequence"]: row for row in dataset.base_rows_json or []}
        combined = [{**row, "_funding_row_id": (stored_by_sequence.get(row["sequence"], {})
                     .get("_funding_row_id") or f"legacy:{row['sequence']}")}
                    for row in existing_rows]
        added_rows = [{**{key: row[key] for key in _KEYS}, "_funding_row_id": str(uuid4())}
                      for row in additions]
        combined.extend(added_rows)
        prior_by_name = {}
        for item in existing_rows:
            prior_by_name[item["name"]] = item["balance"]
        new_keys = [self._registration_key(request.image.source_sha256, row)
                    for row in request.rows]
        entries = list((current.extraction_policy_json or {}).get("funding_usage_entries") or [])
        for source, row, key, stored in zip(request.rows, additions, new_keys, added_rows):
            entries.append({"registration_key": key, "sequence": row["sequence"],
                "source_row_id": stored["_funding_row_id"],
                "name": row["name"], "image_document_id": str(request.image.document_id),
                "image_document_version_id": str(request.image.document_version_id),
                "image_sha256": request.image.source_sha256,
                "image_extraction_run_id": str(request.image.extraction_run_id),
                "name_ref": source.name_ref.model_dump(),
                "registration_ref": source.registration_ref.model_dump() if source.registration_ref else None,
                "amounts": [candidate.model_dump() for candidate in source.amounts],
                "balance_before": prior_by_name.get(row["name"]),
                "balance_after": row["balance"], "status": row["status"]})
            prior_by_name[row["name"]] = row["balance"]
        with tempfile.TemporaryDirectory(prefix="file-agent-funding-continue-") as directory:
            output = Path(directory) / request.filename
            book = load_workbook(path)
            try:
                append_summary_rows(book, additions)
                rebuild_personal_sheets(book, combined)
                book.save(output)
            finally:
                book.close()
            file_sha = self.datasets.storage.sha256_file(output)
            dataset.base_rows_json = combined
            dataset.target_layout_json = {"header_range": "A1:G1",
                "data_range": f"A2:G{max(2, len(combined) + 1)}"}
            dataset.revision += 1
            export_id = str(uuid4())
            version_number = int(self.db.query(func.max(DocumentVersion.version_number)).filter(
                DocumentVersion.document_id == working.document_id).scalar() or 0) + 1
            storage_path = self.datasets.storage.dataset_working_version_relative_path(
                dataset_id=dataset.id, export_id=export_id, version_number=version_number)
            self.datasets.storage.import_working_copy(source=output, relative_path=storage_path,
                                                       expected_sha256=file_sha)
            working, version = self.datasets._publish_working_version(
                dataset=dataset, working_copy=working, parent_version=current,
                storage_path=storage_path, version_number=version_number,
                sha256=file_sha, size_bytes=output.stat().st_size)
            version.extraction_policy_json = {**dict(current.extraction_policy_json or {}),
                "structured_dataset_export_mode": "funding-usage-v2",
                "funding_usage_history": history,
                "funding_usage_entries": entries,
                "funding_usage_continue_request_key": request.request_key,
                "funding_usage_continue_digest": request_digest,
                "image_document_version_id": str(request.image.document_version_id),
                "image_extraction_run_id": str(request.image.extraction_run_id)}
            export = StructuredDatasetExport(id=export_id, dataset_id=dataset.id,
                dataset_revision=dataset.revision, storage_path=storage_path,
                sha256=file_sha, size_bytes=output.stat().st_size,
                document_version_id=version.id)
            self.db.add(export)
            self.db.flush()
            result = {**self.datasets._export_result(export, reused=False),
                      "status": "PUBLISHING", "excel_published": False,
                      "added_record_count": len(additions)}
            publication = stage_publication(self.db, dataset=dataset,
                working_copy=working, version=version, base_version_id=current.id,
                result=result, agent_run=self.datasets.agent_run)
        return {**result, "people": additions, "_publication_id": publication.id}

    @staticmethod
    def _render(path: Path, baseline_path: Path, people: list[dict], funding=None):
        style_book = load_workbook(baseline_path, read_only=False, data_only=True)
        try:
            style_sheet = style_book["汇总"]
            if funding is not None:
                book = render_book(people, funding, style_sheet=style_sheet)
                try:
                    book.save(path)
                finally:
                    book.close()
                return
            book = Workbook()
            summary = book.active
            summary.title = "使用登记汇总"
            for column, label in enumerate(_HEADERS, 1):
                cell = summary.cell(1, column, label)
                sample = style_sheet.cell(1, min(column, style_sheet.max_column))
                if sample.has_style:
                    _copy_cell_style(sample, cell)
                else:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="DDEBF7")
                summary.column_dimensions[cell.column_letter].width = [8, 16, 48, 18, 22, 20, 24][column - 1]
            summary.freeze_panes = "A2"
            for row in people:
                number = row["sequence"] + 1
                for column, key in enumerate(_KEYS, 1):
                    value = row[key]
                    if isinstance(value, str):
                        value = _safe_spreadsheet_text(value)
                    cell = summary.cell(number, column, value)
                    sample = style_sheet.cell(2, min(column, style_sheet.max_column))
                    if sample.has_style:
                        _copy_cell_style(sample, cell)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if column in {5, 6}:
                        cell.number_format = '#,##0.00;[Red]-#,##0.00'
                lines = max(str(row["registration"] or "").count("\n"),
                            str(row["date"] or "").count("\n")) + 1
                summary.row_dimensions[number].height = min(180, max(28, 18 * lines))
            rebuild_personal_sheets(book, people)
            book.save(path)
            book.close()
        finally:
            style_book.close()

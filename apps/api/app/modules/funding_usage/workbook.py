"""成果资助工作簿的受控首表读取与个人页投影。"""

from __future__ import annotations

from copy import copy
from decimal import Decimal, InvalidOperation

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.modules.structured_datasets.service import fail
from app.modules.structured_extraction.export import _safe_spreadsheet_text


HEADERS = ["序号", "姓名", "使用登记原文", "日期", "本次使用金额合计（元）", "扣减后余额（元）", "核对状态"]
KEYS = ["sequence", "name", "registration", "date", "usage_amount", "balance", "status"]


def money(value, *, cell: str):
    if value is None:
        return None
    try:
        amount = Decimal(str(value))
        if isinstance(value, bool) or not amount.is_finite() or amount.as_tuple().exponent < -2:
            raise InvalidOperation
        return amount
    except (InvalidOperation, ValueError):
        fail("FUNDING_AMOUNT_INVALID", "使用金额或余额必须是最多两位小数的数字。", 422, cell=cell)


def read_summary_rows(book, *, allow_new_rows=False):
    if not book.sheetnames or book.sheetnames[0] != "使用登记汇总":
        fail("FUNDING_SUMMARY_INVALID", "第一张工作表必须是“使用登记汇总”。", 422)
    sheet = book.worksheets[0]
    if [sheet.cell(1, col).value for col in range(1, 8)] != HEADERS:
        fail("FUNDING_SUMMARY_INVALID", "成果资助汇总表 A–G 标题与当前版本不一致。", 422)
    rows = []
    seen_sequences = set()
    for row_number in range(2, sheet.max_row + 1):
        values = [sheet.cell(row_number, col).value for col in range(1, 8)]
        if all(value is None for value in values):
            continue
        sequence, name = values[:2]
        new_row = allow_new_rows and sequence is None
        if ((not new_row and (not isinstance(sequence, int) or isinstance(sequence, bool)
                or sequence < 1 or sequence in seen_sequences)) or not isinstance(name, str)
                or not name.strip()):
            fail("FUNDING_ROW_ID_INVALID",
                 f"汇总页第 {row_number} 行的 A 列序号或 B 列姓名无效；人工新增行的 A 列可留空。", 422,
                 row_number=row_number)
        if not new_row:
            seen_sequences.add(sequence)
        for col in (5, 6):
            amount = money(values[col - 1], cell=f"{get_column_letter(col)}{row_number}")
            if col == 5 and amount is not None and (amount < 0 or amount > Decimal("1000000000")):
                fail("FUNDING_AMOUNT_INVALID", "确认使用金额必须是非负且不超过十亿元的金额。", 422,
                     cell=f"E{row_number}")
        row = dict(zip(KEYS, values))
        if allow_new_rows:
            row["_sheet_row_number"] = row_number
        if new_row:
            row["_manual_input_row"] = row_number
        rows.append(row)
    return rows


def _summary_row_height(row):
    lines = max(str(row["registration"] or "").count("\n"),
                str(row["date"] or "").count("\n")) + 1
    return min(180, max(28, 18 * lines))


def repair_summary_row_styles(book, rows):
    """只补旧追加行缺失的样式，保留已有的人工格式。"""
    summary = book.worksheets[0]
    if not rows:
        fail("FUNDING_SUMMARY_INVALID", "目标汇总表没有可供追加的现有数据行。", 422)
    template_row = next((index for index, row in enumerate(rows, 2)
                         if all(summary.cell(index, column).has_style
                                for column in range(1, 8))), 2)
    for index, row in enumerate(rows, 2):
        for column in range(1, 8):
            source = summary.cell(template_row, column)
            target = summary.cell(index, column)
            if not target.has_style:
                target._style = copy(source._style)
                target.alignment = Alignment(vertical="top", wrap_text=True)
        if summary.row_dimensions[index].height is None:
            summary.row_dimensions[index].height = _summary_row_height(row)


def append_summary_rows(book, rows):
    """新增登记沿用最近一条真实数据行的格式，并按内容设置行高。"""
    summary = book.worksheets[0]
    existing = read_summary_rows(book)
    if existing:
        repair_summary_row_styles(book, existing)
    template_row = next((index for index in range(summary.max_row, 1, -1)
                         if summary.cell(index, 1).value is not None
                         and summary.cell(index, 2).value is not None), None)
    for row in rows:
        summary.append([_safe_spreadsheet_text(row[key]) if isinstance(row[key], str)
                        else row[key] for key in KEYS])
        target_row = summary.max_row
        for column in range(1, 8):
            target = summary.cell(target_row, column)
            if template_row is not None:
                target._style = copy(summary.cell(template_row, column)._style)
            target.alignment = Alignment(vertical="top", wrap_text=True)
            if column in (5, 6) and template_row is None:
                target.number_format = '#,##0.00;[Red]-#,##0.00'
        summary.row_dimensions[target_row].height = _summary_row_height(row)
        template_row = target_row


def rebuild_personal_sheets(book, rows, *, allow_removed_people=False, original_rows=None):
    """按汇总页逐行生成个人页；第一行出现的姓名固定个人页顺序。"""
    people = []
    for row in rows:
        if row["name"] not in people:
            people.append(row["name"])
    if len(book.worksheets) > len(people) + 1 and not allow_removed_people:
        fail("FUNDING_PERSONAL_SHEET_INVALID", "个人工作表数量超出汇总页人员数。", 422)
    original_people = []
    for row in original_rows if original_rows is not None else rows:
        if row["name"] not in original_people:
            original_people.append(row["name"])
    titles = {name: book.worksheets[index + 1].title
              for index, name in enumerate(original_people) if index + 1 < len(book.worksheets)}
    summary = book.worksheets[0]
    source_columns = (1, 3, 4, 5, 6, 7)
    header_styles = [copy(summary.cell(1, column)._style) for column in source_columns]
    data_styles = [copy(summary.cell(2, column)._style) for column in source_columns]
    while len(book.worksheets) > 1:
        book.remove(book.worksheets[-1])
    for name in people:
        title = titles.get(name)
        if title is None:
            title = name[:31]
            if not title or any(char in title for char in '[]:*?/\\') or title in book.sheetnames:
                title = f"人员{next(row['sequence'] for row in rows if row['name'] == name)}"
        sheet = book.create_sheet(title)
        sheet.append(["序号", "使用登记原文", "日期", "确认使用金额（元）", "扣减后余额（元）", "核对说明"])
        own_rows = [row for row in rows if row["name"] == name]
        for row in own_rows:
            sheet.append([row["sequence"], row["registration"], row["date"],
                          row["usage_amount"], row["balance"], row["status"]])
        sheet.append([None, "当前可信余额（元）", None, None, own_rows[-1]["balance"]])
        sheet.freeze_panes = "A2"
        for col, width in enumerate((9, 48, 24, 24, 22, 30), 1):
            sheet.column_dimensions[sheet.cell(1, col).column_letter].width = width
            sheet.cell(1, col)._style = copy(header_styles[col - 1])
            if not sheet.cell(1, col).has_style:
                sheet.cell(1, col).font = Font(bold=True)
        for cells in sheet.iter_rows(min_row=2):
            for col, cell in enumerate(cells, 1):
                cell._style = copy(data_styles[col - 1])
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row_number, row in enumerate(own_rows, 2):
            lines = max(str(row["registration"] or "").count("\n"),
                        str(row["date"] or "").count("\n"),
                        str(row["status"] or "").count("\n")) + 1
            sheet.row_dimensions[row_number].height = min(180, max(28, 18 * lines))
        sheet.row_dimensions[sheet.max_row].height = 28
        for row_number in range(2, sheet.max_row + 1):
            sheet.cell(row_number, 4).number_format = '#,##0.00;[Red]-#,##0.00'
            sheet.cell(row_number, 5).number_format = '#,##0.00;[Red]-#,##0.00'

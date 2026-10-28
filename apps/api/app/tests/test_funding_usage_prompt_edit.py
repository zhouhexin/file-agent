"""提示词修改余额表与人工回传的共同金额口径。"""

from hashlib import sha256
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app.core.config import get_settings
from app.db.models import DocumentExtractionRun, DocumentVersion
from app.tests.test_funding_usage import _request
from app.tests.test_structured_datasets import demo, seed_excel_target, seed_workbuddy_document


def _published_book(client, sessions, user, monkeypatch):
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    text = "张三 使用登记 报销差旅30元"
    image_id = seed_workbuddy_document(sessions, filename="提示词首图.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1, "text": text}])
    header = [None] * 14
    header[1], header[13] = "姓名", "资助总额（元）"
    people = []
    for name, amount in (("张三", 100), ("李四", 200)):
        row = [None] * 14
        row[1], row[13] = name, amount
        people.append(row)
    workbook_id = seed_excel_target(sessions, filename="提示词来源.xlsx",
        sheets={"汇总": [header, *people]})
    with sessions() as db:
        image_version = db.query(DocumentVersion).filter_by(document_id=image_id).one()
        workbook_version = db.query(DocumentVersion).filter_by(document_id=workbook_id).one()
        image_run = db.query(DocumentExtractionRun).filter_by(document_id=image_id).one()
    payload = _request(text, amount="30", registration="报销差旅30元").model_dump(mode="json")
    payload["image"] = {"document_id": image_id,
        "document_version_id": image_version.id, "source_sha256": image_version.sha256,
        "extraction_run_id": image_run.id}
    payload["workbook"] = {"document_id": workbook_id,
        "document_version_id": workbook_version.id, "source_sha256": workbook_version.sha256}
    payload["filename"] = "提示词余额表.xlsx"
    payload["request_key"] = str(uuid4())
    preview = client.post("/api/funding-usage/preview", headers=user, json=payload)
    assert preview.status_code == 200, preview.text
    result = client.post("/api/funding-usage/publish", headers=user,
        json={**payload, "preview_digest": preview.json()["preview_digest"]})
    assert result.status_code == 200, result.text
    return result.json()


def _target(client, user, filename):
    response = client.post("/api/funding-usage/prompt/target", headers=user,
                           json={"filename": filename})
    assert response.status_code == 200, response.text
    return response.json()


def _prompt_payload(target, *, edits=(), totals=(), names=(), instruction="请修改余额表"):
    return {"filename": target["filename"],
        "base_document_version_id": target["base_document_version_id"],
        "base_sha256": target["base_sha256"], "base_revision": target["base_revision"],
        "request_key": str(uuid4()), "user_request": instruction,
        "row_edits": list(edits), "total_edits": list(totals),
        "name_edits": list(names)}


def _prompt_publish(client, user, payload):
    checked = client.post("/api/funding-usage/prompt/preview", headers=user, json=payload)
    assert checked.status_code == 200, checked.text
    saved = client.post("/api/funding-usage/prompt/publish", headers=user,
        json={**payload, "preview_digest": checked.json()["preview_digest"]})
    assert saved.status_code == 200, saved.text
    assert saved.json()["status"] == "COMPLETED"
    return checked.json(), saved.json()


def test_prompt_add_update_delete_and_other_funding_manual_return(demo, monkeypatch):
    client, sessions, users = demo
    user = users[0]
    initial = _published_book(client, sessions, user, monkeypatch)
    target = _target(client, user, initial["filename"])
    old = target["rows"][0]
    payload = _prompt_payload(target, edits=[
        {"action": "UPDATE", "name": "张三", "source_row_id": old["_funding_row_id"],
         "usage_amount": 40},
        {"action": "ADD", "name": "张三", "after_row_id": old["_funding_row_id"],
         "registration": "新增会议支出10元", "date": "2026-10-26", "usage_amount": 10},
    ], totals=[{"name": "张三", "field": "other_total", "value": 25}],
        instruction="张三原差旅支出改为40元，新增会议支出10元，其他资助25元")
    preview, saved = _prompt_publish(client, user, payload)
    assert (preview["added_record_count"], preview["updated_record_count"],
            preview["total_change_count"]) == (1, 1, 1)
    book = load_workbook(BytesIO(client.get(saved["download_url"], headers=user).content))
    assert [book["使用登记汇总"].cell(2, col).value for col in (3, 4, 5, 6)] == [100, 25, 50, 75]
    assert [book["张三"].cell(row, 4).value for row in (2, 3)] == [40, 10]
    assert book["张三"]["B6"].value == 25
    book.close()
    repeated = client.post("/api/funding-usage/prompt/publish", headers=user,
        json={**payload, "preview_digest": preview["preview_digest"]})
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["document_version_id"] == saved["document_version_id"]
    stale = client.post("/api/funding-usage/prompt/preview", headers=user,
        json={**payload, "request_key": str(uuid4())})
    assert stale.status_code == 409, stale.text
    assert stale.json()["error"]["code"] == "FUNDING_TARGET_VERSION_CONFLICT"

    target = _target(client, user, initial["filename"])
    delete = _prompt_payload(target, edits=[{"action": "DELETE", "name": "张三",
        "source_row_id": old["_funding_row_id"]}], instruction="删除张三的原差旅登记")
    checked, removed = _prompt_publish(client, user, delete)
    assert checked["deleted_record_count"] == 1
    book = load_workbook(BytesIO(client.get(removed["download_url"], headers=user).content))
    assert [book["使用登记汇总"].cell(2, col).value for col in (4, 5, 6)] == [25, 10, 115]
    assert book["张三"]["A2"].value == 1
    assert book["张三"]["D2"].value == 10
    assert book["张三"]["B5"].value == 25
    book["张三"]["B5"] = 30
    output = BytesIO()
    book.save(output)
    book.close()
    edited = output.getvalue()
    returned = client.post("/api/excel-writeback/manual-revisions/publish", headers=user,
        data={"filename": removed["filename"], "request_key": str(uuid4()),
              "base_document_version_id": removed["document_version_id"],
              "base_sha256": removed["sha256"], "base_revision": str(removed["revision"]),
              "file_sha256": sha256(edited).hexdigest()},
        files={"file": (removed["filename"], edited,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert returned.status_code == 200, returned.text
    book = load_workbook(BytesIO(client.get(returned.json()["download_url"], headers=user).content))
    assert [book["使用登记汇总"].cell(2, col).value for col in (3, 4, 5, 6)] == [100, 30, 10, 120]
    assert [book["张三"].cell(row, 2).value for row in (4, 5, 6, 7)] == [100, 30, 10, 120]
    book.close()

    next_text = "张三 使用登记 报销新差旅5元"
    next_id = seed_workbuddy_document(sessions, filename="提示词后续图.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1,
                                           "text": next_text}])
    with sessions() as db:
        next_version = db.query(DocumentVersion).filter_by(document_id=next_id).one()
        next_run = db.query(DocumentExtractionRun).filter_by(document_id=next_id).one()
    follow = {"image": {"document_id": next_id, "document_version_id": next_version.id,
                        "source_sha256": next_version.sha256, "extraction_run_id": next_run.id},
              "filename": removed["filename"],
              "base_document_version_id": returned.json()["document_version_id"],
              "base_sha256": returned.json()["sha256"],
              "base_revision": returned.json()["revision"],
              "request_key": str(uuid4()),
              "rows": _request(next_text, amount="5", registration="报销新差旅5元")
                      .model_dump(mode="json")["rows"]}
    checked = client.post("/api/funding-usage/continue/preview", headers=user, json=follow)
    assert checked.status_code == 200, checked.text
    appended = client.post("/api/funding-usage/continue/publish", headers=user,
        json={**follow, "preview_digest": checked.json()["preview_digest"]})
    assert appended.status_code == 200, appended.text
    book = load_workbook(BytesIO(client.get(appended.json()["download_url"], headers=user).content))
    assert [book["使用登记汇总"].cell(2, col).value for col in (4, 5, 6)] == [30, 15, 115]
    book.close()


def test_prompt_edit_rejects_ambiguous_or_invalid_actions(demo, monkeypatch):
    client, sessions, users = demo
    initial = _published_book(client, sessions, users[0], monkeypatch)
    target = _target(client, users[0], initial["filename"])
    row_id = target["rows"][0]["_funding_row_id"]
    for edit in (
        {"action": "UPDATE", "name": "张三", "usage_amount": 5},
        {"action": "DELETE", "name": "张三", "source_row_id": row_id, "usage_amount": 5},
        {"action": "ADD", "name": "不存在", "registration": "支出5元", "usage_amount": 5},
        {"action": "UPDATE", "name": "李四", "source_row_id": row_id, "usage_amount": 5},
    ):
        response = client.post("/api/funding-usage/prompt/preview", headers=users[0],
            json=_prompt_payload(target, edits=[edit]))
        assert response.status_code == 422, response.text


def test_prompt_add_and_remove_first_detail_keeps_teacher_on_summary(demo, monkeypatch):
    client, sessions, users = demo
    user = users[0]
    initial = _published_book(client, sessions, user, monkeypatch)
    target = _target(client, user, initial["filename"])
    assert [row["name"] for row in target["rows"]] == ["张三"]
    added_payload = _prompt_payload(target, edits=[{
        "action": "ADD", "name": "李四", "registration": "会议支出5元",
        "usage_amount": 5}], instruction="给李四新增会议支出5元")
    _, added = _prompt_publish(client, user, added_payload)
    book = load_workbook(BytesIO(client.get(added["download_url"], headers=user).content))
    assert book.sheetnames == ["使用登记汇总", "张三", "李四"]
    assert [book["使用登记汇总"].cell(3, col).value for col in (3, 4, 5, 6)] == [200, None, 5, 195]
    book.close()
    target = _target(client, user, initial["filename"])
    row_id = next(row["_funding_row_id"] for row in target["rows"] if row["name"] == "李四")
    delete_payload = _prompt_payload(target, edits=[{
        "action": "DELETE", "name": "李四", "source_row_id": row_id}],
        instruction="删除李四的会议支出5元")
    _, removed = _prompt_publish(client, user, delete_payload)
    book = load_workbook(BytesIO(client.get(removed["download_url"], headers=user).content))
    assert book.sheetnames == ["使用登记汇总", "张三"]
    assert [book["使用登记汇总"].cell(3, col).value for col in (3, 4, 5, 6)] == [200, None, 0, 200]
    book.close()


def test_prompt_rename_updates_teacher_identity_and_keeps_manual_return(demo, monkeypatch):
    client, sessions, users = demo
    user = users[0]
    initial = _published_book(client, sessions, user, monkeypatch)
    target = _target(client, user, initial["filename"])
    source_row_id = target["rows"][0]["_funding_row_id"]
    with sessions() as db:
        from app.db.models import StructuredDataset
        before_people = db.get(StructuredDataset, initial["dataset_id"]).base_rows_json
        before_person_ids = [row["_funding_row_id"] for row in before_people]
    payload = _prompt_payload(target, names=[
        {"old_name": "张三", "new_name": "张山"},
        {"old_name": "李四", "new_name": "李司"},
    ], instruction="纠正两位教师的手写姓名")
    preview, saved = _prompt_publish(client, user, payload)
    assert preview["renamed_person_count"] == saved["renamed_person_count"] == 2
    assert [person["name"] for person in preview["people"]] == ["张山", "李司"]
    repeated = client.post("/api/funding-usage/prompt/publish", headers=user,
        json={**payload, "preview_digest": preview["preview_digest"]})
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["document_version_id"] == saved["document_version_id"]
    book = load_workbook(BytesIO(client.get(saved["download_url"], headers=user).content))
    assert book.sheetnames == ["使用登记汇总", "张山"]
    assert [book["使用登记汇总"].cell(row, 2).value for row in (2, 3)] == ["张山", "李司"]
    assert [book["使用登记汇总"].cell(2, col).value for col in (3, 5, 6)] == [100, 30, 70]
    assert book["张山"].cell(2, 2).value == "报销差旅30元"
    book["张山"]["B5"] = 9
    output = BytesIO()
    book.save(output)
    book.close()

    with sessions() as db:
        from app.db.models import StructuredDataset
        version = db.get(DocumentVersion, saved["document_version_id"])
        dataset = db.get(StructuredDataset, saved["dataset_id"])
        policy = version.extraction_policy_json
        assert set(policy["funding_usage_funding_v3"]) == {"张山", "李司"}
        assert policy["funding_usage_person_order_v3"] == ["张山", "李司"]
        assert policy["funding_usage_rows_v3"][0]["_funding_row_id"] == source_row_id
        assert policy["funding_usage_entries"][0]["name"] == "张山"
        assert policy["funding_usage_entries"][0]["name_ref"]["quote"] == "张三"
        assert [row["name"] for row in dataset.base_rows_json] == ["张山", "李司"]
        assert [row["_funding_row_id"] for row in dataset.base_rows_json] == before_person_ids
    new_target = _target(client, user, initial["filename"])
    assert new_target["rows"][0]["name"] == "张山"
    assert new_target["rows"][0]["_funding_row_id"] == source_row_id
    edited = output.getvalue()
    returned = client.post("/api/excel-writeback/manual-revisions/publish", headers=user,
        data={"filename": saved["filename"], "request_key": str(uuid4()),
              "base_document_version_id": saved["document_version_id"],
              "base_sha256": saved["sha256"], "base_revision": str(saved["revision"]),
              "file_sha256": sha256(edited).hexdigest()},
        files={"file": (saved["filename"], edited,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert returned.status_code == 200, returned.text
    book = load_workbook(BytesIO(client.get(returned.json()["download_url"], headers=user).content))
    assert book["使用登记汇总"].cell(2, 2).value == "张山"
    assert [book["使用登记汇总"].cell(2, col).value for col in (4, 6)] == [9, 79]
    book.close()

    next_text = "张山 使用登记 报销车费5元"
    next_id = seed_workbuddy_document(sessions, filename="更名后追加图片.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1,
                                           "text": next_text}])
    with sessions() as db:
        next_version = db.query(DocumentVersion).filter_by(document_id=next_id).one()
        next_run = db.query(DocumentExtractionRun).filter_by(document_id=next_id).one()
    follow = {"image": {"document_id": next_id,
                        "document_version_id": next_version.id,
                        "source_sha256": next_version.sha256,
                        "extraction_run_id": next_run.id},
              "filename": saved["filename"],
              "base_document_version_id": returned.json()["document_version_id"],
              "base_sha256": returned.json()["sha256"],
              "base_revision": returned.json()["revision"],
              "request_key": str(uuid4()),
              "rows": _request(next_text, amount="5", registration="报销车费5元",
                               name="张山").model_dump(mode="json")["rows"]}
    checked = client.post("/api/funding-usage/continue/preview", headers=user, json=follow)
    assert checked.status_code == 200, checked.text
    appended = client.post("/api/funding-usage/continue/publish", headers=user,
        json={**follow, "preview_digest": checked.json()["preview_digest"]})
    assert appended.status_code == 200, appended.text
    book = load_workbook(BytesIO(client.get(appended.json()["download_url"], headers=user).content))
    assert [book["使用登记汇总"].cell(2, col).value for col in (2, 4, 5, 6)] == ["张山", 9, 35, 74]
    assert book["张山"].cell(3, 2).value == "报销车费5元"
    book.close()


def test_prompt_rename_rejects_missing_duplicate_and_invalid_names(demo, monkeypatch):
    client, sessions, users = demo
    initial = _published_book(client, sessions, users[0], monkeypatch)
    target = _target(client, users[0], initial["filename"])
    for old_name, new_name in (
        ("不存在", "张山"), ("张三", "张三"), ("张三", "李四"),
        ("张三", "使用登记汇总"), ("张三", "张/山"), ("张三", " 张山"),
        ("张三", "李 四"),
    ):
        response = client.post("/api/funding-usage/prompt/preview", headers=users[0],
            json=_prompt_payload(target, names=[{"old_name": old_name, "new_name": new_name}]))
        assert response.status_code == 422, response.text

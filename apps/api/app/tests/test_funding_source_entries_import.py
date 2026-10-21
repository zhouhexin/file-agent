"""单人资助 Excel 追加到成果明细汇总表的端到端校验。"""

from io import BytesIO
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from app.core.config import get_settings
from app.db.models import DocumentVersion
from app.modules.funding_usage.source_entries import (PERSON_HEADERS, SUMMARY_HEADERS,
    PAPER_HEADERS, MEETING_HEADERS, _candidates, _duplicate_status, apply)
from app.tests.test_structured_datasets import demo, seed_excel_target


def _row(size, values):
    row = [None] * size
    for index, value in values.items():
        row[index - 1] = value
    return row


def test_untitled_meeting_does_not_duplicate_existing_titled_meeting():
    book = Workbook()
    sheet = book.active
    sheet.append(["教师成果明细"])
    sheet.append(["年份", "序号", "成果类型", "期刊/会议名称", "论文/成果名称",
                  "论文类型", "资助类型", "发表/召开时间", "资助类别",
                  "资助金额（元）", "地点", "状态", "备注"])
    sheet.append(_row(13, {1: "2020", 3: "会议", 4: "已有会议名称",
                           8: "2020年10月16日-10月18日", 9: "差旅", 10: 1000}))
    sheet.append(_row(13, {1: "合计", 10: 1000}))
    item = {"kind": "会议", "cells": ("2020", None, "会议", None, None, None, None,
            "2020年10月16日-10月18日", "差旅", 1000, None, None, None)}
    try:
        assert _duplicate_status(sheet, item) == "POSSIBLE_DUPLICATE"
    finally:
        book.close()


def test_missing_paper_owner_is_not_inferred_from_meeting_or_filename():
    book = Workbook()
    paper = book.active
    paper.title = "论文"
    paper.append(PAPER_HEADERS)
    paper.append([1, None, "期刊甲", "缺姓名论文", None, "版面费", "2025.3.8", "CCF-A", 1000])
    book.create_sheet("Sheet2")
    meeting = book.create_sheet("会议")
    meeting.append(MEETING_HEADERS)
    meeting.append([1, "温营坤", "会议甲", None, "2020年10月16日", "否", "差旅", 500, None])
    try:
        rows = _candidates(book, "科研成果资助汇总表-温营坤.xlsx")
        assert rows[0]["name"] == "" and not rows[0]["owner_inferred"]
        assert rows[1]["name"] == "温营坤"
    finally:
        book.close()


def test_source_text_and_sequence_are_saved_as_literals(tmp_path):
    source = Workbook()
    paper = source.active
    paper.title = "论文"
    paper.append(PAPER_HEADERS)
    paper.append(["=1+1", "张三", "期刊甲", "=HYPERLINK(\"https://example.com\")",
                  None, "版面费", "2025.3.8", "CCF-A", 1000])
    source.create_sheet("Sheet2")
    source.create_sheet("会议").append(MEETING_HEADERS)
    target = Workbook()
    summary = target.active
    summary.title = "汇总"
    summary.append(SUMMARY_HEADERS)
    summary.append(_row(14, {1: 1, 2: "张三", 3: 0, 4: 0, 5: 0, 14: 0}))
    detail = target.create_sheet("张三")
    detail.append(["张三 · 科研成果资助明细"])
    detail.append(PERSON_HEADERS)
    detail.append(_row(13, {1: "合计", 10: 0}))
    try:
        entries = _candidates(source, "科研成果资助汇总表-张三.xlsx")
        entries[0]["status"] = "ADDED"
        apply(target, entries, "张三")
        output = tmp_path / "safe.xlsx"
        target.save(output)
    finally:
        source.close()
        target.close()
    saved = load_workbook(output)
    try:
        assert saved["张三"]["B3"].value == "'=1+1"
        assert saved["张三"]["E3"].value == "'=HYPERLINK(\"https://example.com\")"
        assert saved["张三"]["B3"].data_type == "s"
        assert saved["张三"]["E3"].data_type == "s"
    finally:
        saved.close()


def test_single_teacher_import_skips_existing_papers_and_requires_owner_confirmation(demo, monkeypatch):
    client, sessions, users = demo
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    name = "温营坤"
    title_one = "第一篇已有论文"
    title_two = "第二篇已有论文"
    target_id = seed_excel_target(sessions, filename="教师成果明细汇总表.xlsx", sheets={
        "汇总": [list(SUMMARY_HEADERS), _row(14, {1: 1, 2: name, 3: 2, 5: 2,
                                                     12: 1, 13: 1, 14: 21000}),
                 _row(14, {1: "总计", 2: "1 人", 3: 2, 4: 0, 5: 2,
                           7: 0, 12: 1, 13: 1, 14: 21000})],
        name: [[name + " · 科研成果资助明细"], list(PERSON_HEADERS),
               _row(13, {1: "2025", 2: 1, 3: "论文", 4: "期刊甲", 5: title_one,
                         8: "2025.3.8", 10: 10000}),
               _row(13, {1: "2026", 2: 2, 3: "论文", 4: "期刊乙", 5: title_two,
                         8: "2026.1.24", 10: 11000}),
               _row(13, {1: "合计", 4: "共 2 项", 10: 21000})]})
    source_id = seed_excel_target(sessions, filename="科研成果资助汇总表-温营坤.xlsx", sheets={
        "论文": [list(PAPER_HEADERS),
                 [1, name, "期刊甲", title_one, None, "版面费", "2025.3.8", "CCF-A", 10000],
                 [2, name, "期刊乙", title_two, None, "版面费", "2026.1.24", "CCF-B", 11000]],
        "Sheet2": [[None]],
        "会议": [list(MEETING_HEADERS) + [None],
                 [None, None, None, None, "2020年10月16日-10月18日", "否", "差旅", 1000, "有", None]]})
    with sessions() as db:
        source = db.query(DocumentVersion).filter_by(document_id=source_id).one()
        old_target = db.query(DocumentVersion).filter_by(document_id=target_id).one()
        old_version_id, old_sha = old_target.id, old_target.sha256
    payload = {"workbook": {"document_id": source_id,
        "document_version_id": source.id, "source_sha256": source.sha256},
        "filename": "教师成果明细汇总表.xlsx", "request_key": str(uuid4())}

    pending = client.post("/api/funding-usage/source-entries/preview", headers=users[0], json=payload)
    assert pending.status_code == 200, pending.text
    assert pending.json()["added_count"] == 0
    assert [row["status"] for row in pending.json()["rows"]] == [
        "DUPLICATE", "DUPLICATE", "OWNER_CONFIRMATION_REQUIRED"]
    assert client.post("/api/funding-usage/source-entries/publish", headers=users[0],
        json={**payload, "preview_digest": pending.json()["preview_digest"]}).status_code == 422

    confirmed = {**payload, "confirm_inferred_owner": True}
    preview = client.post("/api/funding-usage/source-entries/preview", headers=users[0], json=confirmed)
    assert preview.status_code == 200, preview.text
    assert (preview.json()["added_count"], preview.json()["duplicate_count"],
            preview.json()["proposed_totals"]["research_total"]) == (1, 2, 22000)
    assert preview.json()["rows"][2]["missing_fields"] == ["申请人", "会议名称", "地点"]
    stale = {**confirmed, "request_key": str(uuid4())}
    stale_preview = client.post("/api/funding-usage/source-entries/preview", headers=users[0], json=stale)
    assert stale_preview.status_code == 200, stale_preview.text
    published = client.post("/api/funding-usage/source-entries/publish", headers=users[0],
        json={**confirmed, "preview_digest": preview.json()["preview_digest"]})
    assert published.status_code == 200, published.text
    result = published.json()
    assert result["status"] == "COMPLETED" and result["excel_published"]
    assert result["version_number"] == 2 and result["added_count"] == 1
    book = load_workbook(BytesIO(client.get(result["download_url"], headers=users[0]).content))
    try:
        assert [book["汇总"].cell(2, col).value for col in (3, 4, 5, 7, 14)] == [2, 1, 3, 1, 22000]
        assert [book["汇总"].cell(3, col).value for col in (3, 4, 5, 7, 14)] == [2, 1, 3, 1, 22000]
        assert [book[name].cell(5, col).value for col in (1, 3, 4, 8, 9, 10, 12)] == [
            "2020", "会议", None, "2020年10月16日-10月18日", "差旅", 1000, "待核对"]
        assert book[name]["J6"].value == 22000
    finally:
        book.close()
    with sessions() as db:
        old = db.get(DocumentVersion, old_version_id)
        assert old.sha256 == old_sha
    repeated = client.post("/api/funding-usage/source-entries/publish", headers=users[0],
        json={**confirmed, "preview_digest": preview.json()["preview_digest"]})
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["document_version_id"] == result["document_version_id"]
    stale_publish = client.post("/api/funding-usage/source-entries/publish", headers=users[0],
        json={**stale, "preview_digest": stale_preview.json()["preview_digest"]})
    assert stale_publish.status_code == 409, stale_publish.text
    assert stale_publish.json()["error"]["code"] == "FUNDING_PREVIEW_CHANGED"
    synchronized = client.post("/api/funding-usage/source-summary/preview", headers=users[0],
        json={"workbook": {"document_id": target_id,
              "document_version_id": result["document_version_id"],
              "source_sha256": result["sha256"]}, "request_key": str(uuid4())})
    assert synchronized.status_code == 200, synchronized.text
    assert synchronized.json()["changed_count"] == 0
    rerun = client.post("/api/funding-usage/source-entries/preview", headers=users[0],
        json={**confirmed, "request_key": str(uuid4())})
    assert rerun.status_code == 200, rerun.text
    assert rerun.json()["added_count"] == 0
    assert rerun.json()["duplicate_count"] == 3


def test_unknown_teacher_is_reported_without_creating_person_sheet(demo, monkeypatch):
    client, sessions, users = demo
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    target_id = seed_excel_target(sessions, filename="不含该教师的成果表.xlsx", sheets={
        "汇总": [list(SUMMARY_HEADERS), _row(14, {1: 1, 2: "张三", 3: 0, 4: 0, 5: 0, 14: 0})],
        "张三": [["张三 · 科研成果资助明细"], list(PERSON_HEADERS),
                 _row(13, {1: "合计", 10: 0})]})
    source_id = seed_excel_target(sessions, filename="科研成果资助汇总表-温营坤.xlsx", sheets={
        "论文": [list(PAPER_HEADERS),
                 [1, "温营坤", "期刊甲", "新论文", None, "版面费", "2025.3.8", "CCF-A", 1000]],
        "Sheet2": [[None]], "会议": [list(MEETING_HEADERS)]})
    with sessions() as db:
        source = db.query(DocumentVersion).filter_by(document_id=source_id).one()
    payload = {"workbook": {"document_id": source_id,
        "document_version_id": source.id, "source_sha256": source.sha256},
        "filename": "不含该教师的成果表.xlsx", "request_key": str(uuid4())}
    preview = client.post("/api/funding-usage/source-entries/preview", headers=users[0], json=payload)
    assert preview.status_code == 200, preview.text
    assert preview.json()["added_count"] == 0
    assert preview.json()["proposed_totals"] is None
    assert preview.json()["rows"][0]["status"] == "TARGET_PERSON_MISSING"

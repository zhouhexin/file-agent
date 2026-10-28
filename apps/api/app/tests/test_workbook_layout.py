"""复杂汇总模板的实际布局、保真填充和失败原子性回归。"""

import zipfile

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from app.modules.structured_datasets.service import StructuredDatasetService
from app.modules.structured_extraction.export import _append_xlsx_rows_atomic
from app.modules.structured_extraction.workbook_layout import WorkbookLayout, read_layout


LAYOUT = {"header_range": "A2:AA4", "data_range": "A5:AA6"}


def test_layout_keeps_untitled_column_at_physical_position():
    book = Workbook()
    sheet = book.active
    sheet.append(["姓名", None, "单位"])
    sheet.append(["已有人员", None, "原单位"])
    labels, rows = read_layout(sheet, WorkbookLayout(
        header_range="A1:C1", data_range="A2:C2"))
    assert labels == ["姓名", "〔空白列:B〕", "单位"]
    assert rows == [["已有人员", None, "原单位"]]


def test_layout_rejects_real_header_using_blank_column_marker():
    book = Workbook()
    sheet = book.active
    sheet.append(["姓名", "〔空白列:B〕"])
    with pytest.raises(ValueError, match="系统保留的空白列标记"):
        read_layout(sheet, WorkbookLayout(header_range="A1:B1", data_range="A2:B2"))


def make_template(path):
    """构造与实际模板同类的 27 列三层表头，不使用用户个人信息。"""
    book = Workbook()
    sheet = book.active
    sheet.title = "登记"
    sheet.merge_cells("A1:AA1")
    sheet["A1"] = "教学科研人员业绩汇总"
    for col in range(1, 14):
        sheet.merge_cells(start_row=2, start_column=col, end_row=4, end_column=col)
        sheet.cell(2, col, ["序号", "姓名", "金额", "取得时间", "取得时间"][col - 1] if col <= 5 else f"字段{col}")
    sheet.merge_cells("N2:Y2")
    sheet["N2"] = "业绩"
    sheet.merge_cells("N3:Y3")
    sheet["N3"] = "具体条件"
    for col in range(14, 28):
        sheet.cell(4, col, f"字段{col}" + ("完整条件" * 50 if col == 20 else ""))
    sheet["A5"] = 1
    sheet["A6"] = 2
    sheet["B5"].fill = PatternFill("solid", fgColor="FFFF00")
    sheet.merge_cells("A7:AA7")
    sheet["A7"] = "填写说明，禁止在本行下追加"
    book.create_sheet("附表")["B2"] = "=1+2"
    book.save(path)


def test_merged_headers_reserved_rows_and_package_preservation(tmp_path):
    """只填预留行，序号/尾注/样式/其他工作表与所有非目标 ZIP part 保持不变。"""
    original, output = tmp_path / "原表.xlsx", tmp_path / "结果.xlsx"
    make_template(original)
    service = object.__new__(StructuredDatasetService)
    name, fields, base = service._read_excel_schema(original, "登记", WorkbookLayout(**LAYOUT))
    assert len(fields) == 27 and base == []
    assert fields[3].label == "取得时间 [D]" and fields[4].label == "取得时间 [E]"
    assert len(fields[19].label) > 80
    labels = [f.label for f in fields]
    row = [None, "测试人员", 20] + [None] * 24
    _append_xlsx_rows_atomic(template_path=original, target_path=output, headers=labels,
                             rows=[row], sheet_name=name, layout=LAYOUT)
    with zipfile.ZipFile(original) as before, zipfile.ZipFile(output) as after:
        assert before.namelist() == after.namelist()
        assert [p for p in before.namelist() if before.read(p) != after.read(p)] == ["xl/worksheets/sheet1.xml"]
    workbook = load_workbook(output)
    assert workbook["登记"]["B5"].value == "测试人员"
    assert workbook["登记"]["A5"].value == 1
    assert workbook["登记"]["A7"].value.startswith("填写说明")
    assert workbook["登记"]["B5"].fill.fgColor.rgb == "00FFFF00"
    assert workbook["附表"]["B2"].value == "=1+2"
    workbook.close()
    # 第二次只占第二个预留行，第三次不允许越界到说明下方。
    _append_xlsx_rows_atomic(template_path=output, target_path=output, headers=labels,
                             rows=[row], sheet_name=name, layout=LAYOUT)
    previous = output.read_bytes()
    with pytest.raises(ValueError, match="TARGET_EXCEL_DATA_RANGE_FULL"):
        _append_xlsx_rows_atomic(template_path=output, target_path=output, headers=labels,
                                 rows=[row], sheet_name=name, layout=LAYOUT)
    assert output.read_bytes() == previous


def test_complex_header_requires_explicit_layout_and_rejects_overlaps(tmp_path):
    """不能把大标题误当成单列；填写范围不得侵入表头和合并注释。"""
    path = tmp_path / "表.xlsx"
    make_template(path)
    with pytest.raises(HTTPException) as caught:
        object.__new__(StructuredDatasetService)._read_excel_schema(path, "登记")
    assert caught.value.detail["code"] == "TARGET_EXCEL_LAYOUT_REQUIRED"
    with pytest.raises(ValueError):
        WorkbookLayout(header_range="A2:AA4", data_range="A4:AA6")
    workbook = load_workbook(path)
    with pytest.raises(ValueError, match="合并"):
        read_layout(workbook.active, WorkbookLayout(header_range="A2:AA4", data_range="A5:AA7"))
    workbook.close()


def test_existing_formula_is_not_an_empty_row(tmp_path):
    """无缓存值的公式不能被 data_only 读取误判为空行后覆盖。"""
    path = tmp_path / "表.xlsx"
    make_template(path)
    workbook = load_workbook(path)
    workbook.active["B5"] = "=1+1"
    workbook.save(path)
    labels, _ = read_layout(workbook.active, WorkbookLayout(**LAYOUT))
    workbook.close()
    _append_xlsx_rows_atomic(template_path=path, target_path=path, headers=labels,
                             rows=[[None, "第二行"] + [None] * 25], sheet_name="登记", layout=LAYOUT)
    result = load_workbook(path)
    assert result.active["B5"].value == "=1+1" and result.active["B6"].value == "第二行"
    result.close()


def test_workbuddy_grid_must_match_every_nonempty_header_cell(tmp_path):
    """逐格核对须拒绝缺格或虚构的额外单元格，错误只给数量不泄露模板正文。"""

    from types import SimpleNamespace
    path = tmp_path / "表.xlsx"
    make_template(path)
    service = object.__new__(StructuredDatasetService)

    class FakeQuery:
        """只模拟本测试所需的固定提取元素查询。"""
        def __init__(self, elements):
            self.elements = elements

        def filter_by(self, **_kwargs):
            """提取运行 ID 已由调用者固定。"""
            return self

        def all(self):
            """返回测试的真实坐标候选。"""
            return self.elements

    def element(row, column, text):
        """构造一个 WorkBuddy 提交的表格单元格。"""
        return SimpleNamespace(label="table_cell", text_content=text,
                               metadata_json={"sheet_name": "登记", "row_index": row,
                                              "column_index": column})

    elements = [element(1, 0, "序号")]
    service.db = SimpleNamespace(query=lambda _model: FakeQuery(elements))
    name, fields, _ = service._read_excel_schema(path, "登记", WorkbookLayout(**LAYOUT))
    with pytest.raises(HTTPException) as caught:
        service._assert_workbuddy_schema_matches(
            expected={"run_id": "test", "sheets": {"登记": object()}}, sheet_name=name,
            actual_labels=[field.label for field in fields], layout=WorkbookLayout(**LAYOUT),
            converted_path=path)
    assert caught.value.detail["code"] == "TARGET_EXCEL_CONVERSION_SCHEMA_MISMATCH"
    assert caught.value.detail["missing_cells"] > 0

    # 有实际标题之外的假单元格也不能被无条件忽略。
    workbook = load_workbook(path)
    elements[:] = [element(cell.row - 1, cell.column - 1, str(cell.value))
                   for row in workbook["登记"].iter_rows(min_row=2, max_row=4, min_col=1, max_col=27)
                   for cell in row if cell.value is not None and str(cell.value).strip()]
    workbook.close()
    elements.append(element(3, 0, "虚构标题"))
    with pytest.raises(HTTPException) as caught:
        service._assert_workbuddy_schema_matches(
            expected={"run_id": "test", "sheets": {"登记": object()}}, sheet_name=name,
            actual_labels=[field.label for field in fields], layout=WorkbookLayout(**LAYOUT),
            converted_path=path)
    assert caught.value.detail["unexpected_cells"] == 1

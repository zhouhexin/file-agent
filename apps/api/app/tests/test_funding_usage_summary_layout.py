"""新版成果资助表以个人页为登记来源，首页只汇总金额。"""

from hashlib import sha256
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.core.config import get_settings
from app.db.models import DocumentExtractionRun, DocumentVersion, WorkingCopy
from app.modules.file_lifecycle.storage import FileLifecycleStorageService
from app.tests.test_funding_usage import _request
from app.tests.test_structured_datasets import demo, seed_excel_target, seed_workbuddy_document
from app.modules.funding_usage.summary_workbook import (DETAIL_HEADERS, TOTAL_LABELS,
                                                         read_detail_book, render_book)


def test_personal_sequences_and_totals_are_independent():
    rows = [
        {"name": "张三", "registration": "支出10元", "date": None,
         "usage_amount": 10, "status": "已核对"},
        {"name": "李四", "registration": "支出20元", "date": None,
         "usage_amount": 20, "status": "已核对"},
        {"name": "张三", "registration": "支出30元", "date": None,
         "usage_amount": 30, "status": "已核对"},
    ]
    funding = {"张三": {"research_total": 100, "other_total": 5},
               "李四": {"research_total": 200, "other_total": None}}
    book = render_book(rows, funding)
    try:
        assert [book["张三"].cell(row, 1).value for row in (2, 3)] == [1, 2]
        assert book["李四"]["A2"].value == 1
        assert [book.worksheets[0].cell(2, col).value for col in (3, 4, 5, 6)] == [
            100, 5, 40, 65]
        assert [book.worksheets[0].cell(3, col).value for col in (3, 4, 5, 6)] == [
            200, None, 20, 180]
        read_rows, read_funding = read_detail_book(book, ["张三", "李四"])
        assert [row["name"] for row in read_rows] == ["张三", "张三", "李四"]
        assert read_funding == funding
    finally:
        book.close()


def test_previously_published_empty_personal_sheet_can_still_be_read():
    rows = [{"name": "张三", "registration": "支出10元", "date": None,
             "usage_amount": 10, "status": "已核对"}]
    funding = {"张三": {"research_total": 100, "other_total": None},
               "李四": {"research_total": 200, "other_total": None}}
    book = render_book(rows, funding)
    try:
        assert book.sheetnames == ["使用登记汇总", "张三"]
        empty_sheet = book.create_sheet("李四")  # 兼容旧版已发布的空白个人页。
        empty_sheet.append(DETAIL_HEADERS)
        for index, label in enumerate(TOTAL_LABELS, 3):
            empty_sheet.cell(index, 1, label)
        empty_sheet["B3"] = 200
        empty_sheet["B5"] = 0
        empty_sheet["B6"] = 200
        read_rows, read_funding = read_detail_book(
            book, ["张三", "李四"], required_names={"张三"})
        assert [row["name"] for row in read_rows] == ["张三"]
        assert read_funding == funding
    finally:
        book.close()


def test_refresh_research_totals_keeps_usage_and_supports_next_append(demo, monkeypatch):
    client, sessions, users = demo
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    original_text = "张三 使用登记 报销差旅30元"
    next_text = "张三 使用登记 报销会议10元"
    image_id = seed_workbuddy_document(sessions, filename="资助刷新首图.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1,
                                           "text": original_text}])
    next_image_id = seed_workbuddy_document(sessions, filename="资助刷新续图.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1,
                                           "text": next_text}])

    def workbook_rows(amounts):
        header = [None] * 14
        header[1], header[13] = "姓名", "资助总额（元）"
        result = [header]
        for name, amount in amounts:
            row = [None] * 14
            row[1], row[13] = name, amount
            result.append(row)
        return result

    old_id = seed_excel_target(sessions, filename="刷新前成果表.xlsx", sheets={"汇总": workbook_rows([
        ("张三", 100), ("李四", 200), ("王五", 300), ("赵六", 400), ("钱七", 500)])})
    detail_header = [None] * 13
    detail_header[9] = "资助金额（元）"
    detail_first, detail_second, detail_total = ([None] * 13 for _ in range(3))
    detail_first[9], detail_second[9] = 50, 70
    detail_total[0], detail_total[9] = "合计", 120
    mismatch_item, mismatch_total = ([None] * 13 for _ in range(2))
    mismatch_item[9] = 20
    mismatch_total[0], mismatch_total[9] = "合计", 30
    new_id = seed_excel_target(sessions, filename="刷新后成果表.xlsx", sheets={
        "汇总": workbook_rows([
            ("张三", 100), ("李四", 250), ("王五", 310), ("王五", 320),
            ("赵六", "待核对"), ("额外人员", 600)]),
        "张三": [["张三 · 科研成果资助明细"], detail_header,
                 detail_first, detail_second, detail_total],
        "赵六": [["赵六 · 科研成果资助明细"], detail_header,
                 mismatch_item, mismatch_total]})

    def snapshot(document_id):
        with sessions() as db:
            version = db.query(DocumentVersion).filter_by(document_id=document_id).one()
            return {"document_id": document_id, "document_version_id": version.id,
                    "source_sha256": version.sha256}

    with sessions() as db:
        image_run = db.query(DocumentExtractionRun).filter_by(document_id=image_id).one()
    first = _request(original_text, amount="30", registration="报销差旅30元").model_dump(mode="json")
    first["image"] = {**snapshot(image_id), "extraction_run_id": image_run.id}
    first["workbook"] = snapshot(old_id)
    first["filename"] = "资助刷新余额表.xlsx"
    preview = client.post("/api/funding-usage/preview", headers=users[0], json=first)
    assert preview.status_code == 200, preview.text
    initial = client.post("/api/funding-usage/publish", headers=users[0],
        json={**first, "preview_digest": preview.json()["preview_digest"]})
    assert initial.status_code == 200, initial.text
    initial = initial.json()

    refresh = {"workbook": snapshot(new_id), "filename": initial["filename"],
        "base_document_version_id": initial["document_version_id"],
        "base_sha256": initial["sha256"], "base_revision": initial["revision"],
        "request_key": str(uuid4())}
    checked = client.post("/api/funding-usage/refresh/preview", headers=users[0], json=refresh)
    assert checked.status_code == 200, checked.text
    checked = checked.json()
    assert checked["changed_count"] == 2
    assert checked["review_count"] == 3
    assert checked["source_only_names"] == ["额外人员"]
    personal = next(item for item in checked["people"] if item["name"] == "张三")
    assert (personal["source_basis"], personal["source_summary_total"],
            personal["new_research_total"]) == ("PERSONAL_DETAIL", 100, 120)
    assert {item["name"]: item["status"] for item in checked["people"]} == {
        "张三": "UPDATED", "李四": "UPDATED", "王五": "SOURCE_DUPLICATE",
        "赵六": "DETAIL_TOTAL_MISMATCH", "钱七": "SOURCE_MISSING"}
    assert next(item for item in checked["people"] if item["name"] == "张三")["total_spent"] == 30
    published = client.post("/api/funding-usage/refresh/publish", headers=users[0],
        json={**refresh, "preview_digest": checked["preview_digest"]})
    assert published.status_code == 200, published.text
    published = published.json()
    assert published["status"] == "COMPLETED"
    assert published["changed_count"] == 2
    book = load_workbook(BytesIO(client.get(published["download_url"], headers=users[0]).content))
    assert book.sheetnames == ["使用登记汇总", "张三"]
    assert [book["使用登记汇总"].cell(2, col).value for col in (3, 4, 5, 6)] == [120, None, 30, 90]
    assert [book["使用登记汇总"].cell(3, col).value for col in (3, 5, 6)] == [250, 0, 250]
    assert [book["使用登记汇总"].cell(row, 3).value for row in (4, 5, 6)] == [300, 400, 500]
    assert [book["张三"].cell(row, 2).value for row in (4, 5, 6, 7)] == [120, None, 30, 90]
    assert book["张三"]["D2"].value == 30
    book.close()


    repeated = client.post("/api/funding-usage/refresh/publish", headers=users[0],
        json={**refresh, "preview_digest": checked["preview_digest"]})
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["document_version_id"] == published["document_version_id"]
    stale = client.post("/api/funding-usage/refresh/preview", headers=users[0],
        json={**refresh, "request_key": str(uuid4())})
    assert stale.status_code == 409, stale.text
    assert stale.json()["error"]["code"] == "FUNDING_TARGET_VERSION_CONFLICT"

    with sessions() as db:
        next_run = db.query(DocumentExtractionRun).filter_by(document_id=next_image_id).one()
    append = {"image": {**snapshot(next_image_id), "extraction_run_id": next_run.id},
        "filename": published["filename"],
        "base_document_version_id": published["document_version_id"],
        "base_sha256": published["sha256"], "base_revision": published["revision"],
        "request_key": str(uuid4()),
        "rows": _request(next_text, amount="10", registration="报销会议10元")
                .model_dump(mode="json")["rows"]}
    append_preview = client.post("/api/funding-usage/continue/preview", headers=users[0], json=append)
    assert append_preview.status_code == 200, append_preview.text
    appended = client.post("/api/funding-usage/continue/publish", headers=users[0],
        json={**append, "preview_digest": append_preview.json()["preview_digest"]})
    assert appended.status_code == 200, appended.text
    appended = appended.json()
    book = load_workbook(BytesIO(client.get(appended["download_url"], headers=users[0]).content))
    assert [book["使用登记汇总"].cell(2, col).value for col in (3, 5, 6)] == [120, 40, 80]
    book["张三"]["D2"] = 35
    output = BytesIO()
    book.save(output)
    book.close()
    edited = output.getvalue()
    revised = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": appended["filename"], "request_key": str(uuid4()),
              "base_document_version_id": appended["document_version_id"],
              "base_sha256": appended["sha256"], "base_revision": str(appended["revision"]),
              "file_sha256": sha256(edited).hexdigest()},
        files={"file": (appended["filename"], edited,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert revised.status_code == 200, revised.text
    book = load_workbook(BytesIO(client.get(revised.json()["download_url"], headers=users[0]).content))
    assert [book["使用登记汇总"].cell(2, col).value for col in (3, 5, 6)] == [120, 45, 75]
    assert [book["张三"].cell(row, 2).value for row in (5, 6, 7, 8)] == [120, None, 45, 75]
    book.close()


def test_source_summary_sync_publishes_new_version_and_preserves_detail(demo, monkeypatch):
    client, sessions, users = demo
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    header = [None] * 14
    header[1], header[13] = "姓名", "资助总额（元）"
    summary_row = [None] * 14
    summary_row[1], summary_row[13] = "张三", 100
    detail_header = [None] * 10
    detail_header[9] = "资助金额（元）"
    item = [None] * 10
    item[9] = 120
    footer = [None] * 10
    footer[0], footer[9] = "合计", 120
    source_id = seed_excel_target(sessions, filename="待同步教师成果明细表.xlsx", sheets={
        "汇总": [header, summary_row], "张三": [["张三"], detail_header, item, footer]})
    with sessions() as db:
        original = db.query(DocumentVersion).filter_by(document_id=source_id).one()
        old_version_id, old_sha = original.id, original.sha256
    payload = {"workbook": {"document_id": source_id,
        "document_version_id": old_version_id, "source_sha256": old_sha},
        "request_key": str(uuid4())}
    checked = client.post("/api/funding-usage/source-summary/preview", headers=users[0], json=payload)
    assert checked.status_code == 200, checked.text
    assert checked.json()["changed_count"] == 1
    published = client.post("/api/funding-usage/source-summary/publish", headers=users[0],
        json={**payload, "preview_digest": checked.json()["preview_digest"]})
    assert published.status_code == 200, published.text
    result = published.json()
    assert result["status"] == "COMPLETED" and result["excel_published"]
    assert result["version_number"] == 2
    workbook = load_workbook(BytesIO(client.get(result["download_url"], headers=users[0]).content))
    assert workbook["汇总"]["N2"].value == 120
    assert workbook["张三"]["J4"].value == 120
    workbook.close()
    with sessions() as db:
        original = db.get(DocumentVersion, old_version_id)
        current = db.get(DocumentVersion, result["document_version_id"])
        working = db.query(WorkingCopy).filter_by(document_id=source_id).one()
        assert original.sha256 == old_sha and original.storage_path != current.storage_path
        assert FileLifecycleStorageService().sha256_file(
            FileLifecycleStorageService().working_copy_path(original.storage_path)) == old_sha
        assert working.current_version_id == current.id
    repeated = client.post("/api/funding-usage/source-summary/publish", headers=users[0],
        json={**payload, "preview_digest": checked.json()["preview_digest"]})
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["document_version_id"] == result["document_version_id"]

    image_text = "张三 使用登记 报销差旅30元"
    image_id = seed_workbuddy_document(sessions, filename="来源同步登记.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1,
                                           "text": image_text}])
    old_source_id = seed_excel_target(sessions, filename="旧来源.xlsx", sheets={
        "汇总": [header, summary_row]})
    with sessions() as db:
        old_source = db.query(DocumentVersion).filter_by(document_id=old_source_id).one()
        image = db.query(DocumentVersion).filter_by(document_id=image_id).one()
        image_run = db.query(DocumentExtractionRun).filter_by(document_id=image_id).one()
    first = _request(image_text, amount="30", registration="报销差旅30元").model_dump(mode="json")
    first["image"] = {"document_id": image_id, "document_version_id": image.id,
                      "source_sha256": image.sha256, "extraction_run_id": image_run.id}
    first["workbook"] = {"document_id": old_source_id,
                         "document_version_id": old_source.id,
                         "source_sha256": old_source.sha256}
    first["filename"] = "同步后的余额表.xlsx"
    first_preview = client.post("/api/funding-usage/preview", headers=users[0], json=first)
    assert first_preview.status_code == 200, first_preview.text
    balance = client.post("/api/funding-usage/publish", headers=users[0],
        json={**first, "preview_digest": first_preview.json()["preview_digest"]})
    assert balance.status_code == 200, balance.text
    balance = balance.json()
    refresh = {"workbook": {"document_id": result["document_id"],
        "document_version_id": result["document_version_id"],
        "source_sha256": result["sha256"]},
        "filename": balance["filename"],
        "base_document_version_id": balance["document_version_id"],
        "base_sha256": balance["sha256"], "base_revision": balance["revision"],
        "request_key": str(uuid4())}
    preview = client.post("/api/funding-usage/refresh/preview", headers=users[0], json=refresh)
    assert preview.status_code == 200, preview.text
    assert preview.json()["changed_count"] == 1
    balance_updated = client.post("/api/funding-usage/refresh/publish", headers=users[0],
        json={**refresh, "preview_digest": preview.json()["preview_digest"]})
    assert balance_updated.status_code == 200, balance_updated.text
    balance_book = load_workbook(BytesIO(client.get(
        balance_updated.json()["download_url"], headers=users[0]).content))
    assert balance_book["使用登记汇总"]["C2"].value == 120
    assert balance_book["使用登记汇总"]["E2"].value == 30
    assert balance_book["使用登记汇总"]["F2"].value == 90
    balance_book.close()


@pytest.mark.parametrize("registration_prefix", ["", "-"])
def test_jsonb_funding_order_does_not_change_append_or_manual_return(
        demo, monkeypatch, registration_prefix):
    client, sessions, users = demo
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    first_text = f"王五 使用登记 {registration_prefix}报销会议30元\n李四 使用登记 报销办公20元"
    next_text = "李四 使用登记 报销资料10元"
    first_id = seed_workbuddy_document(sessions, filename="顺序首图.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1,
                                           "text": first_text}])
    next_id = seed_workbuddy_document(sessions, filename="顺序追加图.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1,
                                           "text": next_text}])
    headings, no_use, second, third = ([None] * 14 for _ in range(4))
    headings[1], headings[13] = "姓名", "资助总额（元）"
    no_use[1], no_use[13] = "张三", 100
    second[1], second[13] = "李四", 200
    third[1], third[13] = "王五", 300
    old_id = seed_excel_target(sessions, filename="顺序旧表.xlsx",
                               sheets={"汇总": [headings, no_use, second, third]})

    def source(document_id):
        with sessions() as db:
            version = db.query(DocumentVersion).filter_by(document_id=document_id).one()
            run = db.query(DocumentExtractionRun).filter_by(document_id=document_id).one()
            return {"document_id": document_id, "document_version_id": version.id,
                    "source_sha256": version.sha256, "extraction_run_id": run.id}

    first = _request(first_text, amount="30", name="王五",
                     registration=f"{registration_prefix}报销会议30元").model_dump(mode="json")
    first["rows"] += _request(first_text, amount="20", name="李四",
                              registration="报销办公20元").model_dump(mode="json")["rows"]
    first["image"] = source(first_id)
    first["filename"] = "顺序余额表.xlsx"
    with sessions() as db:
        old = db.query(DocumentVersion).filter_by(document_id=old_id).one()
        first["workbook"] = {"document_id": old_id,
                             "document_version_id": old.id, "source_sha256": old.sha256}
    preview = client.post("/api/funding-usage/preview", headers=users[0], json=first)
    assert preview.status_code == 200, preview.text
    initial = client.post("/api/funding-usage/publish", headers=users[0],
        json={**first, "preview_digest": preview.json()["preview_digest"]})
    assert initial.status_code == 200, initial.text
    initial = initial.json()

    def simulate_old_jsonb_order(version_id):
        with sessions() as db:
            version = db.get(DocumentVersion, version_id)
            policy = dict(version.extraction_policy_json)
            policy.pop("funding_usage_person_order_v3", None)
            funding = policy["funding_usage_funding_v3"]
            policy["funding_usage_funding_v3"] = {
                name: funding[name] for name in ("王五", "张三", "李四")}
            version.extraction_policy_json = policy
            db.commit()

    simulate_old_jsonb_order(initial["document_version_id"])
    baseline_bytes = client.get(initial["download_url"], headers=users[0]).content
    book = load_workbook(BytesIO(baseline_bytes))
    assert book.sheetnames == ["使用登记汇总", "李四", "王五"]
    unchanged = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": initial["filename"], "request_key": str(uuid4()),
              "base_document_version_id": initial["document_version_id"],
              "base_sha256": initial["sha256"], "base_revision": str(initial["revision"]),
              "file_sha256": sha256(baseline_bytes).hexdigest()},
        files={"file": (initial["filename"], baseline_bytes,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert unchanged.status_code == 422, unchanged.text
    assert unchanged.json()["error"]["code"] == "MANUAL_EXCEL_UNCHANGED"
    assert "个人明细页" in unchanged.json()["error"]["message"]
    book["李四"]["D2"] = 25
    output = BytesIO()
    book.save(output)
    book.close()
    edited = output.getvalue()
    revised = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": initial["filename"], "request_key": str(uuid4()),
              "base_document_version_id": initial["document_version_id"],
              "base_sha256": initial["sha256"], "base_revision": str(initial["revision"]),
              "file_sha256": sha256(edited).hexdigest()},
        files={"file": (initial["filename"], edited,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert revised.status_code == 200, revised.text
    revised = revised.json()
    book = load_workbook(BytesIO(client.get(revised["download_url"], headers=users[0]).content))
    assert [book["使用登记汇总"].cell(row, 2).value for row in (2, 3, 4)] == [
        "张三", "李四", "王五"]
    assert book["使用登记汇总"]["F3"].value == 175
    assert book["王五"]["B2"].value == (
        "'" if registration_prefix else "") + f"{registration_prefix}报销会议30元"
    book.close()
    with sessions() as db:
        version = db.get(DocumentVersion, revised["document_version_id"])
        recorded = next(row for row in version.extraction_policy_json["funding_usage_rows_v3"]
                        if row["name"] == "王五")
        assert recorded["registration"] == f"{registration_prefix}报销会议30元"

    simulate_old_jsonb_order(revised["document_version_id"])
    payload = {"image": source(next_id), "filename": revised["filename"],
        "base_document_version_id": revised["document_version_id"],
        "base_sha256": revised["sha256"], "base_revision": revised["revision"],
        "request_key": str(uuid4()),
        "rows": _request(next_text, amount="10", name="李四",
                         registration="报销资料10元").model_dump(mode="json")["rows"]}
    checked = client.post("/api/funding-usage/continue/preview", headers=users[0], json=payload)
    assert checked.status_code == 200, checked.text
    appended = client.post("/api/funding-usage/continue/publish", headers=users[0],
        json={**payload, "preview_digest": checked.json()["preview_digest"]})
    assert appended.status_code == 200, appended.text
    book = load_workbook(BytesIO(client.get(
        appended.json()["download_url"], headers=users[0]).content))
    assert [book["使用登记汇总"].cell(row, 2).value for row in (2, 3, 4)] == [
        "张三", "李四", "王五"]
    assert book["使用登记汇总"]["F3"].value == 165
    assert book.sheetnames == ["使用登记汇总", "李四", "王五"]
    book.close()


def test_personal_edits_recompute_summary_and_next_image_balance(demo, monkeypatch):
    client, sessions, users = demo
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    first_text = "张三 使用登记 报销差旅300元"
    second_text = "张三 使用登记 报销会议100元"
    third_text = "张三 使用登记 报销资料25元"
    fourth_text = "李四 使用登记 报销办公50元"
    image_ids = [seed_workbuddy_document(sessions, filename=f"新版来源{index}.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1, "text": text}])
        for index, text in enumerate((first_text, second_text, third_text, fourth_text), 1)]
    headings, record, unused, still_unused, footer = ([None] * 14 for _ in range(5))
    headings[1], headings[13] = "姓名", "资助总额（元）"
    record[1], record[13] = "张三", 1000
    unused[1], unused[13] = "李四", 2000
    still_unused[1], still_unused[13] = "王五", 3000
    footer[0], footer[1], footer[13] = "总计", "50 人", 6000
    old_id = seed_excel_target(sessions, filename="新版成果基准.xlsx",
                               sheets={"汇总": [headings, record, unused, still_unused, footer]})

    def image(document_id):
        with sessions() as db:
            version = db.query(DocumentVersion).filter_by(document_id=document_id).one()
            run = db.query(DocumentExtractionRun).filter_by(document_id=document_id).one()
            return {"document_id": document_id, "document_version_id": version.id,
                    "source_sha256": version.sha256, "extraction_run_id": run.id}

    first = _request(first_text, amount="300").model_dump(mode="json")
    first["image"] = image(image_ids[0])
    first["filename"] = "新版成果资助余额.xlsx"
    with sessions() as db:
        old = db.query(DocumentVersion).filter_by(document_id=old_id).one()
        first["workbook"] = {"document_id": old_id, "document_version_id": old.id,
                             "source_sha256": old.sha256}
    preview = client.post("/api/funding-usage/preview", headers=users[0], json=first)
    assert preview.status_code == 200, preview.text
    initial = client.post("/api/funding-usage/publish", headers=users[0],
        json={**first, "preview_digest": preview.json()["preview_digest"]})
    assert initial.status_code == 200, initial.text
    initial = initial.json()
    book = load_workbook(BytesIO(client.get(initial["download_url"], headers=users[0]).content))
    assert [book.worksheets[0].cell(1, col).value for col in range(1, 7)] == [
        "序号", "姓名", "总科研成果资助金额", "其他资助金额（院优青、专委、其他资助）",
        "总支出", "剩余总资助金额"]
    assert [book.worksheets[0].cell(2, col).value for col in range(1, 7)] == [
        1, "张三", 1000, None, 300, 700]
    assert [book.worksheets[0].cell(3, col).value for col in range(1, 7)] == [
        2, "李四", 2000, None, 0, 2000]
    assert [book.worksheets[0].cell(4, col).value for col in range(1, 7)] == [
        3, "王五", 3000, None, 0, 3000]
    assert book.sheetnames == ["使用登记汇总", "张三"]
    assert [book["张三"].cell(1, col).value for col in range(1, 6)] == [
        "序号", "使用登记原文", "日期", "确认使用金额（元）", "核对说明"]
    assert book["张三"]["A2"].value == 1
    assert book["张三"]["A4"].value == "总科研成果资助金额"
    book.close()

    def append(text, amount, registration, document_id, base, *, name="张三"):
        payload = {"image": image(document_id), "filename": base["filename"],
            "base_document_version_id": base["document_version_id"],
            "base_sha256": base["sha256"], "base_revision": base["revision"],
            "request_key": str(uuid4()),
            "rows": _request(text, amount=amount, registration=registration, name=name)
                    .model_dump(mode="json")["rows"]}
        checked = client.post("/api/funding-usage/continue/preview", headers=users[0], json=payload)
        assert checked.status_code == 200, checked.text
        result = client.post("/api/funding-usage/continue/publish", headers=users[0],
            json={**payload, "preview_digest": checked.json()["preview_digest"]})
        assert result.status_code == 200, result.text
        return result.json()

    appended = append(second_text, "100", "报销会议100元", image_ids[1], initial)
    book = load_workbook(BytesIO(client.get(appended["download_url"], headers=users[0]).content))
    assert book["使用登记汇总"].max_row == 4
    assert book.sheetnames == ["使用登记汇总", "张三"]
    assert [book["张三"].cell(row, 1).value for row in (2, 3)] == [1, 2]
    book["张三"]["D2"] = 350
    book["张三"]["B6"] = 50  # 其他资助金额；用户核对后填写。
    output = BytesIO()
    book.save(output)
    book.close()
    edited = output.getvalue()
    revised = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": appended["filename"], "request_key": str(uuid4()),
              "base_document_version_id": appended["document_version_id"],
              "base_sha256": appended["sha256"], "base_revision": str(appended["revision"]),
              "file_sha256": sha256(edited).hexdigest()},
        files={"file": (appended["filename"], edited,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert revised.status_code == 200, revised.text
    revised = revised.json()
    book = load_workbook(BytesIO(client.get(revised["download_url"], headers=users[0]).content))
    assert [book["使用登记汇总"].cell(2, col).value for col in (3, 4, 5, 6)] == [
        1000, 50, 450, 600]
    assert [book["张三"].cell(row, 2).value for row in (5, 6, 7, 8)] == [
        1000, 50, 450, 600]
    assert [book["使用登记汇总"].cell(3, col).value for col in (3, 4, 5, 6)] == [
        2000, None, 0, 2000]
    assert [book["使用登记汇总"].cell(4, col).value for col in (3, 4, 5, 6)] == [
        3000, None, 0, 3000]
    book.close()

    final = append(third_text, "25", "报销资料25元", image_ids[2], revised)
    book = load_workbook(BytesIO(client.get(final["download_url"], headers=users[0]).content))
    assert book["使用登记汇总"].max_row == 4
    assert book["使用登记汇总"]["F2"].value == 575
    assert [book["张三"].cell(row, 1).value for row in (2, 3, 4)] == [1, 2, 3]
    book["张三"].delete_rows(4)
    output = BytesIO()
    book.save(output)
    book.close()
    edited = output.getvalue()
    removed = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": final["filename"], "request_key": str(uuid4()),
              "base_document_version_id": final["document_version_id"],
              "base_sha256": final["sha256"], "base_revision": str(final["revision"]),
              "file_sha256": sha256(edited).hexdigest()},
        files={"file": (final["filename"], edited,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert removed.status_code == 200, removed.text
    book = load_workbook(BytesIO(client.get(removed.json()["download_url"], headers=users[0]).content))
    assert book["使用登记汇总"]["F2"].value == 600
    assert [book["张三"].cell(row, 1).value for row in (2, 3)] == [1, 2]
    book.close()
    readded = append(third_text, "25", "报销资料25元", image_ids[2], removed.json())
    book = load_workbook(BytesIO(client.get(readded["download_url"], headers=users[0]).content))
    assert book["使用登记汇总"]["F2"].value == 575
    book.close()
    first_for_unused = append(fourth_text, "50", "报销办公50元", image_ids[3],
                              readded, name="李四")
    book = load_workbook(BytesIO(client.get(first_for_unused["download_url"], headers=users[0]).content))
    assert [book["使用登记汇总"].cell(3, col).value for col in (3, 4, 5, 6)] == [
        2000, None, 50, 1950]
    assert book.sheetnames == ["使用登记汇总", "张三", "李四"]
    assert book["李四"]["A2"].value == 1
    assert book["使用登记汇总"]["F4"].value == 3000
    book["李四"].delete_rows(2)
    output = BytesIO()
    book.save(output)
    book.close()
    edited = output.getvalue()
    removed_unused = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": first_for_unused["filename"], "request_key": str(uuid4()),
              "base_document_version_id": first_for_unused["document_version_id"],
              "base_sha256": first_for_unused["sha256"],
              "base_revision": str(first_for_unused["revision"]),
              "file_sha256": sha256(edited).hexdigest()},
        files={"file": (first_for_unused["filename"], edited,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert removed_unused.status_code == 200, removed_unused.text
    book = load_workbook(BytesIO(client.get(
        removed_unused.json()["download_url"], headers=users[0]).content))
    assert book.sheetnames == ["使用登记汇总", "张三"]
    assert book["使用登记汇总"]["E3"].value == 0
    assert book["使用登记汇总"]["F3"].value == 2000
    book.close()

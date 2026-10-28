"""完整文件回传路径与旧表格路径并存的端到端回归。"""

import hashlib
import io
import json
import asyncio
import zipfile
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from app.core.config import get_settings
from app.core.request_limits import ExcelWritebackBodyLimitMiddleware
from app.db.models import DocumentSearchProfile, DocumentVersion, StructuredDatasetMember, StructuredDatasetPublication, StructuredDatasetTask, WorkingCopy, WorkingCopyRoot
from app.modules.file_lifecycle.storage import FileLifecycleStorageService
from app.modules.structured_datasets.task_schemas import TaskPayload
from app.modules.structured_datasets.writeback import writeback_digest
from app.modules.structured_datasets.publication import (
    finalize_publication, recover_pending_publications, repair_needs_repair_publication,
)
from app.scripts.rebuild_dataset_search_profile import rebuild_profile
from app.tests.test_dataset_tasks import make_payload
from app.tests.test_structured_datasets import BASE, demo, seed_workbuddy_document, seed_excel_target


def _ready(client, sessions, headers, monkeypatch, tmp_path, *, review_mode="DIRECT_SAVE"):
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    payload = make_payload(client, sessions, headers, template=True)
    target_id = payload["manifest"]["target"]["uploaded"]["document_id"]
    response = client.post(f"{BASE}/from-excel", headers=headers,
        json={"document_id": target_id, "sheet_name": "登记", "idempotency_key": "writeback-target"})
    assert response.status_code == 200, response.text
    target = response.json()
    checkout = client.post("/api/excel-writeback/checkout", headers=headers,
        json={"filename": "教学汇总.xlsx", "request_key": "first", "review_mode": review_mode})
    assert checkout.status_code == 200, checkout.text
    task = checkout.json()
    payload["manifest"]["target"] = {
        "dataset_id": target["dataset_id"], "document_version_id": task["document_version_id"],
        "filename": task["filename"], "sheet_name": task["sheet_name"],
        "labels": task["labels"], "layout": task["layout"],
    }
    payload = TaskPayload.model_validate(payload).model_dump(mode="json")
    baseline = tmp_path / "baseline.xlsx"
    downloaded = client.get(task["download_url"], headers=headers)
    assert downloaded.status_code == 200
    baseline.write_bytes(downloaded.content)
    assert hashlib.sha256(baseline.read_bytes()).hexdigest() == task["base_sha256"]
    return task, payload, baseline


def _candidate(monkeypatch, task, payload, baseline, tmp_path):
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3] / "mcp"))
    from file_agent_mcp.xlsx_writeback import patch_workbook
    output = tmp_path / "candidate.xlsx"
    patch_workbook(baseline=baseline, output=output, sheet_name=task["sheet_name"],
                   labels=task["labels"], rows=[["教学工作量：120学时 √", None]], layout=task["layout"])
    sha = hashlib.sha256(output.read_bytes()).hexdigest()
    return output, sha, writeback_digest(payload, sha)


def _publish(client, headers, task, payload, output, sha, digest):
    return client.post(f"/api/excel-writeback/tasks/{task['task_id']}/publish", headers=headers,
        data={"result_digest": digest, "file_sha256": sha, "payload_json": json.dumps(payload, ensure_ascii=False)},
        files={"file": (task["filename"], output.read_bytes(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_generic_writeback_numbers_multiple_evidenced_items(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    payload["manifest"]["sealed_only"] = False
    payload["records"][0]["fields"][0] = {
        "fragments": [
            {"source_unit_id": "unit-0", "quote": "教学工作量：120学时\n本栏情况属实"},
            {"source_unit_id": "unit-0", "quote": "年度考核未提供"},
        ],
        "numbered_items": True,
        "confidence": 1,
    }
    payload = TaskPayload.model_validate(payload).model_dump(mode="json")
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3] / "mcp"))
    from file_agent_mcp.xlsx_writeback import patch_workbook

    output = tmp_path / "numbered-items.xlsx"
    shown = "1. 教学工作量：120学时\n本栏情况属实\n2. 年度考核未提供"
    patch_workbook(baseline=baseline, output=output, sheet_name=task["sheet_name"],
                   labels=task["labels"], rows=[[shown, None]], layout=task["layout"])
    sha = hashlib.sha256(output.read_bytes()).hexdigest()
    digest = writeback_digest(payload, sha)
    preview = client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
                          json={"result_digest": digest, "file_sha256": sha})
    assert preview.status_code == 200, preview.text
    published = _publish(client, users[0], task, payload, output, sha, digest)
    assert published.status_code == 200, published.text
    book = load_workbook(io.BytesIO(client.get(published.json()["download_url"], headers=users[0]).content))
    assert book["登记"]["A2"].value == shown
    book.close()


def test_generic_paper_titles_must_be_numbered_before_submit(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    _, payload, _ = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    payload["manifest"]["target"]["labels"][0] = "论文题目"
    payload["manifest"]["sealed_only"] = False
    payload["records"][0]["fields"][0] = {
        "fragments": [
            {"source_unit_id": "unit-0", "quote": "第一篇\n续行"},
            {"source_unit_id": "unit-0", "quote": "第二篇"},
        ],
    }
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3] / "mcp"))
    from file_agent_mcp.dataset_task_contracts import TaskPayload as McpTaskPayload
    for contract in (TaskPayload, McpTaskPayload):
        with pytest.raises(ValueError, match="论文题目有多条原文片段时必须逐篇编号"):
            contract.model_validate(payload)
        payload["records"][0]["fields"][0]["numbered_items"] = True
        contract.model_validate(payload)
        del payload["records"][0]["fields"][0]["numbered_items"]


def test_generic_writeback_keeps_blank_column_and_publishes_row_sequence(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    source_id = seed_workbuddy_document(sessions, filename="人员.pdf", content_type="application/pdf", units=[
        {"kind": "PDF_PAGE", "page_number": 1, "text": "张三\n计算机学院"},
    ])
    source = client.get(f"{BASE}/sources/{source_id}", headers=users[0]).json()
    target_id = seed_excel_target(sessions, filename="序号登记.xlsx", sheets={
        "登记": [["序号", "姓名", None, "单位"], [1, "已有人员", None, "原单位"]],
    })
    created = client.post(f"{BASE}/from-excel", headers=users[0], json={
        "document_id": target_id, "sheet_name": "登记", "idempotency_key": "numbered-target",
    })
    assert created.status_code == 200, created.text
    task_response = client.post("/api/excel-writeback/checkout", headers=users[0], json={
        "filename": "序号登记.xlsx", "request_key": "numbered-row", "review_mode": "DIRECT_SAVE",
    })
    assert task_response.status_code == 200, task_response.text
    task = task_response.json()
    assert task["labels"] == ["序号", "姓名", "〔空白列:C〕", "单位"]
    payload = TaskPayload.model_validate({
        "manifest": {
            "sources": [{key: source[key] for key in (
                "document_id", "document_version_id", "extraction_run_id", "source_sha256")},
            ],
            "target": {"dataset_id": created.json()["dataset_id"],
                       "document_version_id": task["document_version_id"],
                       "filename": task["filename"], "sheet_name": task["sheet_name"],
                       "labels": task["labels"], "layout": task["layout"]},
        },
        "records": [{"source_index": 0, "fields": [
            {"blank_reason": "TARGET_DERIVED"},
            {"raw_text": "张三", "evidence_refs": [{"source_unit_id": "unit-0", "quote": "张三"}]},
            {"blank_reason": "TARGET_BLANK_COLUMN"},
            {"raw_text": "计算机学院", "evidence_refs": [{"source_unit_id": "unit-0", "quote": "计算机学院"}]},
        ]}],
    }).model_dump(mode="json")
    filled_blank = json.loads(json.dumps(payload, ensure_ascii=False))
    filled_blank["records"][0]["fields"][2] = {
        "raw_text": "误填", "evidence_refs": [{"source_unit_id": "unit-0", "quote": "张三"}],
    }
    with pytest.raises(ValueError, match="目标空白列必须原样保留"):
        TaskPayload.model_validate(filled_blank)
    baseline = tmp_path / "numbered-baseline.xlsx"
    baseline.write_bytes(client.get(task["download_url"], headers=users[0]).content)
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3] / "mcp"))
    from file_agent_mcp.dataset_task_contracts import TaskPayload as McpTaskPayload
    from file_agent_mcp.xlsx_writeback import patch_workbook
    McpTaskPayload.model_validate(payload)
    with pytest.raises(ValueError, match="目标空白列必须原样保留"):
        McpTaskPayload.model_validate(filled_blank)

    output = tmp_path / "numbered-result.xlsx"
    patch_workbook(baseline=baseline, output=output, sheet_name=task["sheet_name"],
                   labels=task["labels"], rows=[[None, "张三", None, "计算机学院"]], layout=task["layout"])
    sha = hashlib.sha256(output.read_bytes()).hexdigest()
    digest = writeback_digest(payload, sha)
    preview = client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
                          json={"result_digest": digest, "file_sha256": sha})
    assert preview.status_code == 200, preview.text
    published = _publish(client, users[0], task, payload, output, sha, digest)
    assert published.status_code == 200, published.text
    book = load_workbook(io.BytesIO(client.get(published.json()["download_url"], headers=users[0]).content))
    assert [book["登记"][f"A{row}"].value for row in (2, 3)] == [1, 2]
    assert book["登记"]["B3"].value == "张三"
    assert book["登记"]["C1"].value is None and book["登记"]["C3"].value is None
    assert book["登记"]["D3"].value == "计算机学院"
    book.close()


def test_writeback_publishes_workbuddy_bytes_and_preserves_old_path(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    output, sha, digest = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    preview = client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
                          json={"result_digest": digest, "file_sha256": sha})
    assert preview.status_code == 200, preview.text
    result = _publish(client, users[0], task, payload, output, sha, digest)
    assert result.status_code == 200, result.text
    saved = result.json()
    assert saved["excel_published"] and saved["task_record_count"] == 1
    assert saved["filename"] == "教学汇总.xlsx"
    searched = client.post("/api/search", headers=users[0], json={
        "query": "教学汇总.xlsx", "top_k": 10,
    })
    assert searched.status_code == 200, searched.text
    assert any(item["working_copy_id"] == saved["working_copy_id"]
               for item in searched.json()["files"])
    downloaded = client.get(saved["download_url"], headers=users[0])
    assert downloaded.status_code == 200
    assert downloaded.content == output.read_bytes()
    book = load_workbook(io.BytesIO(downloaded.content))
    assert book["登记"]["A2"].value == "教学工作量：120学时 √"
    book.close()
    again = _publish(client, users[0], task, payload, output, sha, digest)
    assert again.status_code == 200, again.text
    assert again.json()["document_version_id"] == saved["document_version_id"]
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 1
        assert db.query(StructuredDatasetTask).filter_by(id=task["task_id"]).one().status == "COMPLETED"
        version = db.get(DocumentVersion, saved["document_version_id"])
        assert version.version_number == saved["version_number"]
        working = db.get(WorkingCopy, saved["working_copy_id"])
        profile = db.query(DocumentSearchProfile).filter_by(working_copy_id=working.id).one()
        assert profile.status == "ACTIVE"
        assert profile.document_version_id == saved["document_version_id"]
        db.delete(profile)
        db.flush()
        repaired = rebuild_profile(
            db, working_copy_id=working.id,
            expected_version_id=saved["document_version_id"],
        )
        assert repaired["profile_id"]
        assert db.query(DocumentSearchProfile).filter_by(working_copy_id=working.id).one().document_version_id == saved["document_version_id"]
        root = db.get(WorkingCopyRoot, working.working_copy_root_id)
        assert working.filename == "教学汇总.xlsx"
        assert Path(working.relative_path).name == "教学汇总.xlsx"
        assert root.relative_storage_path == "structured-datasets"
        visible = FileLifecycleStorageService().working_copy_path(
            f"{root.relative_storage_path}/{working.relative_path}")
        assert visible.read_bytes() == output.read_bytes()


def test_full_professor_row_publishes_once_with_separate_ocr_markers(demo, monkeypatch, tmp_path):
    """同一 PDF 的 B–AA 一次落盘；待核对格可见但不参与规范化事实。"""
    client, sessions, users = demo
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    capabilities = client.get(f"{BASE}/tasks/capabilities", headers=users[0])
    assert capabilities.status_code == 200
    assert capabilities.json()["full_row_write_scope"] == "A:AA"
    assert capabilities.json()["ocr_uncertain_markers"] is True
    labels = ["序号", *[f"字段{number}" for number in range(2, 28)]]
    target_id = seed_excel_target(sessions, filename="教授汇总.xlsx", sheets={"汇总": [labels, [1]]})
    layout = {"header_range": "A1:AA1", "data_range": "A2:AA3"}
    registered = client.post(f"{BASE}/from-excel", headers=users[0], json={
        "document_id": str(target_id), "sheet_name": "汇总", "layout": layout,
        "idempotency_key": "full-professor-target"})
    assert registered.status_code == 200, registered.text
    checkout = client.post("/api/excel-writeback/checkout", headers=users[0], json={
        "filename": "教授汇总.xlsx", "request_key": "full-professor-row", "review_mode": "DIRECT_SAVE"})
    assert checkout.status_code == 200, checkout.text
    task = checkout.json()
    source_id = seed_workbuddy_document(sessions, filename="申报.pdf", content_type="application/pdf",
        units=[{"kind": "PDF_PAGE", "page_number": 1,
                "text": "单位 学院甲 姓名 张三 性别 男 出生 1980 学历 博士 职称 教授 年份 2026 部门 教学\n论文 12?篇\n师德评价 优秀"}])
    source = client.get(f"{BASE}/sources/{source_id}", headers=users[0]).json()
    unit_id = source["source_units"][0]["source_unit_id"]
    clear = [("学院甲", "单位 学院甲"), ("张三", "姓名 张三"), ("男", "性别 男"),
             ("1980", "出生 1980"), ("博士", "学历 博士"), ("教授", "职称 教授"),
             ("2026", "年份 2026"), ("教学", "部门 教学")]
    fields = [{"blank_reason": "TARGET_DERIVED"}]
    fields.extend({"raw_text": value, "page_number": 1,
                   "evidence_refs": [{"source_unit_id": unit_id, "quote": quote}]}
                  for value, quote in clear)
    fields.append({"raw_text": "12?篇", "page_number": 1, "ocr_status": "UNCERTAIN",
                   "ocr_unit_id": unit_id,
                   "evidence_refs": [{"source_unit_id": unit_id, "quote": "论文 12?篇"}]})
    fields.append({"blank_reason": "OCR_UNREADABLE", "page_number": 1,
                   "ocr_status": "UNREADABLE", "ocr_unit_id": unit_id})
    fields.extend({"blank_reason": "来源未提供对应事实"} for _ in range(16))
    source_quote = {"source_unit_id": unit_id, "quote": "师德评价 优秀"}
    fields[12] = {"raw_text": "师德评价 优秀", "page_number": 1,
                  "fragments": [source_quote], "source_only_fragments": [source_quote]}
    payload = TaskPayload.model_validate({"manifest": {
        "sources": [{key: source[key] for key in (
            "document_id", "document_version_id", "extraction_run_id", "source_sha256")}],
        "target": {"dataset_id": registered.json()["dataset_id"],
                   "document_version_id": task["document_version_id"], "filename": task["filename"],
                   "sheet_name": task["sheet_name"], "layout": task["layout"], "labels": task["labels"]},
        "sealed_only": True, "identity_column_range": "B:I", "source_fact_column_range": "J:P",
        "write_scope": "A:AA"}, "records": [{"source_index": 0, "fields": fields}]}).model_dump(mode="json")
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3] / "mcp"))
    from file_agent_mcp.xlsx_writeback import patch_workbook
    from file_agent_mcp.excel_writeback_tools import _full_professor_rows
    from file_agent_mcp.dataset_task_contracts import TaskPayload as HostPayload
    assert HostPayload.model_validate(payload).model_dump(mode="json") == payload
    baseline = tmp_path / "full-baseline.xlsx"
    baseline.write_bytes(client.get(task["download_url"], headers=users[0]).content)
    output = tmp_path / "full-candidate.xlsx"
    cells = {chr(66 + index): value for index, (value, _) in enumerate(clear)}
    cells.update(J="〔OCR 待核对：12?篇〕", K="〔OCR 未识别，待核对〕",
                 M="师德评价 优秀〔PDF 原文，未见对应印章〕")
    row = _full_professor_rows([cells], len(task["labels"]))[0]
    patch_workbook(baseline=baseline, output=output, sheet_name="汇总", labels=task["labels"],
                   rows=[row], layout=task["layout"], scoped_a_i=True)
    sha = hashlib.sha256(output.read_bytes()).hexdigest()
    digest = writeback_digest(payload, sha)
    preview = client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
                          json={"result_digest": digest, "file_sha256": sha})
    assert preview.status_code == 200, preview.text
    published = _publish(client, users[0], task, payload, output, sha, digest)
    assert published.status_code == 200, published.text
    result = published.json()
    assert result["excel_published"] and result["task_record_count"] == 1
    book = load_workbook(io.BytesIO(client.get(result["download_url"], headers=users[0]).content))
    assert book["汇总"]["A2"].value == 1
    assert book["汇总"]["B2"].value == "学院甲"
    assert book["汇总"]["J2"].value == "〔OCR 待核对：12?篇〕"
    assert book["汇总"]["K2"].value == "〔OCR 未识别，待核对〕"
    assert book["汇总"]["M2"].value == "师德评价 优秀〔PDF 原文，未见对应印章〕"
    book.close()
    query = client.post(f"{BASE}/{result['dataset_id']}/query", headers=users[0], json={})
    assert query.status_code == 200, query.text
    current = query.json()["records"][-1]
    uncertain = [field for field in current["fields"].values()
                 if field.get("assessment", {}).get("ocr_status") == "UNCERTAIN"]
    assert len(uncertain) == 1 and uncertain[0]["value"] is None


def test_writeback_publication_failure_is_recoverable_without_duplicate_append(demo, monkeypatch, tmp_path):
    """文件同步失败后不得报告成功；数据库提交后的候选只补发，不重复追加。"""
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    output, sha, digest = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    assert client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": digest, "file_sha256": sha}).status_code == 200
    original = FileLifecycleStorageService.publish_dataset_current_file
    calls = 0

    def fail_once(self, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise OSError("injected visible sync failure")
        return original(self, **kwargs)

    monkeypatch.setattr(FileLifecycleStorageService, "publish_dataset_current_file", fail_once)
    first = _publish(client, users[0], task, payload, output, sha, digest)
    assert first.status_code == 200, first.text
    assert first.json()["status"] == "PUBLISHING"
    assert first.json()["excel_published"] is False
    assert client.get(task["download_url"], headers=users[0]).status_code == 423
    listed = client.get("/api/working-copies", headers=users[0])
    assert listed.status_code == 200
    assert next(item for item in listed.json() if item["id"] == task["working_copy_id"])[
        "publication_status"] == "DB_COMMITTED"
    organized = client.get("/api/classification/organization/files", headers=users[0])
    assert organized.status_code == 200, organized.text
    entry = next(item for item in organized.json()["files"] if item["working_copy_id"] == task["working_copy_id"])
    assert entry["publication_status"] == "DB_COMMITTED"
    assert entry["download_url"] is None
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 1
        journal = db.query(StructuredDatasetPublication).filter_by(task_id=task["task_id"]).one()
        assert journal.status == "DB_COMMITTED"
        assert db.get(StructuredDatasetTask, task["task_id"]).status == "PUBLISHING"
    assert recover_pending_publications(sessions)["recovered"] == 1
    assert recover_pending_publications(sessions)["recovered"] == 0
    completed = client.get(f"/api/excel-writeback/tasks/{task['task_id']}", headers=users[0])
    assert completed.status_code == 200 and completed.json()["excel_published"] is True
    assert client.get(task["download_url"], headers=users[0]).content == output.read_bytes()
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 1
        assert db.query(StructuredDatasetPublication).filter_by(task_id=task["task_id"]).one().status == "COMPLETED"


def test_needs_repair_only_republishes_known_baseline_without_duplicate_append(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    output, sha, digest = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    assert client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": digest, "file_sha256": sha}).status_code == 200
    original = FileLifecycleStorageService.publish_dataset_current_file

    def deny_replace(self, **kwargs):
        raise PermissionError("injected file lock")

    monkeypatch.setattr(FileLifecycleStorageService, "publish_dataset_current_file", deny_replace)
    result = _publish(client, users[0], task, payload, output, sha, digest)
    assert result.status_code == 200 and result.json()["excel_published"] is False
    with sessions() as db:
        publication = db.query(StructuredDatasetPublication).filter_by(task_id=task["task_id"]).one()
        publication_id = publication.id
    for _ in range(4):
        with sessions() as db:
            assert finalize_publication(db, publication_id) is False
    monkeypatch.setattr(FileLifecycleStorageService, "publish_dataset_current_file", original)
    with sessions() as db:
        publication = db.get(StructuredDatasetPublication, publication_id)
        assert publication.status == "NEEDS_REPAIR" and publication.attempt_count == 5
        visible = FileLifecycleStorageService().working_copy_path(publication.visible_relative_path)
        assert visible.read_bytes() == baseline.read_bytes()
        visible.write_bytes(b"unexpected external edit")
        with pytest.raises(ValueError, match="拒绝覆盖"):
            repair_needs_repair_publication(
                db, publication_id, expected_visible_sha256=task["base_sha256"])
        db.rollback()
        assert visible.read_bytes() == b"unexpected external edit"
        visible.write_bytes(baseline.read_bytes())
    with sessions() as db:
        assert repair_needs_repair_publication(
            db, publication_id, expected_visible_sha256=task["base_sha256"])
    with sessions() as db:
        publication = db.get(StructuredDatasetPublication, publication_id)
        assert publication.status == "COMPLETED"
        assert db.get(StructuredDatasetTask, task["task_id"]).status == "COMPLETED"
        assert db.query(StructuredDatasetMember).count() == 1
        assert FileLifecycleStorageService().working_copy_path(
            publication.visible_relative_path).read_bytes() == output.read_bytes()
    assert client.get(task["download_url"], headers=users[0]).content == output.read_bytes()


def test_manual_excel_revision_updates_same_target_and_query_without_appending(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    output, sha, digest = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    assert client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": digest, "file_sha256": sha}).status_code == 200
    first = _publish(client, users[0], task, payload, output, sha, digest)
    assert first.status_code == 200 and first.json()["excel_published"]
    previous = first.json()
    book = load_workbook(output)
    data_row = range_boundaries(task["layout"]["data_range"])[1] if task["layout"] else 2
    sheet = book[task["sheet_name"]]
    sheet.cell(data_row, 2).value = "手工补全"
    edited = tmp_path / "manual.xlsx"
    book.save(edited)
    book.close()
    edited_sha = hashlib.sha256(edited.read_bytes()).hexdigest()
    form = {"filename": task["filename"], "request_key": "manual-1",
            "base_document_version_id": previous["document_version_id"],
            "base_sha256": sha, "base_revision": str(previous["revision"]),
            "file_sha256": edited_sha}
    def upload(fields, path=edited):
        return client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
            data=fields, files={"file": (task["filename"], path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    published = upload(form)
    assert published.status_code == 200, published.text
    current = published.json()
    assert current["excel_published"] and current["status"] == "COMPLETED"
    assert current["working_copy_id"] == previous["working_copy_id"]
    assert current["document_version_id"] != previous["document_version_id"]
    assert current["manual_changed_cell_count"] == 1
    with sessions() as db:
        profile = db.query(DocumentSearchProfile).filter_by(
            working_copy_id=current["working_copy_id"]).one()
        assert profile.document_version_id == current["document_version_id"]
    searched = client.post("/api/search", headers=users[0], json={
        "query": task["filename"], "top_k": 10,
    })
    assert searched.status_code == 200, searched.text
    assert any(item["document_version_id"] == current["document_version_id"]
               for item in searched.json()["files"])
    assert client.get(current["download_url"], headers=users[0]).content == edited.read_bytes()
    assert client.get(previous["download_url"], headers=users[0]).content == output.read_bytes()
    target = client.get(f"{BASE}/targets/by-filename", headers=users[0],
        params={"filename": task["filename"]}).json()
    field_key = target["fields"][1]["key"]
    query = client.post(f"{BASE}/{current['dataset_id']}/query", headers=users[0], json={}).json()
    assert query["records"][0]["fields"][field_key]["value"] == "手工补全"
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 1
    replay = upload(form)
    assert replay.status_code == 200 and replay.json()["document_version_id"] == current["document_version_id"]
    changed_form = {**form, "request_key": "manual-stale"}
    assert upload(changed_form).status_code == 409
    book = load_workbook(edited)
    book[task["sheet_name"]].cell(data_row, 1).value = None
    cleared = tmp_path / "cleared.xlsx"
    book.save(cleared)
    book.close()
    current_form = {"filename": task["filename"], "request_key": "manual-cleared",
            "base_document_version_id": current["document_version_id"],
            "base_sha256": edited_sha, "base_revision": str(current["revision"]),
            "file_sha256": hashlib.sha256(cleared.read_bytes()).hexdigest()}
    rejected = upload(current_form, cleared)
    assert rejected.status_code == 422 and rejected.json()["error"]["code"] == "MANUAL_EXCEL_CELL_CLEARED"


def test_manual_revision_publishes_matching_direct_visible_edit_without_losing_it(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    output, sha, digest = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    assert client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": digest, "file_sha256": sha}).status_code == 200
    first = _publish(client, users[0], task, payload, output, sha, digest)
    assert first.status_code == 200, first.text
    previous = first.json()
    edited = tmp_path / "direct-visible-edit.xlsx"
    book = load_workbook(output)
    data_row = range_boundaries(task["layout"]["data_range"])[1] if task["layout"] else 2
    book[task["sheet_name"]].cell(data_row, 2).value = "直接补全"
    book.save(edited)
    book.close()
    edited_sha = hashlib.sha256(edited.read_bytes()).hexdigest()
    with sessions() as db:
        from app.db.models import WorkingCopy, WorkingCopyRoot
        from app.modules.file_lifecycle.storage import FileLifecycleStorageService
        working = db.get(WorkingCopy, previous["working_copy_id"])
        root = db.get(WorkingCopyRoot, working.working_copy_root_id)
        visible = FileLifecycleStorageService().working_copy_path(
            f"{root.relative_storage_path}/{working.relative_path}")
        visible.write_bytes(edited.read_bytes())
    assert client.get(f"{BASE}/targets/by-filename", headers=users[0],
        params={"filename": task["filename"]}).status_code == 409
    form = {"filename": task["filename"], "request_key": "manual-direct-visible",
            "base_document_version_id": previous["document_version_id"],
            "base_sha256": sha, "base_revision": str(previous["revision"]),
            "file_sha256": edited_sha}
    different = tmp_path / "different-upload.xlsx"
    book = load_workbook(edited)
    book[task["sheet_name"]].cell(data_row, 3).value = "另一份编辑"
    book.save(different)
    book.close()
    different_sha = hashlib.sha256(different.read_bytes()).hexdigest()
    conflict = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={**form, "request_key": "manual-different", "file_sha256": different_sha},
        files={"file": (task["filename"], different.read_bytes(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert conflict.status_code == 409
    assert conflict.json()["error"]["code"] == "MANUAL_EXCEL_VISIBLE_CONFLICT"
    assert visible.read_bytes() == edited.read_bytes()
    uploaded = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data=form, files={"file": (task["filename"], edited.read_bytes(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert uploaded.status_code == 200, uploaded.text
    current = uploaded.json()
    assert current["status"] == "COMPLETED" and current["excel_published"]
    assert current["working_copy_id"] == previous["working_copy_id"]
    assert client.get(current["download_url"], headers=users[0]).content == edited.read_bytes()
    assert client.get(previous["download_url"], headers=users[0]).content == output.read_bytes()


def test_writeback_rejects_zip_expansion_before_member_append(demo, monkeypatch, tmp_path):
    """小压缩包中的超大 XML 必须在 openpyxl 和成员写入前被拒绝。"""
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    output, _, _ = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    with zipfile.ZipFile(output, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/oversized.xml", b"x" * (17 * 1024 * 1024))
    sha = hashlib.sha256(output.read_bytes()).hexdigest()
    digest = writeback_digest(payload, sha)
    assert client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": digest, "file_sha256": sha}).status_code == 200
    response = _publish(client, users[0], task, payload, output, sha, digest)
    assert response.status_code == 413, response.text
    assert response.json()["error"]["code"] == "XLSX_UNCOMPRESSED_TOO_LARGE"
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 0


def test_writeback_unknown_commit_outcome_rechecks_database(demo, monkeypatch, tmp_path):
    """提交已生效但驱动抛错时，新会话核对发布记录并安全完成。"""
    from sqlalchemy.orm import Session
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    output, sha, digest = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    assert client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": digest, "file_sha256": sha}).status_code == 200
    original_commit = Session.commit
    injected = False

    def commit_then_report_failure(self):
        nonlocal injected
        row = self.get(StructuredDatasetTask, task["task_id"])
        publishing = row is not None and row.status == "PUBLISHING"
        original_commit(self)
        if publishing and not injected:
            injected = True
            raise OSError("simulated lost commit acknowledgement")

    monkeypatch.setattr(Session, "commit", commit_then_report_failure)
    response = _publish(client, users[0], task, payload, output, sha, digest)
    assert injected and response.status_code == 200, response.text
    assert response.json()["excel_published"] is True
    assert client.get(task["download_url"], headers=users[0]).content == output.read_bytes()
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 1


def test_writeback_failed_database_commit_leaves_visible_baseline(demo, monkeypatch, tmp_path):
    """数据库事务未提交时，候选不可变文件不能提前替换同名可见文件。"""
    from sqlalchemy.orm import Session
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    output, sha, digest = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    assert client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": digest, "file_sha256": sha}).status_code == 200
    original_commit = Session.commit
    injected = False

    def fail_before_commit(self):
        nonlocal injected
        row = self.get(StructuredDatasetTask, task["task_id"])
        publishing = row is not None and row.status == "PUBLISHING"
        if publishing and not injected:
            injected = True
            raise OSError("simulated database commit failure")
        return original_commit(self)

    monkeypatch.setattr(Session, "commit", fail_before_commit)
    response = _publish(client, users[0], task, payload, output, sha, digest)
    assert injected and response.status_code == 503, response.text
    assert response.json()["error"]["code"] == "DATASET_COMMIT_UNCONFIRMED"
    assert client.get(task["download_url"], headers=users[0]).content == baseline.read_bytes()
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 0
        assert db.query(StructuredDatasetPublication).count() == 1  # 初次注册已完成


def test_review_writeback_cannot_add_unconfirmed_source_after_preview(demo, monkeypatch, tmp_path):
    """复核 Excel 指纹不等于复核来源正文；新增来源确认引用必须重新预览。"""
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path,
                                     review_mode="REVIEW_REQUIRED")
    payload["manifest"]["extraction_confirmations"] = [{
        "task_id": str(uuid4()), "confirmation_id": str(uuid4()), "result_digest": "a" * 64}]
    output, sha, digest = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    assert client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": digest, "file_sha256": sha}).status_code == 200
    assert client.post(f"/api/excel-writeback/tasks/{task['task_id']}/confirm", headers=users[0],
        json={"result_digest": digest}).status_code == 200
    response = _publish(client, users[0], task, payload, output, sha, digest)
    assert response.status_code == 422, response.text
    assert response.json()["error"]["code"] == "EXTRACTION_CONFIRMATION_STALE"
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 0


def test_writeback_stream_limit_applies_before_multipart_parsing():
    """无 Content-Length 的流式请求也在 multipart 解析之前返回 413。"""
    from starlette.testclient import TestClient
    from app.main import app
    from app.modules.structured_datasets.writeback_router import require_writeback
    from app.modules.auth.dependencies import get_current_user
    from app.core.database import get_db
    client = TestClient(app)
    chunks = (b"x" * (1024 * 1024) for _ in range(37))
    app.dependency_overrides[require_writeback] = lambda: None
    app.dependency_overrides[get_current_user] = lambda: None
    app.dependency_overrides[get_db] = lambda: None
    try:
        response = client.post("/api/excel-writeback/tasks/00000000-0000-0000-0000-000000000000/publish",
                               headers={"Transfer-Encoding": "chunked", "Content-Type": "application/x-www-form-urlencoded"},
                               content=chunks)
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 413, response.text
    assert response.json()["error"]["code"] == "WRITEBACK_REQUEST_TOO_LARGE"


def test_legacy_export_uses_same_recoverable_publication(demo, monkeypatch):
    """旧 export 接口也不得在可见文件同步失败时报告发布完成。"""
    from app.tests.test_structured_datasets import create_dataset, seed_image, candidate
    client, sessions, users = demo
    dataset_id = create_dataset(client, users[0], key="legacy-publication")
    source_id = seed_image(sessions, "姓名：张三\n金额：20")
    request = candidate(client, users[0], source_id, "张三", "20", key="legacy-member")
    assert client.post(f"{BASE}/{dataset_id}/records", headers=users[0], json=request).status_code == 200
    original = FileLifecycleStorageService.publish_dataset_current_file
    calls = 0

    def fail_once(self, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise OSError("injected old export sync failure")
        return original(self, **kwargs)

    monkeypatch.setattr(FileLifecycleStorageService, "publish_dataset_current_file", fail_once)
    first = client.post(f"{BASE}/{dataset_id}/exports", headers=users[0])
    assert first.status_code == 200, first.text
    assert first.json()["status"] == "PUBLISHING" and not first.json()["excel_published"]
    assert client.get(first.json()["download_url"], headers=users[0]).status_code == 423
    with sessions() as db:
        assert db.query(StructuredDatasetPublication).filter_by(dataset_id=dataset_id).one().status == "DB_COMMITTED"
    assert recover_pending_publications(sessions)["recovered"] == 1
    assert client.get(first.json()["download_url"], headers=users[0]).status_code == 200
    repeated = client.post(f"{BASE}/{dataset_id}/exports", headers=users[0])
    assert repeated.status_code == 200 and repeated.json()["reused"] is True


def test_writeback_rejects_unrelated_workbook_changes(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    output, _, _ = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    book = load_workbook(output)
    book["登记"]["D1"] = "未经授权"
    book.save(output)
    book.close()
    sha = hashlib.sha256(output.read_bytes()).hexdigest()
    digest = writeback_digest(payload, sha)
    preview = client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
                          json={"result_digest": digest, "file_sha256": sha})
    assert preview.status_code == 200, preview.text
    response = _publish(client, users[0], task, payload, output, sha, digest)
    assert response.status_code == 422, response.text
    assert any(code in response.text for code in (
        "WRITEBACK_TEMPLATE_CHANGED", "WRITEBACK_ROWS_MISMATCH", "TARGET_EXCEL_SCHEMA_INVALID")), response.text


def test_writeback_repeated_append_uses_same_filename_and_new_version(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    output, sha, digest = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    assert client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": digest, "file_sha256": sha}).status_code == 200
    first = _publish(client, users[0], task, payload, output, sha, digest)
    assert first.status_code == 200, first.text
    first_saved = first.json()
    second_checkout = client.post("/api/excel-writeback/checkout", headers=users[0],
        json={"filename": "教学汇总.xlsx", "request_key": "second", "review_mode": "DIRECT_SAVE"})
    assert second_checkout.status_code == 200, second_checkout.text
    second_task = second_checkout.json()
    assert second_task["document_id"] == first_saved["document_id"]
    assert second_task["document_version_id"] == first_saved["document_version_id"]
    source_id = seed_workbuddy_document(sessions, filename="第二份.pdf", content_type="application/pdf",
        units=[{"kind": "PDF_PAGE", "page_number": 1,
                "text": "教学工作量：120学时\n本栏情况属实\n教务处圆章"}])
    source = client.get(f"{BASE}/sources/{source_id}", headers=users[0]).json()
    payload["manifest"]["sources"] = [{key: source[key] for key in
        ("document_id", "document_version_id", "extraction_run_id", "source_sha256")}]
    payload["manifest"]["target"]["document_version_id"] = second_task["document_version_id"]
    payload = TaskPayload.model_validate(payload).model_dump(mode="json")
    second_base = tmp_path / "second-base.xlsx"
    second_base.write_bytes(client.get(second_task["download_url"], headers=users[0]).content)
    second_output, second_sha, second_digest = _candidate(
        monkeypatch, second_task, payload, second_base, tmp_path)
    assert client.post(f"/api/excel-writeback/tasks/{second_task['task_id']}/preview", headers=users[0],
        json={"result_digest": second_digest, "file_sha256": second_sha}).status_code == 200
    second = _publish(client, users[0], second_task, payload, second_output, second_sha, second_digest)
    assert second.status_code == 200, second.text
    second_saved = second.json()
    assert second_saved["document_id"] == first_saved["document_id"]
    assert second_saved["working_copy_id"] == first_saved["working_copy_id"]
    assert second_saved["version_number"] == first_saved["version_number"] + 1
    assert second_saved["filename"] == "教学汇总.xlsx"
    current = load_workbook(io.BytesIO(client.get(second_saved["download_url"], headers=users[0]).content))
    assert [current["登记"]["A2"].value, current["登记"]["A3"].value] == [
        "教学工作量：120学时 √", "教学工作量：120学时 √"]
    current.close()


def test_writeback_requires_confirmation_of_same_file_digest(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path,
                                     review_mode="REVIEW_REQUIRED")
    output, sha, digest = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    assert client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": digest, "file_sha256": sha}).status_code == 200
    refused = _publish(client, users[0], task, payload, output, sha, digest)
    assert refused.status_code == 409, refused.text
    assert "WRITEBACK_CONFIRMATION_REQUIRED" in refused.text
    confirmed = client.post(f"/api/excel-writeback/tasks/{task['task_id']}/confirm", headers=users[0],
        json={"result_digest": digest})
    assert confirmed.status_code == 200, confirmed.text
    saved = _publish(client, users[0], task, payload, output, sha, digest)
    assert saved.status_code == 200, saved.text
    assert saved.json()["excel_published"]


def test_writeback_rejects_stale_checkout_without_overwriting_current(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    stale = client.post("/api/excel-writeback/checkout", headers=users[0],
        json={"filename": "教学汇总.xlsx", "request_key": "stale", "review_mode": "DIRECT_SAVE"})
    assert stale.status_code == 200, stale.text
    stale_task = stale.json()
    output, sha, digest = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    for current in (task, stale_task):
        assert client.post(f"/api/excel-writeback/tasks/{current['task_id']}/preview", headers=users[0],
            json={"result_digest": digest, "file_sha256": sha}).status_code == 200
    first = _publish(client, users[0], task, payload, output, sha, digest)
    assert first.status_code == 200, first.text
    rejected = _publish(client, users[0], stale_task, payload, output, sha, digest)
    assert rejected.status_code == 409, rejected.text
    assert "WRITEBACK_VERSION_CONFLICT" in rejected.text
    current = client.get(first.json()["download_url"], headers=users[0])
    assert current.content == output.read_bytes()


def test_writeback_request_limit_counts_stream_without_content_length():
    """新增行、更正和人工回传都须在 multipart 解析前限制实际请求体。"""
    for path in ("/api/excel-writeback/tasks/task-id/publish",
                 "/api/excel-writeback/corrections/tasks/task-id/publish",
                 "/api/excel-writeback/manual-revisions/publish"):
        called = []
        messages = iter([
            {"type": "http.request", "body": b"1234567890", "more_body": True},
            {"type": "http.request", "body": b"abcdefghij", "more_body": False},
        ])
        sent = []

        async def application(scope, receive, send):
            called.append("entered")
            while (await receive()).get("more_body"):
                pass
            called.append("completed")

        async def receive():
            return next(messages)

        async def send(message):
            sent.append(message)

        middleware = ExcelWritebackBodyLimitMiddleware(application, max_body_bytes=16)
        asyncio.run(middleware({"type": "http", "method": "POST",
            "path": path, "headers": []}, receive, send))
        assert called[0] == "entered"
        assert next(message["status"] for message in sent if message["type"] == "http.response.start") == 413

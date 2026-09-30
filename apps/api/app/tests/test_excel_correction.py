"""已发布 Excel 行更正的端到端测试：不重复来源 OCR，也不新增人员行。"""

import hashlib
import io
import json
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.db.models import (DocumentExtractionRun, DocumentVersion, StructuredDatasetCorrection,
                           StructuredDatasetMember, StructuredDataset)
from app.modules.structured_datasets.correction import ExcelCorrectionService
from app.modules.structured_datasets.writeback import writeback_digest
from app.modules.structured_datasets.correction_schemas import CorrectionPayload
from app.tests.test_excel_writeback import _candidate, _publish, _ready
from app.tests.test_structured_datasets import demo


@pytest.mark.parametrize("review_mode", ["DIRECT_SAVE", "REVIEW_REQUIRED"])
def test_correction_changes_only_published_cell_and_preserves_history(demo, monkeypatch, tmp_path, review_mode):
    """直写默认无需确认；显式复核仍须确认，两者都保留来源和旧版本。"""
    client, sessions, users = demo
    task, payload, baseline = _ready(client, sessions, users[0], monkeypatch, tmp_path)
    first_file, first_sha, first_digest = _candidate(monkeypatch, task, payload, baseline, tmp_path)
    assert client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": first_digest, "file_sha256": first_sha}).status_code == 200
    initial = _publish(client, users[0], task, payload, first_file, first_sha, first_digest)
    assert initial.status_code == 200, initial.text
    saved = initial.json()
    query = client.post(f"/api/structured-datasets/{saved['dataset_id']}/query", headers=users[0], json={})
    assert query.status_code == 200, query.text
    record = query.json()["records"][0]
    with sessions() as db:
        extraction_count = db.query(DocumentExtractionRun).count()
        member_count = db.query(StructuredDatasetMember).count()
    checkout = client.post("/api/excel-writeback/corrections/checkout", headers=users[0], json={
        "filename": task["filename"], "request_key": "fix-one-cell", "record_id": record["record_id"],
        "row_number": 2, **({"review_mode": review_mode} if review_mode == "REVIEW_REQUIRED" else {})})
    assert checkout.status_code == 200, checkout.text
    correction = checkout.json()
    field_key = correction["field_keys"][0]
    assert correction["base_values"][field_key] == "教学工作量：120学时 √"
    change_field = dict(payload["records"][0]["fields"][0])
    change_field.pop("assessment")
    candidate = {"record_id": correction["record_id"], "row_number": 2,
        "source": payload["manifest"]["sources"][0], "changes": [{
            "field_key": field_key, "before": correction["base_values"][field_key],
            "after": "教学工作量：120学时", "field": change_field}]}
    candidate = CorrectionPayload.model_validate(candidate).model_dump(mode="json")
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3] / "mcp"))
    from file_agent_mcp.xlsx_writeback import patch_existing_cells
    old_file = tmp_path / "correction-base.xlsx"
    old_file.write_bytes(client.get(correction["download_url"], headers=users[0]).content)
    new_file = tmp_path / "correction-result.xlsx"
    patch_existing_cells(baseline=old_file, output=new_file, sheet_name=correction["sheet_name"],
        row_number=2, changes={"A2": ("教学工作量：120学时 √", "教学工作量：120学时")})
    # 服务端不能因一格得到授权就接受同一行另一格的静默改写。
    tampered = tmp_path / "tampered.xlsx"
    patch_existing_cells(baseline=old_file, output=tampered, sheet_name=correction["sheet_name"],
        row_number=2, changes={"A2": ("教学工作量：120学时 √", "教学工作量：120学时"),
                               "B2": (None, "未授权新增")})
    with sessions() as db:
        dataset = db.get(StructuredDataset, saved["dataset_id"])
        with pytest.raises(HTTPException) as blocked:
            ExcelCorrectionService._verify_correction_workbook(
                object.__new__(ExcelCorrectionService), old_file, tampered, dataset, correction,
                CorrectionPayload.model_validate(candidate))
    assert blocked.value.detail["code"] == "WRITEBACK_TEMPLATE_CHANGED"
    sha = hashlib.sha256(new_file.read_bytes()).hexdigest()
    digest = writeback_digest(candidate, sha)
    preview = client.post(f"/api/excel-writeback/corrections/tasks/{correction['task_id']}/preview",
        headers=users[0], json={"result_digest": digest, "file_sha256": sha})
    expected_status = "WRITEBACK_READY" if review_mode == "DIRECT_SAVE" else "WRITEBACK_WAITING_CONFIRMATION"
    assert preview.status_code == 200 and preview.json()["status"] == expected_status

    def publish():
        return client.post(f"/api/excel-writeback/corrections/tasks/{correction['task_id']}/publish",
            headers=users[0], data={"result_digest": digest, "file_sha256": sha,
                "payload_json": json.dumps(candidate, ensure_ascii=False)},
            files={"file": (task["filename"], new_file.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})

    if review_mode == "REVIEW_REQUIRED":
        assert publish().status_code == 409
        confirmed = client.post(f"/api/excel-writeback/corrections/tasks/{correction['task_id']}/confirm",
            headers=users[0], json={"result_digest": digest})
        assert confirmed.status_code == 200, confirmed.text
    result = publish()
    assert result.status_code == 200, result.text
    published = result.json()
    assert published["excel_published"] and published["corrected_cell_count"] == 1
    assert published["document_version_id"] != saved["document_version_id"]
    result_book = load_workbook(io.BytesIO(client.get(published["download_url"], headers=users[0]).content))
    assert result_book["登记"]["A1"].value == "教学工作量不少于80学时"
    assert result_book["登记"]["A2"].value == "教学工作量：120学时"
    result_book.close()
    old_book = load_workbook(io.BytesIO(client.get(saved["download_url"], headers=users[0]).content))
    assert old_book["登记"]["A2"].value == "教学工作量：120学时 √"
    old_book.close()
    current_query = client.post(f"/api/structured-datasets/{saved['dataset_id']}/query", headers=users[0], json={})
    assert current_query.json()["total_records"] == 1
    assert current_query.json()["records"][0]["fields"][field_key]["display_value"] == "教学工作量：120学时"
    assert publish().json()["document_version_id"] == published["document_version_id"]
    with sessions() as db:
        assert db.query(StructuredDatasetCorrection).count() == 1
        assert db.query(StructuredDatasetMember).count() == member_count
        assert db.query(DocumentExtractionRun).count() == extraction_count
        assert db.get(DocumentVersion, published["document_version_id"]).parent_version_id == saved["document_version_id"]

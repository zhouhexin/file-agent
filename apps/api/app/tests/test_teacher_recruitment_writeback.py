"""教师引进模板：两份正式材料共证一人一行，并保护旧单元格。"""

import hashlib
import io
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree as ET

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from app.core.config import get_settings
from app.db.models import StructuredDatasetMember
from app.modules.structured_datasets.task_schemas import TaskPayload
from app.modules.structured_datasets.assessment import display_value
from app.modules.structured_datasets.writeback import ExcelWritebackService, writeback_digest
from app.tests.test_structured_datasets import BASE, demo, seed_excel_target, seed_workbuddy_document


def test_excel_br_display_and_writeback_ignores_layout_only(tmp_path, monkeypatch):
    baseline = tmp_path / "baseline.xlsx"
    output = tmp_path / "result.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "应届"
    sheet.append(["序号", "姓名", "学习经历"])
    book.save(baseline)
    book.close()

    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3] / "mcp"))
    from file_agent_mcp.xlsx_writeback import patch_workbook

    raw = "博士<BR>硕士<br/>本科<br />论文"
    assert display_value(raw, {}) == "博士\n硕士\n本科\n论文"
    patch_workbook(baseline=baseline, output=output, sheet_name="应届",
                   labels=["序号", "姓名", "学习经历"],
                   rows=[[None, "李四", raw]],
                   layout={"header_range": "A1:C1", "data_range": "A2:C2"}, scoped_a_i=True)
    dataset = SimpleNamespace(target_layout_json={"header_range": "A1:C1", "data_range": "A2:C2"},
                              target_sheet_name="应届", schema_json=[
                                  {"key": "A", "label": "序号"},
                                  {"key": "B", "label": "姓名"},
                                  {"key": "C", "label": "学习经历"}])
    assert ExcelWritebackService._verify_workbook(
        None, base_path=baseline, candidate_path=output, dataset=dataset,
        expected_rows=[[None, "李四", "博士  硕士 本科 论文"]], scoped_a_i=True) == 0

    low_height = tmp_path / "low-height.xlsx"
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(low_height, "w") as target:
        for info in source.infolist():
            contents = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(contents)
                row = next(node for node in root.iter()
                           if node.tag.endswith("}row") and node.attrib.get("r") == "2")
                row.set("ht", "14")
                contents = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(info, contents)
    assert ExcelWritebackService._verify_workbook(
        None, base_path=baseline, candidate_path=low_height, dataset=dataset,
        expected_rows=[[None, "李四", raw]], scoped_a_i=True) == 0

    with pytest.raises(HTTPException) as error:
        ExcelWritebackService._verify_workbook(
            None, base_path=baseline, candidate_path=output, dataset=dataset,
            expected_rows=[[None, "李四", "博士  硕士 本科 A++类"]], scoped_a_i=True)
    assert error.value.detail["code"] == "WRITEBACK_ROWS_MISMATCH"

    unclean = tmp_path / "unclean.xlsx"
    book = load_workbook(output)
    book["应届"]["C2"] = "博士<br>硕士<br/>本科<br />论文"
    book.save(unclean)
    book.close()
    with pytest.raises(HTTPException) as error:
        ExcelWritebackService._verify_workbook(
            None, base_path=baseline, candidate_path=unclean, dataset=dataset,
            expected_rows=[[None, "李四", raw]], scoped_a_i=True)
    assert error.value.detail["code"] == "WRITEBACK_BREAK_TAG_REMAINING"


def _grouped_task(client, sessions, headers, monkeypatch, tmp_path, *, supplemental_has_identity=True):
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    target_id = seed_excel_target(sessions, filename="教师引进汇总.xlsx",
        sheets={"应届": [["序号", "姓名", "学院", "政治面貌"],
                       [None, None, None, None], [None, None, ".", None]]})
    layout = {"header_range": "A1:D1", "data_range": "A2:D3"}
    registered = client.post(f"{BASE}/from-excel", headers=headers, json={
        "document_id": str(target_id), "sheet_name": "应届", "layout": layout,
        "idempotency_key": "recruitment-target"})
    assert registered.status_code == 200, registered.text
    checkout = client.post("/api/excel-writeback/checkout", headers=headers, json={
        "filename": "教师引进汇总.xlsx", "request_key": "recruitment-grouped", "review_mode": "DIRECT_SAVE"})
    assert checkout.status_code == 200, checkout.text
    task = checkout.json()
    first_id = seed_workbuddy_document(sessions, filename="面试.pdf", content_type="application/pdf",
        units=[{"kind": "PDF_PAGE", "page_number": 1, "text": "姓名 张三"}])
    second_text = "姓名 张三 学院 计算机学院" if supplemental_has_identity else "学院 计算机学院"
    second_id = seed_workbuddy_document(sessions, filename="履历.xlsx", content_type=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        units=[{"kind": "SHEET", "sheet_name": "基本信息", "text": second_text}])
    responses = [client.get(f"{BASE}/sources/{source_id}?content_origin=WORKBUDDY", headers=headers)
                 for source_id in (first_id, second_id)]
    assert all(response.status_code == 200 for response in responses), [response.text for response in responses]
    sources = [response.json() for response in responses]
    bindings = [{key: source[key] for key in (
        "document_id", "document_version_id", "extraction_run_id", "source_sha256")}
        for source in sources]
    payload = TaskPayload.model_validate({"manifest": {
        "sources": bindings, "write_scope": "TEMPLATE_ALL",
        "target": {"dataset_id": task["dataset_id"], "document_version_id": task["document_version_id"],
                   "filename": task["filename"], "sheet_name": task["sheet_name"],
                   "layout": task["layout"], "labels": task["labels"]}},
        "records": [{"source_index": 0, "source_indices": [0, 1], "fields": [
            {"blank_reason": "TARGET_DERIVED"},
            {"raw_text": "张三", "evidence_refs": [{"source_unit_id": sources[0]["source_units"][0]["source_unit_id"],
                                                 "quote": "姓名 张三"}]},
            {"raw_text": "计算机学院", "evidence_refs": [{"source_unit_id": sources[1]["source_units"][0]["source_unit_id"],
                                                    "quote": "学院 计算机学院"}]},
            {"blank_reason": "SOURCE_FIELD_NOT_FOUND"}]}]}).model_dump(mode="json")
    baseline = tmp_path / "baseline.xlsx"
    baseline.write_bytes(client.get(task["download_url"], headers=headers).content)
    return task, payload, baseline, sources


def test_grouped_sources_publish_one_row_and_preserve_existing_cell(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    task, payload, baseline, sources = _grouped_task(client, sessions, users[0], monkeypatch, tmp_path)
    monkeypatch.syspath_prepend("apps/mcp")
    from file_agent_mcp.xlsx_writeback import patch_workbook

    output = tmp_path / "result.xlsx"
    patch_workbook(baseline=baseline, output=output, sheet_name="应届", labels=task["labels"],
                   rows=[[None, "张三", "计算机学院", None]], layout=task["layout"], scoped_a_i=True)
    sha = hashlib.sha256(output.read_bytes()).hexdigest()
    digest = writeback_digest(payload, sha)
    preview = client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": digest, "file_sha256": sha})
    assert preview.status_code == 200, preview.text
    saved = client.post(f"/api/excel-writeback/tasks/{task['task_id']}/publish", headers=users[0],
        data={"result_digest": digest, "file_sha256": sha,
              "payload_json": json.dumps(payload, ensure_ascii=False)},
        files={"file": (task["filename"], output.read_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert saved.status_code == 200, saved.text
    receipt = saved.json()
    assert receipt["excel_published"] and receipt["task_record_count"] == 1
    book = load_workbook(io.BytesIO(client.get(receipt["download_url"], headers=users[0]).content))
    assert [book["应届"][f"{column}2"].value for column in "ABC"] == [1, "张三", "计算机学院"]
    assert book["应届"]["D2"].value is None
    assert book["应届"]["C3"].value == "."
    book.close()
    with sessions() as db:
        members = db.query(StructuredDatasetMember).all()
        assert len(members) == 1
        assert members[0].source_versions_json == [source["document_version_id"] for source in sources]
    reused = client.post(f"{BASE}/{receipt['dataset_id']}/records", headers=users[0], json={
        **{key: sources[1][key] for key in (
            "document_id", "document_version_id", "extraction_run_id", "source_sha256")},
        "target_working_copy_id": receipt["working_copy_id"],
        "target_document_id": receipt["document_id"],
        "target_document_version_id": receipt["document_version_id"],
        "content_origin": "WORKBUDDY", "idempotency_key": "reuse-secondary", "records": []})
    assert reused.status_code == 409
    assert reused.json()["error"]["code"] == "SOURCE_ALREADY_APPENDED"

    next_id = seed_workbuddy_document(sessions, filename="新考察.pdf", content_type="application/pdf",
        units=[{"kind": "PDF_PAGE", "page_number": 1, "text": "姓名 李四 学院 数学学院"}])
    next_source = client.get(f"{BASE}/sources/{next_id}?content_origin=WORKBUDDY", headers=users[0]).json()
    checkout = client.post("/api/excel-writeback/checkout", headers=users[0], json={
        "filename": "教师引进汇总.xlsx", "request_key": "recruitment-append", "review_mode": "DIRECT_SAVE"})
    assert checkout.status_code == 200, checkout.text
    next_task = checkout.json()
    next_payload = TaskPayload.model_validate({"manifest": {
        "sources": [{key: next_source[key] for key in (
            "document_id", "document_version_id", "extraction_run_id", "source_sha256")}],
        "write_scope": "TEMPLATE_ALL",
        "target": {"dataset_id": next_task["dataset_id"],
                   "document_version_id": next_task["document_version_id"],
                   "filename": next_task["filename"], "sheet_name": next_task["sheet_name"],
                   "layout": next_task["layout"], "labels": next_task["labels"]}},
        "records": [{"source_index": 0, "source_indices": [0], "fields": [
            {"blank_reason": "TARGET_DERIVED"},
            {"raw_text": "李四", "evidence_refs": [{
                "source_unit_id": next_source["source_units"][0]["source_unit_id"], "quote": "姓名 李四"}]},
            {"raw_text": "数学学院", "evidence_refs": [{
                "source_unit_id": next_source["source_units"][0]["source_unit_id"], "quote": "学院 数学学院"}]},
            {"blank_reason": "SOURCE_FIELD_NOT_FOUND"}]}]}).model_dump(mode="json")
    next_base = tmp_path / "next-base.xlsx"
    next_base.write_bytes(client.get(next_task["download_url"], headers=users[0]).content)
    next_output = tmp_path / "next-result.xlsx"
    patch_workbook(baseline=next_base, output=next_output, sheet_name="应届",
                   labels=next_task["labels"], rows=[[None, "李四", "数学学院", None]],
                   layout=next_task["layout"], scoped_a_i=True)
    next_sha = hashlib.sha256(next_output.read_bytes()).hexdigest()
    next_digest = writeback_digest(next_payload, next_sha)
    assert client.post(f"/api/excel-writeback/tasks/{next_task['task_id']}/preview", headers=users[0],
        json={"result_digest": next_digest, "file_sha256": next_sha}).status_code == 200
    appended = client.post(f"/api/excel-writeback/tasks/{next_task['task_id']}/publish", headers=users[0],
        data={"result_digest": next_digest, "file_sha256": next_sha,
              "payload_json": json.dumps(next_payload, ensure_ascii=False)},
        files={"file": (next_task["filename"], next_output.read_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert appended.status_code == 200, appended.text
    newer = appended.json()
    assert newer["version_number"] == receipt["version_number"] + 1
    assert newer["document_id"] == receipt["document_id"]
    book = load_workbook(io.BytesIO(client.get(newer["download_url"], headers=users[0]).content))
    assert [book["应届"][f"{column}4"].value for column in "ABC"] == [2, "李四", "数学学院"]
    assert book["应届"]["C3"].value == "."
    book.close()


def test_unrelated_supplement_cannot_be_merged_into_person_row(demo, monkeypatch, tmp_path):
    client, sessions, users = demo
    task, payload, baseline, _ = _grouped_task(
        client, sessions, users[0], monkeypatch, tmp_path, supplemental_has_identity=False)
    monkeypatch.syspath_prepend("apps/mcp")
    from file_agent_mcp.xlsx_writeback import patch_workbook

    output = tmp_path / "unrelated.xlsx"
    patch_workbook(baseline=baseline, output=output, sheet_name="应届", labels=task["labels"],
                   rows=[[None, "张三", "计算机学院", None]], layout=task["layout"], scoped_a_i=True)
    sha = hashlib.sha256(output.read_bytes()).hexdigest()
    digest = writeback_digest(payload, sha)
    preview = client.post(f"/api/excel-writeback/tasks/{task['task_id']}/preview", headers=users[0],
        json={"result_digest": digest, "file_sha256": sha})
    assert preview.status_code == 200, preview.text
    rejected = client.post(f"/api/excel-writeback/tasks/{task['task_id']}/publish", headers=users[0],
        data={"result_digest": digest, "file_sha256": sha,
              "payload_json": json.dumps(payload, ensure_ascii=False)},
        files={"file": (task["filename"], output.read_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert rejected.status_code == 422
    assert rejected.json()["error"]["code"] == "TEMPLATE_PERSON_SOURCE_MISMATCH"
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 0

"""Shared 2019 catalog indexing and exact-match behavior."""

import hashlib
import importlib.util
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from app.db.models import ReferenceCatalogVersion
from app.modules.reference_catalogs.service import CATALOG_KEY, build_index, normalize_issn, query_catalog
from app.modules.structured_datasets.schemas import FieldCandidate
from app.modules.structured_datasets.teacher_render import render_teacher_cell


def test_index_uses_only_published_visible_sheets(tmp_path):
    book = Workbook()
    rules = book.active
    rules.title = "Sheet1"
    rules["A1"] = "SCI 一区 A++；二区 A+"
    entries = book.create_sheet("A+、A、B类期刊目录")
    entries.append([None, "期刊名称", None, "刊号", "所属类别"])
    entries.append([None, "期刊名称", None, "刊号", "所属类别"])
    for index in range(621):
        entries.append([None, f"Journal {index}", None,
                        f"1234-{index:04d}", ("A+", "第三类 A", "第四类 B")[index % 3]])
    hidden = book.create_sheet("待确定")
    hidden.sheet_state = "hidden"
    hidden.append([None, "Hidden Journal", None, "9999-9999", "A+"])
    path = tmp_path / "converted.xlsx"
    book.save(path)

    index = build_index(path)
    assert index["entry_count"] == 621
    assert len(index["entries"]) == 621
    assert [index["entries"][i]["school_grade"] for i in range(3)] == ["A+", "A", "B"]
    assert all(item["journal_name"] != "Hidden Journal" for item in index["entries"])
    assert index["entries"][0]["issn_keys"] == ["12340000"]


def test_query_rejects_issn_name_disagreement(tmp_path, monkeypatch):
    from app.modules.reference_catalogs import service

    content = {"source_sha256": "a" * 64, "document_version_id": "version-1",
               "entries": [
                   {"row": 3, "journal_name": "First Journal", "issn_keys": ["12345678"], "school_grade": "A+"},
                   {"row": 4, "journal_name": "Second Journal", "issn_keys": ["87654321"], "school_grade": "B"}],
               "rules_sheet": "Sheet1", "rules": []}
    root = tmp_path / "storage"
    root.mkdir()
    target = root / "index.json"
    encoded = json.dumps(content).encode()
    target.write_bytes(encoded)
    record = ReferenceCatalogVersion(catalog_key=CATALOG_KEY, version_number=1,
        source_sha256="a" * 64, document_version_id="version-1", entry_count=2,
        index_storage_path="index.json", index_sha256=hashlib.sha256(encoded).hexdigest(), status="ACTIVE")
    record.id = "catalog-version-1"
    monkeypatch.setattr(service, "active_catalog", lambda _db: record)
    monkeypatch.setattr(service, "get_settings", lambda: type("Settings", (), {"file_storage_root": root})())

    result = query_catalog(None, issn="1234-5678", journal_name="First Journal")
    assert result["match_status"] == "FOUND"
    assert result["matches"][0]["school_grade"] == "A+"
    assert query_catalog(None, issn="1234-5678", journal_name="Second Journal")["match_status"] == "CONFLICT"
    assert query_catalog(None, issn="1234-5678", journal_name="Missing Journal")["match_status"] == "CONFLICT"
    assert query_catalog(None, journal_name="Missing Journal")["match_status"] == "NOT_FOUND"
    assert normalize_issn("ISSN:0736-587X") == "0736587X"


def test_teacher_render_uses_only_original_fragments():
    field = FieldCandidate.model_validate({
        "fragments": [{"source_unit_id": "00000000-0000-0000-0000-000000000001",
                       "quote": "博士毕业 2026.09 西安理工大学 计算机科学与技术"},
                      {"source_unit_id": "00000000-0000-0000-0000-000000000001",
                       "quote": "本科毕业 2017.07 西南民族大学 网络工程"}],
        "teacher_render": {"kind": "EDUCATION", "items": [
            {"stage": "博士", "parts": [
                {"fragment_index": 0, "role": "GRADUATION", "text": "2026.09"},
                {"fragment_index": 0, "role": "SCHOOL", "text": "西安理工大学"},
                {"fragment_index": 0, "role": "MAJOR", "text": "计算机科学与技术"}]},
            {"stage": "本科", "parts": [
                {"fragment_index": 1, "role": "GRADUATION", "text": "2017.07"},
                {"fragment_index": 1, "role": "SCHOOL", "text": "西南民族大学"},
                {"fragment_index": 1, "role": "MAJOR", "text": "网络工程"}]}]}})
    rendered, audit = render_teacher_cell(field, column="I", db=None)
    assert rendered == ("博士：2026.09\n西安理工大学\n计算机科学与技术\n\n"
                        "硕士：硕博连读\n\n本科：2017.07\n西南民族大学\n网络工程")
    assert audit[0]["reason"] == "博士资料存在且硕士栏未填写"


def test_teacher_paper_grade_is_derived_from_frozen_catalog(monkeypatch):
    import uuid
    from app.modules.structured_datasets import teacher_render

    version = str(uuid.uuid4())
    calls = []
    def fake_query(_db, **kwargs):
        calls.append(kwargs)
        return {"catalog_version_id": version, "match_status": "FOUND", "matches": [{
            "school_grade": "A+", "sheet": "A+、A、B类期刊目录", "row": 13,
            "category": "A+"}]}
    monkeypatch.setattr(teacher_render, "query_catalog", fake_query)
    monkeypatch.setattr(teacher_render, "catalog_info", lambda _db: {"catalog_version_id": version})
    field = FieldCandidate.model_validate({
        "fragments": [{"source_unit_id": "00000000-0000-0000-0000-000000000001",
                       "quote": "Journal One ISSN 1234-5678 SCI-2区 IF 2.1"}],
        "teacher_render": {"kind": "PAPER", "catalog_version_id": version, "items": [{
            "ordinal": 1, "parts": [
                {"fragment_index": 0, "role": "JOURNAL_NAME", "text": "Journal One"},
                {"fragment_index": 0, "role": "ISSN", "text": "1234-5678"},
                {"fragment_index": 0, "role": "ZONE", "text": "SCI-2区"},
                {"fragment_index": 0, "role": "TEXT", "text": "IF 2.1"}]}]}})
    rendered, audit = render_teacher_cell(field, column="L", db=None)
    assert rendered == "1.A+类,SCI-2区,IF 2.1"
    assert audit[0]["row"] == 13
    assert calls == [{"issn": "1234-5678", "journal_name": "Journal One"}]


@pytest.mark.parametrize(("zone", "grade"), [
    ("SCI/A&HCI/SSCI中科院1 区", "A++"),
    ("SCI/A&HCI/SSCI中科院2 区", "A+"),
    ("SCI-2区", "A+"),
])
def test_teacher_paper_combined_zone_uses_catalog_rule(monkeypatch, zone, grade):
    import uuid
    from app.modules.structured_datasets import teacher_render

    version = str(uuid.uuid4())
    monkeypatch.setattr(teacher_render, "catalog_info", lambda _db: {"catalog_version_id": version})
    field = FieldCandidate.model_validate({
        "fragments": [{"source_unit_id": "00000000-0000-0000-0000-000000000001",
                       "quote": f"{zone}，IF:9.098，被引0 次"}],
        "teacher_render": {"kind": "PAPER", "catalog_version_id": version, "items": [{
            "ordinal": 1, "parts": [
                {"fragment_index": 0, "role": "ZONE", "text": zone},
                {"fragment_index": 0, "role": "TEXT", "text": "IF:9.098，被引0 次"}]}]}})

    rendered, audit = render_teacher_cell(field, column="L", db=None)
    assert rendered == f"1.{grade}类,{zone},IF:9.098，被引0 次"
    assert audit[0]["zone_grade"] == grade


def test_teacher_paper_explicit_zone_requires_zone_role():
    field = FieldCandidate.model_validate({
        "fragments": [{"source_unit_id": "00000000-0000-0000-0000-000000000001",
                       "quote": "SCI/A&HCI/SSCI中科院1 区，IF:9.098"}],
        "teacher_render": {"kind": "PAPER", "items": [{"ordinal": 1, "parts": [
            {"fragment_index": 0, "role": "TEXT", "text": "SCI/A&HCI/SSCI中科院1 区，IF:9.098"}]}]}})

    with pytest.raises(ValueError, match="TEACHER_RENDER_ZONE_REQUIRED"):
        render_teacher_cell(field, column="L", db=None)


def test_teacher_paper_catalog_conflict_keeps_source_without_grade(monkeypatch):
    import uuid
    from app.modules.structured_datasets import teacher_render

    version = str(uuid.uuid4())
    monkeypatch.setattr(teacher_render, "catalog_info", lambda _db: {"catalog_version_id": version})
    monkeypatch.setattr(teacher_render, "query_catalog", lambda _db, **_kwargs: {
        "match_status": "CONFLICT", "matches": []})
    field = FieldCandidate.model_validate({
        "fragments": [{"source_unit_id": "00000000-0000-0000-0000-000000000001",
                       "quote": "SCI/A&HCI/SSCI中科院1 区 ISSN 1234-5678"}],
        "teacher_render": {"kind": "PAPER", "catalog_version_id": version, "items": [{
            "ordinal": 1, "parts": [
                {"fragment_index": 0, "role": "ZONE", "text": "SCI/A&HCI/SSCI中科院1 区"},
                {"fragment_index": 0, "role": "ISSN", "text": "1234-5678"}]}]}})

    rendered, audit = render_teacher_cell(field, column="L", db=None)
    assert rendered == "1.SCI/A&HCI/SSCI中科院1 区"
    assert audit[0]["school_grade_status"] == "CONFLICT"


def test_teacher_paper_ccf_skips_catalog_even_with_zone_and_issn(monkeypatch):
    from app.modules.structured_datasets import teacher_render

    def unexpected_catalog(*_args, **_kwargs):
        raise AssertionError("CCF 条目不应读取 2019 版目录")

    monkeypatch.setattr(teacher_render, "catalog_info", unexpected_catalog)
    monkeypatch.setattr(teacher_render, "query_catalog", unexpected_catalog)
    field = FieldCandidate.model_validate({
        "fragments": [{"source_unit_id": "00000000-0000-0000-0000-000000000001",
                       "quote": "CCF C类期刊，SCI-2区，ISSN 1234-5678，被引3次"}],
        "teacher_render": {"kind": "PAPER", "items": [{"ordinal": 1, "parts": [
            {"fragment_index": 0, "role": "TEXT", "text": "CCF C类期刊"},
            {"fragment_index": 0, "role": "ZONE", "text": "SCI-2区"},
            {"fragment_index": 0, "role": "ISSN", "text": "1234-5678"},
            {"fragment_index": 0, "role": "TEXT", "text": "被引3次"}]}]}})

    rendered, audit = render_teacher_cell(field, column="L", db=None)
    assert rendered == "1.CCF C类期刊,SCI-2区,被引3次"
    assert audit == [{"ordinal": 1, "school_grade_status": "CCF_SOURCE"}]


def test_teacher_paper_ccf_grade_moves_before_other_facts(monkeypatch):
    from app.modules.structured_datasets import teacher_render

    monkeypatch.setattr(teacher_render, "catalog_info", lambda _db: (_ for _ in ()).throw(
        AssertionError("CCF 条目不应读取目录")))
    field = FieldCandidate.model_validate({
        "fragments": [{"source_unit_id": "00000000-0000-0000-0000-000000000001",
                       "quote": "EI索引，CCF A类会议，被引3次"}],
        "teacher_render": {"kind": "PAPER", "items": [{"ordinal": 1, "parts": [
            {"fragment_index": 0, "role": "TEXT", "text": "EI索引，CCF A类会议，被引3次"}]}]}})

    rendered, audit = render_teacher_cell(field, column="L", db=None)
    assert rendered == "1.CCF A类会议,EI索引,被引3次"
    assert audit == [{"ordinal": 1, "school_grade_status": "CCF_SOURCE"}]


def test_teacher_paper_catalog_b_grade_moves_before_source_facts(monkeypatch):
    import uuid
    from app.modules.structured_datasets import teacher_render

    version = str(uuid.uuid4())
    monkeypatch.setattr(teacher_render, "catalog_info", lambda _db: {"catalog_version_id": version})
    monkeypatch.setattr(teacher_render, "query_catalog", lambda _db, **_kwargs: {
        "catalog_version_id": version, "match_status": "FOUND", "matches": [{
            "school_grade": "B", "sheet": "A+、A、B类期刊目录", "row": 15,
            "category": "第四类 B"}]})
    field = FieldCandidate.model_validate({
        "fragments": [{"source_unit_id": "00000000-0000-0000-0000-000000000001",
                       "quote": "Journal One ISSN 1234-5678，被引3次"}],
        "teacher_render": {"kind": "PAPER", "catalog_version_id": version, "items": [{
            "ordinal": 1, "parts": [
                {"fragment_index": 0, "role": "ISSN", "text": "1234-5678"},
                {"fragment_index": 0, "role": "TEXT", "text": "被引3次"}]}]}})

    rendered, _ = render_teacher_cell(field, column="L", db=None)
    assert rendered == "1.B类,被引3次"


def test_teacher_paper_ccf_priority_is_per_item(monkeypatch):
    import uuid
    from app.modules.structured_datasets import teacher_render

    version = str(uuid.uuid4())
    calls = []
    monkeypatch.setattr(teacher_render, "catalog_info", lambda _db: {"catalog_version_id": version})
    def fake_query(_db, **kwargs):
        calls.append(kwargs)
        return {"catalog_version_id": version, "match_status": "FOUND", "matches": [{
            "school_grade": "A+", "sheet": "A+、A、B类期刊目录", "row": 13,
            "category": "A+"}]}
    monkeypatch.setattr(teacher_render, "query_catalog", fake_query)
    field = FieldCandidate.model_validate({
        "fragments": [
            {"source_unit_id": "00000000-0000-0000-0000-000000000001",
             "quote": "CCF C类期刊，SCI-1区，ISSN 1111-1111"},
            {"source_unit_id": "00000000-0000-0000-0000-000000000001",
             "quote": "Journal One ISSN 1234-5678 SCI-2区"}],
        "teacher_render": {"kind": "PAPER", "catalog_version_id": version, "items": [
            {"ordinal": 1, "parts": [
                {"fragment_index": 0, "role": "TEXT", "text": "CCF C类期刊"},
                {"fragment_index": 0, "role": "ZONE", "text": "SCI-1区"},
                {"fragment_index": 0, "role": "ISSN", "text": "1111-1111"}]},
            {"ordinal": 2, "parts": [
                {"fragment_index": 1, "role": "JOURNAL_NAME", "text": "Journal One"},
                {"fragment_index": 1, "role": "ISSN", "text": "1234-5678"},
                {"fragment_index": 1, "role": "ZONE", "text": "SCI-2区"}]}]}})

    rendered, audit = render_teacher_cell(field, column="L", db=None)
    assert rendered == "1.CCF C类期刊,SCI-1区\n2.A+类,SCI-2区"
    assert audit[0] == {"ordinal": 1, "school_grade_status": "CCF_SOURCE"}
    assert calls == [{"issn": "1234-5678", "journal_name": "Journal One"}]


def test_multiline_writeback_keeps_wrap_and_sets_visible_height(tmp_path):
    module_path = Path(__file__).resolve().parents[3] / "mcp" / "file_agent_mcp" / "xlsx_writeback.py"
    spec = importlib.util.spec_from_file_location("workbuddy_xlsx_writeback", module_path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    book = Workbook()
    sheet = book.active
    sheet.title = "应届"
    sheet.append(["序号", "姓名", "学习经历"])
    sheet["C2"].alignment = Alignment(wrap_text=True)
    baseline = tmp_path / "baseline.xlsx"
    result = tmp_path / "result.xlsx"
    book.save(baseline)
    value = "博士：2026.09\n西安理工大学\n\n硕士：硕博连读\n\n本科：2017.07"
    module.patch_workbook(baseline=baseline, output=result, sheet_name="应届",
                          labels=["序号", "姓名", "学习经历"],
                          rows=[[None, "张三", value]],
                          layout={"header_range": "A1:C1", "data_range": "A2:C2"}, scoped_a_i=True)
    published = load_workbook(result)
    assert published["应届"]["C2"].value == value
    assert published["应届"]["C2"].alignment.wrap_text is True
    assert published["应届"].row_dimensions[2].height >= 108
    published.close()

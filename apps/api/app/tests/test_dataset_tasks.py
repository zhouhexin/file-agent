"""完整表格任务的真实发布、无编号审批、逐列覆盖、标记和恢复回归。"""

import io
from uuid import uuid4

from openpyxl import load_workbook

from app.db.models import StructuredDatasetTask, StructuredDataset, StructuredDatasetMember, DocumentVersion, DocumentExtractionRun
from app.modules.structured_datasets.task_schemas import TaskPayload
from app.modules.structured_datasets.tasks import DatasetTaskService, payload_digest
from app.modules.structured_datasets.service import StructuredDatasetService
from app.tests.test_structured_datasets import demo, seed_workbuddy_document, seed_excel_target, BASE


def make_payload(client, sessions, headers, *, template=False, filename="test.pdf",
                 content_type="application/pdf", unit_kind="PDF_PAGE"):
    """测试文本只有整栏声明，没有任何项号，且有一列明确留空。"""
    document_id = seed_workbuddy_document(sessions, filename=filename, content_type=content_type,
        units=[{"kind": unit_kind, "page_number": 1 if unit_kind != "TEXT_SECTION" else None,
                "text": "教学工作量：120学时\n本栏情况属实\n教务处圆章\n年度考核未提供"}])
    source = client.get(f"{BASE}/sources/{document_id}", headers=headers).json()
    title = "教学工作量不少于80学时"
    fields = [{"key": "hours", "label": title, "field_type": "string"},
              {"key": "review", "label": "考核", "field_type": "string"}]
    def ref(text):
        """使用真实宿主 unit 标识，由服务端映射到正式 Page UUID。"""
        return {"source_unit_id": "unit-0", "quote": text}
    value = {"raw_text": "教学工作量：120学时", "confidence": 1,
        "evidence_refs": [ref("教学工作量：120学时")],
        "approval": {"binding_method": "WHOLE_SECTION", "declaration": ref("本栏情况属实"),
                     "seal": ref("教务处圆章"), "scope": [ref("教学工作量：120学时")], "explanation": "同栏盖章审批"},
        "assessment": {"title": title, "groups": [{"checks": [{"criterion": title, "operator": "GTE",
            "expected": "80", "observed": "120", "evidence": [ref("教学工作量：120学时")]}]}]}}
    target = {"filename": "教学汇总.xlsx", "labels": [title, "考核"],
              "new": {"name": "教学汇总", "fields": fields, "idempotency_key": "target"}}
    if template:
        target_id = seed_excel_target(sessions, filename="教学汇总.xlsx", sheets={"登记": [[title, "考核"]]})
        target.pop("new")
        target["uploaded"] = {"document_id": target_id, "idempotency_key": "target"}
        target["sheet_name"] = "登记"
    return TaskPayload.model_validate({"manifest": {"sources": [{key: source[key] for key in
        ("document_id", "document_version_id", "extraction_run_id", "source_sha256")}], "target": target, "sealed_only": True},
        "records": [{"source_index": 0, "fields": [value, {"blank_reason": "来源未提供考核情况"}]}]}).model_dump(mode="json")


def prepare(client, headers, payload, mode="REVIEW_REQUIRED", request_key="task-1"):
    """只传指纹和目标元数据，预览不得保存正文。"""
    digest = payload_digest(payload)
    response = client.post(f"{BASE}/tasks/preview", headers=headers,
        json={"request_key": request_key, "result_digest": digest, "review_mode": mode, "manifest": payload["manifest"]})
    assert response.status_code == 200, response.text
    return response.json()["task_id"], digest


def commit(client, headers, task_id, digest, payload=None):
    """执行一次确定阶段；测试直接观察服务器状态而非模拟成功字典。"""
    return client.post(f"{BASE}/tasks/{task_id}/commit", headers=headers,
        json={"result_digest": digest, **({"payload": payload} if payload else {})})


def test_merged_template_labels_match_only_same_physical_columns() -> None:
    """旧 XLS 转换的层级分隔符与重复叶列号可兼容，文字变化或错列仍必须拒绝。"""

    layout = {"header_range": "G2:J4", "data_range": "G5:J6"}
    expected = ["取得时间", "学历", "取得时间", "正常申报业绩 子项"]
    actual = ["取得时间 [G]", "学历", "取得时间 [I]", "正常申报业绩 / 子项"]
    assert DatasetTaskService._target_labels_match(expected, actual, layout)
    assert not DatasetTaskService._target_labels_match(expected, actual, None)
    assert not DatasetTaskService._target_labels_match(expected, ["取得时间 [H]", *actual[1:]], layout)
    assert not DatasetTaskService._target_labels_match(expected, [*actual[:3], "正常申报业绩 / 其他"], layout)
    assert not DatasetTaskService._target_labels_match(expected, [actual[2], actual[1], actual[0], actual[3]], layout)


def test_task_confirmation_publishes_real_file_and_preserves_headers(demo):
    """整栏审批无需项号；正文、表头、数据格末尾勾和当前文件分别核验。"""
    client, sessions, users = demo
    payload = make_payload(client, sessions, users[0], template=True)
    task_id, digest = prepare(client, users[0], payload)
    with sessions() as db:
        assert db.get(StructuredDatasetTask, task_id).payload_json == {}
        assert db.query(StructuredDataset).count() == 0
    assert commit(client, users[0], task_id, digest, payload).status_code == 409
    assert client.get(f"{BASE}/tasks/{task_id}", headers=users[1]).status_code == 404
    response = client.post(f"{BASE}/tasks/{task_id}/confirm", headers=users[0], json={"result_digest": digest})
    assert response.status_code == 200, response.text
    for step in range(3):
        response = commit(client, users[0], task_id, digest, payload if step == 0 else None)
        assert response.status_code == 200, response.text
    result = response.json()
    assert result["status"] == "COMPLETED"
    assert result["excel_published"] and result["task_record_count"] == 1 and result["checkmark_count"] == 1
    download = client.get(result["download_url"], headers=users[0])
    assert download.status_code == 200
    workbook = load_workbook(io.BytesIO(download.content))
    assert workbook.active["A1"].value == "教学工作量不少于80学时"
    assert workbook.active["A2"].value == "教学工作量：120学时 √"
    assert workbook.active["B2"].value is None
    query = client.post(f"{BASE}/{result['dataset_id']}/query", headers=users[0], json={}).json()
    field = query["records"][0]["fields"]["field_1"]
    assert field["raw_text"] == "教学工作量：120学时"
    assert field["display_value"].endswith(" √")
    repeat = commit(client, users[0], task_id, digest).json()
    assert repeat["document_version_id"] == result["document_version_id"]
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 1


def test_word_and_image_sources_use_real_workbuddy_text(demo):
    """Word 和图片来源都复用 WorkBuddy 固定正文，不回退到程序内部解析或 OCR。"""

    client, sessions, users = demo
    for index, (filename, content_type, kind) in enumerate((
        ("材料.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "TEXT_SECTION"),
        ("扫描件.png", "image/png", "IMAGE_PAGE"),
    )):
        payload = make_payload(client, sessions, users[0], filename=filename,
                               content_type=content_type, unit_kind=kind)
        if index:
            payload["manifest"]["target"]["new"]["name"] = "教学汇总2"
            payload["manifest"]["target"]["new"]["filename"] = "教学汇总2.xlsx"
            payload["manifest"]["target"]["filename"] = "教学汇总2.xlsx"
        payload = TaskPayload.model_validate(payload).model_dump(mode="json")
        task_id, digest = prepare(client, users[0], payload, mode="DIRECT_SAVE", request_key=f"source-kind-{index}")
        for step in range(3):
            response = commit(client, users[0], task_id, digest, payload if step == 0 else None)
            assert response.status_code == 200, response.text
        assert response.json()["excel_published"] is True


def test_template_identity_columns_b_to_i_are_saved_without_seal_or_checkmark(demo, monkeypatch):
    """身份列 B:I 只免盖章与资格判定；仍以原文证据、确认指纹和原模板发布。"""

    import pytest
    from pydantic import ValidationError
    client, sessions, users = demo
    labels = ["序号", "姓名", "性别", "出生年月", "单位", "岗位", "职称", "工号", "电话", "教学工作量不少于80学时"]
    target_id = seed_excel_target(sessions, filename="身份业绩表.xlsx", sheets={"登记": [labels]})
    items = ["姓名：张三", "性别：女", "出生年月：1990年1月", "单位：某学院",
             "岗位：教师", "职称：讲师", "工号：123", "电话：456"]
    fact = "教学工作量：120学时"
    source_id = seed_workbuddy_document(sessions, filename="身份材料.pdf", content_type="application/pdf",
        units=[{"kind": "PDF_PAGE", "page_number": 1,
                "text": "\n".join([*items, fact, "本栏情况属实", "教务处圆章"])}])
    source = client.get(f"{BASE}/sources/{source_id}", headers=users[0]).json()

    def quote(value):
        """定位到 WorkBuddy 已保存的同一真实页，不把免审批误作免证据。"""
        return {"source_unit_id": "unit-0", "quote": value}

    fields = [{"blank_reason": "模板序号列不由来源填写"}]
    fields.extend({"raw_text": value.split("：", 1)[1], "evidence_refs": [quote(value)]} for value in items)
    fields.append({"raw_text": fact, "evidence_refs": [quote(fact)],
        "approval": {"binding_method": "WHOLE_SECTION", "declaration": quote("本栏情况属实"),
                     "seal": quote("教务处圆章"), "scope": [quote(fact)], "explanation": "同栏盖章审批"},
        "assessment": {"title": labels[-1], "groups": [{"checks": [{"criterion": labels[-1],
            "operator": "GTE", "expected": "80", "observed": "120", "evidence": [quote(fact)]}]}]}})
    payload = {"manifest": {"sources": [{key: source[key] for key in
        ("document_id", "document_version_id", "extraction_run_id", "source_sha256")}],
        "target": {"filename": "身份业绩表.xlsx", "sheet_name": "登记", "labels": labels,
                   "uploaded": {"document_id": target_id, "sheet_name": "登记", "idempotency_key": "identity-target"}},
        "sealed_only": True, "identity_column_range": "B:I"},
        "records": [{"source_index": 0, "fields": fields}]}
    without_exemption = {**payload, "manifest": {key: value for key, value in payload["manifest"].items()
                                                   if key != "identity_column_range"}}
    with pytest.raises(ValidationError, match="审批证据"):
        TaskPayload.model_validate(without_exemption)
    normalized = TaskPayload.model_validate(payload).model_dump(mode="json")
    from pathlib import Path
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3] / "mcp"))
    from file_agent_mcp.dataset_task_contracts import TaskPayload as HostPayload
    assert HostPayload.model_validate(payload).model_dump(mode="json") == normalized
    without_identity_evidence = TaskPayload.model_validate(payload).model_dump(mode="json")
    without_identity_evidence["records"][0]["fields"][1]["evidence_refs"] = []
    with pytest.raises(ValidationError, match="原文引用"):
        TaskPayload.model_validate(without_identity_evidence)
    task_id, digest = prepare(client, users[0], normalized, mode="DIRECT_SAVE", request_key="identity-b-i")
    for step in range(3):
        response = commit(client, users[0], task_id, digest, normalized if step == 0 else None)
        assert response.status_code == 200, response.text
    result = response.json()
    assert result["status"] == "COMPLETED" and result["checkmark_count"] == 1
    workbook = load_workbook(io.BytesIO(client.get(result["download_url"], headers=users[0]).content))
    sheet = workbook["登记"]
    assert [sheet.cell(2, column).value for column in range(2, 10)] == [item.split("：", 1)[1] for item in items]
    assert all("√" not in str(sheet.cell(2, column).value) for column in range(2, 10))
    assert sheet["J2"].value == "教学工作量：120学时 √"
    workbook.close()


def test_task_rejects_changed_payload_and_recovers_after_export_failure(demo, monkeypatch):
    """导出失败后持久保留追加进度，重试不生成重复行。"""
    client, sessions, users = demo
    payload = make_payload(client, sessions, users[0])
    task_id, digest = prepare(client, users[0], payload, mode="DIRECT_SAVE")
    changed = TaskPayload.model_validate(payload).model_dump(mode="json")
    changed["records"][0]["fields"][1]["blank_reason"] = "不同内容"
    assert commit(client, users[0], task_id, digest, changed).status_code == 409
    assert commit(client, users[0], task_id, digest, payload).status_code == 200
    assert commit(client, users[0], task_id, digest).status_code == 200
    original = StructuredDatasetService.export
    def failed_export(*args, **kwargs):
        """只模拟发布失败，不触碰实际文件。"""
        raise OSError("test failure")
    monkeypatch.setattr(StructuredDatasetService, "export", failed_export)
    assert commit(client, users[0], task_id, digest).status_code == 500
    state = client.get(f"{BASE}/tasks/{task_id}", headers=users[0]).json()
    assert not state["excel_published"] and state["status"] == "PUBLISHING"
    monkeypatch.setattr(StructuredDatasetService, "export", original)
    resumed = commit(client, users[0], task_id, digest)
    assert resumed.status_code == 200, resumed.text
    assert resumed.json()["status"] == "COMPLETED"
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 1


def test_missing_column_or_blank_reason_is_rejected(demo):
    """未解释的漏列不能通过统一提交契约。"""
    import pytest
    from pydantic import ValidationError
    client, sessions, users = demo
    payload = make_payload(client, sessions, users[0])
    payload["records"][0]["fields"].pop()
    with pytest.raises(ValidationError):
        TaskPayload.model_validate(payload)


def test_failed_target_job_reports_terminal_parent_state(demo):
    """异步转换失败必须在原 task_id 上给出终态、job ID 和重新预览动作。"""

    from app.modules.managed_files.jobs import FilesystemJobQueue
    from app.modules.structured_datasets.tasks import mark_target_preparation_failed
    client, sessions, users = demo
    payload = make_payload(client, sessions, users[0])
    task_id, _ = prepare(client, users[0], payload, mode="DIRECT_SAVE")
    with sessions() as db:
        user_id = db.get(StructuredDatasetTask, task_id).user_id
        job = FilesystemJobQueue(db).create_job(
            job_type="STRUCTURED_DATASET_TARGET_PREPARE", queue_name="STRUCTURED_EXTRACTION",
            root_id=None, created_by=user_id, deduplication_key=f"test-target-{task_id}",
            priority=20, max_attempts=1,
            payload={"structured_dataset_task_id": task_id},
        )
        row = db.get(StructuredDatasetTask, task_id)
        row.status = "PREPARING_TARGET"
        row.progress_json = {"filesystem_job_id": job.id}
        mark_target_preparation_failed(db, job, code="TARGET_EXCEL_CONVERSION_SCHEMA_MISMATCH",
                                       message="模板网格不匹配，请重新预览。", needs_input=True)
        db.commit()
        job_id = job.id
    state = client.get(f"{BASE}/tasks/{task_id}", headers=users[0]).json()
    assert state["status"] == "NEEDS_INPUT"
    assert state["error"]["code"] == "TARGET_EXCEL_CONVERSION_SCHEMA_MISMATCH"
    assert state["filesystem_job_id"] == job_id
    assert state["next_action"] == "create_new_preview"
    assert state["excel_published"] is False


def test_complex_xls_task_missing_grid_needs_input_without_queue(demo):
    """统一任务在复杂 XLS 缺少真实表格坐标时终止于父 task，不生成仅表头文件。"""

    from app.db.models import FilesystemJob
    client, sessions, users = demo
    payload = make_payload(client, sessions, users[0])
    template_id = seed_workbuddy_document(
        sessions, filename="教学汇总.xls", content_type="application/vnd.ms-excel",
        units=[{"kind": "SHEET", "sheet_name": "登记", "text": "姓名\t考核"}],
    )
    with sessions() as db:
        version = db.query(DocumentVersion).filter_by(document_id=template_id).one()
        run = db.query(DocumentExtractionRun).filter_by(document_id=template_id).one()
        before_jobs = db.query(FilesystemJob).count()
        version_id, run_id = version.id, run.id
    target = payload["manifest"]["target"]
    target.pop("new")
    target.update({"filename": "教学汇总.xlsx", "sheet_name": "登记",
        "layout": {"header_range": "A2:B2", "data_range": "A3:B3"},
        "uploaded": {"document_id": template_id, "document_version_id": version_id,
                     "extraction_run_id": run_id, "sheet_name": "登记",
                     "layout": {"header_range": "A2:B2", "data_range": "A3:B3"},
                     "legacy_target_policy": "CONVERT_TO_XLSX", "idempotency_key": "template"}})
    payload = TaskPayload.model_validate(payload).model_dump(mode="json")
    task_id, digest = prepare(client, users[0], payload, mode="DIRECT_SAVE", request_key="missing-grid-task")
    response = commit(client, users[0], task_id, digest, payload)
    assert response.status_code == 200, response.text
    state = response.json()
    assert state["status"] == "NEEDS_INPUT"
    assert state["error"]["code"] == "TARGET_EXCEL_HEADER_CELLS_MISSING"
    assert state["next_action"] == "create_new_preview"
    with sessions() as db:
        assert db.query(FilesystemJob).count() == before_jobs
        assert db.query(StructuredDataset).count() == 0


def test_mcp_payload_fingerprint_matches_server_and_target_is_reserved(demo, monkeypatch):
    """真实宿主 schema 的指纹与 API 一致；旧导出不能穿透进行中的任务。"""
    from pathlib import Path
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3] / "mcp"))
    from file_agent_mcp.dataset_task_contracts import TaskPayload as HostPayload
    from file_agent_mcp.extraction_drafts import result_digest
    client, sessions, users = demo
    payload = make_payload(client, sessions, users[0])
    host = HostPayload.model_validate(payload).model_dump(mode="json")
    assert host == payload
    assert result_digest(host) == payload_digest(payload)
    task_id, digest = prepare(client, users[0], payload, mode="DIRECT_SAVE")
    assert commit(client, users[0], task_id, digest, host).status_code == 200
    with sessions() as db:
        target = db.query(StructuredDataset).one()
        assert target.active_task_id == task_id
        dataset_id = target.id
    blocked = client.post(f"{BASE}/{dataset_id}/exports", headers=users[1])
    assert blocked.status_code == 409 and blocked.json()["error"]["code"] == "DATASET_TASK_IN_PROGRESS"
    assert commit(client, users[0], task_id, digest).status_code == 200
    assert commit(client, users[0], task_id, digest).json()["excel_published"]
    with sessions() as db:
        assert db.get(StructuredDataset, dataset_id).active_task_id is None


def test_fragment_field_keeps_quotes_and_allows_whole_section(demo):
    """多个已验证片段可合并同一格，不要求原文含完整拼接后的字符串。"""
    client, sessions, users = demo
    payload = make_payload(client, sessions, users[0])
    field = payload["records"][0]["fields"][0]
    field["raw_text"] = None
    field["fragments"] = [{"source_unit_id": "unit-0", "quote": "教学工作量：120学时"},
                          {"source_unit_id": "unit-0", "quote": "本栏情况属实"}]
    field["approval"]["scope"] = list(field["fragments"])
    payload = TaskPayload.model_validate(payload).model_dump(mode="json")
    task_id, digest = prepare(client, users[0], payload, mode="DIRECT_SAVE")
    for step in range(3):
        response = commit(client, users[0], task_id, digest, payload if step == 0 else None)
        assert response.status_code == 200, response.text
    result = response.json()
    wb = load_workbook(io.BytesIO(client.get(result["download_url"], headers=users[0]).content))
    assert wb.active["A2"].value == "教学工作量：120学时\n本栏情况属实 √"


def test_special_template_source_facts_are_marked_per_fragment(demo, monkeypatch):
    """专用模板 J 与 P 的无章事实按片段落表；P 的印章不能覆盖同格无章教学评价。"""
    client, sessions, users = demo
    labels = [f"列{index}" for index in range(1, 17)]
    target_id = seed_excel_target(sessions, filename="来源状态表.xlsx", sheets={"登记": [labels]})
    source_id = seed_workbuddy_document(sessions, filename="评审材料.pdf", content_type="application/pdf",
        units=[{"kind": "PDF_PAGE", "page_number": 1,
            "text": "班主任三年\n教学工作量120学时\n教学效果优秀\n本栏情况属实\n教务处圆章"}])
    source = client.get(f"{BASE}/sources/{source_id}", headers=users[0]).json()
    def ref(value):
        """片段与审批均引用相同的固定来源页面。"""
        return {"source_unit_id": "unit-0", "quote": value}
    fields = [{"blank_reason": "来源没有该列资料"} for _ in labels]
    fields[9] = {"fragments": [ref("班主任三年")], "source_only_fragments": [ref("班主任三年")],
                  "confidence": 1}
    fields[15] = {"fragments": [ref("教学工作量120学时"), ref("教学效果优秀")],
                  "source_only_fragments": [ref("教学效果优秀")], "confidence": 1,
                  "approval": {"binding_method": "WHOLE_SECTION", "declaration": ref("本栏情况属实"),
                      "seal": ref("教务处圆章"), "scope": [ref("教学工作量120学时")],
                      "explanation": "工作量同栏盖章，教学效果无章"}}
    manifest = {"sources": [{key: source[key] for key in
        ("document_id", "document_version_id", "extraction_run_id", "source_sha256")}],
        "target": {"filename": "来源状态表.xlsx", "sheet_name": "登记", "labels": labels,
            "uploaded": {"document_id": target_id, "sheet_name": "登记", "idempotency_key": "source-facts"}},
        "sealed_only": True, "identity_column_range": "B:I", "source_fact_column_range": "J:P"}
    payload = TaskPayload.model_validate({"manifest": manifest,
        "records": [{"source_index": 0, "fields": fields}]}).model_dump(mode="json")
    from pathlib import Path
    import pytest
    from pydantic import ValidationError
    monkey_payload = {**payload, "manifest": {**payload["manifest"], "source_fact_column_range": None}}
    with pytest.raises(ValidationError):
        TaskPayload.model_validate(monkey_payload)
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3] / "mcp"))
    from file_agent_mcp.dataset_task_contracts import TaskPayload as HostPayload
    assert HostPayload.model_validate(payload).model_dump(mode="json") == payload
    task_id, digest = prepare(client, users[0], payload, mode="DIRECT_SAVE", request_key="source-facts")
    for step in range(3):
        response = commit(client, users[0], task_id, digest, payload if step == 0 else None)
        assert response.status_code == 200, response.text
    result = response.json()
    workbook = load_workbook(io.BytesIO(client.get(result["download_url"], headers=users[0]).content))
    sheet = workbook["登记"]
    assert sheet["J2"].value == "班主任三年〔PDF 原文，未见对应印章〕"
    assert sheet["P2"].value == "教学工作量120学时\n教学效果优秀〔PDF 原文，未见对应印章〕"
    workbook.close()


def test_mcp_confirm_drives_real_api_publication_and_resume(demo, monkeypatch, tmp_path):
    """MCP 确认实际调用 API 并下载 XLSX；重复确认只恢复已发布结果。"""
    import asyncio
    from pathlib import Path
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[3] / "mcp"))
    from mcp.server.fastmcp import FastMCP
    from file_agent_mcp.dataset_tools import register_dataset_tools
    from file_agent_mcp.dataset_task_contracts import TaskPayload as HostPayload
    client, sessions, users = demo
    payload = make_payload(client, sessions, users[0], template=True)
    payload["records"][0]["fields"][0]["confidence"] = None
    calls = []

    class ApiAdapter:
        """只替换网络传输，数据库、API、模板和导出都是真实实现。"""
        async def dataset_request(self, method, path, payload=None):
            """进入现有受控 API 路由。"""
            calls.append((method, path, payload))
            response = client.request(method, BASE + path, headers=users[0], json=payload)
            assert response.status_code == 200, response.text
            return response.json()

        async def download_dataset_export(self, *, dataset_id, export_id, output_dir):
            """验证真实下载内容；测试缓存只写 pytest 临时目录。"""
            response = client.get(f"{BASE}/{dataset_id}/exports/{export_id}/content", headers=users[0])
            assert response.status_code == 200
            path = tmp_path / "教学汇总.xlsx"
            path.write_bytes(response.content)
            return {"filename": path.name, "resource_uri": path.as_uri(),
                "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "size_bytes": len(response.content)}

        async def close(self):
            """测试没有网络连接。"""

    async def scenario():
        """使用真实注册函数及本地草稿存储跑完整确认流程。"""
        mcp = FastMCP("task-integration")
        register_dataset_tools(mcp, ApiAdapter, lambda: tmp_path / "cache")
        tools = mcp._tool_manager._tools
        preview = await tools["structured_dataset_preview"].fn(HostPayload.model_validate(payload), "host-1")
        assert preview["preview_result"]["records"] == payload["records"]
        assert "records" not in calls[0][2]
        from uuid import UUID
        task_id = UUID(preview["task_id"])
        result = await tools["structured_dataset_confirm"].fn(task_id, preview["result_digest"])
        assert result.structuredContent["excel_published"] is True
        assert result.structuredContent["downloaded"] is True
        first_version = result.structuredContent["document_version_id"]
        result = await tools["structured_dataset_commit"].fn(task_id, preview["result_digest"])
        assert result.structuredContent["document_version_id"] == first_version
        workbook = load_workbook(tmp_path / "教学汇总.xlsx")
        assert workbook.active["A2"].value == "教学工作量：120学时 √"
        with sessions() as db:
            assert db.query(StructuredDatasetMember).count() == 1
    asyncio.run(scenario())


def test_second_source_failure_rolls_back_whole_batch_and_releases_target(demo):
    """后一来源错误不留下第一来源的半批记录，也不永久占用目标。"""
    import copy
    client, sessions, users = demo
    payload = make_payload(client, sessions, users[0])
    second = make_payload(client, sessions, users[0])
    payload["manifest"]["sources"].append(second["manifest"]["sources"][0])
    record = copy.deepcopy(second["records"][0])
    record["source_index"] = 1
    record["fields"][0]["evidence_refs"][0]["quote"] = "不在原文中的内容"
    payload["records"].append(record)
    task_id, digest = prepare(client, users[0], payload, mode="DIRECT_SAVE")
    assert commit(client, users[0], task_id, digest, payload).status_code == 200
    rejected = commit(client, users[0], task_id, digest)
    assert rejected.status_code == 422, rejected.text
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 0
        target = db.query(StructuredDataset).one()
        assert target.revision == 0 and target.active_task_id is None
        task = db.get(StructuredDatasetTask, task_id)
        assert task.progress_json["next_source"] == 0 and task.progress_json["revisions"] == []


def test_invalid_approval_identifies_column_and_requires_new_preview(demo):
    """审批引文不成立时指出具体列，已确认草稿不能反复重试或静默改写。"""
    client, sessions, users = demo
    payload = make_payload(client, sessions, users[0])
    value = payload["records"][0]["fields"][0]
    value["approval"]["declaration"] = value["evidence_refs"][0]
    payload = TaskPayload.model_validate(payload).model_dump(mode="json")
    task_id, digest = prepare(client, users[0], payload, mode="DIRECT_SAVE", request_key="invalid-approval-column")
    rejected = commit(client, users[0], task_id, digest, payload)
    assert rejected.status_code == 422, rejected.text
    detail = rejected.json()["error"]
    assert detail["code"] == "APPROVAL_DECLARATION_REQUIRED"
    assert (detail["details"]["record_index"], detail["details"]["column"]) == (1, "A")
    receipt = client.get(f"{BASE}/tasks/{task_id}", headers=users[0]).json()
    assert receipt["status"] == "NEEDS_INPUT"
    assert receipt["error"]["column"] == "A"
    assert commit(client, users[0], task_id, digest).json()["status"] == "NEEDS_INPUT"
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 0
        assert db.query(StructuredDataset).count() == 0

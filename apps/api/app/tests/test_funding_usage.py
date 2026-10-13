"""成果资助使用登记金额来自手写引文，余额仅作确定性扣减。"""

from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4
from zipfile import ZipFile
from hashlib import sha256
from decimal import Decimal
from datetime import datetime

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side

from app.modules.funding_usage.schemas import FundingUsageRequest
from app.modules.funding_usage.service import FundingUsageService, _registration_dates
from app.modules.funding_usage.workbook import append_summary_rows, rebuild_personal_sheets
from app.core.config import get_settings
from app.db.models import DocumentVersion, DocumentExtractionRun, StructuredDatasetPublication
from app.tests.test_structured_datasets import demo, seed_excel_target, seed_workbuddy_document


def _request(text: str, *, amount: str, decision: str = "CONFIRMED_SPEND",
             registration: str = "报销差旅300元"):
    name = "张三"
    name_start = text.index(name)
    registration_start = text.index(registration)
    amount_start = text.index(amount, registration_start)
    source = {"document_id": str(uuid4()), "document_version_id": str(uuid4()),
              "source_sha256": "a" * 64}
    return FundingUsageRequest.model_validate({
        "image": {**source, "extraction_run_id": str(uuid4())},
        "workbook": {**source, "source_sha256": "b" * 64},
        "filename": "使用登记与余额.xlsx", "request_key": "test-one",
        "rows": [{"page_number": 1, "name": name,
                  "name_ref": {"quote": name, "start": name_start, "end": name_start + len(name)},
                  "registration_ref": {"quote": registration, "start": registration_start,
                                       "end": registration_start + len(registration)},
                  "amounts": [{"quote": amount, "start": amount_start,
                               "end": amount_start + len(amount),
                               "decision": decision, "reason": "测试手写金额"}]}]})


def _service(text: str, amount_in_old_book: int = 1000):
    service = object.__new__(FundingUsageService)
    service._sources = lambda request: ({1: SimpleNamespace(text_content=text)}, Path("old.xlsx"))
    service._baseline = lambda path: {"张三": [amount_in_old_book]}
    return service


def test_only_handwritten_registration_amount_is_deducted():
    text = "张三 学院资助经费5000元 论文注册等3000元 使用登记 报销差旅300元"
    service = _service(text)
    result, _ = service._evaluate(_request(text, amount="300"))
    assert result[0]["usage_amount"] == 300
    assert result[0]["balance"] == 700
    assert result[0]["status"] == "已核对"


def test_handwritten_addition_expression_is_summed_without_rewriting_ocr():
    expression = "269+6560+2030.+3885+17250"
    registration = "报销费用" + expression
    text = "张三 学院资助经费80000元 使用登记 " + registration
    result, _ = _service(text, 50000)._evaluate(
        _request(text, amount=expression, registration=registration))
    assert result[0]["registration"] == registration
    assert result[0]["usage_amount"] == 29994
    assert result[0]["balance"] == 20006
    assert result[0]["status"] == "已核对"


def test_handwritten_addition_terms_can_be_quoted_separately():
    expression = "269+6560+2030.+3885+17250"
    registration = "报销费用" + expression
    text = "张三 使用登记 " + registration
    payload = _request(text, amount="269", registration=registration)
    for term in ("6560", "2030", "3885", "17250"):
        start = text.index(term, text.index(registration))
        payload.rows[0].amounts.append(type(payload.rows[0].amounts[0])(
            quote=term, start=start, end=start + len(term),
            decision="CONFIRMED_SPEND", reason="手写加法式中的加数"))
    result, _ = _service(text, 50000)._evaluate(payload)
    assert result[0]["usage_amount"] == 29994
    assert result[0]["balance"] == 20006


@pytest.mark.parametrize("expression", ["269++6560", "269-6560", "269+6560.", "269+2026年"])
def test_invalid_or_ambiguous_addition_expression_is_rejected(expression):
    text = "张三 使用登记 " + expression
    with pytest.raises(HTTPException) as exc:
        _service(text, 50000)._evaluate(
            _request(text, amount=expression, registration=expression))
    assert exc.value.detail["code"] == "FUNDING_AMOUNT_INVALID"


def test_uncertain_amount_leaves_balance_blank():
    text = "张三 学院资助经费5000元 使用登记 报销差旅300元"
    result, _ = _service(text)._evaluate(_request(text, amount="300", decision="UNCERTAIN"))
    assert result[0]["usage_amount"] is None
    assert result[0]["balance"] is None
    assert result[0]["status"] == "手写金额待核对"


def test_continue_keeps_later_balance_blank_until_earlier_amount_is_confirmed():
    text = "张三 使用登记 报销差旅300元；张三 使用登记 报销资料50元"
    request = _request(text, amount="300").model_dump(mode="json")
    second_name = text.rindex("张三")
    second_registration = text.index("报销资料50元")
    second_amount = text.index("50", second_registration)
    request["rows"].append({"page_number": 1, "name": "张三",
        "name_ref": {"quote": "张三", "start": second_name, "end": second_name + 2},
        "registration_ref": {"quote": "报销资料50元", "start": second_registration,
                             "end": second_registration + len("报销资料50元")},
        "amounts": [{"quote": "50", "start": second_amount, "end": second_amount + 2,
                     "decision": "CONFIRMED_SPEND", "reason": "手写支出"}]})
    request["rows"][0]["amounts"][0]["decision"] = "UNCERTAIN"
    payload = FundingUsageRequest.model_validate(request)
    service = _service(text)
    rows, _ = service._evaluate(payload, totals_override={"张三": [Decimal("700.00")]},
                                allow_duplicate_names=True, starting_sequence=1)
    assert [row["balance"] for row in rows] == [None, None]
    assert rows[1]["usage_amount"] == 50
    payload.rows[0].amounts[0].decision = "CONFIRMED_SPEND"
    confirmed, _ = service._evaluate(payload, totals_override={"张三": [Decimal("700.00")]},
                                     allow_duplicate_names=True, starting_sequence=1)
    assert [row["balance"] for row in confirmed] == [400, 350]


def test_continue_uses_each_persons_own_latest_balance():
    text = "张三 使用登记 报销差旅100元；李四 使用登记 报销资料50元"
    request = _request(text, amount="100", registration="报销差旅100元").model_dump(mode="json")
    name_start = text.index("李四")
    registration_start = text.index("报销资料50元")
    amount_start = text.index("50", registration_start)
    request["rows"].append({"page_number": 1, "name": "李四",
        "name_ref": {"quote": "李四", "start": name_start, "end": name_start + 2},
        "registration_ref": {"quote": "报销资料50元", "start": registration_start,
                             "end": registration_start + len("报销资料50元")},
        "amounts": [{"quote": "50", "start": amount_start, "end": amount_start + 2,
                     "decision": "CONFIRMED_SPEND", "reason": "手写支出"}]})
    rows, _ = _service(text)._evaluate(FundingUsageRequest.model_validate(request),
        totals_override={"张三": [Decimal("700.00")], "李四": [Decimal("500.00")]},
        allow_duplicate_names=True, starting_sequence=2)
    assert [(row["sequence"], row["name"], row["balance"]) for row in rows] == [
        (3, "张三", 600), (4, "李四", 450)]


def test_unclassified_number_in_registration_blocks_final_balance():
    text = "张三 使用登记 报销差旅300元，2026年复核"
    registration = "报销差旅300元，2026年复核"
    payload = _request("张三 使用登记 报销差旅300元", amount="300")
    payload.rows[0].registration_ref.quote = registration
    payload.rows[0].registration_ref.end = payload.rows[0].registration_ref.start + len(registration)
    result, _ = _service(text)._evaluate(payload)
    assert result[0]["balance"] is None
    assert result[0]["status"] == "手写金额待核对"


def test_explicit_no_usage_allows_zero_but_unpriced_registration_does_not():
    text = "张三 使用登记 未使用"
    source = _request("张三 使用登记 报销差旅300元", amount="300").model_dump(mode="json")
    start = text.index("未使用")
    source["rows"][0].update({"registration_ref": {"quote": "未使用", "start": start, "end": start + 3},
                              "no_usage": True, "amounts": []})
    result, _ = _service(text)._evaluate(FundingUsageRequest.model_validate(source))
    assert result[0]["usage_amount"] == 0
    assert result[0]["balance"] == 1000
    source["rows"][0]["registration_ref"]["quote"] = "报销差旅"
    source["rows"][0]["registration_ref"]["end"] = start + 4
    source["rows"][0]["no_usage"] = False
    other_text = "张三 使用登记 报销差旅"
    result, _ = _service(other_text)._evaluate(FundingUsageRequest.model_validate(source))
    assert result[0]["balance"] is None


def test_amount_outside_registration_is_rejected():
    text = "张三 学院资助经费5000元 使用登记 报销差旅300元"
    payload = _request(text, amount="300")
    printed_start = text.index("5000")
    payload.rows[0].amounts[0].quote = "5000"
    payload.rows[0].amounts[0].start = printed_start
    payload.rows[0].amounts[0].end = printed_start + 4
    with pytest.raises(HTTPException) as exc:
        _service(text)._evaluate(payload)
    assert exc.value.detail["code"] == "FUNDING_AMOUNT_SOURCE_INVALID"


def test_year_in_handwritten_date_cannot_be_spend():
    text = "张三 使用登记 报销差旅300元，2026年核对"
    payload = _request("张三 使用登记 报销差旅300元", amount="300")
    payload.rows[0].registration_ref.quote = "报销差旅300元，2026年核对"
    payload.rows[0].registration_ref.end = payload.rows[0].registration_ref.start + len(payload.rows[0].registration_ref.quote)
    year_start = text.index("2026")
    payload.rows[0].amounts.append(type(payload.rows[0].amounts[0])(
        quote="2026", start=year_start, end=year_start + 4,
        decision="CONFIRMED_SPEND", reason="误判日期"))
    with pytest.raises(HTTPException) as exc:
        _service(text)._evaluate(payload)
    assert exc.value.detail["code"] == "FUNDING_DATE_NOT_SPEND"


def test_registration_dates_use_image_year_without_guessing_ambiguous_year():
    registration = "12.8 报销差旅300元"
    ref = SimpleNamespace(quote=registration, start=10)
    date_candidate = SimpleNamespace(decision="EXCLUDED", reason="手写日期",
                                     start=10, end=14)
    assert _registration_dates(ref, "2019年度科研成果资助 " + registration,
                               [date_candidate]) == ("2019-12-08", False)
    assert _registration_dates(ref, "2019年度 2021.12.16 " + registration,
                               [date_candidate]) == ("待核对：12.8（图片年份不唯一）", True)
    assert _registration_dates(ref, "报销2021元 " + registration,
                               [date_candidate]) == ("待核对：12.8（图片未识别到年份）", True)
    assert _registration_dates(ref, "2019年度 " + registration, []) == (None, False)
    assert _registration_dates(SimpleNamespace(quote="2月30日", start=0),
                               "2021年度 2月30日", []) == ("待核对：2月30日（日期无效）", True)


def test_multiple_full_dates_from_registration_are_preserved_in_order():
    registration = "2021.12.16 报销差旅300元；2023年1月5日 报销100元；2021.12.16"
    ref = SimpleNamespace(quote=registration, start=0)
    assert _registration_dates(ref, "2019年度 " + registration, []) == (
        "2021-12-16\n2023-01-05", False)


def test_evaluate_adds_registration_date_without_changing_amount():
    registration = "2021.12.16 报销差旅300元"
    text = "2019年度 张三 使用登记 " + registration
    payload = _request(text, amount="300", registration=registration)
    start = text.index("2021.12.16")
    payload.rows[0].amounts.append(type(payload.rows[0].amounts[0])(
        quote="2021.12.16", start=start, end=start + len("2021.12.16"),
        decision="EXCLUDED", reason="日期"))
    result, _ = _service(text)._evaluate(payload)
    assert result[0]["date"] == "2021-12-16"
    assert result[0]["usage_amount"] == 300
    assert result[0]["balance"] == 700


def test_ambiguous_date_is_visible_without_erasing_confirmed_balance():
    registration = "12.8 报销差旅300元"
    text = "2019年度 2021年其他登记 张三 使用登记 " + registration
    payload = _request(text, amount="300", registration=registration)
    start = text.index("12.8")
    payload.rows[0].amounts.append(type(payload.rows[0].amounts[0])(
        quote="12.8", start=start, end=start + 4,
        decision="EXCLUDED", reason="日期"))
    result, _ = _service(text)._evaluate(payload)
    assert result[0]["date"] == "待核对：12.8（图片年份不唯一）"
    assert result[0]["status"] == "金额已核对；日期待核对"
    assert result[0]["usage_amount"] == 300
    assert result[0]["balance"] == 700


def test_unclear_date_candidate_is_reported_as_pending():
    registration = "2021? 报销差旅300元"
    text = "张三 使用登记 " + registration
    payload = _request(text, amount="300", registration=registration)
    start = text.index("2021?")
    payload.rows[0].amounts.append(type(payload.rows[0].amounts[0])(
        quote="2021?", start=start, end=start + 5,
        decision="UNCERTAIN", reason="日期识别不清"))
    result, _ = _service(text)._evaluate(payload)
    assert result[0]["date"] == "待核对：2021?（日期识别不清）"
    assert result[0]["status"] == "手写金额待核对；日期待核对"


def test_render_contains_no_old_total(tmp_path):
    source = tmp_path / "old.xlsx"
    old = Workbook()
    sheet = old.active
    sheet.title = "汇总"
    sheet["B1"] = "姓名"
    sheet["N1"] = "资助总额（元）"
    sheet["B2"] = "张三"
    sheet["N2"] = 918273
    sheet["B1"].font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    sheet["B1"].fill = PatternFill("solid", fgColor="123456")
    sheet["B2"].border = Border(bottom=Side(style="thin", color="654321"))
    old.save(source)
    old.close()
    target = tmp_path / "new.xlsx"
    FundingUsageService._render(target, source, [{
        "sequence": 1, "name": "张三", "registration": "12.8 报销差旅300元",
        "date": "待核对：12.8（图片年份不唯一）", "usage_amount": 300.0,
        "balance": 700.0, "status": "金额已核对；日期待核对"}])
    book = load_workbook(target, read_only=True, data_only=True)
    try:
        assert book.sheetnames == ["使用登记汇总", "张三"]
        assert book["使用登记汇总"]["D2"].value == "待核对：12.8（图片年份不唯一）"
        assert book["使用登记汇总"]["E2"].value == 300
        assert book["使用登记汇总"]["F2"].value == 700
        assert book["使用登记汇总"]["G2"].value == "金额已核对；日期待核对"
        assert book["张三"]["C1"].value == "日期"
        assert book["张三"]["C2"].value == "待核对：12.8（图片年份不唯一）"
        assert book["张三"]["E3"].value == 700
        assert book["使用登记汇总"]["B1"].fill.fgColor.rgb == "00123456"
        assert book["使用登记汇总"]["B2"].border.bottom.style == "thin"
    finally:
        book.close()
    with ZipFile(target) as archive:
        assert b"918273" not in b"".join(archive.read(name) for name in archive.namelist())


def test_append_formats_multiline_row_and_upgrades_legacy_personal_sheet():
    book = Workbook()
    summary = book.active
    summary.title = "使用登记汇总"
    summary.append(["序号", "姓名", "使用登记原文", "日期", "本次使用金额合计（元）",
                    "扣减后余额（元）", "核对状态"])
    summary.append([1, "张三", "首次支出300元", "2026-10-10", 300, 700, "已核对"])
    summary["C2"].font = Font(name="Microsoft YaHei", color="123456")
    summary["E2"].number_format = '#,##0.00;[Red]-#,##0.00'
    legacy = book.create_sheet("张三")
    legacy.append(["序号", "使用登记原文", "确认使用金额（元）", "核对说明"])
    legacy.append([1, "首次支出300元", 300, "已核对"])
    legacy.append([None, "扣减后余额（元）", 700])
    second = {"sequence": 2, "name": "张三", "registration": "会议费100元\n资料费50元",
              "date": "2026-10-11", "usage_amount": 150, "balance": 550, "status": "已核对"}
    append_summary_rows(book, [second])
    assert summary["C3"].font.name == summary["C2"].font.name
    assert summary["C3"].font.color.rgb == summary["C2"].font.color.rgb
    assert summary["E3"].number_format == summary["E2"].number_format
    assert summary["C3"].alignment.wrap_text is True
    assert summary.row_dimensions[3].height == 36
    first = {"sequence": 1, "name": "张三", "registration": "首次支出300元",
             "date": "2026-10-10", "usage_amount": 300, "balance": 700, "status": "已核对"}
    rebuild_personal_sheets(book, [first, second])
    assert [book["张三"].cell(1, col).value for col in range(1, 7)] == [
        "序号", "使用登记原文", "日期", "确认使用金额（元）", "扣减后余额（元）", "核对说明"]
    assert book["张三"]["C3"].value == "2026-10-11"
    assert book["张三"]["E3"].value == 550
    assert book["张三"].row_dimensions[3].height == 36
    book.close()


def test_delete_registration_recalculates_balance_and_releases_source(demo, monkeypatch):
    from io import BytesIO

    client, sessions, users = demo
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    first_text = "张三 使用登记 报销差旅300元"
    second_text = "张三 使用登记 报销会议100元"
    first_id = seed_workbuddy_document(sessions, filename="删除来源.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1, "text": first_text}])
    second_id = seed_workbuddy_document(sessions, filename="保留来源.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1, "text": second_text}])
    headers, record = [None] * 14, [None] * 14
    headers[1], headers[13] = "姓名", "资助总额（元）"
    record[1], record[13] = "张三", 1000
    old_id = seed_excel_target(sessions, filename="删除测试基准.xlsx", sheets={"汇总": [headers, record]})

    def image(document_id):
        with sessions() as db:
            version = db.query(DocumentVersion).filter_by(document_id=document_id).one()
            run = db.query(DocumentExtractionRun).filter_by(document_id=document_id).one()
            return {"document_id": document_id, "document_version_id": version.id,
                    "source_sha256": version.sha256, "extraction_run_id": run.id}

    first = _request(first_text, amount="300").model_dump(mode="json")
    first["image"] = image(first_id)
    with sessions() as db:
        old = db.query(DocumentVersion).filter_by(document_id=old_id).one()
        first["workbook"] = {"document_id": old_id, "document_version_id": old.id,
                             "source_sha256": old.sha256}
    first["filename"] = "删除回传测试.xlsx"
    preview = client.post("/api/funding-usage/preview", headers=users[0], json=first)
    assert preview.status_code == 200, preview.text
    initial = client.post("/api/funding-usage/publish", headers=users[0],
        json={**first, "preview_digest": preview.json()["preview_digest"]})
    assert initial.status_code == 200, initial.text
    initial = initial.json()

    def continuation(text, amount, registration, source, base, key):
        return {"image": image(source), "filename": base["filename"],
            "base_document_version_id": base["document_version_id"],
            "base_sha256": base["sha256"], "base_revision": base["revision"],
            "request_key": key,
            "rows": _request(text, amount=amount, registration=registration).model_dump(mode="json")["rows"]}

    second = continuation(second_text, "100", "报销会议100元", second_id, initial, "delete-append")
    preview = client.post("/api/funding-usage/continue/preview", headers=users[0], json=second)
    assert preview.status_code == 200, preview.text
    added = client.post("/api/funding-usage/continue/publish", headers=users[0],
        json={**second, "preview_digest": preview.json()["preview_digest"]})
    assert added.status_code == 200, added.text
    added = added.json()
    downloaded = client.get(added["download_url"], headers=users[0])
    book = load_workbook(BytesIO(downloaded.content))
    book["使用登记汇总"].delete_rows(2)
    output = BytesIO()
    book.save(output)
    book.close()
    edited = output.getvalue()
    deleted = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": added["filename"], "request_key": "delete-one-registration",
              "base_document_version_id": added["document_version_id"],
              "base_sha256": added["sha256"], "base_revision": str(added["revision"]),
              "file_sha256": sha256(edited).hexdigest()},
        files={"file": (added["filename"], edited,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert deleted.status_code == 200, deleted.text
    deleted = deleted.json()
    assert deleted["deleted_row_sequences"] == [1]
    assert deleted["deleted_record_count"] == 1
    assert [(item["before"], item["after"]) for item in deleted["renumbered_sequences"]] == [(2, 1)]
    revised = client.get(deleted["download_url"], headers=users[0])
    book = load_workbook(BytesIO(revised.content), read_only=True, data_only=True)
    assert [book["使用登记汇总"].cell(2, col).value for col in (1, 5, 6)] == [1, 100, 900]
    assert book["使用登记汇总"].max_row == 2
    assert book["张三"]["E2"].value == 900
    book.close()
    with sessions() as db:
        version = db.get(DocumentVersion, deleted["document_version_id"])
        policy = version.extraction_policy_json
        assert len(policy["funding_usage_history"]) == 1
        assert [entry["sequence"] for entry in policy["funding_usage_entries"]] == [1]
        assert policy["funding_usage_entries"][0]["balance_before"] == 1000
        assert [entry["sequence"] for entry in policy["funding_usage_deleted_entries"]] == [1]
        assert (policy["funding_usage_entries"][0]["source_row_id"]
                != policy["funding_usage_deleted_entries"][0]["source_row_id"])
        assert policy["funding_usage_deleted_entries"][0]["registration_key"] not in policy["funding_usage_history"]
        assert policy["funding_usage_deleted_entries"][0]["image_document_version_id"] == image(first_id)["document_version_id"]
    readd = continuation(first_text, "300", "报销差旅300元", first_id, deleted, "readd-deleted")
    preview = client.post("/api/funding-usage/continue/preview", headers=users[0], json=readd)
    assert preview.status_code == 200, preview.text
    assert preview.json()["people"][0]["sequence"] == 2
    assert preview.json()["people"][0]["balance"] == 600
    republished = client.post("/api/funding-usage/continue/publish", headers=users[0],
        json={**readd, "preview_digest": preview.json()["preview_digest"]})
    assert republished.status_code == 200, republished.text
    republished = republished.json()
    book = load_workbook(BytesIO(client.get(republished["download_url"], headers=users[0]).content))
    book["使用登记汇总"].delete_rows(2, 2)
    output = BytesIO()
    book.save(output)
    book.close()
    emptied = output.getvalue()
    all_deleted = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": republished["filename"], "request_key": "delete-all-registration",
              "base_document_version_id": republished["document_version_id"],
              "base_sha256": republished["sha256"], "base_revision": str(republished["revision"]),
              "file_sha256": sha256(emptied).hexdigest()},
        files={"file": (republished["filename"], emptied,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert all_deleted.status_code == 200, all_deleted.text
    all_deleted = all_deleted.json()
    assert all_deleted["deleted_row_sequences"] == [1, 2]
    book = load_workbook(BytesIO(client.get(all_deleted["download_url"], headers=users[0]).content),
                         read_only=True, data_only=True)
    assert book.sheetnames == ["使用登记汇总"]
    assert all(book["使用登记汇总"].cell(2, col).value is None for col in range(1, 8))
    book.close()
    restore = continuation(second_text, "100", "报销会议100元", second_id,
                           all_deleted, "restore-after-all-deleted")
    preview = client.post("/api/funding-usage/continue/preview", headers=users[0], json=restore)
    assert preview.status_code == 200, preview.text
    assert preview.json()["people"][0]["sequence"] == 1
    assert preview.json()["people"][0]["balance"] == 900
    restored = client.post("/api/funding-usage/continue/publish", headers=users[0],
        json={**restore, "preview_digest": preview.json()["preview_digest"]})
    assert restored.status_code == 200, restored.text
    book = load_workbook(BytesIO(client.get(restored.json()["download_url"], headers=users[0]).content),
                         read_only=True, data_only=True)
    assert book["使用登记汇总"]["A2"].value == 1
    assert book["使用登记汇总"]["F2"].value == 900
    assert book["张三"]["E2"].value == 900
    book.close()


def test_manual_return_adds_rows_in_middle_and_end_then_continues(demo, monkeypatch):
    from io import BytesIO

    client, sessions, users = demo
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    first_text = "张三 使用登记 报销差旅300元"
    second_text = "张三 使用登记 报销会议100元"
    first_id = seed_workbuddy_document(sessions, filename="人工新增首笔.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1, "text": first_text}])
    second_id = seed_workbuddy_document(sessions, filename="人工新增次笔.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1, "text": second_text}])
    headers, record = [None] * 14, [None] * 14
    headers[1], headers[13] = "姓名", "资助总额（元）"
    record[1], record[13] = "张三", 1000
    old_id = seed_excel_target(sessions, filename="人工新增基准.xlsx", sheets={"汇总": [headers, record]})

    def image(document_id):
        with sessions() as db:
            version = db.query(DocumentVersion).filter_by(document_id=document_id).one()
            run = db.query(DocumentExtractionRun).filter_by(document_id=document_id).one()
            return {"document_id": document_id, "document_version_id": version.id,
                    "source_sha256": version.sha256, "extraction_run_id": run.id}

    first = _request(first_text, amount="300").model_dump(mode="json")
    first["image"] = image(first_id)
    with sessions() as db:
        old = db.query(DocumentVersion).filter_by(document_id=old_id).one()
        first["workbook"] = {"document_id": old_id, "document_version_id": old.id,
                             "source_sha256": old.sha256}
    first["filename"] = "人工新增回传测试.xlsx"
    preview = client.post("/api/funding-usage/preview", headers=users[0], json=first)
    assert preview.status_code == 200, preview.text
    initial = client.post("/api/funding-usage/publish", headers=users[0],
        json={**first, "preview_digest": preview.json()["preview_digest"]})
    assert initial.status_code == 200, initial.text
    initial = initial.json()

    second = {"image": image(second_id), "filename": initial["filename"],
        "base_document_version_id": initial["document_version_id"],
        "base_sha256": initial["sha256"], "base_revision": initial["revision"],
        "request_key": "manual-add-source-second",
        "rows": _request(second_text, amount="100", registration="报销会议100元")
                .model_dump(mode="json")["rows"]}
    preview = client.post("/api/funding-usage/continue/preview", headers=users[0], json=second)
    assert preview.status_code == 200, preview.text
    previous = client.post("/api/funding-usage/continue/publish", headers=users[0],
        json={**second, "preview_digest": preview.json()["preview_digest"]})
    assert previous.status_code == 200, previous.text
    previous = previous.json()
    with sessions() as db:
        old_entries = db.get(DocumentVersion, previous["document_version_id"])
        old_ids = [item["source_row_id"] for item in
                   old_entries.extraction_policy_json["funding_usage_entries"]]

    book = load_workbook(BytesIO(client.get(previous["download_url"], headers=users[0]).content))
    sheet = book["使用登记汇总"]
    sheet.insert_rows(3)
    for column, value in enumerate([None, "张三", "独立差旅费50元", "2026-10-11",
                                    50, None, "已核对"], 1):
        sheet.cell(3, column).value = value
    sheet["E4"] = 125  # 同轮也更正一条已有登记金额，旧 F 值由服务端重算。
    sheet["C4"] = 6789
    sheet["G4"] = True
    sheet.append([None, "张三", 12345, datetime(2026, 10, 12), 25, None, 7])
    output = BytesIO()
    book.save(output)
    book.close()
    edited = output.getvalue()
    added = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": previous["filename"], "request_key": "manual-add-middle-and-end",
              "base_document_version_id": previous["document_version_id"],
              "base_sha256": previous["sha256"], "base_revision": str(previous["revision"]),
              "file_sha256": sha256(edited).hexdigest()},
        files={"file": (previous["filename"], edited,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert added.status_code == 200, added.text
    added = added.json()
    assert added["added_record_count"] == 2
    assert added["added_row_sequences"] == [2, 4]
    assert [(item["before"], item["after"]) for item in added["renumbered_sequences"]] == [(2, 3)]
    result_book = load_workbook(BytesIO(client.get(added["download_url"], headers=users[0]).content),
                                read_only=True, data_only=True)
    summary = result_book["使用登记汇总"]
    assert [summary.cell(row, 1).value for row in range(2, 6)] == [1, 2, 3, 4]
    assert [summary.cell(row, 6).value for row in range(2, 6)] == [700, 650, 525, 500]
    assert [summary.cell(4, col).value for col in (3, 7)] == ["6789", "TRUE"]
    assert [summary.cell(5, col).value for col in (3, 4, 7)] == ["12345", "2026-10-12", "7"]
    assert [result_book["张三"].cell(row, 5).value for row in range(2, 6)] == [700, 650, 525, 500]
    result_book.close()
    with sessions() as db:
        version = db.get(DocumentVersion, added["document_version_id"])
        entries = version.extraction_policy_json["funding_usage_entries"]
        assert entries[2]["source_row_id"] == old_ids[0]
        assert entries[0]["source_row_id"] not in {
            entries[1]["source_row_id"], entries[3]["source_row_id"]}
        assert [entries[index]["source_kind"] for index in (1, 3)] == [
            "MANUAL_EXCEL_REVISION", "MANUAL_EXCEL_REVISION"]
        assert [entries[index]["manual_upload_row"] for index in (1, 3)] == [3, 5]
        assert len({item["source_row_id"] for item in entries}) == 4
        manual_ids = [entries[index]["source_row_id"] for index in (1, 3)]

    next_text = "张三 使用登记 报销材料5元"
    next_id = seed_workbuddy_document(sessions, filename="人工新增后继续.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1, "text": next_text}])
    continuation = {"image": image(next_id), "filename": added["filename"],
        "base_document_version_id": added["document_version_id"],
        "base_sha256": added["sha256"], "base_revision": added["revision"],
        "request_key": "continue-after-manual-add",
        "rows": _request(next_text, amount="5", registration="报销材料5元")
                .model_dump(mode="json")["rows"]}
    preview = client.post("/api/funding-usage/continue/preview", headers=users[0], json=continuation)
    assert preview.status_code == 200, preview.text
    assert preview.json()["people"][0]["sequence"] == 5
    assert preview.json()["people"][0]["balance"] == 495
    continued = client.post("/api/funding-usage/continue/publish", headers=users[0],
        json={**continuation, "preview_digest": preview.json()["preview_digest"]})
    assert continued.status_code == 200, continued.text
    continued = continued.json()
    final_book = load_workbook(BytesIO(client.get(continued["download_url"], headers=users[0]).content),
                               read_only=True, data_only=True)
    assert final_book["使用登记汇总"]["A6"].value == 5
    assert final_book["使用登记汇总"]["F6"].value == 495
    assert final_book["张三"]["E6"].value == 495
    final_book.close()
    book = load_workbook(BytesIO(client.get(continued["download_url"], headers=users[0]).content))
    book["使用登记汇总"]["D3"] = "2026-10-13"
    output = BytesIO()
    book.save(output)
    book.close()
    edited_again = output.getvalue()
    revised = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": continued["filename"], "request_key": "manual-added-row-second-return",
              "base_document_version_id": continued["document_version_id"],
              "base_sha256": continued["sha256"], "base_revision": str(continued["revision"]),
              "file_sha256": sha256(edited_again).hexdigest()},
        files={"file": (continued["filename"], edited_again,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert revised.status_code == 200, revised.text
    with sessions() as db:
        version = db.get(DocumentVersion, revised.json()["document_version_id"])
        entries = version.extraction_policy_json["funding_usage_entries"]
        assert [entries[index]["source_row_id"] for index in (1, 3)] == manual_ids
        assert entries[1]["current_row"]["date"] == "2026-10-13"


def test_manual_return_renumbers_existing_gaps_without_changing_sources(tmp_path, monkeypatch):
    from io import BytesIO

    from app.modules.funding_usage.workbook import HEADERS
    from app.modules.structured_datasets.manual_revision import ManualRevisionService

    book = Workbook()
    sheet = book.active
    sheet.title = "使用登记汇总"
    sheet.append(HEADERS)
    rows = [
        {"sequence": 1, "name": "张三", "registration": "第一次使用100元", "date": None,
         "usage_amount": 100, "balance": 900, "status": "已核对"},
        {"sequence": 37, "name": "张三", "registration": "后续使用200元", "date": None,
         "usage_amount": 200, "balance": 700, "status": "已核对"},
    ]
    for row in rows:
        sheet.append([row[key] for key in ("sequence", "name", "registration", "date",
                                            "usage_amount", "balance", "status")])
    rebuild_personal_sheets(book, rows)
    path = tmp_path / "existing.xlsx"
    book.save(path)
    book.close()
    source_version = SimpleNamespace(storage_path=str(path), sha256=sha256(path.read_bytes()).hexdigest())
    service = object.__new__(ManualRevisionService)
    service.db = SimpleNamespace(get=lambda model, key: source_version if key == "source-version" else None)
    service.user = SimpleNamespace(id="user")
    service.datasets = SimpleNamespace(
        _records=lambda dataset: [{"record_id": f"base:dataset:{row['sequence']}",
                                   "fields": {"sequence": {"value": row["sequence"]}}}
                                  for row in rows],
        storage=SimpleNamespace(working_copy_path=lambda value: Path(value),
                                sha256_file=lambda value: sha256(Path(value).read_bytes()).hexdigest()))
    monkeypatch.setattr(FundingUsageService, "_baseline",
                        lambda self, value: {"张三": [Decimal(1000)]})
    dataset = SimpleNamespace(target_source_version_id="source-version",
                              base_rows_json=[{**row, "_funding_row_id": f"source-{row['sequence']}"}
                                              for row in rows])
    current = SimpleNamespace(extraction_policy_json={
        "structured_dataset_export_mode": "funding-usage-v2",
        "funding_usage_source_active_complete": True,
        "funding_usage_entries": [
            {"sequence": row["sequence"], "name": row["name"],
             "registration_key": f"key-{row['sequence']}",
             "source_row_id": f"source-{row['sequence']}"} for row in rows],
        "funding_usage_history": ["key-1", "key-37"],
    })
    origin = SimpleNamespace(extraction_policy_json={"funding_usage_evidence": []})
    changes, output, normalized, overrides, policy, deleted = service._prepare_funding_revision(
        dataset, path, path, current, origin, "renumber-only")
    assert deleted == []
    assert [row["sequence"] for row in normalized] == [1, 2]
    assert [row["_funding_row_id"] for row in normalized] == ["source-1", "source-37"]
    assert [(item["before"], item["after"]) for item in policy["funding_usage_sequence_changes"]] == [(37, 2)]
    assert [entry["sequence"] for entry in policy["funding_usage_entries"]] == [1, 2]
    assert policy["funding_usage_history"] == ["key-1", "key-37"]
    assert policy["funding_usage_next_sequence"] == 3
    assert any(item["field_key"] == "__sequence__" for item in changes)
    result = load_workbook(BytesIO(output), read_only=True, data_only=True)
    assert result["使用登记汇总"]["A3"].value == 2
    assert result["张三"]["A3"].value == 2
    result.close()


def test_append_repairs_previously_unstyled_summary_rows():
    book = Workbook()
    summary = book.active
    summary.title = "使用登记汇总"
    summary.append(["序号", "姓名", "使用登记原文", "日期", "本次使用金额合计（元）",
                    "扣减后余额（元）", "核对状态"])
    summary.append([1, "张三", "首次支出300元", None, 300, 700, "已核对"])
    summary["C2"].font = Font(name="Microsoft YaHei")
    summary["E2"].number_format = '#,##0.00;[Red]-#,##0.00'
    summary.append([2, "张三", "会议费100元\n资料费50元", None, 150, 550, "已核对"])
    assert summary["C3"].has_style is False
    append_summary_rows(book, [{"sequence": 3, "name": "张三", "registration": "打印费20元",
                                "date": None, "usage_amount": 20, "balance": 530,
                                "status": "已核对"}])
    for row_number in (3, 4):
        assert summary.cell(row_number, 3).font.name == "Microsoft YaHei"
        assert summary.cell(row_number, 3).alignment.wrap_text is True
        assert summary.cell(row_number, 5).number_format == summary["E2"].number_format
    assert summary.row_dimensions[3].height == 36
    assert summary.row_dimensions[4].height == 28
    book.close()


def test_workbuddy_preview_publish_and_download(demo, monkeypatch):
    client, sessions, users = demo
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    text = "张三 学院资助经费5000元 使用登记 报销差旅300元"
    image_id = seed_workbuddy_document(sessions, filename="登记.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1, "text": text}])
    headers = [None] * 14
    headers[1], headers[13] = "姓名", "资助总额（元）"
    record = [None] * 14
    record[1], record[13] = "张三", 1000
    workbook_id = seed_excel_target(sessions, filename="旧成果汇总.xlsx",
        sheets={"汇总": [headers, record]})
    with sessions() as db:
        image_version = db.query(DocumentVersion).filter_by(document_id=image_id).one()
        image_run = db.query(DocumentExtractionRun).filter_by(document_id=image_id).one()
        workbook_version = db.query(DocumentVersion).filter_by(document_id=workbook_id).one()
        source = {"document_id": image_id, "document_version_id": image_version.id,
                  "source_sha256": image_version.sha256}
        old_source = {"document_id": workbook_id, "document_version_id": workbook_version.id,
                      "source_sha256": workbook_version.sha256}
        image_run_id = image_run.id
    request = _request(text, amount="300").model_dump(mode="json")
    request["image"] = {**source, "extraction_run_id": image_run_id}
    request["workbook"] = old_source
    before = client.post("/api/funding-usage/preview", headers=users[0], json=request)
    assert before.status_code == 200, before.text
    preview = before.json()
    assert preview["people"][0]["balance"] == 700
    published = client.post("/api/funding-usage/publish", headers=users[0],
                            json={**request, "preview_digest": preview["preview_digest"]})
    assert published.status_code == 200, published.text
    receipt = published.json()
    assert receipt["status"] == "COMPLETED"
    assert receipt["excel_published"] is True
    assert receipt["version_number"] == 1
    repeated = client.post("/api/funding-usage/publish", headers=users[0],
                           json={**request, "preview_digest": preview["preview_digest"]})
    assert repeated.status_code == 200, repeated.text
    assert repeated.json()["document_version_id"] == receipt["document_version_id"]
    assert repeated.json()["reused"] is True
    downloaded = client.get(receipt["download_url"], headers=users[0])
    assert downloaded.status_code == 200
    from io import BytesIO
    book = load_workbook(BytesIO(downloaded.content), read_only=True, data_only=True)
    try:
        assert book["使用登记汇总"]["D1"].value == "日期"
        assert book["使用登记汇总"]["D2"].value is None
        assert book["使用登记汇总"]["E2"].value == 300
        assert book["使用登记汇总"]["F2"].value == 700
    finally:
        book.close()
    with sessions() as db:
        assert db.query(StructuredDatasetPublication).filter_by(
            dataset_id=receipt["dataset_id"], status="COMPLETED").count() == 1
        assert db.get(DocumentVersion, workbook_version.id).sha256 == old_source["source_sha256"]

    # 用户下载后补全登记文字，仍可走已有同名手工修订通道。
    from io import BytesIO
    editable = load_workbook(BytesIO(downloaded.content))
    editable["使用登记汇总"]["C2"] = "报销差旅300元，票据已核实"
    editable["使用登记汇总"]["D2"] = "2021-12-16"
    output = BytesIO()
    editable.save(output)
    editable.close()
    corrected = output.getvalue()
    manual = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": receipt["filename"], "request_key": "funding-manual-1",
              "base_document_version_id": receipt["document_version_id"],
              "base_sha256": receipt["sha256"], "base_revision": str(receipt["revision"]),
              "file_sha256": sha256(corrected).hexdigest()},
        files={"file": (receipt["filename"], corrected,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert manual.status_code == 200, manual.text
    assert manual.json()["status"] == "COMPLETED"
    assert manual.json()["version_number"] == 2
    latest = client.get(manual.json()["download_url"], headers=users[0])
    assert latest.status_code == 200
    personal = load_workbook(BytesIO(latest.content))
    assert personal["使用登记汇总"]["D2"].value == "2021-12-16"
    assert personal["张三"]["B2"].value == "报销差旅300元，票据已核实"
    assert personal["张三"]["C2"].value == "2021-12-16"
    personal["张三"]["B2"] = "个人页私改"
    output = BytesIO()
    personal.save(output)
    personal.close()
    changed = output.getvalue()
    rejected = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": receipt["filename"], "request_key": "funding-personal-reject",
              "base_document_version_id": manual.json()["document_version_id"],
              "base_sha256": manual.json()["sha256"],
              "base_revision": str(manual.json()["revision"]),
              "file_sha256": sha256(changed).hexdigest()},
        files={"file": (receipt["filename"], changed,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert rejected.status_code == 422
    assert "FUNDING_PERSONAL_SHEET_REVISION_UNSUPPORTED" in rejected.text


def test_continue_deducts_latest_published_balance_and_manual_summary_sync(demo, monkeypatch):
    from io import BytesIO

    client, sessions, users = demo
    monkeypatch.setenv("WORKBUDDY_EXCEL_WRITEBACK_ENABLED", "true")
    get_settings.cache_clear()
    first_text = "张三 使用登记 报销差旅300元"
    first_id = seed_workbuddy_document(sessions, filename="首次.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1, "text": first_text}])
    second_text = "张三 使用登记 报销会议100元"
    second_id = seed_workbuddy_document(sessions, filename="追加.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1, "text": second_text}])
    third_text = "张三 使用登记 报销资料50元"
    third_id = seed_workbuddy_document(sessions, filename="再次追加.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1, "text": third_text}])
    headers = [None] * 14
    headers[1], headers[13] = "姓名", "资助总额（元）"
    record = [None] * 14
    record[1], record[13] = "张三", 1000
    old_id = seed_excel_target(sessions, filename="初始资助.xlsx", sheets={"汇总": [headers, record]})

    def image_snapshot(document_id):
        with sessions() as db:
            version = db.query(DocumentVersion).filter_by(document_id=document_id).one()
            run = db.query(DocumentExtractionRun).filter_by(document_id=document_id).one()
            return {"document_id": document_id, "document_version_id": version.id,
                    "source_sha256": version.sha256, "extraction_run_id": run.id}

    first = _request(first_text, amount="300").model_dump(mode="json")
    first["image"] = image_snapshot(first_id)
    with sessions() as db:
        old_version = db.query(DocumentVersion).filter_by(document_id=old_id).one()
        first["workbook"] = {"document_id": old_id,
            "document_version_id": old_version.id, "source_sha256": old_version.sha256}
    first["filename"] = "累计使用登记.xlsx"
    p = client.post("/api/funding-usage/preview", headers=users[0], json=first)
    assert p.status_code == 200, p.text
    initial = client.post("/api/funding-usage/publish", headers=users[0],
        json={**first, "preview_digest": p.json()["preview_digest"]})
    assert initial.status_code == 200, initial.text
    initial = initial.json()
    initial_download = client.get(initial["download_url"], headers=users[0])
    assert initial_download.status_code == 200
    initial_book = load_workbook(BytesIO(initial_download.content))
    initial_summary = initial_book["使用登记汇总"]
    initial_personal_headers = [initial_book["张三"].cell(1, col).value for col in range(1, 7)]
    assert initial_personal_headers == ["序号", "使用登记原文", "日期",
        "确认使用金额（元）", "扣减后余额（元）", "核对说明"]
    initial_style = [initial_summary.cell(2, col)._style for col in range(1, 8)]
    initial_personal_style = [initial_book["张三"].cell(2, col)._style for col in range(1, 7)]
    initial_book.close()

    def continue_request(text, amount, registration, image_id, base, key):
        row = _request(text, amount=amount, registration=registration).model_dump(mode="json")["rows"]
        return {"image": image_snapshot(image_id), "filename": base["filename"],
                "base_document_version_id": base["document_version_id"],
                "base_sha256": base["sha256"], "base_revision": base["revision"],
                "request_key": key, "rows": row}

    second = continue_request(second_text, "100", "报销会议100元", second_id, initial, "append-1")
    preview = client.post("/api/funding-usage/continue/preview", headers=users[0], json=second)
    assert preview.status_code == 200, preview.text
    assert preview.json()["people"][0]["balance"] == 600
    added = client.post("/api/funding-usage/continue/publish", headers=users[0],
        json={**second, "preview_digest": preview.json()["preview_digest"]})
    assert added.status_code == 200, added.text
    added = added.json()
    assert added["version_number"] == 2
    assert added["added_record_count"] == 1
    stale = continue_request(third_text, "50", "报销资料50元", third_id,
                             initial, "append-stale")
    conflict = client.post("/api/funding-usage/continue/preview", headers=users[0], json=stale)
    assert conflict.status_code == 409
    assert "FUNDING_TARGET_VERSION_CONFLICT" in conflict.text
    repeat = client.post("/api/funding-usage/continue/publish", headers=users[0],
        json={**second, "preview_digest": preview.json()["preview_digest"]})
    assert repeat.status_code == 200, repeat.text
    assert repeat.json()["reused"] is True
    downloaded = client.get(added["download_url"], headers=users[0])
    assert downloaded.status_code == 200
    book = load_workbook(BytesIO(downloaded.content))
    assert book["使用登记汇总"]["F2"].value == 700
    assert book["使用登记汇总"]["F3"].value == 600
    for col in range(1, 8):
        assert book["使用登记汇总"].cell(3, col)._style == initial_style[col - 1]
    assert book["使用登记汇总"].row_dimensions[3].height == 28
    assert [book["张三"].cell(1, col).value for col in range(1, 7)] == initial_personal_headers
    for col in range(1, 7):
        assert book["张三"].cell(3, col)._style == initial_personal_style[col - 1]
    assert book["张三"]["E3"].value == 600
    book["使用登记汇总"]["E3"] = 120
    output = BytesIO()
    book.save(output)
    book.close()
    corrected = output.getvalue()
    manual = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": added["filename"], "request_key": "funding-summary-only",
              "base_document_version_id": added["document_version_id"],
              "base_sha256": added["sha256"], "base_revision": str(added["revision"]),
              "file_sha256": sha256(corrected).hexdigest()},
        files={"file": (added["filename"], corrected,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert manual.status_code == 200, manual.text
    manual = manual.json()
    assert manual["version_number"] == 3
    revised = client.get(manual["download_url"], headers=users[0])
    book = load_workbook(BytesIO(revised.content), read_only=True, data_only=True)
    assert book["使用登记汇总"]["F3"].value == 580
    assert book["张三"]["D3"].value == 120
    assert book["张三"]["E3"].value == 580
    book.close()

    edited = load_workbook(BytesIO(revised.content))
    edited["使用登记汇总"]["F3"] = 575
    output = BytesIO()
    edited.save(output)
    missing_reason = output.getvalue()
    rejected = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": manual["filename"], "request_key": "funding-balance-no-reason",
              "base_document_version_id": manual["document_version_id"],
              "base_sha256": manual["sha256"], "base_revision": str(manual["revision"]),
              "file_sha256": sha256(missing_reason).hexdigest()},
        files={"file": (manual["filename"], missing_reason,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert rejected.status_code == 422
    assert "FUNDING_BALANCE_CORRECTION_REASON_REQUIRED" in rejected.text
    edited["使用登记汇总"]["G3"] = "人工更正：票据核对后余额调整"
    output = BytesIO()
    edited.save(output)
    edited.close()
    override_bytes = output.getvalue()
    override = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": manual["filename"], "request_key": "funding-balance-with-reason",
              "base_document_version_id": manual["document_version_id"],
              "base_sha256": manual["sha256"], "base_revision": str(manual["revision"]),
              "file_sha256": sha256(override_bytes).hexdigest()},
        files={"file": (manual["filename"], override_bytes,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert override.status_code == 200, override.text
    manual = override.json()
    assert manual["version_number"] == 4

    third = continue_request(third_text, "50", "报销资料50元", third_id, manual, "append-2")
    preview = client.post("/api/funding-usage/continue/preview", headers=users[0], json=third)
    assert preview.status_code == 200, preview.text
    assert preview.json()["people"][0]["balance"] == 525
    published = client.post("/api/funding-usage/continue/publish", headers=users[0],
        json={**third, "preview_digest": preview.json()["preview_digest"]})
    assert published.status_code == 200, published.text
    assert published.json()["version_number"] == 5
    duplicate = continue_request(second_text, "100", "报销会议100元", second_id,
        published.json(), "append-duplicate")
    rejected = client.post("/api/funding-usage/continue/preview", headers=users[0], json=duplicate)
    assert rejected.status_code == 409
    assert "FUNDING_REGISTRATION_DUPLICATE" in rejected.text
    reuploaded_id = seed_workbuddy_document(sessions, filename="追加.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1,
                                          "text": second_text}])
    reuploaded = continue_request(second_text, "100", "报销会议100元", reuploaded_id,
                                  published.json(), "append-reuploaded")
    rejected = client.post("/api/funding-usage/continue/preview", headers=users[0], json=reuploaded)
    assert rejected.status_code == 409
    assert "FUNDING_REGISTRATION_DUPLICATE" in rejected.text

    pending_text = "张三 使用登记 报销杂费40元"
    pending_id = seed_workbuddy_document(sessions, filename="待核对.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1,
                                          "text": pending_text}])
    pending = continue_request(pending_text, "40", "报销杂费40元", pending_id,
                               published.json(), "append-pending")
    pending["rows"][0]["amounts"][0]["decision"] = "UNCERTAIN"
    preview = client.post("/api/funding-usage/continue/preview", headers=users[0], json=pending)
    assert preview.status_code == 200, preview.text
    assert preview.json()["people"][0]["balance"] is None
    pending_result = client.post("/api/funding-usage/continue/publish", headers=users[0],
        json={**pending, "preview_digest": preview.json()["preview_digest"]})
    assert pending_result.status_code == 200, pending_result.text
    pending_result = pending_result.json()
    later_text = "张三 使用登记 报销打印20元"
    later_id = seed_workbuddy_document(sessions, filename="后续.png",
        content_type="image/png", units=[{"kind": "IMAGE_PAGE", "page_number": 1,
                                          "text": later_text}])
    later = continue_request(later_text, "20", "报销打印20元", later_id,
                             pending_result, "append-after-pending")
    preview = client.post("/api/funding-usage/continue/preview", headers=users[0], json=later)
    assert preview.status_code == 200, preview.text
    assert preview.json()["people"][0]["usage_amount"] == 20
    assert preview.json()["people"][0]["balance"] is None
    later_result = client.post("/api/funding-usage/continue/publish", headers=users[0],
        json={**later, "preview_digest": preview.json()["preview_digest"]})
    assert later_result.status_code == 200, later_result.text
    later_result = later_result.json()
    latest_bytes = client.get(later_result["download_url"], headers=users[0]).content
    book = load_workbook(BytesIO(latest_bytes))
    book["使用登记汇总"]["E5"] = 40
    output = BytesIO()
    book.save(output)
    book.close()
    confirmed_bytes = output.getvalue()
    confirmed = client.post("/api/excel-writeback/manual-revisions/publish", headers=users[0],
        data={"filename": later_result["filename"], "request_key": "funding-confirm-pending",
              "base_document_version_id": later_result["document_version_id"],
              "base_sha256": later_result["sha256"],
              "base_revision": str(later_result["revision"]),
              "file_sha256": sha256(confirmed_bytes).hexdigest()},
        files={"file": (later_result["filename"], confirmed_bytes,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert confirmed.status_code == 200, confirmed.text
    latest = client.get(confirmed.json()["download_url"], headers=users[0])
    book = load_workbook(BytesIO(latest.content), read_only=True, data_only=True)
    assert book["使用登记汇总"]["F5"].value == 485
    assert book["使用登记汇总"]["F6"].value == 465
    assert "当前余额待核对" not in book["使用登记汇总"]["G6"].value
    assert book["张三"]["E5"].value == 485
    assert book["张三"]["E6"].value == 465
    book.close()

"""共享累计表格 Demo 的端到端数据正确性、跨用户与旧事实保护测试。"""

import hashlib
import io
import zipfile
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image

from app.core.config import get_settings
from app.db.models import (
    ChangeItem, Document, DocumentVersion, DocumentPage, DocumentElement, DocumentExtractionRun,
    FileObject, ManagedRoot, ManagedFile, WorkingCopyRoot, WorkingCopy, WorkingCopyPathRecord, User,
    StructuredDataset, StructuredDatasetExport, StructuredDatasetMember, StructuredExtractionRun,
    ToolInvocation,
)
from app.modules.file_lifecycle.shared_workspace import get_shared_workspace_id
from app.modules.file_lifecycle.storage import FileLifecycleStorageService
from app.modules.files.office_conversion import ConvertedDocumentArtifact
from app.modules.managed_files.worker import process_next_filesystem_job
from app.tests.helpers import client_with_database, clear_overrides


BASE = "/api/structured-datasets"
FIELDS = [{"key": "name", "label": "姓名", "field_type": "string", "required": True},
          {"key": "amount", "label": "金额", "field_type": "decimal", "required": True}]


@pytest.fixture
def demo(monkeypatch, tmp_path):
    """使用隔离数据库和存储，不访问真实 OCR、LLM 或部署数据。"""
    monkeypatch.setenv("WORKBUDDY_STRUCTURED_DATASET_DEMO_ENABLED", "true")
    monkeypatch.setenv("FILE_STORAGE_ROOT", str(tmp_path / "storage"))
    monkeypatch.setenv("WORKING_COPY_STORAGE_ROOT", str(tmp_path / "working"))
    get_settings.cache_clear()
    client, sessions = client_with_database()
    # SQLite 默认关闭外键；显式开启，避免假测试掩盖真实版本和审计归属问题。
    with sessions.kw["bind"].connect() as connection:
        connection.exec_driver_sql("PRAGMA foreign_keys=ON")
    users = []
    for username in ("dataset-user-a", "dataset-user-b"):
        assert client.post("/api/auth/register", json={"username": username, "password": "password123", "display_name": username}).status_code < 300
        token = client.post("/api/auth/login", json={"username": username, "password": "password123"}).json()["access_token"]
        users.append({"Authorization": f"Bearer {token}"})
    yield client, sessions, users
    client.close()
    clear_overrides()


def seed_image(sessions, text, filename="demo.png"):
    """直接固定真实业务模型和解析事实，避免测试依赖 OCR 服务。"""
    image_buffer = io.BytesIO()
    Image.new("RGB", (8, 8), tuple(hashlib.sha256(text.encode()).digest()[:3])).save(image_buffer, format="PNG")
    content = image_buffer.getvalue()
    with sessions() as db:
        user = db.query(User).filter_by(username="dataset-user-a").one()
        workspace_id = get_shared_workspace_id(db)
        document = Document(user_id=user.id, workspace_id=workspace_id, original_filename=filename,
            content_type="image/png", size_bytes=len(content), sha256=hashlib.sha256(content).hexdigest())
        db.add(document)
        db.flush()
        version = DocumentVersion(document_id=document.id, storage_path=f"shared/{document.id}.png",
            filename=filename, content_type="image/png", size_bytes=123, sha256=document.sha256,
            storage_tier="WORKING_COPY", source_type="WORKING_COPY", created_by=user.id)
        root = ManagedRoot(root_key=str(uuid4()), display_name="测试原图", container_path="/test/source")
        db.add_all([version, root])
        db.flush()
        managed = ManagedFile(root_id=root.id, relative_path=filename, relative_path_hash=hashlib.sha256(filename.encode()).hexdigest(),
            filename=filename, extension=".png", size_bytes=123, content_sha256=document.sha256)
        working_root = WorkingCopyRoot(workspace_id=workspace_id, managed_root_id=root.id,
            root_key=str(uuid4()), relative_storage_path="shared", status="ACTIVE")
        db.add_all([managed, working_root])
        db.flush()
        working = WorkingCopy(working_copy_root_id=working_root.id, workspace_id=workspace_id,
            managed_file_id=managed.id, document_id=document.id, current_version_id=version.id,
            relative_path=version.storage_path, relative_path_hash=hashlib.sha256(version.storage_path.encode()).hexdigest(),
            filename=filename, extension=".png", size_bytes=123, content_sha256=document.sha256,
            imported_source_sha256=document.sha256, status="ACTIVE")
        extraction = DocumentExtractionRun(document_id=document.id, document_version_id=version.id,
            status="COMPLETED", extractor="workbuddy-external-ocr", parser_name="external-ocr", parser_version="1", parser_config_hash="f" * 64)
        db.add_all([working, extraction])
        db.flush()
        db.add(DocumentPage(document_id=document.id, extraction_run_id=extraction.id, page_number=1, text_content=text))
        image_path = FileLifecycleStorageService().working_copy_path(version.storage_path)
        image_path.parent.mkdir(parents=True, exist_ok=True)
        image_path.write_bytes(content)
        db.commit()
        return document.id


def seed_workbuddy_document(sessions, *, filename, content_type, units):
    """建立已由 WorkBuddy 验收并发布的共享工作副本，不调用服务器解析器。"""

    content = f"immutable:{filename}".encode()
    suffix = f".{filename.rsplit('.', 1)[-1].lower()}"
    with sessions() as db:
        user = db.query(User).filter_by(username="dataset-user-a").one()
        workspace_id = get_shared_workspace_id(db)
        source_hash = hashlib.sha256(content).hexdigest()
        document = Document(user_id=user.id, workspace_id=workspace_id, original_filename=filename,
            content_type=content_type, size_bytes=len(content), sha256=source_hash)
        db.add(document)
        db.flush()
        relative_path = f"{document.id}/{filename}"
        storage_path = f"shared/{relative_path}"
        version = DocumentVersion(document_id=document.id, storage_path=storage_path,
            filename=filename, content_type=content_type, size_bytes=len(content), sha256=source_hash,
            storage_tier="WORKING_COPY", source_type="WORKING_COPY", created_by=user.id,
            extraction_policy_json={"mode": "WORKBUDDY_DOCUMENT", "contract_version": "workbuddy-document-v1",
                "origin_task_id": str(uuid4()), "origin_upload_version_id": str(uuid4()), "source_sha256": source_hash})
        root = ManagedRoot(root_key=str(uuid4()), display_name="宿主文档", container_path="/test/workbuddy")
        db.add_all([version, root])
        db.flush()
        managed = ManagedFile(root_id=root.id, relative_path=relative_path,
            relative_path_hash=hashlib.sha256(relative_path.encode()).hexdigest(), filename=filename,
            extension=suffix, size_bytes=len(content), content_sha256=source_hash)
        working_root = WorkingCopyRoot(workspace_id=workspace_id, managed_root_id=root.id,
            root_key=str(uuid4()), relative_storage_path="shared", status="ACTIVE")
        db.add_all([managed, working_root])
        db.flush()
        working = WorkingCopy(working_copy_root_id=working_root.id, workspace_id=workspace_id,
            managed_file_id=managed.id, document_id=document.id, current_version_id=version.id,
            relative_path=relative_path, relative_path_hash=hashlib.sha256(relative_path.encode()).hexdigest(),
            filename=filename, extension=suffix, size_bytes=len(content), content_sha256=source_hash,
            imported_source_sha256=source_hash, status="ACTIVE")
        extraction = DocumentExtractionRun(document_id=document.id, document_version_id=version.id,
            status="COMPLETED", extractor="workbuddy-external-document", parser_name="workbuddy-test-reader",
            parser_version="1", parser_config_hash="e" * 64)
        db.add_all([working, extraction])
        db.flush()
        for order, unit in enumerate(units):
            external_unit_id = f"unit-{order}"
            page = DocumentPage(document_id=document.id, extraction_run_id=extraction.id,
                page_number=unit.get("page_number"), sheet_name=unit.get("sheet_name"),
                text_content=unit["text"], metadata_json={
                    "content_origin": "WORKBUDDY", "contract_version": "workbuddy-document-v1",
                    "source_unit_id": external_unit_id, "unit_order": order, "unit_kind": unit["kind"],
                    "sheet_name": unit.get("sheet_name"),
                    "unit_status": "READ", "coverage_status": "COMPLETE", "provider_name": "workbuddy-test-reader",
                    "coverage": {"status": "COMPLETE", "unit_count": len(units), "warnings": []},
                })
            element = DocumentElement(document_id=document.id, extraction_run_id=extraction.id,
                element_index=order, label="text", text_content=unit["text"],
                page_number=unit.get("page_number"), metadata_json={"source_unit_id": external_unit_id,
                    "sheet_name": unit.get("sheet_name"),
                    "start": 0, "end": len(unit["text"])})
            db.add_all([page, element])
        path = FileLifecycleStorageService().working_copy_path(storage_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        db.commit()
        return document.id


def seed_excel_target(
    sessions,
    *,
    filename="2026年人才材料登记表.xlsx",
    sheets=None,
    advanced_formatting=False,
    configure=None,
):
    """建立真实 XLSX 活动工作副本，供目标 schema 与物理文件名测试使用。"""

    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    for sheet_name, rows in (sheets or {"登记": [["姓名", "金额"], ["已有人员", 10]]}).items():
        sheet = workbook.create_sheet(sheet_name)
        for row in rows:
            sheet.append(row)
        if rows:
            sheet.freeze_panes = "A2"
            sheet.column_dimensions["A"].width = 24
            sheet["A1"].font = Font(bold=True)
            sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    if advanced_formatting:
        # 高级模板用于验证 OOXML 包级追加不会重存并丢失图表、校验、公式或合并区域。
        business_sheet = next(sheet for sheet in workbook.worksheets if sheet.max_row > 1)
        business_sheet["D2"] = "=SUM(B2:B2)"
        business_sheet.auto_filter.ref = "A1:B2"
        validation = DataValidation(type="whole", operator="greaterThan", formula1="0")
        business_sheet.add_data_validation(validation)
        validation.add("B2:B100")
        table = Table(displayName="TalentRecords", ref="A1:B2")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        business_sheet.add_table(table)
        chart = BarChart()
        chart.add_data(Reference(business_sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        business_sheet.add_chart(chart, "G2")
        business_sheet.sheet_properties.pageSetUpPr.fitToPage = True
        business_sheet.page_setup.orientation = "landscape"
        support_sheet = next(sheet for sheet in workbook.worksheets if sheet is not business_sheet)
        support_sheet.merge_cells("D1:E1")
        support_sheet["D1"] = "保留说明"
    if configure:
        configure(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()
    with sessions() as db:
        user = db.query(User).filter_by(username="dataset-user-a").one()
        workspace_id = get_shared_workspace_id(db)
        sha256 = hashlib.sha256(content).hexdigest()
        document = Document(
            user_id=user.id,
            workspace_id=workspace_id,
            original_filename=filename,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size_bytes=len(content),
            sha256=sha256,
            status="READY",
            ingest_status="READY",
        )
        root = ManagedRoot(
            root_key=str(uuid4()),
            display_name="测试 Excel",
            container_path="/test/excel",
        )
        db.add_all([document, root])
        db.flush()
        relative_path = f"{document.id}/{filename}"
        storage_path = f"excel-targets/{relative_path}"
        version = DocumentVersion(
            document_id=document.id,
            version_number=1,
            storage_path=storage_path,
            filename=filename,
            content_type=document.content_type,
            size_bytes=len(content),
            sha256=sha256,
            storage_tier="WORKING_COPY",
            source_type="WORKING_COPY",
            created_by=user.id,
        )
        managed = ManagedFile(
            root_id=root.id,
            relative_path=relative_path,
            relative_path_hash=hashlib.sha256(relative_path.encode()).hexdigest(),
            filename=filename,
            extension=".xlsx",
            size_bytes=len(content),
            content_sha256=sha256,
        )
        working_root = WorkingCopyRoot(
            workspace_id=workspace_id,
            managed_root_id=root.id,
            root_key=str(uuid4()),
            relative_storage_path="excel-targets",
            status="ACTIVE",
        )
        db.add_all([version, managed, working_root])
        db.flush()
        working_copy = WorkingCopy(
            working_copy_root_id=working_root.id,
            workspace_id=workspace_id,
            managed_file_id=managed.id,
            document_id=document.id,
            current_version_id=version.id,
            relative_path=relative_path,
            relative_path_hash=hashlib.sha256(relative_path.encode()).hexdigest(),
            filename=filename,
            extension=".xlsx",
            size_bytes=len(content),
            content_sha256=sha256,
            imported_source_sha256=sha256,
            status="ACTIVE",
        )
        db.add(working_copy)
        db.flush()
        version.working_copy_id = working_copy.id
        db.add(FileObject(
            document_id=document.id,
            storage_backend="working_copy_local",
            storage_path=storage_path,
            size_bytes=len(content),
            sha256=sha256,
        ))
        path = FileLifecycleStorageService().working_copy_path(storage_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        db.commit()
        return document.id


def create_dataset(client, headers, key="table-one"):
    """创建共享固定 schema。"""
    response = client.post(BASE, headers=headers, json={"name": "累计登记", "fields": FIELDS, "idempotency_key": key})
    assert response.status_code == 200, response.text
    return response.json()["dataset_id"]


def candidate(client, headers, document_id, name, amount, key="append-one", confidence=0.98):
    """严格使用 source 返回的版本，不猜测对象范围。"""
    response = client.get(f"{BASE}/sources/{document_id}", headers=headers)
    assert response.status_code == 200, response.text
    source = response.json()
    return {k: source[k] for k in ("document_id", "document_version_id", "extraction_run_id", "source_sha256")} | {
        "idempotency_key": key, "records": [{"fields": {
            "name": {"raw_text": name, "confidence": confidence},
            "amount": {"raw_text": amount, "confidence": confidence},
        }}],
    }


def workbuddy_candidate(client, headers, document_id, rows, *, key):
    """从正式 source_units 生成带稳定 UUID 和原文范围的字段证据。"""

    response = client.get(f"{BASE}/sources/{document_id}?content_origin=WORKBUDDY", headers=headers)
    assert response.status_code == 200, response.text
    source = response.json()
    records = []
    for unit_index, name, amount in rows:
        unit = source["source_units"][unit_index]
        fields = {}
        for field_key, raw_text in (("name", name), ("amount", amount)):
            start = unit["text"].index(raw_text)
            fields[field_key] = {"raw_text": raw_text, "confidence": 0.99, "page_number": unit["page_number"],
                "evidence_refs": [{"source_unit_id": unit["source_unit_id"], "quote": raw_text,
                    "start": start, "end": start + len(raw_text)}]}
        records.append({"fields": fields})
    return {k: source[k] for k in ("document_id", "document_version_id", "extraction_run_id", "source_sha256")} | {
        "content_origin": "WORKBUDDY", "idempotency_key": key, "records": records,
    }


def test_complex_template_api_roundtrip_keeps_layout_and_publishes_data(demo):
    """真实 API 从上传模板到追加、版本发布、普通文件和按文件名恢复布局形成闭环。"""
    client, sessions, (headers, _) = demo

    def configure(workbook):
        """仅构造无隐私的多层表头及序号预留行。"""
        sheet = workbook["登记"]
        sheet.merge_cells("A1:C1")
        sheet.merge_cells("A2:A3")
        sheet.merge_cells("B2:B3")
        sheet.merge_cells("C2:C3")
        sheet.merge_cells("A6:C6")

    document_id = seed_excel_target(sessions, filename="复杂模板.xlsx", configure=configure,
        sheets={"登记": [["汇总表"], ["序号", "姓名", "金额"], [], [1], [2], ["填写说明"]]})
    layout = {"header_range": "A2:C3", "data_range": "A4:C5"}
    response = client.post(f"{BASE}/from-excel", headers=headers, json={
        "document_id": document_id, "sheet_name": "登记", "layout": layout, "idempotency_key": "complex"})
    assert response.status_code == 200, response.text
    target = response.json()
    assert target["layout"] == layout and target["base_record_count"] == 0
    for i in (1, 2):
        source_id = seed_image(sessions, f"姓名：人员{i} 金额：{i}", filename=f"材料{i}.png")
        request = candidate(client, headers, source_id, f"人员{i}", str(i), key=f"row-{i}")
        target = client.get(f"{BASE}/targets/by-filename", headers=headers, params={"filename": "复杂模板.xlsx"}).json()
        keys = [f["key"] for f in target["fields"]]
        request["records"][0]["fields"] = dict(zip(keys[1:], request["records"][0]["fields"].values()))
        request.update(target_working_copy_id=target["working_copy_id"], target_document_id=target["document_id"],
                       target_document_version_id=target["document_version_id"])
        appended = client.post(f"{BASE}/{target['dataset_id']}/records", headers=headers, json=request)
        assert appended.status_code == 200, appended.text
        exported = client.post(f"{BASE}/{target['dataset_id']}/exports", headers=headers)
        assert exported.status_code == 200, exported.text
        receipt = exported.json()
        assert receipt["saved_to_server"] and receipt["confirmed_record_count"] == i
        assert receipt["working_copy_id"] == target["working_copy_id"]
    downloaded = client.get(receipt["download_url"], headers=headers)
    workbook = load_workbook(io.BytesIO(downloaded.content))
    assert workbook.active["A4"].value == 1 and workbook.active["B4"].value == "人员1"
    assert workbook.active["A5"].value == 2 and workbook.active["B5"].value == "人员2"
    assert workbook.active["A6"].value == "填写说明"
    workbook.close()
    with sessions() as db:
        dataset = db.get(StructuredDataset, target["dataset_id"])
        working = db.get(WorkingCopy, dataset.working_copy_id)
        root = db.get(WorkingCopyRoot, working.working_copy_root_id)
        visible = FileLifecycleStorageService().working_copy_path(f"{root.relative_storage_path}/{working.relative_path}")
        assert visible.name == receipt["filename"] == "复杂模板.xlsx"
        assert visible.read_bytes() == downloaded.content


def test_shared_append_query_export_and_idempotent_retry(demo):
    """跨用户累计后只保留一个普通文件，revision 形成同一文档的不可变版本。"""
    client, sessions, (user_a, user_b) = demo
    dataset_id = create_dataset(client, user_a)
    original_ids = []
    first_export = None
    for index, (name, amount) in enumerate((("张三", "20"), ("李四", "100")), 1):
        document_id = seed_image(sessions, f"姓名：{name}\n金额：{amount}")
        original_ids.append(document_id)
        request = candidate(client, user_b, document_id, name, amount, key=f"image-{index}")
        if index > 1:
            binding = client.get(f"{BASE}/{dataset_id}", headers=user_b).json()
            request.update({
                "target_working_copy_id": binding["working_copy_id"],
                "target_document_id": binding["document_id"],
                "target_document_version_id": binding["document_version_id"],
            })
        response = client.post(f"{BASE}/{dataset_id}/records", headers=user_b, json=request)
        assert response.status_code == 200, response.text
        assert response.json()["revision"] == index
        repeat = client.post(f"{BASE}/{dataset_id}/records", headers=user_a, json=request)
        assert repeat.status_code == 200 and repeat.json()["reused"]
        assert repeat.json()["record_count"] == index
        if index == 1:
            response = client.post(f"{BASE}/{dataset_id}/exports", headers=user_a)
            assert response.status_code == 200, response.text
            first_export = response.json()
    query = client.post(f"{BASE}/{dataset_id}/query", headers=user_a, json={"sum_field": "amount", "limit": 1}).json()
    assert query["sum"] == "120" and query["matched_records"] == 2 and query["truncated"]
    filtered = client.post(f"{BASE}/{dataset_id}/query", headers=user_b,
        json={"filters": [{"field": "amount", "op": "gte", "value": 50}]}).json()
    assert filtered["matched_records"] == 1 and filtered["records"][0]["fields"]["name"]["value"] == "李四"
    exported = client.post(f"{BASE}/{dataset_id}/exports", headers=user_b)
    assert exported.status_code == 200, exported.text
    assert exported.json()["revision"] == 2 and exported.json()["export_id"] != first_export["export_id"]
    download = client.get(exported.json()["download_url"], headers=user_a)
    sheet = load_workbook(io.BytesIO(download.content)).active
    assert list(sheet.values) == [("姓名", "金额"), ("张三", "20"), ("李四", "100")]
    old = client.get(first_export["download_url"], headers=user_b)
    assert len(list(load_workbook(io.BytesIO(old.content)).active.values)) == 2
    # 当前普通文件意外缺失时，幂等重导出必须从不可变版本重新发布，不能只返回成功元数据。
    with sessions() as db:
        dataset = db.get(StructuredDataset, dataset_id)
        working_copy = db.get(WorkingCopy, dataset.working_copy_id)
        root = db.get(WorkingCopyRoot, working_copy.working_copy_root_id)
        current_path = FileLifecycleStorageService().working_copy_path(
            f"{root.relative_storage_path}/{working_copy.relative_path}"
        )
        current_path.unlink()
        assert not current_path.exists()
    repeated_export = client.post(f"{BASE}/{dataset_id}/exports", headers=user_a).json()
    assert repeated_export["reused"]
    assert repeated_export["working_copy_id"] == exported.json()["working_copy_id"]
    assert repeated_export["document_version_id"] == exported.json()["document_version_id"]
    assert current_path.is_file()
    assert client.get(BASE, headers=user_b).json()["datasets"][0]["dataset_id"] == dataset_id
    ordinary_files = [
        item for item in client.get("/api/working-copies", headers=user_b).json()
        if item["id"] == exported.json()["working_copy_id"]
    ]
    assert len(ordinary_files) == 1
    assert ordinary_files[0]["filename"] == "累计登记.xlsx"
    assert ordinary_files[0]["current_version_id"] == exported.json()["document_version_id"]
    file_versions = client.get(
        f"/api/working-copies/{ordinary_files[0]['id']}/versions",
        headers=user_a,
    ).json()
    assert [item["version_number"] for item in file_versions] == [1, 2]
    current_download = client.get(
        f"/api/working-copies/{ordinary_files[0]['id']}/download",
        headers=user_b,
    )
    assert list(load_workbook(io.BytesIO(current_download.content)).active.values) == [
        ("姓名", "金额"), ("张三", "20"), ("李四", "100")
    ]
    with sessions() as db:
        dataset = db.get(StructuredDataset, dataset_id)
        working_copy = db.get(WorkingCopy, dataset.working_copy_id)
        versions = (
            db.query(DocumentVersion)
            .filter_by(document_id=working_copy.document_id)
            .order_by(DocumentVersion.version_number)
            .all()
        )
        exports = (
            db.query(StructuredDatasetExport)
            .filter_by(dataset_id=dataset_id)
            .order_by(StructuredDatasetExport.dataset_revision)
            .all()
        )
        # 普通文件列表以 WorkingCopy 为逻辑对象，因此历史版本不会增加第二条文件。
        assert db.query(WorkingCopy).filter_by(id=dataset.working_copy_id).count() == 1
        assert working_copy.relative_path == f"{dataset_id}/累计登记.xlsx"
        assert working_copy.filename == "累计登记.xlsx"
        assert len(versions) == 2
        assert [version.version_number for version in versions] == [1, 2]
        assert versions[1].parent_version_id == versions[0].id
        assert working_copy.current_version_id == versions[1].id
        assert first_export["working_copy_id"] == exported.json()["working_copy_id"]
        assert first_export["document_id"] == exported.json()["document_id"]
        assert first_export["version_number"] == 1
        assert exported.json()["version_number"] == 2
        assert [item.document_version_id for item in exports] == [version.id for version in versions]
        # 当前文件对象只有一条；历史内容只由两个 DocumentVersion 表达。
        file_object = db.query(FileObject).filter_by(document_id=working_copy.document_id).one()
        assert file_object.storage_path == f"structured-datasets/{working_copy.relative_path}"
        assert db.query(WorkingCopyPathRecord).filter_by(working_copy_id=working_copy.id).count() == 2
        storage = FileLifecycleStorageService()
        current_path = storage.working_copy_path(file_object.storage_path)
        assert current_path.is_file()
        assert current_path.name == working_copy.filename == "累计登记.xlsx"
        assert storage.sha256_file(current_path) == working_copy.content_sha256
        for version in versions:
            assert version.storage_path.startswith(f"structured-datasets/{dataset_id}/.versions/")
            assert storage.working_copy_path(version.storage_path).is_file()
        # 新实现不再在 FILE_STORAGE_ROOT 额外保留一份 derivative 导出。
        assert not (storage.upload_path("derivatives") / "structured-datasets").exists()
        assert db.query(StructuredDatasetMember).count() == 2
        assert db.query(StructuredExtractionRun).count() == 2
        assert db.query(ChangeItem).filter_by(change_type="STRUCTURED_DATASET_APPENDED").count() == 2
        assert db.query(ToolInvocation).filter_by(status="FAILED").count() == 0
        for document_id in original_ids:
            assert db.query(DocumentExtractionRun).filter_by(document_id=document_id, extractor="workbuddy-external-ocr").count() == 1
            version = db.query(DocumentVersion).filter_by(document_id=document_id).one()
            storage = FileLifecycleStorageService()
            assert storage.sha256_file(storage.working_copy_path(version.storage_path)) == version.sha256


def test_word_and_pdf_append_use_workbuddy_units_and_export_together(demo):
    """Word 空页码和 PDF 真实页码均只复用宿主全文，并累计到同一 Excel。"""

    client, sessions, (user_a, user_b) = demo
    dataset_id = create_dataset(client, user_a, key="mixed-documents")
    word_id = seed_workbuddy_document(sessions, filename="登记.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        units=[{"kind": "TEXT_SECTION", "page_number": None, "text": "姓名：样本丙 金额：50"},
               {"kind": "TEXT_SECTION", "page_number": None, "text": "姓名：样本丁 金额：150"}])
    pdf_id = seed_workbuddy_document(sessions, filename="登记.pdf", content_type="application/pdf",
        units=[{"kind": "PDF_PAGE", "page_number": 1, "text": "姓名：样本己 金额：80"},
               {"kind": "PDF_PAGE", "page_number": 2, "text": "姓名：样本庚 金额：120"}])

    word = workbuddy_candidate(client, user_b, word_id,
        [(0, "样本丙", "50"), (1, "样本丁", "150")], key="word-1")
    pdf = workbuddy_candidate(client, user_a, pdf_id,
        [(0, "样本己", "80"), (1, "样本庚", "120")], key="pdf-1")
    assert client.post(f"{BASE}/{dataset_id}/records", headers=user_b, json=word).status_code == 200
    assert client.post(f"{BASE}/{dataset_id}/records", headers=user_a, json=pdf).status_code == 200
    summary = client.get(f"{BASE}/{dataset_id}", headers=user_b).json()
    assert summary["document_count"] == 2
    assert summary["image_count"] == 0
    assert summary["format_counts"] == {"image": 0, "word": 1, "pdf": 1}
    query = client.post(f"{BASE}/{dataset_id}/query", headers=user_a,
        json={"sum_field": "amount", "limit": 10}).json()
    assert query["matched_records"] == 4 and query["sum"] == "400"
    exported = client.post(f"{BASE}/{dataset_id}/exports", headers=user_a).json()
    rows = list(load_workbook(io.BytesIO(client.get(exported["download_url"], headers=user_b).content)).active.values)
    assert rows == [("姓名", "金额"), ("样本丙", "50"), ("样本丁", "150"), ("样本己", "80"), ("样本庚", "120")]
    with sessions() as db:
        derived = db.query(DocumentElement).filter(DocumentElement.metadata_json.is_not(None)).all()
        traced = [element for element in derived if (element.metadata_json or {}).get("source_evidence_refs")]
        assert len(traced) == 4
        assert all(item["source_element_ids"] for element in traced
                   for item in element.metadata_json["source_evidence_refs"])


def test_workbuddy_source_index_and_selected_units_keep_evidence_ids(demo):
    """索引不回传正文；按页选取仍使用原始区段 UUID，提交校验保留完整来源。"""
    client, sessions, (headers, _) = demo
    document_id = seed_workbuddy_document(sessions, filename="逐页材料.pdf", content_type="application/pdf",
        units=[{"kind": "PDF_PAGE", "page_number": 1, "text": "封面姓名：样本甲"},
               {"kind": "PDF_PAGE", "page_number": 2, "text": "项目：已核准金额 20"}])
    base = f"{BASE}/sources/{document_id}?content_origin=WORKBUDDY"
    index_response = client.get(base + "&mode=INDEX", headers=headers)
    assert index_response.status_code == 200, index_response.text
    index = index_response.json()
    assert index["total_source_unit_count"] == 2
    assert [unit["order"] for unit in index["source_units"]] == [0, 1]
    assert all("text" not in unit for unit in index["source_units"])
    assert index["pages"] == []
    selected_response = client.get(base + "&unit_orders=1", headers=headers)
    assert selected_response.status_code == 200, selected_response.text
    selected = selected_response.json()
    assert len(selected["source_units"]) == 1
    assert selected["source_units"][0]["source_unit_id"] == index["source_units"][1]["source_unit_id"]
    assert selected["source_units"][0]["text"] == "项目：已核准金额 20"
    assert selected["source_sha256"] == index["source_sha256"]
    assert client.get(base + "&unit_orders=7", headers=headers).status_code == 422
    assert client.get(base + "&unit_orders=1&unit_orders=1", headers=headers).status_code == 422


def test_invalid_evidence_is_retained_but_excluded_from_totals(demo):
    """找不到原值不得被当成成功数字；待复核记录仍可检查。"""
    client, sessions, (headers, _) = demo
    dataset_id = create_dataset(client, headers)
    document_id = seed_image(sessions, "张三 金额 20")
    request = candidate(client, headers, document_id, "张三", "9999")
    response = client.post(f"{BASE}/{dataset_id}/records", headers=headers, json=request)
    assert response.status_code == 200, response.text
    assert response.json()["review_field_count"] == 1
    query = client.post(f"{BASE}/{dataset_id}/query", headers=headers, json={"sum_field": "amount"}).json()
    assert query["review_records"] == 1 and query["matched_records"] == 0 and query["sum"] == "0"
    review = client.post(f"{BASE}/{dataset_id}/query", headers=headers, json={"include_review": True, "sum_field": "amount"}).json()
    assert len(review["records"]) == 1 and review["sum"] == "0"
    exported = client.post(f"{BASE}/{dataset_id}/exports", headers=headers)
    assert exported.status_code == 409
    assert exported.json()["error"]["code"] == "DATASET_EXPORT_NO_CONFIRMED_RECORDS"


def test_export_rejects_revision_member_drift(demo):
    """来源清理导致成员链缺失时必须关闭式失败，不能发布标题行或复用旧成功。"""

    client, sessions, (headers, _) = demo
    dataset_id = create_dataset(client, headers, key="revision-drift")
    document_id = seed_image(sessions, "姓名：张三 金额：20")
    request = candidate(client, headers, document_id, "张三", "20", key="revision-drift-row")
    assert client.post(f"{BASE}/{dataset_id}/records", headers=headers, json=request).status_code == 200
    with sessions() as db:
        db.query(StructuredDatasetMember).filter_by(dataset_id=dataset_id).delete()
        db.commit()
    exported = client.post(f"{BASE}/{dataset_id}/exports", headers=headers)
    assert exported.status_code == 409
    assert exported.json()["error"]["code"] == "DATASET_REVISION_INCONSISTENT"


def test_conflicts_keep_facts_atomic_and_record_failed_invocation(demo):
    """版本错误、同键异载荷和同版本新键均不能重复增加成员。"""
    client, sessions, (headers, _) = demo
    dataset_id = create_dataset(client, headers)
    document_id = seed_image(sessions, "张三 20")
    request = candidate(client, headers, document_id, "张三", "20")
    wrong = dict(request, source_sha256="0" * 64)
    failed = client.post(f"{BASE}/{dataset_id}/records", headers=headers, json=wrong)
    assert failed.status_code == 409 and failed.json()["error"]["code"] == "SOURCE_SNAPSHOT_CHANGED"
    assert client.get(f"{BASE}/{dataset_id}", headers=headers).json()["revision"] == 0
    assert client.post(f"{BASE}/{dataset_id}/records", headers=headers, json=request).status_code == 200
    failed = client.post(f"{BASE}/{dataset_id}/records", headers=headers, json=wrong)
    assert failed.json()["error"]["code"] == "IDEMPOTENCY_CONFLICT"
    failed = client.post(f"{BASE}/{dataset_id}/records", headers=headers, json=dict(request, idempotency_key="new-key"))
    assert failed.json()["error"]["code"] == "SOURCE_ALREADY_APPENDED"
    with sessions() as db:
        assert db.query(StructuredDatasetMember).count() == 1
        assert db.query(StructuredExtractionRun).count() == 1
        assert db.query(ToolInvocation).filter_by(status="FAILED").count() == 3


def test_formula_protection_and_export_failure_does_not_undo_append(demo, monkeypatch):
    """复用现有安全 writer；导出失败后已提交记录和 revision 仍保留。"""
    client, sessions, (headers, _) = demo
    dataset_id = create_dataset(client, headers)
    document_id = seed_image(sessions, "=1+1 金额 20")
    request = candidate(client, headers, document_id, "=1+1", "20")
    assert client.post(f"{BASE}/{dataset_id}/records", headers=headers, json=request).status_code == 200
    from app.modules.structured_datasets import service
    writer = service._write_xlsx_atomic
    def failed_writer(**kwargs):
        """模拟磁盘失败，不写入真实外部目录。"""
        raise OSError("fake disk failure")
    monkeypatch.setattr(service, "_write_xlsx_atomic", failed_writer)
    assert client.post(f"{BASE}/{dataset_id}/exports", headers=headers).status_code == 500
    assert client.get(f"{BASE}/{dataset_id}", headers=headers).json()["record_count"] == 1
    monkeypatch.setattr(service, "_write_xlsx_atomic", writer)
    exported = client.post(f"{BASE}/{dataset_id}/exports", headers=headers).json()
    sheet = load_workbook(io.BytesIO(client.get(exported["download_url"], headers=headers).content)).active
    assert sheet["A2"].data_type == "s" and sheet["A2"].value == "'=1+1"


def test_feature_gate_and_strict_input(demo, monkeypatch):
    """保留登录和独立开关；不允许自由路径或未声明字段进入写入。"""
    client, _, (headers, _) = demo
    assert client.get(BASE).status_code == 401
    response = client.post(BASE, headers=headers, json={"name": "bad", "fields": FIELDS,
        "idempotency_key": "bad", "owner_user_id": "someone"})
    assert response.status_code == 422
    monkeypatch.setenv("WORKBUDDY_STRUCTURED_DATASET_DEMO_ENABLED", "false")
    get_settings.cache_clear()
    assert client.get(BASE, headers=headers).status_code == 503


def test_unknown_confidence_and_type_conflict_are_reviewed(demo):
    """未提供评分和数字解析失败不得被默认当成高质量字段。"""
    client, sessions, (headers, _) = demo
    dataset_id = create_dataset(client, headers)
    document_id = seed_image(sessions, "张三 金额 未知")
    request = candidate(client, headers, document_id, "张三", "未知", confidence=None)
    response = client.post(f"{BASE}/{dataset_id}/records", headers=headers, json=request)
    assert response.status_code == 200, response.text
    assert response.json()["review_field_count"] == 2
    query = client.post(f"{BASE}/{dataset_id}/query", headers=headers, json={"include_review": True}).json()
    assert query["records"][0]["fields"]["amount"]["value"] is None


def test_date_filters_reuse_normalization_and_reject_invalid_date(demo):
    """日期比较先规范化，不能对任意输入做字典序比较。"""
    client, sessions, (headers, _) = demo
    response = client.post(BASE, headers=headers, json={"name": "日期登记", "idempotency_key": "dates",
        "fields": [{"key": "date", "label": "日期", "field_type": "date", "required": True}]})
    dataset_id = response.json()["dataset_id"]
    document_id = seed_image(sessions, "2026年9月18日")
    request = candidate(client, headers, document_id, "unused", "unused")
    request["records"] = [{"fields": {"date": {"raw_text": "2026年9月18日", "confidence": 0.98}}}]
    assert client.post(f"{BASE}/{dataset_id}/records", headers=headers, json=request).status_code == 200
    query = client.post(f"{BASE}/{dataset_id}/query", headers=headers,
        json={"filters": [{"field": "date", "op": "gte", "value": "2026年9月1日"}]})
    assert query.json()["matched_records"] == 1
    assert client.post(f"{BASE}/{dataset_id}/query", headers=headers,
        json={"filters": [{"field": "date", "op": "gte", "value": "昨天"}]}).status_code == 422


def test_demo_migration_roundtrip_leaves_existing_tables():
    """新迁移可独立升级回退，且不会删除旧的文件事实表。"""
    import importlib.util
    from pathlib import Path
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from sqlalchemy import create_engine, inspect

    path = Path(__file__).resolve().parents[2] / "alembic/versions/20260920_0001_create_structured_datasets.py"
    spec = importlib.util.spec_from_file_location("dataset_migration_test", path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    engine = create_engine("sqlite+pysqlite:///:memory:")
    with engine.begin() as connection:
        connection.exec_driver_sql("CREATE TABLE documents (id TEXT PRIMARY KEY)")
        migration.op = Operations(MigrationContext.configure(connection))
        migration.upgrade()
        inspector = inspect(connection)
        assert {"structured_datasets", "structured_dataset_members", "structured_dataset_exports"}.issubset(inspector.get_table_names())
        assert len(inspector.get_unique_constraints("structured_dataset_members")) == 2
        migration.downgrade()
        assert inspect(connection).get_table_names() == ["documents"]
    # 同时编译真实部署使用的 PostgreSQL DDL，不连接或变更部署数据库。
    sql = io.StringIO()
    migration.op = Operations(MigrationContext.configure(dialect_name="postgresql", opts={"as_sql": True, "output_buffer": sql}))
    migration.upgrade()
    assert "schema_json JSONB" in sql.getvalue()
    assert "TIMESTAMP WITH TIME ZONE" in sql.getvalue()
    assert "ON DELETE RESTRICT" in sql.getvalue()


def test_working_copy_publication_migration_compiles_postgresql_ddl():
    """发布关联迁移必须同时固定唯一工作副本和唯一文档版本，且回退不删除文件事实。"""

    import importlib.util
    from pathlib import Path
    from alembic.migration import MigrationContext
    from alembic.operations import Operations

    path = Path(__file__).resolve().parents[2] / "alembic/versions/20260920_0003_publish_structured_dataset_working_copy.py"
    spec = importlib.util.spec_from_file_location("dataset_working_copy_migration_test", path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    sql = io.StringIO()
    migration.op = Operations(MigrationContext.configure(
        dialect_name="postgresql",
        opts={"as_sql": True, "output_buffer": sql},
    ))
    migration.upgrade()
    ddl = sql.getvalue()
    assert "working_copy_id" in ddl
    assert "document_version_id" in ddl
    assert "CREATE UNIQUE INDEX ix_structured_datasets_working_copy_id" in ddl
    assert "CREATE UNIQUE INDEX ix_structured_dataset_exports_document_version_id" in ddl
    assert "ON DELETE SET NULL" in ddl


def test_new_uploaded_excel_registers_new_target_without_appending_old(demo):
    """新上传目标走独立注册；只有无新目标时才可用文件名解析旧表。"""
    client, sessions, (headers, _) = demo
    old_document = seed_excel_target(sessions, filename="旧登记表.xlsx")
    old = client.post(f"{BASE}/from-excel", headers=headers, json={
        "document_id": old_document, "idempotency_key": "old-target"})
    assert old.status_code == 200, old.text
    old_target = old.json()
    new_document = seed_excel_target(sessions, filename="新登记表.xlsx")
    new = client.post(f"{BASE}/from-excel", headers=headers, json={
        "document_id": new_document, "idempotency_key": "new-uploaded-target"})
    assert new.status_code == 200, new.text
    new_target = new.json()
    assert new_target["dataset_id"] != old_target["dataset_id"]
    assert new_target["document_id"] != old_target["document_id"]
    assert new_target["working_copy_id"] != old_target["working_copy_id"]
    old_after = client.get(f"{BASE}/targets/by-filename", headers=headers,
        params={"filename": "旧登记表.xlsx"})
    assert old_after.status_code == 200, old_after.text
    assert old_after.json()["dataset_id"] == old_target["dataset_id"]
    assert old_after.json()["document_version_id"] == old_target["document_version_id"]


def test_uploaded_excel_becomes_named_target_and_keeps_existing_rows(demo):
    """上传 Excel 优先建立目标，后续按文件名解析并在同一工作副本生成新版本。"""

    client, sessions, (headers, _) = demo
    target_document_id = seed_excel_target(
        sessions,
        sheets={"登记": [["姓名", "金额"], ["已有人员", 10]], "空白说明": []},
        advanced_formatting=True,
    )
    with sessions() as db:
        original_working_copy = db.query(WorkingCopy).filter_by(
            document_id=target_document_id
        ).one()
        original_root = db.get(WorkingCopyRoot, original_working_copy.working_copy_root_id)
        old_visible = FileLifecycleStorageService().working_copy_path(
            f"{original_root.relative_storage_path}/{original_working_copy.relative_path}"
        )
        assert old_visible.is_file()
        with zipfile.ZipFile(old_visible, "r") as archive:
            original_parts = {name: archive.read(name) for name in archive.namelist()}
    created = client.post(
        f"{BASE}/from-excel",
        headers=headers,
        json={
            "document_id": target_document_id,
            "sheet_name": "登记",
            "idempotency_key": "uploaded-target-1",
        },
    )
    assert created.status_code == 200, created.text
    target = created.json()
    assert target["filename"] == "2026年人才材料登记表.xlsx"
    assert target["sheet_name"] == "登记"
    assert [field["label"] for field in target["fields"]] == ["姓名", "金额"]
    assert target["base_record_count"] == target["record_count"] == 1
    resolved = client.get(
        f"{BASE}/targets/by-filename",
        headers=headers,
        params={"filename": "2026年人才材料登记表.xlsx"},
    )
    assert resolved.status_code == 200, resolved.text
    binding = resolved.json()
    assert binding["dataset_id"] == target["dataset_id"]
    assert binding["resolved_by"] == "EXACT_FILENAME"

    source_document_id = seed_image(sessions, "姓名：新增人员 金额：20", filename="新增材料.png")
    source = client.get(f"{BASE}/sources/{source_document_id}", headers=headers).json()
    name_key, amount_key = [field["key"] for field in binding["fields"]]
    append = client.post(
        f"{BASE}/{target['dataset_id']}/records",
        headers=headers,
        json={
            **{key: source[key] for key in (
                "document_id", "document_version_id", "extraction_run_id", "source_sha256"
            )},
            "target_working_copy_id": binding["working_copy_id"],
            "target_document_id": binding["document_id"],
            "target_document_version_id": binding["document_version_id"],
            "idempotency_key": "append-to-uploaded-target",
            "records": [{"fields": {
                name_key: {"raw_text": "新增人员", "confidence": 0.99},
                amount_key: {"raw_text": "20", "confidence": 0.99},
            }}],
        },
    )
    assert append.status_code == 200, append.text
    exported = client.post(f"{BASE}/{target['dataset_id']}/exports", headers=headers)
    assert exported.status_code == 200, exported.text
    result = exported.json()
    assert result["working_copy_id"] == binding["working_copy_id"]
    assert result["document_id"] == binding["document_id"]
    assert result["filename"] == "2026年人才材料登记表.xlsx"
    downloaded = client.get(result["download_url"], headers=headers)
    workbook = load_workbook(io.BytesIO(downloaded.content))
    sheet = workbook["登记"]
    assert sheet.title == "登记"
    assert [row[:2] for row in sheet.values] == [
        ("姓名", "金额"), ("已有人员", 10), ("新增人员", 20)
    ]
    assert workbook.sheetnames == ["登记", "空白说明"]
    assert sheet.freeze_panes == "A2"
    assert sheet.column_dimensions["A"].width == 24
    assert sheet["A1"].font.bold is True
    assert sheet["A1"].fill.fgColor.rgb == "00FFF2CC"
    assert sheet["A3"].style_id == sheet["A2"].style_id
    assert "D1:E1" in workbook["空白说明"].merged_cells
    assert sheet["D2"].value == "=SUM(B2:B2)"
    assert sheet.auto_filter.ref == "A1:B3"
    assert sheet.tables["TalentRecords"].ref == "A1:B3"
    assert str(next(iter(sheet.data_validations.dataValidation)).sqref) == "B2:B100"
    assert len(sheet._charts) == 1
    with zipfile.ZipFile(io.BytesIO(downloaded.content), "r") as archive:
        exported_parts = {name: archive.read(name) for name in archive.namelist()}
    # 目标 worksheet/table 允许增加行和扩展范围，其他高级 OOXML part 必须逐字节保持。
    for part_name in (
        "xl/workbook.xml",
        "xl/styles.xml",
        "xl/worksheets/sheet2.xml",
        "xl/charts/chart1.xml",
        "xl/drawings/drawing1.xml",
    ):
        assert exported_parts[part_name] == original_parts[part_name]
    with sessions() as db:
        working_copy = db.get(WorkingCopy, binding["working_copy_id"])
        versions = db.query(DocumentVersion).filter_by(
            document_id=working_copy.document_id
        ).order_by(DocumentVersion.version_number).all()
        assert len(versions) == 2
        assert versions[0].storage_path.startswith(
            f"structured-datasets/{target['dataset_id']}/.versions/"
        )
        assert working_copy.current_version_id == versions[1].id
        root = db.get(WorkingCopyRoot, working_copy.working_copy_root_id)
        assert root.relative_storage_path == "structured-datasets"
        assert working_copy.relative_path == (
            f"{target['dataset_id']}/2026年人才材料登记表.xlsx"
        )
        visible = FileLifecycleStorageService().working_copy_path(
            f"{root.relative_storage_path}/{working_copy.relative_path}"
        )
        assert visible.name == working_copy.filename
        assert visible.is_file()
        assert not old_visible.exists()


def test_existing_target_relocation_and_new_export_removes_old_visible_file(demo):
    """修复前目标从旧目录迁移并同时生成新版本时，提交后只能留下新入口。"""

    client, sessions, (headers, _) = demo
    document_id = seed_excel_target(sessions, filename="旧位置登记表.xlsx")
    created = client.post(f"{BASE}/from-excel", headers=headers, json={
        "document_id": document_id,
        "idempotency_key": "legacy-visible-location",
    })
    assert created.status_code == 200, created.text
    target = created.json()
    storage = FileLifecycleStorageService()
    with sessions() as db:
        dataset = db.get(StructuredDataset, target["dataset_id"])
        working_copy = db.get(WorkingCopy, dataset.working_copy_id)
        managed_file = db.get(ManagedFile, working_copy.managed_file_id)
        current_root = db.get(WorkingCopyRoot, working_copy.working_copy_root_id)
        legacy_root = db.query(WorkingCopyRoot).filter(
            WorkingCopyRoot.workspace_id == working_copy.workspace_id,
            WorkingCopyRoot.id != current_root.id,
            WorkingCopyRoot.relative_storage_path == "excel-targets",
        ).one()
        current_path = storage.working_copy_path(
            f"{current_root.relative_storage_path}/{working_copy.relative_path}"
        )
        legacy_relative_path = f"{working_copy.document_id}/{working_copy.filename}"
        legacy_storage_path = f"{legacy_root.relative_storage_path}/{legacy_relative_path}"
        legacy_path = storage.import_working_copy(
            source=current_path,
            relative_path=legacy_storage_path,
            expected_sha256=working_copy.content_sha256,
        )
        current_path.unlink()
        working_copy.working_copy_root_id = legacy_root.id
        working_copy.relative_path = legacy_relative_path
        working_copy.relative_path_hash = hashlib.sha256(legacy_relative_path.encode()).hexdigest()
        managed_file.root_id = legacy_root.managed_root_id
        managed_file.relative_path = legacy_relative_path
        managed_file.relative_path_hash = hashlib.sha256(legacy_relative_path.encode()).hexdigest()
        dataset.target_filename = ""
        file_object = db.query(FileObject).filter_by(document_id=working_copy.document_id).one()
        file_object.storage_path = legacy_storage_path
        db.commit()
        assert legacy_path.is_file() and not current_path.exists()

    binding = client.get(f"{BASE}/{target['dataset_id']}", headers=headers).json()
    source_document_id = seed_image(sessions, "姓名：迁移人员 金额：30", filename="迁移材料.png")
    source = client.get(f"{BASE}/sources/{source_document_id}", headers=headers).json()
    name_key, amount_key = [field["key"] for field in binding["fields"]]
    request = {
        **{key: source[key] for key in (
            "document_id", "document_version_id", "extraction_run_id", "source_sha256"
        )},
        "idempotency_key": "relocate-and-export",
        "records": [{"fields": {
            name_key: {"raw_text": "迁移人员", "confidence": 0.99},
            amount_key: {"raw_text": "30", "confidence": 0.99},
        }}],
        "target_working_copy_id": binding["working_copy_id"],
        "target_document_id": binding["document_id"],
        "target_document_version_id": binding["document_version_id"],
    }
    appended = client.post(f"{BASE}/{target['dataset_id']}/records", headers=headers, json=request)
    assert appended.status_code == 200, appended.text
    exported = client.post(f"{BASE}/{target['dataset_id']}/exports", headers=headers)
    assert exported.status_code == 200, exported.text
    assert not legacy_path.exists()
    with sessions() as db:
        dataset = db.get(StructuredDataset, target["dataset_id"])
        working_copy = db.get(WorkingCopy, dataset.working_copy_id)
        root = db.get(WorkingCopyRoot, working_copy.working_copy_root_id)
        new_path = storage.working_copy_path(f"{root.relative_storage_path}/{working_copy.relative_path}")
        assert root.relative_storage_path == "structured-datasets"
        assert dataset.target_filename == working_copy.filename
        assert new_path.is_file()
        assert storage.sha256_file(new_path) == working_copy.content_sha256


def test_existing_flattened_export_is_rebuilt_from_original_template(demo):
    """旧 writer 生成的扁平工作簿必须从 revision 0 模板自修复，且不重复累计行。"""

    client, sessions, (headers, _) = demo
    document_id = seed_excel_target(
        sessions,
        filename="待修复登记表.xlsx",
        sheets={"登记": [["姓名", "金额"], ["已有人员", 10]], "模板说明": []},
        advanced_formatting=True,
    )
    target = client.post(f"{BASE}/from-excel", headers=headers, json={
        "document_id": document_id,
        "sheet_name": "登记",
        "idempotency_key": "flattened-target",
    }).json()
    binding = client.get(f"{BASE}/{target['dataset_id']}", headers=headers).json()
    source_document_id = seed_image(sessions, "姓名：修复人员 金额：40", filename="修复材料.png")
    source = client.get(f"{BASE}/sources/{source_document_id}", headers=headers).json()
    name_key, amount_key = [field["key"] for field in binding["fields"]]
    append_request = {
        **{key: source[key] for key in (
            "document_id", "document_version_id", "extraction_run_id", "source_sha256"
        )},
        "target_working_copy_id": binding["working_copy_id"],
        "target_document_id": binding["document_id"],
        "target_document_version_id": binding["document_version_id"],
        "idempotency_key": "flattened-target-row",
        "records": [{"fields": {
            name_key: {"raw_text": "修复人员", "confidence": 0.99},
            amount_key: {"raw_text": "40", "confidence": 0.99},
        }}],
    }
    assert client.post(
        f"{BASE}/{target['dataset_id']}/records", headers=headers, json=append_request
    ).status_code == 200
    first_export = client.post(f"{BASE}/{target['dataset_id']}/exports", headers=headers).json()

    # 模拟升级前 writer 已经发布的扁平文件：只剩标题和数据，没有原模板高级结构。
    flattened = Workbook()
    flattened_sheet = flattened.active
    flattened_sheet.title = "登记"
    flattened_sheet.append(["姓名", "金额"])
    flattened_sheet.append(["已有人员", 10])
    flattened_sheet.append(["修复人员", 40])
    flattened_buffer = io.BytesIO()
    flattened.save(flattened_buffer)
    flattened_bytes = flattened_buffer.getvalue()
    flattened_sha = hashlib.sha256(flattened_bytes).hexdigest()
    storage = FileLifecycleStorageService()
    with sessions() as db:
        dataset = db.get(StructuredDataset, target["dataset_id"])
        export = db.query(StructuredDatasetExport).filter_by(
            dataset_id=dataset.id,
            dataset_revision=dataset.revision,
        ).one()
        version = db.get(DocumentVersion, export.document_version_id)
        working_copy = db.get(WorkingCopy, dataset.working_copy_id)
        managed_file = db.get(ManagedFile, working_copy.managed_file_id)
        document = db.get(Document, working_copy.document_id)
        file_object = db.query(FileObject).filter_by(document_id=document.id).one()
        version_path = storage.working_copy_path(version.storage_path)
        visible_path = storage.working_copy_path(file_object.storage_path)
        version_path.write_bytes(flattened_bytes)
        visible_path.write_bytes(flattened_bytes)
        version.extraction_policy_json = {
            "structured_dataset_export_mode": "template-append-v2"
        }
        version.sha256 = export.sha256 = working_copy.content_sha256 = flattened_sha
        working_copy.imported_source_sha256 = flattened_sha
        managed_file.content_sha256 = managed_file.fingerprint = flattened_sha
        document.sha256 = file_object.sha256 = flattened_sha
        version.size_bytes = export.size_bytes = working_copy.size_bytes = len(flattened_bytes)
        managed_file.size_bytes = document.size_bytes = file_object.size_bytes = len(flattened_bytes)
        db.commit()

    repaired = client.post(f"{BASE}/{target['dataset_id']}/exports", headers=headers)
    assert repaired.status_code == 200, repaired.text
    repaired_result = repaired.json()
    assert repaired_result["reused"] is False
    assert repaired_result["export_id"] == first_export["export_id"]
    assert repaired_result["version_number"] == first_export["version_number"] + 1
    content = client.get(repaired_result["download_url"], headers=headers).content
    workbook = load_workbook(io.BytesIO(content), data_only=False)
    sheet = workbook["登记"]
    assert [row[:2] for row in sheet.values] == [
        ("姓名", "金额"), ("已有人员", 10), ("修复人员", 40)
    ]
    assert workbook.sheetnames == ["登记", "模板说明"]
    assert sheet["D2"].value == "=SUM(B2:B2)"
    assert len(sheet._charts) == 1
    assert sheet.tables["TalentRecords"].ref == "A1:B3"


@pytest.mark.parametrize("failure_mode", [None, "publish", "cleanup"])
@pytest.mark.parametrize("replace_existing", [False, True])
def test_legacy_xls_target_is_prepared_asynchronously_in_same_working_copy(
    demo, monkeypatch, failure_mode, replace_existing,
):
    """XLS 先返回任务；worker 校验 WorkBuddy schema 后发布同一逻辑文件的 XLSX 新版本。"""

    client, sessions, (headers, _) = demo
    old_target = None
    if replace_existing:
        old_document = seed_excel_target(sessions, filename="人才材料登记表.xlsx")
        old_response = client.post(f"{BASE}/from-excel", headers=headers, json={
            "document_id": old_document, "idempotency_key": "old-same-name-target",
        })
        assert old_response.status_code == 200, old_response.text
        old_target = old_response.json()
    document_id = seed_workbuddy_document(
        sessions,
        filename="人才材料登记表.xls",
        content_type="application/vnd.ms-excel",
        units=[{"kind": "SHEET", "sheet_name": "登记", "text": "姓名\t金额\n已有人员\t10"}],
    )
    with sessions() as db:
        source_version = db.query(DocumentVersion).filter_by(document_id=document_id).one()
        extraction = db.query(DocumentExtractionRun).filter_by(
            document_id=document_id,
            document_version_id=source_version.id,
        ).one()
        working_copy = db.query(WorkingCopy).filter_by(document_id=document_id).one()
        # The real submit run belongs to an upload staging version. Archiving
        # clones its pages into the current working-copy version with a new run ID.
        upload_document = Document(
            user_id=db.get(Document, document_id).user_id,
            workspace_id=source_version.document.workspace_id,
            original_filename=source_version.filename,
            content_type=source_version.content_type,
            size_bytes=source_version.size_bytes,
            sha256=source_version.sha256,
        )
        db.add(upload_document)
        db.flush()
        upload_version = DocumentVersion(
            document_id=upload_document.id, storage_path=f"staging/{upload_document.id}.xls",
            filename=source_version.filename, content_type=source_version.content_type,
            size_bytes=source_version.size_bytes, sha256=source_version.sha256,
            storage_tier="UPLOAD", source_type="UPLOAD",
        )
        db.add(upload_version)
        db.flush()
        upload_run = DocumentExtractionRun(
            document_id=upload_document.id, document_version_id=upload_version.id,
            status="COMPLETED", extractor="workbuddy-external-document",
            parser_name=extraction.parser_name, parser_version=extraction.parser_version,
            parser_config_hash=extraction.parser_config_hash,
        )
        db.add(upload_run)
        db.flush()
        for cloned_page in db.query(DocumentPage).filter_by(extraction_run_id=extraction.id).all():
            upload_page = DocumentPage(
                document_id=upload_document.id, extraction_run_id=upload_run.id,
                page_number=cloned_page.page_number, sheet_name=cloned_page.sheet_name,
                text_content=cloned_page.text_content,
                metadata_json=dict(cloned_page.metadata_json or {}),
            )
            db.add(upload_page)
            db.flush()
            cloned_page.metadata_json = {
                **dict(cloned_page.metadata_json or {}),
                "reused_from_upload_extraction": upload_run.id,
                "reused_from_upload_page_id": upload_page.id,
            }
        db.get(ManagedFile, working_copy.managed_file_id).source_upload_version_id = upload_version.id
        source_version.extraction_policy_json = {
            **dict(source_version.extraction_policy_json or {}),
            "origin_upload_version_id": upload_version.id,
            "origin_upload_extraction_run_id": upload_run.id,
        }
        upload_run_id = upload_run.id
        original_working_copy_id = working_copy.id
        source_version_id = source_version.id
        old_visible_path = FileLifecycleStorageService().working_copy_path(
            f"shared/{working_copy.relative_path}"
        )
        assert old_visible_path.is_file()
        db.commit()

    converted_path = FileLifecycleStorageService().upload_path("test-artifacts/converted.xlsx")
    converted_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "登记"
    sheet.append(["姓名", "金额"])
    sheet.append(["已有人员", 10])
    workbook.save(converted_path)
    converted_sha = hashlib.sha256(converted_path.read_bytes()).hexdigest()

    def fake_convert(_service, *, document, document_version, source_path, spec, force_reconvert=False):
        assert document.id == document_id
        assert document_version.id == source_version_id
        assert source_path == old_visible_path
        return ConvertedDocumentArtifact(
            artifact_id=str(uuid4()),
            file_path=converted_path,
            storage_path="test-artifacts/converted.xlsx",
            sha256=converted_sha,
            source_sha256=document_version.sha256,
            converter_name="libreoffice",
            converter_version="test",
            converter_config_hash="c" * 64,
            artifact_type="CONVERTED_XLSX",
            source_format=".xls",
            parsed_format=".xlsx",
            reused=False,
        )

    monkeypatch.setattr(
        "app.modules.structured_datasets.service.OfficeDerivativeService.get_or_create",
        fake_convert,
    )
    request = {
        "document_id": document_id,
        "document_version_id": source_version_id,
        "extraction_run_id": upload_run_id,
        "sheet_name": "登记",
        "legacy_target_policy": "CONVERT_TO_XLSX",
        "idempotency_key": "legacy-xls-target-1",
    }
    if replace_existing and failure_mode is None:
        unrelated_id = seed_workbuddy_document(
            sessions, filename="unrelated.xls", content_type="application/vnd.ms-excel",
            units=[{"kind": "SHEET", "sheet_name": "other", "text": "unrelated"}],
        )
        with sessions() as db:
            unrelated_run_id = db.query(DocumentExtractionRun).filter_by(document_id=unrelated_id).one().id
        rejected = client.post(f"{BASE}/from-excel", headers=headers, json={
            **request, "extraction_run_id": unrelated_run_id,
            "idempotency_key": "unrelated-upload-run",
        })
        assert rejected.status_code == 409, rejected.text
        assert rejected.json()["error"]["code"] == "TARGET_EXCEL_EXTRACTION_INVALID"
    accepted = client.post(f"{BASE}/from-excel", headers=headers, json=request)
    assert accepted.status_code == 202, accepted.text
    queued = accepted.json()
    assert queued["status"] == "PREPARING_TARGET"
    assert queued["target_filename"] == "人才材料登记表.xlsx"

    if failure_mode:
        method_name = ("publish_dataset_current_file" if failure_mode == "publish"
                       else "retire_dataset_legacy_current_file")
        original_action = getattr(FileLifecycleStorageService, method_name)
        calls = 0

        def fail_once(self, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 1:
                raise OSError("injected legacy XLS publication failure")
            return original_action(self, **kwargs)

        monkeypatch.setattr(FileLifecycleStorageService, method_name, fail_once)

    processed_job_id = process_next_filesystem_job(
        session_factory=sessions,
        worker_id="structured-dataset-test",
    )
    assert processed_job_id == queued["filesystem_job_id"]
    job = client.get(f"/api/jobs/{processed_job_id}", headers=headers)
    assert job.status_code == 200, job.text
    if failure_mode:
        from app.modules.structured_datasets.publication import recover_pending_publications
        assert job.json()["status"] == "PUBLISHING"
        assert old_visible_path.is_file()
        if old_target:
            with sessions() as db:
                assert db.get(WorkingCopy, old_target["working_copy_id"]).status == "ACTIVE"
        assert recover_pending_publications(sessions)["recovered"] == 1
        job = client.get(f"/api/jobs/{processed_job_id}", headers=headers)
        assert job.json()["status"] == "COMPLETED"
    result = job.json()["result"]
    assert result["status"] == "READY"
    if old_target:
        assert result["archived_dataset_id"] == old_target["dataset_id"]
    assert result["converted_from_legacy"] is True
    assert result["filename"] == "人才材料登记表.xlsx"
    assert result["working_copy_id"] == original_working_copy_id
    assert result["source_document_version_id"] == source_version_id
    assert result["document_version_id"] != source_version_id

    recalled = client.post(f"{BASE}/from-excel", headers=headers, json=request)
    assert recalled.status_code == 200, recalled.text
    assert recalled.json()["dataset_id"] == result["dataset_id"]
    assert recalled.json()["reused"] is True
    with sessions() as db:
        working_copy = db.get(WorkingCopy, original_working_copy_id)
        assert working_copy.filename == "人才材料登记表.xlsx"
        assert working_copy.extension == ".xlsx"
        assert working_copy.document_id == document_id
        versions = db.query(DocumentVersion).filter_by(document_id=document_id).order_by(
            DocumentVersion.version_number
        ).all()
        assert [version.filename for version in versions] == [
            "人才材料登记表.xls",
            "人才材料登记表.xlsx",
        ]
        root = db.get(WorkingCopyRoot, working_copy.working_copy_root_id)
        visible = FileLifecycleStorageService().working_copy_path(
            f"{root.relative_storage_path}/{working_copy.relative_path}"
        )
        assert visible.is_file()
        assert visible.name == "人才材料登记表.xlsx"
        assert not old_visible_path.exists()
        if old_target:
            old_copy = db.get(WorkingCopy, old_target["working_copy_id"])
            assert old_copy.status == "ARCHIVED"
            assert db.get(ManagedFile, old_copy.managed_file_id).status == "ARCHIVED"
            assert db.get(StructuredDataset, old_target["dataset_id"]) is not None
    if old_target:
        current = client.get(f"{BASE}/targets/by-filename", headers=headers,
            params={"filename": "人才材料登记表.xlsx"})
        assert current.status_code == 200, current.text
        assert current.json()["dataset_id"] == result["dataset_id"]
        assert old_target["dataset_id"] not in {
            item["dataset_id"] for item in client.get(BASE, headers=headers).json()["datasets"]}


def test_legacy_complex_template_without_cell_grid_fails_before_queue(demo):
    """复杂 XLS 表头只有 Sheet 正文时必须立即拒绝，不花时间转换或留下悬挂任务。"""

    from app.db.models import FilesystemJob
    client, sessions, (headers, _) = demo
    document_id = seed_workbuddy_document(
        sessions, filename="无坐标模板.xls", content_type="application/vnd.ms-excel",
        units=[{"kind": "SHEET", "sheet_name": "登记", "text": "姓名\t金额"}],
    )
    with sessions() as db:
        version = db.query(DocumentVersion).filter_by(document_id=document_id).one()
        run = db.query(DocumentExtractionRun).filter_by(document_id=document_id).one()
        before_jobs = db.query(FilesystemJob).count()
        version_id, run_id = version.id, run.id
    response = client.post(f"{BASE}/from-excel", headers=headers, json={
        "document_id": document_id, "document_version_id": version_id,
        "extraction_run_id": run_id, "sheet_name": "登记",
        "layout": {"header_range": "A2:B2", "data_range": "A3:B3"},
        "legacy_target_policy": "CONVERT_TO_XLSX", "idempotency_key": "missing-grid",
    })
    assert response.status_code == 422, response.text
    assert response.json()["error"]["code"] == "TARGET_EXCEL_HEADER_CELLS_MISSING"
    with sessions() as db:
        assert db.query(FilesystemJob).count() == before_jobs


def test_legacy_offset_header_requires_layout_before_conversion(demo):
    """真实网格的表头从第二行开始时，缺少 layout 不应浪费一次转换任务。"""

    from app.db.models import FilesystemJob

    client, sessions, (headers, _) = demo
    document_id = seed_workbuddy_document(
        sessions, filename="多层表头.xls", content_type="application/vnd.ms-excel",
        units=[{"kind": "SHEET", "sheet_name": "登记", "text": "序号\t姓名"}],
    )
    with sessions() as db:
        version = db.query(DocumentVersion).filter_by(document_id=document_id).one()
        run = db.query(DocumentExtractionRun).filter_by(document_id=document_id).one()
        for index, label in enumerate(("序号", "姓名"), 1):
            db.add(DocumentElement(
                document_id=document_id, extraction_run_id=run.id,
                element_index=index, label="table_cell", text_content=label,
                metadata_json={"sheet_name": "登记", "row_index": 1, "column_index": index - 1},
            ))
        version_id, run_id = version.id, run.id
        before_jobs = db.query(FilesystemJob).count()
        db.commit()
    request = {
        "document_id": document_id, "document_version_id": version_id,
        "extraction_run_id": run_id, "sheet_name": "登记",
        "legacy_target_policy": "CONVERT_TO_XLSX", "idempotency_key": "offset-header",
    }
    rejected = client.post(f"{BASE}/from-excel", headers=headers, json=request)
    assert rejected.status_code == 422, rejected.text
    assert rejected.json()["error"]["code"] == "TARGET_EXCEL_LAYOUT_REQUIRED"
    with sessions() as db:
        assert db.query(FilesystemJob).count() == before_jobs
    accepted = client.post(f"{BASE}/from-excel", headers=headers, json={
        **request, "layout": {"header_range": "A2:B2", "data_range": "A3:B3"},
    })
    assert accepted.status_code == 202, accepted.text


def test_legacy_target_layout_error_is_actionable_without_server_log():
    """确定性布局失败应向 WorkBuddy 返回安全错误码，而非泛化的转换失败。"""

    from types import SimpleNamespace
    from fastapi import HTTPException
    from app.modules.managed_files.worker import _public_job_error_message

    message = _public_job_error_message(
        job=SimpleNamespace(job_type="STRUCTURED_DATASET_TARGET_PREPARE"),
        error=HTTPException(422, detail={"code": "TARGET_EXCEL_LAYOUT_REQUIRED"}),
    )
    assert "TARGET_EXCEL_LAYOUT_REQUIRED" in message
    assert "原文件未修改" in message


def test_target_filename_conflict_sheet_ambiguity_and_version_binding(demo):
    """同名、工作表和确认期间版本变化必须关闭式失败，不能猜测目标。"""

    client, sessions, (headers, _) = demo
    invalid_name = client.post(BASE, headers=headers, json={
        "name": "非法目标",
        "filename": "../续写目标.xlsx",
        "fields": FIELDS,
        "idempotency_key": "invalid-target-path",
    })
    assert invalid_name.status_code == 422
    assert invalid_name.json()["error"]["code"] == "TARGET_EXCEL_FILENAME_INVALID"
    missing = client.get(
        f"{BASE}/targets/by-filename",
        headers=headers,
        params={"filename": "不存在的登记表.xlsx"},
    )
    assert missing.status_code == 404
    assert missing.json()["error"]["code"] == "TARGET_EXCEL_NOT_FOUND"
    target_document_id = seed_excel_target(sessions, filename="续写目标.xlsx")
    created = client.post(f"{BASE}/from-excel", headers=headers, json={
        "document_id": target_document_id,
        "idempotency_key": "target-version-test",
    }).json()
    conflict = client.post(BASE, headers=headers, json={
        "name": "另一个目标",
        "filename": "续写目标.xlsx",
        "fields": FIELDS,
        "idempotency_key": "same-visible-name",
    })
    assert conflict.status_code == 409
    assert conflict.json()["error"]["code"] == "TARGET_EXCEL_FILENAME_CONFLICT"

    ambiguous_document_id = seed_excel_target(
        sessions,
        filename="多表目标.xlsx",
        sheets={"表一": [["姓名"]], "表二": [["金额"]]},
    )
    ambiguous = client.post(f"{BASE}/from-excel", headers=headers, json={
        "document_id": ambiguous_document_id,
        "idempotency_key": "ambiguous-sheets",
    })
    assert ambiguous.status_code == 409
    assert ambiguous.json()["error"]["code"] == "TARGET_EXCEL_SHEET_AMBIGUOUS"

    old_binding = client.get(f"{BASE}/targets/by-filename", headers=headers,
        params={"filename": "续写目标.xlsx"}).json()
    first_source = seed_image(sessions, "姓名：甲 金额：1", filename="甲.png")
    first = candidate(client, headers, first_source, "甲", "1", key="target-first")
    first.update({
        "target_working_copy_id": old_binding["working_copy_id"],
        "target_document_id": old_binding["document_id"],
        "target_document_version_id": old_binding["document_version_id"],
    })
    # 上传模板字段 key 由中文标题确定生成；把候选映射到真实 key。
    target_keys = [field["key"] for field in created["fields"]]
    first["records"][0]["fields"] = dict(zip(target_keys, first["records"][0]["fields"].values()))
    assert client.post(f"{BASE}/{created['dataset_id']}/records", headers=headers, json=first).status_code == 200
    assert client.post(f"{BASE}/{created['dataset_id']}/exports", headers=headers).status_code == 200

    second_source = seed_image(sessions, "姓名：乙 金额：2", filename="乙.png")
    stale = candidate(client, headers, second_source, "乙", "2", key="target-stale")
    stale["records"][0]["fields"] = dict(zip(target_keys, stale["records"][0]["fields"].values()))
    stale.update({
        "target_working_copy_id": old_binding["working_copy_id"],
        "target_document_id": old_binding["document_id"],
        "target_document_version_id": old_binding["document_version_id"],
    })
    rejected = client.post(f"{BASE}/{created['dataset_id']}/records", headers=headers, json=stale)
    assert rejected.status_code == 409
    assert rejected.json()["error"]["code"] == "TARGET_EXCEL_VERSION_CONFLICT"


def test_uploaded_xlsx_replaces_same_named_target_only_after_successful_publication(demo, monkeypatch):
    """同名上传表成为新目标；发布失败时旧目标仍活动，恢复后旧版本可下载。"""

    from app.modules.structured_datasets.publication import recover_pending_publications

    client, sessions, (headers, _) = demo
    filename = "同名上传目标.xlsx"
    old_document = seed_excel_target(sessions, filename=filename)
    old_response = client.post(f"{BASE}/from-excel", headers=headers, json={
        "document_id": old_document, "idempotency_key": "old-upload-target",
    })
    assert old_response.status_code == 200, old_response.text
    old_target = old_response.json()
    with sessions() as db:
        old_export = db.query(StructuredDatasetExport).filter_by(dataset_id=old_target["dataset_id"]).one()
        old_export_id = old_export.id

    new_document = seed_excel_target(sessions, filename=filename,
        sheets={"登记": [["姓名", "金额"], ["新上传模板", 20]]})
    original_publish = FileLifecycleStorageService.publish_dataset_current_file
    calls = 0

    def fail_once(self, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise OSError("injected publication failure")
        return original_publish(self, **kwargs)

    monkeypatch.setattr(FileLifecycleStorageService, "publish_dataset_current_file", fail_once)
    new_request = {"document_id": new_document, "idempotency_key": "new-upload-target"}
    pending = client.post(f"{BASE}/from-excel", headers=headers, json=new_request)
    assert pending.status_code == 200, pending.text
    assert pending.json()["status"] == "PUBLISHING"
    with sessions() as db:
        assert db.get(WorkingCopy, old_target["working_copy_id"]).status == "ACTIVE"
        pending_copy = db.get(WorkingCopy, pending.json()["working_copy_id"])
        assert pending_copy.status == "PUBLISHING"
        assert db.get(ManagedFile, pending_copy.managed_file_id).status == "PUBLISHING"
    blocked = client.get(f"{BASE}/targets/by-filename", headers=headers,
        params={"filename": filename})
    assert blocked.status_code == 423
    assert recover_pending_publications(sessions)["recovered"] == 1

    current = client.get(f"{BASE}/targets/by-filename", headers=headers,
        params={"filename": filename})
    assert current.status_code == 200, current.text
    assert current.json()["dataset_id"] == pending.json()["dataset_id"]
    assert current.json()["base_record_count"] == 1
    with sessions() as db:
        old_copy = db.get(WorkingCopy, old_target["working_copy_id"])
        assert old_copy.status == "ARCHIVED"
        assert db.get(ManagedFile, old_copy.managed_file_id).status == "ARCHIVED"
        assert db.get(WorkingCopy, pending.json()["working_copy_id"]).status == "ACTIVE"
        assert db.get(ManagedFile, pending_copy.managed_file_id).status == "ACTIVE"
        archived = db.query(ChangeItem).filter_by(
            change_type="STRUCTURED_DATASET_TARGET_ARCHIVED",
            target_id=old_target["dataset_id"],
        ).one()
        assert archived.after_value_json["replaced_by_dataset_id"] == pending.json()["dataset_id"]
        new_export_id = db.query(StructuredDatasetExport).filter_by(
            dataset_id=pending.json()["dataset_id"]).one().id
    old_download = client.get(f"{BASE}/{old_target['dataset_id']}/exports/{old_export_id}/content",
        headers=headers)
    new_download = client.get(f"{BASE}/{pending.json()['dataset_id']}/exports/{new_export_id}/content",
        headers=headers)
    assert old_download.status_code == new_download.status_code == 200
    assert load_workbook(io.BytesIO(old_download.content), read_only=True).active["A2"].value == "已有人员"
    assert load_workbook(io.BytesIO(new_download.content), read_only=True).active["A2"].value == "新上传模板"
    listed = client.get(BASE, headers=headers).json()["datasets"]
    assert [item["dataset_id"] for item in listed if item["filename"] == filename] == [pending.json()["dataset_id"]]
    retry = client.post(f"{BASE}/from-excel", headers=headers, json=new_request)
    assert retry.status_code == 200
    assert retry.json()["reused"] is True


def test_dataset_rename_moves_only_visible_file_and_keeps_continuation_binding(demo):
    """累计表改名必须同步可见文件名，且不能移动或覆盖隐藏的历史版本。"""

    client, sessions, (headers, _) = demo
    dataset_id = create_dataset(client, headers, key="rename-dataset-target")
    first_source = seed_image(sessions, "姓名：甲 金额：1", filename="改名前来源.png")
    first = candidate(client, headers, first_source, "甲", "1", key="rename-first")
    assert client.post(f"{BASE}/{dataset_id}/records", headers=headers, json=first).status_code == 200
    first_export = client.post(f"{BASE}/{dataset_id}/exports", headers=headers).json()
    old_name = first_export["filename"]
    new_name = "2026年人才材料登记表.xlsx"

    plan = client.post(
        "/api/operations/plans",
        headers=headers,
        json={
            "conversation_id": "dataset-rename-regression",
            "operation_type": "RENAME_WORKING_COPIES",
            "reason": "用户修改累计表文件名",
            "items": [{
                "working_copy_id": first_export["working_copy_id"],
                "after": {"filename": new_name},
            }],
        },
    )
    assert plan.status_code == 200, plan.text
    confirmed = client.post(
        f"/api/operations/plans/{plan.json()['id']}/confirm",
        headers=headers,
        json={"confirmation": "确认重命名"},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["status"] == "EXECUTED"

    assert client.get(
        f"{BASE}/targets/by-filename", headers=headers, params={"filename": old_name}
    ).status_code == 404
    resolved = client.get(
        f"{BASE}/targets/by-filename", headers=headers, params={"filename": new_name}
    )
    assert resolved.status_code == 200, resolved.text
    binding = resolved.json()
    assert binding["dataset_id"] == dataset_id
    assert binding["working_copy_id"] == first_export["working_copy_id"]
    assert binding["document_id"] == first_export["document_id"]

    second_source = seed_image(sessions, "姓名：乙 金额：2", filename="改名后来源.png")
    second = candidate(client, headers, second_source, "乙", "2", key="rename-second")
    unbound = client.post(f"{BASE}/{dataset_id}/records", headers=headers, json=second)
    assert unbound.status_code == 409
    assert unbound.json()["error"]["code"] == "TARGET_EXCEL_BINDING_REQUIRED"
    second.update({
        "target_working_copy_id": binding["working_copy_id"],
        "target_document_id": binding["document_id"],
        "target_document_version_id": binding["document_version_id"],
    })
    appended = client.post(f"{BASE}/{dataset_id}/records", headers=headers, json=second)
    assert appended.status_code == 200, appended.text
    assert appended.json()["record_count_added"] == 1
    second_export = client.post(f"{BASE}/{dataset_id}/exports", headers=headers).json()
    assert second_export["working_copy_id"] == first_export["working_copy_id"]
    assert second_export["document_id"] == first_export["document_id"]
    assert second_export["filename"] == new_name
    assert second_export["version_number"] == 2

    with sessions() as db:
        dataset = db.get(StructuredDataset, dataset_id)
        working_copy = db.get(WorkingCopy, dataset.working_copy_id)
        root = db.get(WorkingCopyRoot, working_copy.working_copy_root_id)
        versions = db.query(DocumentVersion).filter_by(
            document_id=working_copy.document_id
        ).order_by(DocumentVersion.version_number).all()
        storage = FileLifecycleStorageService()
        visible = storage.working_copy_path(
            f"{root.relative_storage_path}/{working_copy.relative_path}"
        )
        assert dataset.target_filename == new_name
        assert working_copy.filename == new_name
        assert visible.name == new_name and visible.is_file()
        assert len(versions) == 2
        assert all("/.versions/" in version.storage_path for version in versions)
        missing_versions = [
            version.storage_path
            for version in versions
            if not storage.working_copy_path(version.storage_path).is_file()
        ]
        assert not missing_versions, missing_versions


def test_excel_target_migrations_compile_postgresql_ddl():
    """目标文件名、工作表、基础行和上传版本字段必须形成连续迁移。"""

    import importlib.util
    from pathlib import Path
    from alembic.migration import MigrationContext
    from alembic.operations import Operations

    path = Path(__file__).resolve().parents[2] / "alembic/versions/20260920_0005_add_structured_dataset_excel_targets.py"
    spec = importlib.util.spec_from_file_location("dataset_excel_target_migration_test", path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    sql = io.StringIO()
    migration.op = Operations(MigrationContext.configure(
        dialect_name="postgresql",
        opts={"as_sql": True, "output_buffer": sql},
    ))
    migration.upgrade()
    ddl = sql.getvalue()
    assert migration.down_revision == "20260920_0004"
    assert "target_filename" in ddl
    assert "target_sheet_name" in ddl
    assert "base_rows_json JSONB" in ddl
    assert "target_source_version_id" in ddl

    backfill_path = (
        Path(__file__).resolve().parents[2]
        / "alembic/versions/20260920_0006_backfill_structured_dataset_target_filename.py"
    )
    spec = importlib.util.spec_from_file_location("dataset_target_filename_backfill_test", backfill_path)
    backfill = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(backfill)
    sql = io.StringIO()
    backfill.op = Operations(MigrationContext.configure(
        dialect_name="postgresql",
        opts={"as_sql": True, "output_buffer": sql},
    ))
    backfill.upgrade()
    ddl = sql.getvalue()
    assert backfill.down_revision == "20260920_0005"
    assert "UPDATE structured_datasets AS dataset" in ddl
    assert "dataset.name || '.xlsx'" in ddl

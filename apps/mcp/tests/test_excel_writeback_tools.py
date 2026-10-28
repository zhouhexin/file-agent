"""新 MCP 工具只编辑固定本地副本，旧工具不受注册影响。"""

import asyncio
import hashlib
import zipfile
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import pytest
from mcp.server.fastmcp import FastMCP
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from file_agent_mcp.excel_writeback_tools import (
    _scoped_a_i_rows, _scoped_identity_review_rows, _template_rows, register_excel_writeback_tools,
)


def test_template_rows_leave_unfound_cells_empty_without_client_nulls():
    row = _template_rows([{"E": "张三", "T": "123"}], 20)[0]
    assert len(row) == 20
    assert row[0] is None and row[4] == "张三" and row[19] == "123"
    assert all(value is None for offset, value in enumerate(row) if offset not in {4, 19})
    with pytest.raises(ValueError, match="TEMPLATE_ROWS_INVALID"):
        _template_rows([{"E": "张三", "F": None}], 20)


@pytest.mark.parametrize("layout", [None, {"header_range": "A1:B1", "data_range": "A2:B4"}])
def test_general_writeback_numbers_new_rows_when_target_has_sequence_column(tmp_path, layout):
    baseline = tmp_path / "numbered.xlsx"
    output = tmp_path / "result.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "登记"
    sheet.append(["序号", "事项"])
    sheet.append([1, "原记录"])
    book.save(baseline)
    book.close()

    patch_workbook(baseline=baseline, output=output, sheet_name="登记",
                   labels=["序号", "事项"], rows=[[None, "第一条"], [None, "第二条"]],
                   layout=layout)
    result = load_workbook(output)
    assert [(result.active.cell(row, 1).value, result.active.cell(row, 2).value)
            for row in (2, 3, 4)] == [(1, "原记录"), (2, "第一条"), (3, "第二条")]
    result.close()


def test_general_writeback_does_not_fill_untitled_blank_column(tmp_path):
    baseline = tmp_path / "blank-column.xlsx"
    output = tmp_path / "result.xlsx"
    book = Workbook()
    book.active.title = "登记"
    book.active.append(["姓名", None, "单位"])
    book.save(baseline)
    book.close()
    labels = ["姓名", "〔空白列:B〕", "单位"]
    patch_workbook(baseline=baseline, output=output, sheet_name="登记",
                   labels=labels, rows=[["张三", None, "计算机学院"]], layout=None)
    result = load_workbook(output)
    assert [result.active.cell(2, col).value for col in (1, 2, 3)] == ["张三", None, "计算机学院"]
    result.close()
    with pytest.raises(ValueError, match="TARGET_BLANK_COLUMN_CHANGED"):
        patch_workbook(baseline=baseline, output=output, sheet_name="登记",
                       labels=labels, rows=[["张三", "误填", "计算机学院"]], layout=None)


def test_general_writeback_accepts_host_empty_string_for_middle_blank_column(tmp_path):
    """宿主空串不能把中间无标题列或业务留空列写成有值单元格。"""
    baseline = tmp_path / "23-column-template.xlsx"
    output = tmp_path / "result.xlsx"
    labels = ["序号", "〔空白列:B〕", "姓名", "性别",
              *[f"字段{column}" for column in range(5, 24)]]
    book = Workbook()
    book.active.title = "登记"
    book.active.append(["序号", None, *labels[2:]])
    book.save(baseline)
    book.close()

    patch_workbook(baseline=baseline, output=output, sheet_name="登记",
                   labels=labels, rows=[[None, "", "张三", "男", "", *([None] * 18)]], layout=None)
    result = load_workbook(output)
    assert result.active.max_column == 23
    assert result.active["A2"].value == 1
    assert result.active["B2"].value is None
    assert result.active["C2"].value == "张三"
    assert result.active["D2"].value == "男"
    assert result.active["E2"].value is None
    result.close()


def test_apply_23_column_template_with_host_empty_string_in_blank_b(tmp_path):
    """真实 MCP apply 入口应接受 23 列模板，并保留所有空列。"""
    baseline = tmp_path / "source.xlsx"
    labels = ["序号", "〔空白列:B〕", "姓名", "性别",
              *[f"字段{column}" for column in range(5, 24)]]
    book = Workbook()
    book.active.title = "登记"
    book.active.append(["序号", None, *labels[2:]])
    book.save(baseline)
    book.close()
    target_id, version_id, task_id = (str(uuid4()) for _ in range(3))
    source_id, source_version, run_id = (str(uuid4()) for _ in range(3))
    sha = hashlib.sha256(baseline.read_bytes()).hexdigest()
    task = {"task_id": task_id, "dataset_id": target_id, "document_version_id": version_id,
            "working_copy_id": str(uuid4()), "filename": "23列登记.xlsx", "sheet_name": "登记",
            "labels": labels, "layout": None, "base_sha256": sha, "review_mode": "DIRECT_SAVE"}

    class FakeClient:
        async def excel_writeback_request(self, method, path, payload=None):
            return task

        async def download_working_copy(self, *, working_copy_id, output_dir):
            return {"sha256": sha, "resource_uri": baseline.as_uri()}

        async def close(self):
            pass

    payload = {"manifest": {"sources": [{"document_id": source_id,
        "document_version_id": source_version, "extraction_run_id": run_id,
        "source_sha256": "a" * 64}], "target": {"dataset_id": target_id,
        "document_version_id": version_id, "filename": task["filename"],
        "sheet_name": task["sheet_name"], "labels": labels}},
        "records": [{"source_index": 0, "fields": [
            {"blank_reason": "TARGET_DERIVED"}, {"blank_reason": "TARGET_BLANK_COLUMN"},
            {"raw_text": "张三"}, {"raw_text": "男"},
            *[{"blank_reason": "SOURCE_FIELD_NOT_FOUND"} for _ in range(19)]]}]}

    async def scenario():
        mcp = FastMCP("blank-b-apply-test")
        register_excel_writeback_tools(mcp, FakeClient, lambda: tmp_path / "cache")
        await mcp.call_tool("excel_writeback_checkout", {"filename": task["filename"], "request_key": "blank-b"})
        await mcp.call_tool("excel_writeback_apply", {"task_id": task_id,
            "payload": payload, "rows": [[None, "", "张三", "男", *([""] * 19)]]})
        editable = tmp_path / "cache" / "excel-writeback" / task_id / task["filename"]
        current = load_workbook(editable)
        assert [current.active.cell(2, column).value for column in (1, 2, 3, 4)] == [
            1, None, "张三", "男"]
        assert all(current.active.cell(2, column).value is None for column in range(5, 24))
        current.close()

    asyncio.run(scenario())


def test_scoped_identity_review_rows_only_populates_b_i_and_l():
    """桥接层内部补空位，避免宿主把 JSON null 改写为非空值。"""
    values = {column: column for column in "BCDEFGHI"}
    rows = _scoped_identity_review_rows([{**values, "L": "2022、2024"}, values], 27)
    assert len(rows) == 2
    assert all(len(row) == 27 for row in rows)
    assert rows[0][1:9] == list("BCDEFGHI")
    assert rows[0][11] == "2022、2024"
    assert rows[1][11] is None
    partial = _scoped_identity_review_rows([{"B": "学院", "I": "2016.6"}], 27)[0]
    assert partial[1] == "学院" and partial[8] == "2016.6"
    assert all(partial[index] is None for index in range(2, 8))
    assert all(row[0] is None and row[9] is None and row[10] is None
               and all(value is None for value in row[12:]) for row in rows)
    with pytest.raises(ValueError, match="SCOPED_WRITE_SCOPE_INVALID"):
        _scoped_identity_review_rows([{**values, "J": "越界"}], 27)
    with pytest.raises(ValueError, match="SCOPED_WRITE_SCOPE_INVALID"):
        _scoped_identity_review_rows([{"B": None}], 27)
from file_agent_mcp.client import WorkBuddyAttachmentRegistry
from file_agent_mcp.workbuddy_submission import WorkBuddySubmissionService
from file_agent_mcp.xlsx_writeback import patch_existing_cells, patch_workbook


def test_patch_existing_cells_changes_only_named_data_cell(tmp_path):
    """更正保留公式、样式和其他部件，旧值不匹配时拒绝覆盖。"""
    baseline = tmp_path / "source.xlsx"
    output = tmp_path / "result.xlsx"
    book = Workbook()
    book.active.title = "登记"
    book.active["A1"] = "姓名"
    book.active["A2"] = "旧值"
    book.active["A2"].font = Font(bold=True, color="FF0000")
    book.active["B2"] = "=1+2"
    book.create_sheet("其他")["A1"] = "不变"
    book.save(baseline)
    book.close()
    patch_existing_cells(baseline=baseline, output=output, sheet_name="登记", row_number=2,
        changes={"A2": ("旧值", "新值 √")})
    result = load_workbook(output)
    assert result["登记"]["A2"].value == "新值 √"
    assert result["登记"]["A2"].font.bold
    assert result["登记"]["B2"].value == "=1+2"
    assert result["其他"]["A1"].value == "不变"
    result.close()
    with zipfile.ZipFile(baseline) as before, zipfile.ZipFile(output) as after:
        for name in before.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert before.read(name) == after.read(name)
    with pytest.raises(ValueError, match="CORRECTION_BEFORE_CHANGED"):
        patch_existing_cells(baseline=baseline, output=output, sheet_name="登记", row_number=2,
            changes={"A2": ("错误旧值", "新值")})
    with pytest.raises(ValueError, match="WRITEBACK_FORMULA_NOT_ALLOWED"):
        patch_existing_cells(baseline=baseline, output=output, sheet_name="登记", row_number=2,
            changes={"B2": ("=1+2", "新值")})


def test_manual_revision_tool_uses_frozen_checkout_and_trusted_attachment(tmp_path, monkeypatch):
    filename = "教授汇总.xlsx"
    path = tmp_path / filename
    book = Workbook()
    book.active["A1"] = "姓名"
    book.save(path)
    book.close()
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    version_id = str(uuid4())
    calls = []

    class FakeClient:
        async def dataset_request(self, method, target_path, payload=None):
            return {"status": "ACTIVE", "filename": filename, "dataset_id": str(uuid4()),
                    "working_copy_id": str(uuid4()), "document_version_id": version_id,
                    "revision": 2}

        async def download_working_copy(self, **kwargs):
            return {"sha256": sha, "resource_uri": path.as_uri(),
                    "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "size_bytes": path.stat().st_size}

        async def upload_manual_excel_revision(self, **kwargs):
            calls.append(kwargs)
            return {"status": "COMPLETED", "excel_published": True,
                    "dataset_id": str(uuid4()), "export_id": str(uuid4()), "filename": filename}

        async def download_dataset_export(self, **kwargs):
            return {"resource_uri": path.as_uri(), "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "size_bytes": path.stat().st_size}

        async def close(self):
            pass

    fake_submission = SimpleNamespace(submission_id="trusted-submission")
    fake_attachment = SimpleNamespace(local_path=path)
    monkeypatch.setattr(WorkBuddyAttachmentRegistry, "from_environment",
                        classmethod(lambda cls: object()))
    monkeypatch.setattr(WorkBuddySubmissionService, "from_environment",
                        classmethod(lambda cls, **kwargs: SimpleNamespace(
                            manual_excel_revision_attachment=lambda **arguments:
                                (fake_submission, fake_attachment))))

    async def scenario():
        mcp = FastMCP("manual-revision-test")
        register_excel_writeback_tools(mcp, FakeClient, lambda: tmp_path)
        checkout = await mcp.call_tool("excel_manual_revision_checkout", {"target_filename": filename})
        assert checkout.structuredContent["base_document_version_id"] == version_id
        await mcp.call_tool("excel_manual_revision_publish", {
            "submission_ref": "fixed-ref", "target_filename": filename,
            "base_document_version_id": version_id, "base_sha256": sha,
            "base_revision": 2})

    asyncio.run(scenario())
    assert len(calls) == 1
    assert calls[0]["path"] == path
    assert calls[0]["base_document_version_id"] == version_id
    assert calls[0]["file_sha256"] == sha


def test_writeback_tools_checkout_and_apply_preserve_baseline(tmp_path):
    baseline = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.title = "登记"
    book.active.append(["姓名"])
    book.active.append(["旧记录"])
    book.active["A2"].font = Font(bold=True, color="FF0000")
    other = book.create_sheet("公式")
    other["A1"] = 1
    other["A2"] = "=A1*2"
    chart = BarChart()
    chart.add_data(Reference(other, min_col=1, min_row=1, max_row=1))
    other.add_chart(chart, "C1")
    book.save(baseline)
    book.close()
    target_id, version_id, task_id = (str(uuid4()) for _ in range(3))
    source_id, source_version, run_id = (str(uuid4()) for _ in range(3))
    sha = hashlib.sha256(baseline.read_bytes()).hexdigest()
    task = {"task_id": task_id, "dataset_id": target_id, "document_version_id": version_id,
            "working_copy_id": str(uuid4()), "filename": "登记表.xlsx", "sheet_name": "登记",
            "labels": ["姓名"], "layout": None, "base_sha256": sha, "review_mode": "DIRECT_SAVE"}
    calls = []

    class FakeClient:
        async def dataset_request(self, method, path, payload=None):
            calls.append((method, path, payload))
            return {"filename": "登记表.xlsx"}

        async def excel_writeback_request(self, method, path, payload=None):
            calls.append((method, path, payload))
            return task

        async def download_working_copy(self, *, working_copy_id, output_dir):
            return {"sha256": sha, "resource_uri": baseline.as_uri()}

        async def close(self):
            pass

    payload = {"manifest": {"sources": [{"document_id": source_id,
        "document_version_id": source_version, "extraction_run_id": run_id,
        "source_sha256": "a" * 64}], "target": {"dataset_id": target_id,
        "document_version_id": version_id, "filename": "登记表.xlsx",
        "sheet_name": "登记", "labels": ["姓名"]}},
        "records": [{"source_index": 0, "fields": [{"raw_text": "张三"}]}]}

    async def scenario():
        mcp = FastMCP("writeback-test")
        register_excel_writeback_tools(mcp, FakeClient, lambda: tmp_path / "cache")
        names = {tool.name for tool in await mcp.list_tools()}
        # 更正工具与新增行工具并存；旧工具名称和行为不改变。
        assert names == {"excel_writeback_checkout", "excel_writeback_apply",
                         "excel_writeback_preview", "excel_writeback_confirm", "excel_writeback_publish",
                         "excel_writeback_correction_checkout", "excel_writeback_correction_apply",
                         "excel_writeback_correction_confirm", "excel_writeback_correction_publish",
                         "excel_manual_revision_checkout", "excel_manual_revision_publish"}
        await mcp.call_tool("excel_writeback_checkout", {"filename": "登记表.xlsx", "request_key": "test"})
        await mcp.call_tool("excel_writeback_apply", {
            "task_id": task_id, "payload": payload, "rows": [["张三"]]})
        editable = tmp_path / "cache" / "excel-writeback" / task_id / "登记表.xlsx"
        assert editable.is_file()
        current = load_workbook(editable)
        assert [current.active["A2"].value, current.active["A3"].value] == ["旧记录", "张三"]
        assert current.active["A2"].font.bold
        assert current["公式"]["A2"].value == "=A1*2"
        current.close()
        assert hashlib.sha256(baseline.read_bytes()).hexdigest() == sha
        with zipfile.ZipFile(baseline) as before, zipfile.ZipFile(editable) as after:
            for name in before.namelist():
                if name != "xl/worksheets/sheet1.xml":
                    assert before.read(name) == after.read(name), name
        assert calls[-1][1].endswith("/preview")
        assert calls[0][2]["review_mode"] == "DIRECT_SAVE"
        # 新上传目标必须先注册，不能仅用同名字符串解析旧目标。
        await mcp.call_tool("excel_writeback_checkout", {
            "filename": "新上传模板.xlsx", "request_key": "new-template",
            "uploaded_document_id": str(uuid4())})
        assert calls[-2][1] == "/from-excel"
        assert calls[-2][2]["document_id"]
        assert calls[-1][1] == "/checkout"
        assert calls[-1][2]["review_mode"] == "DIRECT_SAVE"
        await mcp.call_tool("excel_writeback_checkout", {
            "filename": "新上传模板.xlsx", "request_key": "new-template-layout",
            "uploaded_document_id": str(uuid4()), "sheet_name": "登记",
            "layout": {"header_range": "A2:B2", "data_range": "A3:B3"},
        })
        assert calls[-2][2]["sheet_name"] == "登记"
        assert calls[-2][2]["layout"] == {"header_range": "A2:B2", "data_range": "A3:B3"}

    asyncio.run(scenario())


def test_scoped_apply_accepts_b_to_i_without_null_placeholders(tmp_path):
    baseline = tmp_path / "source.xlsx"
    labels = ["序号", "单位", "姓名", "出生年月", "工作时间", "职称", "职称时间", "学历", "学历时间"]
    values = ["计算机科学与工程学院", "石程", "1980-01", "2001-01", "教授", "2020-01", "博士", "2007-01"]
    book = Workbook()
    book.active.title = "汇总"
    book.active.append(labels)
    book.active["A2"] = 1
    book.save(baseline)
    book.close()
    dataset_id, version_id, task_id = (str(uuid4()) for _ in range(3))
    source_id, source_version, run_id = (str(uuid4()) for _ in range(3))
    layout = {"header_range": "A1:I1", "data_range": "A2:I2"}
    sha = hashlib.sha256(baseline.read_bytes()).hexdigest()
    task = {"task_id": task_id, "dataset_id": dataset_id, "document_version_id": version_id,
            "working_copy_id": str(uuid4()), "filename": "汇总.xlsx", "sheet_name": "汇总",
            "labels": labels, "layout": layout, "base_sha256": sha, "review_mode": "DIRECT_SAVE"}
    calls = []

    class FakeClient:
        async def excel_writeback_request(self, method, path, payload=None):
            calls.append((method, path, payload))
            return task

        async def download_working_copy(self, *, working_copy_id, output_dir):
            return {"sha256": sha, "resource_uri": baseline.as_uri()}

        async def close(self):
            pass

    payload = {"manifest": {"sources": [{"document_id": source_id,
        "document_version_id": source_version, "extraction_run_id": run_id,
        "source_sha256": "a" * 64}], "target": {"dataset_id": dataset_id,
        "document_version_id": version_id, "filename": "汇总.xlsx", "sheet_name": "汇总",
        "layout": layout, "labels": labels}, "write_scope": "A:I",
        "required_pdf_columns": "B:I"}, "records": [{"source_index": 0,
        "fields": [{"blank_reason": "TARGET_DERIVED"}] + [
            {"raw_text": value, "evidence_refs": [{"source_unit_id": "ocr-unit", "quote": value}]}
            for value in values]}]}

    async def scenario():
        mcp = FastMCP("scoped-writeback-test")
        register_excel_writeback_tools(mcp, FakeClient, lambda: tmp_path / "cache")
        await mcp.call_tool("excel_writeback_checkout", {"filename": "汇总.xlsx", "request_key": "scoped"})
        await mcp.call_tool("excel_writeback_apply", {
            "task_id": task_id, "payload": payload, "rows": [values]})
        editable = tmp_path / "cache" / "excel-writeback" / task_id / "汇总.xlsx"
        current = load_workbook(editable, read_only=True)
        assert [current.active.cell(2, col).value for col in range(1, 10)] == [1, *values]
        current.close()
        assert calls[-1][1].endswith("/preview")
        # 真正的 JSON null 仍可按原完整行传输；字符串 "null" 不得被当作空格。
        await mcp.call_tool("excel_writeback_apply", {
            "task_id": task_id, "payload": payload, "rows": [[None, *values]]})
        assert calls[-1][1].endswith("/preview")

    asyncio.run(scenario())
    expanded = _scoped_a_i_rows([values], 27)[0]
    assert len(expanded) == 27 and expanded[0] is None and expanded[1:9] == values
    assert all(value is None for value in expanded[9:])
    with pytest.raises(ValueError, match="SCOPED_WRITE_SCOPE_INVALID"):
        _scoped_a_i_rows([["null", *values]], len(labels))


def test_scoped_append_generates_sequence_and_moves_footer(tmp_path):
    labels = ["序号", *[f"字段{column}" for column in range(2, 28)]]
    book = Workbook()
    sheet = book.active
    sheet.title = "汇总"
    sheet.append(labels)
    sheet["A5"] = 1
    sheet["A7"] = "填表说明"
    sheet.merge_cells("A7:AA7")
    baseline = tmp_path / "baseline.xlsx"
    book.save(baseline)
    book.close()
    layout = {"header_range": "A1:AA1", "data_range": "A5:AA6"}
    current = baseline
    for number in range(1, 5):
        output = tmp_path / f"result-{number}.xlsx"
        patch_workbook(baseline=current, output=output, sheet_name="汇总", labels=labels,
                       rows=[[None, *[f"值{number}-{column}" for column in range(2, 10)],
                              *([None] * 18)]], layout=layout, scoped_a_i=True)
        saved = load_workbook(output)
        sheet = saved["汇总"]
        row = 4 + number
        assert sheet.cell(row, 1).value == number
        assert sheet.cell(row, 2).value == f"值{number}-2"
        assert all(sheet.cell(row, column).value is None for column in range(10, 28))
        assert sheet.cell(max(7, row + 1), 1).value == "填表说明"
        assert f"A{max(7, row + 1)}:AA{max(7, row + 1)}" in {
            str(item) for item in sheet.merged_cells.ranges}
        saved.close()
        if number >= 3:
            layout = {**layout, "data_range": f"A5:AA{row}"}
        current = output


def test_append_multiline_uses_populated_data_row_style(tmp_path):
    """末行只有模板占位符时，新增行仍沿用真实人员行的换行样式。"""
    baseline = tmp_path / "baseline.xlsx"
    output = tmp_path / "result.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "应届"
    sheet.append(["序号", "姓名", "学习经历"])
    sheet["A5"] = 1
    sheet["B5"] = "张三"
    sheet["C5"] = "博士\n硕士"
    sheet["C5"].alignment = Alignment(wrap_text=True)
    sheet["C6"] = "."
    book.save(baseline)
    book.close()
    value = "博士：2026.09\n西安理工大学\n\n硕士：硕博连读"
    patch_workbook(baseline=baseline, output=output, sheet_name="应届",
                   labels=["序号", "姓名", "学习经历"],
                   rows=[[None, "李四", value]],
                   layout={"header_range": "A1:C1", "data_range": "A5:C6"}, scoped_a_i=True)
    result = load_workbook(output)
    assert result["应届"]["C6"].value == "."
    assert result["应届"]["A7"].value == 2
    assert result["应届"]["C7"].value == value
    assert result["应届"]["C7"].alignment.wrap_text is True
    assert result["应届"].row_dimensions[7].height >= 72
    result.close()
    with zipfile.ZipFile(baseline) as before, zipfile.ZipFile(output) as after:
        for part in before.namelist():
            if part != "xl/worksheets/sheet1.xml":
                assert before.read(part) == after.read(part)


def test_append_replaces_case_insensitive_br_tags_before_writing(tmp_path):
    baseline = tmp_path / "baseline.xlsx"
    output = tmp_path / "result.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "应届"
    sheet.append(["序号", "姓名", "学习经历"])
    book.save(baseline)
    book.close()

    patch_workbook(baseline=baseline, output=output, sheet_name="应届",
                   labels=["序号", "姓名", "学习经历"],
                   rows=[[None, "李四", "博士<BR>硕士<br/>本科<br />论文"]],
                   layout={"header_range": "A1:C1", "data_range": "A2:C2"}, scoped_a_i=True)

    result = load_workbook(output)
    assert result["应届"]["C2"].value == "博士\n硕士\n本科\n论文"
    assert result["应届"]["C2"].alignment.wrap_text is not True
    result.close()


def test_empty_template_data_row_does_not_copy_header_style(tmp_path):
    """空模板没有数据样式行时，新数据格不得伪装为加粗着色的标题格。"""
    baseline = tmp_path / "empty-template.xlsx"
    output = tmp_path / "filled.xlsx"
    book = Workbook()
    book.active["A1"] = "姓名"
    book.active["A1"].font = Font(bold=True)
    book.active["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    book.save(baseline)
    book.close()
    patch_workbook(baseline=baseline, output=output, sheet_name="Sheet", labels=["姓名"],
                   rows=[["张三"]], layout=None)
    filled = load_workbook(output)
    assert filled.active["A2"].value == "张三"
    assert not filled.active["A2"].font.bold
    assert filled.active["A2"].fill.patternType is None
    filled.close()

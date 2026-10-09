"""WorkBuddy 本地编辑 Excel 后回传服务器的新 MCP 工具，不修改旧表格工具。"""

from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path
from typing import Annotated, Any, Literal
from urllib.parse import quote, urlparse
from urllib.request import url2pathname
from uuid import UUID

from mcp.types import CallToolResult, ResourceLink, TextContent
from pydantic import Field

from .attachment_transfer import AttachmentTransferService
from .client import WorkBuddyAttachmentRegistry, file_sha256
from .workbuddy_submission import WorkBuddySubmissionService
from .dataset_task_contracts import TaskPayload
from .dataset_tools import DatasetLayout
from .correction_contracts import CorrectionPayload
from .xlsx_writeback import _bounds, _column, patch_existing_cells, patch_workbook
from openpyxl.utils.cell import get_column_letter


def writeback_digest(payload: dict, file_sha256: str) -> str:
    """与服务端一致地冻结候选和完整结果文件。"""
    body = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256((body + ":" + file_sha256).encode("utf-8")).hexdigest()


def _editable(root: Path, task_id: str, filename: str) -> Path:
    """结果只能来自固定任务目录，拒绝客户端传任意本机路径。"""
    if Path(filename).name != filename or not filename.lower().endswith(".xlsx"):
        raise ValueError("WRITEBACK_FILENAME_INVALID: 目标必须是安全的 XLSX 文件名")
    configured = root.expanduser()
    if configured.is_symlink():
        raise ValueError("WRITEBACK_CACHE_INVALID: 缓存根不能是符号链接")
    base = configured.resolve() / "excel-writeback"
    task_dir = base / task_id
    if base.is_symlink() or task_dir.is_symlink():
        raise ValueError("WRITEBACK_CACHE_INVALID: 任务目录不能是符号链接")
    task_dir.mkdir(parents=True, exist_ok=True, mode=0o700)
    result = task_dir / filename
    if result.is_symlink() or task_dir.resolve() not in result.resolve().parents:
        raise ValueError("WRITEBACK_CACHE_INVALID: 结果文件越出固定任务目录")
    return result


def _baseline(editable: Path) -> Path:
    return editable.with_name(".writeback-baseline.xlsx")


def _scoped_a_i_rows(rows: list[list], label_count: int) -> list[list]:
    """允许宿主仅传 B–I 八个值，空列由受控工具补齐。"""
    normalized = []
    for row in rows:
        if len(row) == 8:
            row = [None, *row, *([None] * (label_count - 9))]
        if (len(row) != label_count or row[0] is not None
                or any(value is None for value in row[1:9])
                or any(value is not None for value in row[9:])):
            raise ValueError("SCOPED_WRITE_SCOPE_INVALID: 仅传 B–I 八个值，或传 A/J–末列均为 JSON null 的完整行")
        normalized.append(row)
    return normalized


def _scoped_identity_review_rows(rows: list[dict], label_count: int) -> list[list]:
    """仅接收已确认的 B–I 子集和可选 L；缺列由工具补空位。"""
    if label_count < 12:
        raise ValueError("SCOPED_WRITE_SCOPE_INVALID: 目标必须覆盖 L 列")
    normalized = []
    allowed = set("BCDEFGHI") | {"L"}
    for row in rows:
        if not isinstance(row, dict) or set(row) - allowed:
            raise ValueError("SCOPED_WRITE_SCOPE_INVALID: 每行只提交已确认的 B–I 值与可选 L 年度")
        if any(value is None or isinstance(value, bool)
               or (isinstance(value, str) and not value.strip()) for value in row.values()):
            raise ValueError("SCOPED_WRITE_SCOPE_INVALID: 行映射不得包含显式空值")
        normalized.append([None, *[row.get(column) for column in "BCDEFGHI"], None, None,
                           row.get("L"), *([None] * (label_count - 12))])
    return normalized


def _full_professor_rows(rows: list[list | dict], label_count: int) -> list[list]:
    """完整任务优先按物理列键传非空值，由工具补齐空位与 A 序号。"""
    if label_count != 27:
        raise ValueError("FULL_ROW_SCOPE_INVALID: 目标必须精确覆盖 A–AA")
    normalized = []
    for row in rows:
        if isinstance(row, dict):
            allowed = {get_column_letter(column) for column in range(2, 28)}
            if any(key not in allowed or value is None for key, value in row.items()):
                raise ValueError("FULL_ROW_SCOPE_INVALID: 行映射只能包含 B–AA 的非空展示值")
            row = [None, *[row.get(get_column_letter(column)) for column in range(2, 28)]]
        if len(row) == 26:
            row = [None, *row]
        if len(row) != 27 or row[0] is not None:
            raise ValueError("FULL_ROW_SCOPE_INVALID: 每行只传 B–AA 的 26 个值")
        normalized.append(row)
    return normalized


def register_excel_writeback_tools(mcp, client_factory, download_dir):
    """新工具以完整文件回传发布；旧 structured_dataset_* 清单原样保留。"""

    @mcp.tool(name="excel_writeback_checkout", description=(
        "按完整文件名锁定服务器当前 Excel 版本并下载到 WorkBuddy 固定可编辑缓存。"
        "本轮上传新 Excel 时提供其已归档 document_id/version_id；旧 xls 会先走受控异步提升。"),
        structured_output=False)
    async def checkout(
        filename: Annotated[str, Field(min_length=1, max_length=240)],
        request_key: Annotated[str, Field(min_length=1, max_length=120)],
        review_mode: Literal["DIRECT_SAVE", "REVIEW_REQUIRED"] = "DIRECT_SAVE",
        uploaded_document_id: UUID | None = None,
        uploaded_document_version_id: UUID | None = None,
        uploaded_extraction_run_id: UUID | None = None,
        sheet_name: str | None = None,
        layout: DatasetLayout | None = None,
    ):
        """新上传 Excel 优先；格式转换等待时不谎报已取得可编辑文件。"""
        client = client_factory()
        try:
            if uploaded_document_id:
                target = await client.dataset_request("POST", "/from-excel", {
                    "document_id": str(uploaded_document_id),
                    "document_version_id": str(uploaded_document_version_id) if uploaded_document_version_id else None,
                    "extraction_run_id": str(uploaded_extraction_run_id) if uploaded_extraction_run_id else None,
                    "sheet_name": sheet_name, "layout": layout.model_dump() if layout else None,
                    "legacy_target_policy": "CONVERT_TO_XLSX" if filename.lower().endswith(".xls") else None,
                    "idempotency_key": f"writeback-target-{request_key}"[:120],
                })
                if target.get("kind") == "filesystem_job":
                    return {**target, "next_action": "等待目标格式提升后重试同一 excel_writeback_checkout"}
                filename = target["filename"]
            task = await client.excel_writeback_request("POST", "/checkout", {
                "filename": filename, "request_key": request_key, "review_mode": review_mode})
            resource = await client.download_working_copy(
                working_copy_id=task["working_copy_id"], output_dir=Path(download_dir()))
            if resource["sha256"] != task["base_sha256"]:
                raise ValueError("WRITEBACK_VERSION_CONFLICT: checkout 后目标文件已变化，请重新读取")
            editable = _editable(Path(download_dir()), task["task_id"], task["filename"])
            local_uri = urlparse(resource["resource_uri"])
            if local_uri.scheme != "file" or local_uri.netloc:
                raise ValueError("WRITEBACK_CACHE_INVALID: 下载结果不是本机文件")
            baseline = _baseline(editable)
            if not baseline.exists():
                shutil.copyfile(Path(url2pathname(local_uri.path)), baseline)
            if file_sha256(baseline) != task["base_sha256"]:
                raise ValueError("WRITEBACK_VERSION_CONFLICT: 本地基线已变化，请重新 checkout")
            if not editable.exists():
                shutil.copyfile(baseline, editable)
            receipt = {**task, "editable_resource_uri": editable.as_uri(),
                       "baseline_resource_uri": resource["resource_uri"]}
            return CallToolResult(
                content=[TextContent(type="text", text=json.dumps(receipt, ensure_ascii=False)),
                         ResourceLink(type="resource_link", name=task["filename"], uri=editable.as_uri(),
                                      mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                      size=editable.stat().st_size)],
                structuredContent=receipt,
            )
        finally:
            await client.close()

    @mcp.tool(name="excel_writeback_apply", description=(
        "由 WorkBuddy 提供已提取并判定的每列最终展示值；A–I 定向任务的 rows 每行只传 B–I 八个值，"
        "B–I 加 L 定向任务每行传已确认的 B–I 子集及可选 L 的物理列键映射，由本机补齐其他空位；"
        "A–AA 完整任务每行优先按物理列键传 B–AA 非空值，本机补空位和 A 序号。"
        "本机只在冻结基线的目标工作表新增数据格，"
        "原字节保留其余工作簿部件。完成后自动登记候选和文件共同指纹。"), structured_output=True)
    async def apply(task_id: UUID, payload: TaskPayload,
                    rows: list[list[str | int | float | bool | None] | dict[str, str | int | float | bool]]) -> dict[str, Any]:
        client = client_factory()
        try:
            task = await client.excel_writeback_request("GET", f"/tasks/{task_id}")
            output = _editable(Path(download_dir()), str(task_id), task["filename"])
            baseline = _baseline(output)
            if not baseline.is_file() or file_sha256(baseline) != task["base_sha256"]:
                raise ValueError("WRITEBACK_VERSION_CONFLICT: 本地冻结基线缺失或已变化")
            target = payload.manifest.target
            # 布局在载荷中是 DatasetLayout 模型，任务记录里是等价的 dict；
            # 直接比较两者恒为 False，必须先归一化为 JSON 形态再比。
            target_layout = target.layout.model_dump() if target.layout else None
            if (str(target.dataset_id or "") != task["dataset_id"]
                    or str(target.document_version_id or "") != task["document_version_id"]
                    or target.filename != task["filename"]
                    or (target.sheet_name or "Sheet1") != (task["sheet_name"] or "Sheet1")
                    or target.labels != task["labels"]
                    or target_layout != (task["layout"] or None)):
                raise ValueError("WRITEBACK_TARGET_CHANGED: 候选目标与 checkout 不一致")
            if len(rows) != len(payload.records):
                raise ValueError("WRITEBACK_ROWS_MISMATCH: 数据行与来源候选数量不一致")
            if [record.source_index for record in payload.records] != sorted(
                    record.source_index for record in payload.records):
                raise ValueError("WRITEBACK_ROWS_MISMATCH: 多来源记录必须按来源顺序排列")
            if payload.manifest.write_scope == "A:I":
                left = _bounds(task["layout"]["header_range"])[0] if task["layout"] else 1
                if left != 1:
                    raise ValueError("SCOPED_WRITE_SCOPE_INVALID: 目标数据区必须从 A 列开始")
                rows = _scoped_a_i_rows(rows, len(task["labels"]))
            elif payload.manifest.write_scope == "A:I,L":
                rows = _scoped_identity_review_rows(rows, len(task["labels"]))
            elif payload.manifest.write_scope == "A:AA":
                rows = _full_professor_rows(rows, len(task["labels"]))
            patch_workbook(baseline=baseline, output=output, sheet_name=task["sheet_name"] or "Sheet1",
                           labels=task["labels"], rows=rows, layout=task["layout"],
                           scoped_a_i=payload.manifest.write_scope in {"A:I", "A:I,L", "A:AA"})
            body = payload.model_dump(mode="json")
            sha = file_sha256(output)
            digest = writeback_digest(body, sha)
            result = await client.excel_writeback_request("POST", f"/tasks/{task_id}/preview",
                {"result_digest": digest, "file_sha256": sha,
                 "extraction_confirmations": body["manifest"].get("extraction_confirmations", [])})
            return {**result, "file_sha256": sha, "result_digest": digest,
                    "preview_result": body, "display_required": task["review_mode"] == "REVIEW_REQUIRED"}
        finally:
            await client.close()

    @mcp.tool(name="excel_writeback_preview", description=(
        "WorkBuddy 已在 checkout 返回的固定本地 XLSX 副本中写入数据后，登记完整文件与逐列证据的共同指纹。"
        "复核模式必须向用户展示完整候选并等待确认。"), structured_output=True)
    async def preview(task_id: UUID, payload: TaskPayload) -> dict[str, Any]:
        client = client_factory()
        try:
            task = await client.excel_writeback_request("GET", f"/tasks/{task_id}")
            output = _editable(Path(download_dir()), str(task_id), task["filename"])
            if not output.is_file():
                raise ValueError("WRITEBACK_FILE_MISSING: WorkBuddy 尚未生成结果 Excel")
            body = payload.model_dump(mode="json")
            sha = file_sha256(output)
            digest = writeback_digest(body, sha)
            result = await client.excel_writeback_request("POST", f"/tasks/{task_id}/preview",
                {"result_digest": digest, "file_sha256": sha,
                 "extraction_confirmations": body["manifest"].get("extraction_confirmations", [])})
            return {**result, "file_sha256": sha, "result_digest": digest,
                    "preview_result": body,
                    "display_required": task["review_mode"] == "REVIEW_REQUIRED"}
        finally:
            await client.close()

    @mcp.tool(name="excel_writeback_confirm", description=(
        "用户确认已展示的 Excel 逐列预览后，固定同一结果指纹；不修改本地文件。"), structured_output=True)
    async def confirm(task_id: UUID, result_digest: Annotated[str, Field(pattern=r"^[0-9a-f]{64}$")]) -> dict[str, Any]:
        client = client_factory()
        try:
            return await client.excel_writeback_request("POST", f"/tasks/{task_id}/confirm",
                {"result_digest": result_digest})
        finally:
            await client.close()

    @mcp.tool(name="excel_writeback_publish", description=(
        "只上传 checkout 固定缓存中的完整 XLSX；File Agent 核对目标版本、逐列证据和模板差异后发布同名新版本。"
        "失败时不能宣称 Excel 已写入。"), structured_output=False)
    async def publish(task_id: UUID, result_digest: Annotated[str, Field(pattern=r"^[0-9a-f]{64}$")],
                      payload: TaskPayload):
        client = client_factory()
        try:
            task = await client.excel_writeback_request("GET", f"/tasks/{task_id}")
            output = _editable(Path(download_dir()), str(task_id), task["filename"])
            if not output.is_file():
                raise ValueError("WRITEBACK_FILE_MISSING: 本地结果 Excel 不存在")
            body = payload.model_dump(mode="json")
            sha = file_sha256(output)
            if writeback_digest(body, sha) != result_digest:
                raise ValueError("WRITEBACK_PREVIEW_CHANGED: 结果 Excel 或候选已变化，请重新预览")
            result = await client.upload_excel_writeback(
                task_id=str(task_id), result_digest=result_digest, file_sha256=sha,
                payload=body, path=output, filename=task["filename"])
            if not result.get("excel_published") or result.get("status") != "COMPLETED":
                return result
            resource = await client.download_dataset_export(
                dataset_id=result["dataset_id"], export_id=result["export_id"],
                output_dir=Path(download_dir()))
            return CallToolResult(
                content=[TextContent(type="text", text=json.dumps(result, ensure_ascii=False)),
                         ResourceLink(type="resource_link", name=result["filename"],
                                      uri=resource["resource_uri"], mimeType=resource["content_type"],
                                      size=resource["size_bytes"])],
                structuredContent={**result, "downloaded": True},
            )
        finally:
            await client.close()

    @mcp.tool(name="excel_writeback_correction_checkout", description=(
        "按服务器同名 Excel、稳定记录 ID 和真实行号取得已发布行的更正基线；"
        "复用已归档 PDF 页面结果，不重新 OCR；原追加入口不受影响。"), structured_output=False)
    async def correction_checkout(
        filename: Annotated[str, Field(min_length=1, max_length=240)],
        request_key: Annotated[str, Field(min_length=1, max_length=120)],
        record_id: Annotated[str, Field(min_length=1, max_length=100)],
        row_number: Annotated[int, Field(ge=1, le=100000)],
        source_extraction_run_id: UUID | None = None,
        review_mode: Literal["DIRECT_SAVE", "REVIEW_REQUIRED"] = "DIRECT_SAVE",
    ):
        """更正已有非空格须唯一绑定已发布行；仅显式复核模式等待确认。"""
        client = client_factory()
        try:
            task = await client.excel_writeback_request("POST", "/corrections/checkout", {
                "filename": filename, "request_key": request_key,
                "record_id": record_id, "row_number": row_number,
                "review_mode": review_mode,
                "source_extraction_run_id": str(source_extraction_run_id) if source_extraction_run_id else None})
            resource = await client.download_working_copy(
                working_copy_id=task["working_copy_id"], output_dir=Path(download_dir()))
            if resource["sha256"] != task["base_sha256"]:
                raise ValueError("WRITEBACK_VERSION_CONFLICT: 更正基线下载后已变化")
            output = _editable(Path(download_dir()), task["task_id"], task["filename"])
            uri = urlparse(resource["resource_uri"])
            if uri.scheme != "file" or uri.netloc:
                raise ValueError("WRITEBACK_CACHE_INVALID: 下载结果不是本机文件")
            baseline = _baseline(output)
            if not baseline.exists():
                shutil.copyfile(Path(url2pathname(uri.path)), baseline)
            if file_sha256(baseline) != task["base_sha256"]:
                raise ValueError("WRITEBACK_VERSION_CONFLICT: 本地更正基线已变化")
            if not output.exists():
                shutil.copyfile(baseline, output)
            receipt = {**task, "editable_resource_uri": output.as_uri(),
                       "baseline_resource_uri": resource["resource_uri"]}
            return CallToolResult(
                content=[TextContent(type="text", text=json.dumps(receipt, ensure_ascii=False)),
                         ResourceLink(type="resource_link", name=task["filename"], uri=output.as_uri(),
                                      mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                      size=output.stat().st_size)], structuredContent=receipt)
        finally:
            await client.close()

    @mcp.tool(name="excel_writeback_correction_apply", description=(
        "WorkBuddy 提交逐格旧值、新值和正式来源证据；仅在固定行写本地 XLSX，"
        "登记完整文件指纹；显式复核模式向用户展示更正预览。"), structured_output=True)
    async def correction_apply(task_id: UUID, payload: CorrectionPayload) -> dict[str, Any]:
        """不调用旧追加 writer，禁止更正不属于本任务的行或列。"""
        client = client_factory()
        try:
            task = await client.excel_writeback_request("GET", f"/corrections/tasks/{task_id}")
            output = _editable(Path(download_dir()), str(task_id), task["filename"])
            baseline = _baseline(output)
            if not baseline.is_file() or file_sha256(baseline) != task["base_sha256"]:
                raise ValueError("WRITEBACK_VERSION_CONFLICT: 更正基线缺失或已变化")
            if payload.record_id != task["record_id"] or payload.row_number != task["row_number"]:
                raise ValueError("CORRECTION_RECORD_CHANGED: 更正指向不同记录或行")
            source = payload.source
            if (source.extraction_task_id or str(source.document_id) != task["source_document_id"]
                    or str(source.document_version_id) != task["source_document_version_id"]
                    or str(source.extraction_run_id) != task["source_extraction_run_id"]
                    or source.source_sha256 != task["source_sha256"]):
                raise ValueError("CORRECTION_SOURCE_CHANGED: 更正来源必须是冻结的已归档版本")
            left = _bounds(task["layout"]["header_range"])[0] if task["layout"] else 1
            offsets = {key: left + index for index, key in enumerate(task["field_keys"])}
            changes = {}
            for change in payload.changes:
                if change.field_key not in offsets:
                    raise ValueError("UNKNOWN_FIELD: 更正字段不在目标表")
                before = task["base_values"][change.field_key]
                if (before is None) != (change.before is None) or (before is not None and str(before) != str(change.before)):
                    raise ValueError("CORRECTION_BEFORE_CHANGED: 旧值与冻结行不一致")
                coordinate = f"{_column(offsets[change.field_key])}{task['row_number']}"
                changes[coordinate] = (before, change.after)
            patch_existing_cells(baseline=baseline, output=output,
                sheet_name=task["sheet_name"] or "Sheet1", row_number=task["row_number"], changes=changes)
            body = payload.model_dump(mode="json")
            sha = file_sha256(output)
            digest = writeback_digest(body, sha)
            result = await client.excel_writeback_request("POST", f"/corrections/tasks/{task_id}/preview", {
                "result_digest": digest, "file_sha256": sha, "extraction_confirmations": []})
            return {**result, "file_sha256": sha, "result_digest": digest,
                    "preview_result": body, "display_required": task["review_mode"] == "REVIEW_REQUIRED"}
        finally:
            await client.close()

    @mcp.tool(name="excel_writeback_correction_confirm", description=(
        "显式复核模式下，用户看过逐格差异后确认同一指纹。"), structured_output=True)
    async def correction_confirm(task_id: UUID,
        result_digest: Annotated[str, Field(pattern=r"^[0-9a-f]{64}$")]) -> dict[str, Any]:
        """只确认当前候选，不执行文件发布。"""
        client = client_factory()
        try:
            return await client.excel_writeback_request("POST", f"/corrections/tasks/{task_id}/confirm",
                {"result_digest": result_digest})
        finally:
            await client.close()

    @mcp.tool(name="excel_writeback_correction_publish", description=(
        "只回传已预览登记的固定缓存 XLSX；服务端核验来源、逐格差异及版本后发布同名新版本。"),
        structured_output=False)
    async def correction_publish(task_id: UUID,
        result_digest: Annotated[str, Field(pattern=r"^[0-9a-f]{64}$")], payload: CorrectionPayload):
        """成功以服务器发布回执为准，不能把本地候选当成最终文件。"""
        client = client_factory()
        try:
            task = await client.excel_writeback_request("GET", f"/corrections/tasks/{task_id}")
            output = _editable(Path(download_dir()), str(task_id), task["filename"])
            if not output.is_file():
                raise ValueError("WRITEBACK_FILE_MISSING: 本地更正结果不存在")
            body = payload.model_dump(mode="json")
            sha = file_sha256(output)
            if writeback_digest(body, sha) != result_digest:
                raise ValueError("CORRECTION_PREVIEW_CHANGED: 更正结果或候选已变化")
            result = await client.upload_excel_writeback(task_id=str(task_id), result_digest=result_digest,
                file_sha256=sha, payload=body, path=output, filename=task["filename"], correction=True)
            if result.get("status") != "COMPLETED" or not result.get("excel_published"):
                return result
            resource = await client.download_dataset_export(dataset_id=result["dataset_id"],
                export_id=result["export_id"], output_dir=Path(download_dir()))
            return CallToolResult(
                content=[TextContent(type="text", text=json.dumps(result, ensure_ascii=False)),
                         ResourceLink(type="resource_link", name=result["filename"],
                                      uri=resource["resource_uri"], mimeType=resource["content_type"],
                                      size=resource["size_bytes"])],
                structuredContent={**result, "downloaded": True})
        finally:
            await client.close()

    @mcp.tool(name="excel_manual_revision_checkout", description=(
        "在用户手工编辑前，按完整文件名读取同一张已发布 Excel 的当前版本与文件，"
        "返回必须在回传时原样提交的基线版本 ID 和 SHA。"), structured_output=False)
    async def manual_revision_checkout(
        target_filename: Annotated[str, Field(min_length=1, max_length=240)],
    ):
        client = client_factory()
        try:
            target = await client.dataset_request(
                "GET", f"/targets/by-filename?filename={quote(target_filename, safe='')}")
            if (target.get("status") != "ACTIVE" or target.get("filename") != target_filename
                    or not target.get("working_copy_id") or not target.get("document_version_id")):
                raise ValueError("MANUAL_EXCEL_TARGET_INVALID: 目标必须是唯一已发布的同名 XLSX")
            resource = await client.download_working_copy(
                working_copy_id=target["working_copy_id"], output_dir=Path(download_dir()))
            receipt = {"filename": target_filename, "dataset_id": target["dataset_id"],
                       "working_copy_id": target["working_copy_id"],
                       "base_document_version_id": target["document_version_id"],
                       "base_sha256": resource["sha256"], "base_revision": target["revision"]}
            return CallToolResult(
                content=[TextContent(type="text", text=json.dumps(receipt, ensure_ascii=False)),
                         ResourceLink(type="resource_link", name=target_filename,
                                      uri=resource["resource_uri"], mimeType=resource["content_type"],
                                      size=resource["size_bytes"])], structuredContent=receipt)
        finally:
            await client.close()

    @mcp.tool(name="excel_manual_revision_publish", description=(
        "从本轮受信 WorkBuddy 附件接收一份手工补全的同名 XLSX，"
        "必须提供编辑前 checkout 固定的版本 ID 和 SHA，发布为同一 Excel 的新版本。"),
        structured_output=False)
    async def manual_revision_publish(
        submission_ref: str,
        target_filename: Annotated[str, Field(min_length=1, max_length=240)],
        base_document_version_id: UUID,
        base_sha256: Annotated[str, Field(pattern=r"^[0-9a-f]{64}$")],
        base_revision: Annotated[int, Field(ge=0)],
    ):
        client = client_factory()
        try:
            registry = WorkBuddyAttachmentRegistry.from_environment()
            transfer = AttachmentTransferService(client, registry)
            submission, attachment = WorkBuddySubmissionService.from_environment(
                transfer=transfer, registry=registry).manual_excel_revision_attachment(
                    submission_ref=submission_ref, target_filename=target_filename)
            request_key = hashlib.sha256(submission.submission_id.encode()).hexdigest()[:32]
            result = await client.upload_manual_excel_revision(
                filename=target_filename, request_key=request_key,
                base_document_version_id=str(base_document_version_id),
                base_sha256=base_sha256, base_revision=base_revision,
                file_sha256=file_sha256(attachment.local_path), path=attachment.local_path)
            if result.get("status") != "COMPLETED" or not result.get("excel_published"):
                return result
            resource = await client.download_dataset_export(
                dataset_id=result["dataset_id"], export_id=result["export_id"],
                output_dir=Path(download_dir()))
            return CallToolResult(
                content=[TextContent(type="text", text=json.dumps(result, ensure_ascii=False)),
                         ResourceLink(type="resource_link", name=result["filename"],
                                      uri=resource["resource_uri"], mimeType=resource["content_type"],
                                      size=resource["size_bytes"])],
                structuredContent={**result, "downloaded": True})
        finally:
            await client.close()
